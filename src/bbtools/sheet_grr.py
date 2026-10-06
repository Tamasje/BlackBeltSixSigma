"""Gage R&R sheet: average and range method and ANOVA method for a crossed study (k operators × n parts × r trials).

Course (Ottoy, Les 5, Black Belt in Six Sigma - Measurement System Analysis.pdf):
- Average and range (p. 34-35): EV = R̿ · K1, K1 = 1/d2 (subgroup size r); AV = √((X̄_DIFF · K2)² − EV²/(n r)),
  K2 = 1/d2* (one subgroup of size k); PV = √((R_p · K3)² − EV²/(k r)), K3 = 1/d2* (one subgroup of size n);
  GRR² = EV² + AV², TV² = GRR² + PV², %GRR = GRR / TV. Constants from tabel MSA.pdf: d2 from its g → ∞ row,
  d2* from its g = 1 row (as the course workbook does: d2 1.12838, d2*(3) 1.91155, d2*(5) 2.48124).
- ANOVA (p. 36-37): two-way model without interaction; EV = √MSe, AV = √((MSa − EV²)/(n r)),
  PV = √((MSp − EV²)/(k r)), error df = nk(r − 1) + (k − 1)(n − 1).
- %GRR against tolerance 6·σ_GRR / TOL (p. 24); thresholds 10 % / 30 % (p. 35).
Decision 11: multiplier 6; %GRR against total variation and tolerance; the ANOVA with interaction is shown too, with
its p-value, and no pooling α is chosen; ndc omitted (not in the course).

Row plan: data 8-22; per operator and part 24-32; study size 34-39; average & range 41-56; ANOVA 58-72;
ANOVA with interaction 74-81; thresholds 83-88.
"""
from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.sheet_tables import d2_star_formula, msa_d2_formula
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    constant_row,
    input_cell,
    input_row,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Gage R&R"

HEADER = HeaderBlock(
    tool="Gage R&R: average and range method and ANOVA method (EV, AV, PV, GRR, TV, %GRR)",
    source="source/course/Les 5/20260619_ottoy_Black Belt in Six Sigma - Measurement System Analysis.pdf p. 18, 24, 34-38; "
           "20260619_ottoy_tabel MSA.pdf; 20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx",
    convention="Decision 11: multiplier 6; %GRR against total variation and tolerance; course ANOVA = model without "
               "interaction (interaction also shown); constants from tabel MSA.pdf; ndc not in the course.",
    status=Status.VERIFIED,
    status_detail="tested against course worked examples S10-WE02 (ANOVA) and S10-WE03 (average and range) on the "
                  "course's own study data",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Repeatability (EV), reproducibility (AV), part variation (PV), GRR, TV and %GRR by the average and range "
            "method and by the ANOVA method, with the course's 10 % / 30 % verdict; the ANOVA with the interaction term.",
    inputs="the measurements of a complete crossed study: up to 3 operators × 3 trials (rows) × 10 parts (columns); "
           "optionally the tolerance USL − LSL and α for the F tests.",
    audit="not needed (course worked examples exist)",
    disagreements=(
        "GRR workbook sheet '2way anova' K45 types '=412.5+296.667' (the interaction SS 296.6667 rounded), so its EV² is "
        "30.8333478 instead of 30.8333333; its AV, PV, TV and %GRR shift in the 7th digit. The sheet computes from the "
        "data; the tests compare at a tolerance that allows exactly this.",
    ),
)

PARTS, OPERATORS, TRIALS = 10, 3, 3
FIRST_DATA_ROW = 13
PART_COLUMNS = "BCDEFGHIJK"
OPERATOR_NAMES = ("A", "B", "C")
NUMBER = "0.000000"
PERCENT = "0.00%"
TOL, ALPHA, OK = "$B$9", "$B$10", '$B$39="yes"'
INPUTS = {"tolerance": "B9", "alpha": "B10"}
RESULTS = {
    "k": "B35", "n": "B36", "r": "B37", "count": "B38", "complete": "B39",
    "rbarbar": "B42", "xdiff": "B43", "rp": "B44", "d2": "B45", "d2star_k": "B46", "d2star_n": "B47",
    "ev_ar": "B48", "av_ar": "B49", "pv_ar": "B50", "grr_ar": "B51", "tv_ar": "B52", "pct_ar": "B53", "pct_tol_ar": "B54",
    "ev_anova": "B64", "av_anova": "B65", "pv_anova": "B66", "grr_anova": "B67", "tv_anova": "B68",
    "pct_anova": "B69", "pct_tol_anova": "B70",
}
ANOVA_ROWS = {"parts": 60, "operators": 61, "error": 62, "total": 63}
INTERACTION_ROWS = {"parts": 76, "operators": 77, "interaction": 78, "within": 79, "total": 80}


def data_cell(operator: int, trial: int, part: int) -> str:
    """Input cell of one measurement (0-based operator, trial and part)."""
    return f"{PART_COLUMNS[part]}{FIRST_DATA_ROW + TRIALS * operator + trial}"


def _block(operator: int) -> str:
    """All measurements of one operator (0-based)."""
    top = FIRST_DATA_ROW + TRIALS * operator
    return f"B{top}:K{top + TRIALS - 1}"


def _data(ws: Worksheet) -> None:
    """Section 1: the study data."""
    section_title(ws, 8, "1. Study data: one row per operator and trial, one column per part (coded or real values)")
    input_row(ws, 9, "Tolerance TOL = USL − LSL (optional)", "for %GRR against tolerance")
    input_row(ws, 10, "α for the F tests", "course has no default; most course examples use 5 %", "0.00")
    ws[ALPHA.replace("$", "")] = 0.05
    label(ws, 11, 1, "Fill a complete study: every operator used measures every part used in every trial. "
                     "Up to 3 operators × 3 trials × 10 parts.", italic=True)
    column_titles(ws, 12, ["Operator, trial"] + [f"Part {j + 1}" for j in range(PARTS)])
    for operator in range(OPERATORS):
        for trial in range(TRIALS):
            row = FIRST_DATA_ROW + TRIALS * operator + trial
            label(ws, row, 1, f"Operator {OPERATOR_NAMES[operator]}, trial {trial + 1}", bold=trial == 0)
            for part in range(PARTS):
                input_cell(ws, data_cell(operator, trial, part))
    label(ws, 22, 1, "values per part")
    for part, letter in enumerate(PART_COLUMNS):
        output_cell(ws, f"{letter}22", f"=COUNT({letter}13:{letter}21)", "0")


def _per_cell(ws: Worksheet) -> None:
    """Section 2: range and mean per operator and part, operator means and part means."""
    section_title(ws, 24, "2. Per operator and part (computed)")
    column_titles(ws, 25, ["", *[f"Part {j + 1}" for j in range(PARTS)], "Operator mean"])
    for operator in range(OPERATORS):
        top = FIRST_DATA_ROW + TRIALS * operator
        label(ws, 26 + operator, 1, f"Range over trials, operator {OPERATOR_NAMES[operator]}")
        label(ws, 29 + operator, 1, f"Mean over trials, operator {OPERATOR_NAMES[operator]}")
        for letter in PART_COLUMNS:
            trials = f"{letter}{top}:{letter}{top + TRIALS - 1}"
            output_cell(ws, f"{letter}{26 + operator}", f'=IF(COUNT({trials})>1,MAX({trials})-MIN({trials}),"")', NUMBER)
            output_cell(ws, f"{letter}{29 + operator}", f'=IF(COUNT({trials})>0,AVERAGE({trials}),"")', NUMBER)
        output_cell(ws, f"L{29 + operator}", f'=IF(COUNT({_block(operator)})>0,AVERAGE({_block(operator)}),"")', NUMBER)
    label(ws, 32, 1, "Part mean (all operators and trials)")
    for letter in PART_COLUMNS:
        output_cell(ws, f"{letter}32", f'=IF(COUNT({letter}13:{letter}21)>0,AVERAGE({letter}13:{letter}21),"")', NUMBER)


def _size(ws: Worksheet) -> None:
    """Section 3: k, n, r and the completeness check."""
    section_title(ws, 34, "3. Study size (derived from the data)")
    operators = "+".join(f"(COUNT({_block(i)})>0)" for i in range(OPERATORS))
    trials = "+".join(f"(COUNT(B{FIRST_DATA_ROW + t}:K{FIRST_DATA_ROW + t})>0)" for t in range(TRIALS))
    result_row(ws, 35, "k = number of operators", f"={operators}", "0")
    result_row(ws, 36, "n = number of parts", '=COUNTIF(B22:K22,">0")', "0")
    result_row(ws, 37, "r = number of trials (of operator A)", f"={trials}", "0")
    result_row(ws, 38, "number of measurements", "=COUNT(B13:K21)", "0")
    result_row(ws, 39, "Complete study (k ≥ 2, n ≥ 2, r ≥ 2, k·n·r values)?",
               '=IF(AND(B35>=2,B36>=2,B37>=2,B38=B35*B36*B37),"yes",IF(B38=0,"",'
               '"NO: every operator must measure every part in every trial"))')


def _verdict(value: str) -> str:
    """Course verdict for a %GRR (MSA p. 35), using the thresholds of section 7."""
    return (f'=IF(ISNUMBER({value}),IF({value}<=$B$84,"acceptable (≤ 10 %)",IF({value}<=$B$85,'
            f'"may be acceptable (10 - 30 %)","not acceptable (> 30 %)")),"")')


def _results(ws: Worksheet, first: int, ev: str, av: str, pv: str, source: str) -> None:
    """EV, AV, PV, GRR, TV, %GRR (TV and tolerance) and the two verdicts, from row `first`."""
    r = first
    result_row(ws, r, "EV = repeatability", f'=IF({OK},{ev},"")', NUMBER, source)
    result_row(ws, r + 1, "AV = reproducibility", f'=IF({OK},{av},"")', NUMBER)
    result_row(ws, r + 2, "PV = part variation", f'=IF({OK},{pv},"")', NUMBER)
    result_row(ws, r + 3, "GRR = √(EV² + AV²)", f'=IF(ISNUMBER(B{r}),SQRT(B{r}^2+B{r + 1}^2),"")', NUMBER, "MSA p. 18, 35")
    result_row(ws, r + 4, "TV = √(GRR² + PV²)", f'=IF(ISNUMBER(B{r + 3}),SQRT(B{r + 3}^2+B{r + 2}^2),"")', NUMBER)
    result_row(ws, r + 5, "%GRR of total variation = GRR / TV", f'=IF(ISNUMBER(B{r + 4}),B{r + 3}/B{r + 4},"")', PERCENT,
               "MSA p. 35")
    result_row(ws, r + 6, "%GRR of tolerance = 6 · GRR / TOL",
               f'=IF(AND(ISNUMBER(B{r + 3}),ISNUMBER({TOL}),{TOL}>0),6*B{r + 3}/{TOL},"")', PERCENT, "MSA p. 24")
    # verdicts in column D: column C holds the course source of each row
    output_cell(ws, f"D{r + 5}", _verdict(f"B{r + 5}"))
    output_cell(ws, f"D{r + 6}", _verdict(f"B{r + 6}"))


def _average_range(ws: Worksheet) -> None:
    """Section 4: average and range method."""
    section_title(ws, 41, "4. Average and range method (MSA p. 34-35)")
    result_row(ws, 42, "R̿ = mean of all ranges", f'=IF({OK},AVERAGE(B26:K28),"")', NUMBER)
    result_row(ws, 43, "X̄_DIFF = largest − smallest operator mean", f'=IF({OK},MAX(L29:L31)-MIN(L29:L31),"")', NUMBER)
    result_row(ws, 44, "R_p = largest − smallest part mean", f'=IF({OK},MAX(B32:K32)-MIN(B32:K32),"")', NUMBER)
    result_row(ws, 45, "d2 for m = r (g → ∞): K1 = 1/d2", f'=IF({OK},IFERROR({msa_d2_formula("$B$37")},"r not in table"),"")',
               "0.00000", "tabel MSA.pdf, last row")
    result_row(ws, 46, "d2* for m = k, g = 1: K2 = 1/d2*",
               f'=IF({OK},IFERROR({d2_star_formula("1", "$B$35")},"k not in table"),"")', "0.00000", "tabel MSA.pdf, g = 1")
    result_row(ws, 47, "d2* for m = n, g = 1: K3 = 1/d2*",
               f'=IF({OK},IFERROR({d2_star_formula("1", "$B$36")},"n not in table"),"")', "0.00000", "tabel MSA.pdf, g = 1")
    _results(ws, 48, "B42/B45", "SQRT(MAX(0,(B43/B46)^2-(B42/B45)^2/(B36*B37)))",
             "SQRT(MAX(0,(B44/B47)^2-(B42/B45)^2/(B35*B37)))", "MSA p. 35")


def _anova_table(ws: Worksheet, rows: dict[str, int], entries: dict[str, tuple[str, str, str, str]], denominator: str) -> None:
    """Rows of an ANOVA table: SS, df, MS and for the factors F, p-value and F crit against `denominator`'s MS."""
    for key, (text, ss, df, _) in entries.items():
        r = rows[key]
        label(ws, r, 1, text)
        output_cell(ws, f"B{r}", f'=IF({OK},{ss},"")', "0.0000")
        output_cell(ws, f"C{r}", f'=IF({OK},{df},"")', "0")
        if key != "total":
            output_cell(ws, f"D{r}", f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),N(C{r})>0),B{r}/C{r},"")', "0.0000")
    denominator_row = rows[denominator]
    for key, (_, _, _, is_factor) in entries.items():
        if is_factor:
            r = rows[key]
            ms_e = f"$D${denominator_row}"
            output_cell(ws, f"E{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER({ms_e}),N({ms_e})>0),D{r}/{ms_e},"")', "0.0000")
            output_cell(ws, f"F{r}", f'=IF(ISNUMBER(E{r}),_xlfn.F.DIST.RT(E{r},C{r},$C${denominator_row}),"")', "0.000000")
            output_cell(ws, f"G{r}", f'=IF(ISNUMBER(E{r}),_xlfn.F.INV.RT({ALPHA},C{r},$C${denominator_row}),"")', "0.0000")


def _anova(ws: Worksheet) -> None:
    """Sections 5-6: ANOVA without interaction (the course's method) and with interaction (to judge it)."""
    ss_parts, ss_operators = "B35*B37*DEVSQ(B32:K32)", "B36*B37*DEVSQ(L29:L31)"
    ss_total, ss_cells = "DEVSQ(B13:K21)", "B37*DEVSQ(B29:K31)"
    section_title(ws, 58, "5. ANOVA method (MSA p. 36-37): two-way model without interaction, as the course")
    column_titles(ws, 59, ["Source", "SS", "df", "MS", "F", "p-value", "F crit at α"])
    _anova_table(ws, ANOVA_ROWS, {
        "parts": ("Parts", ss_parts, "B36-1", "factor"),
        "operators": ("Operators", ss_operators, "B35-1", "factor"),
        "error": ("Error (repeatability + interaction)", f"{ss_total}-{ss_parts}-{ss_operators}",
                  "B36*B35*(B37-1)+(B35-1)*(B36-1)", ""),
        "total": ("Total", ss_total, "B36*B35*B37-1", ""),
    }, "error")
    _results(ws, 64, "SQRT(D62)", "SQRT(MAX(0,(D61-D62)/(B36*B37)))", "SQRT(MAX(0,(D60-D62)/(B35*B37)))", "MSA p. 36")

    section_title(ws, 74, "6. ANOVA with the operator × part interaction (as Excel 'Anova: two-factor with replication')")
    column_titles(ws, 75, ["Source", "SS", "df", "MS", "F", "p-value", "F crit at α"])
    _anova_table(ws, INTERACTION_ROWS, {
        "parts": ("Parts", ss_parts, "B36-1", "factor"),
        "operators": ("Operators", ss_operators, "B35-1", "factor"),
        "interaction": ("Interaction operator × part", f"{ss_cells}-{ss_parts}-{ss_operators}", "(B36-1)*(B35-1)", "factor"),
        "within": ("Within (repeatability)", f"{ss_total}-{ss_cells}", "B36*B35*(B37-1)", ""),
        "total": ("Total", ss_total, "B36*B35*B37-1", ""),
    }, "within")
    label(ws, 81, 1, "A large interaction p-value supports pooling it into the error, as section 5 does (the course "
                     "workbook's 'afgeleide ANOVA-tabel zonder interactie'). The course gives EV, AV, PV only for section 5.",
          italic=True)


def _thresholds(ws: Worksheet) -> None:
    """Section 7: the course's acceptance rule and the other thresholds it mentions."""
    section_title(ws, 83, "7. Acceptance of the measurement system")
    constant_row(ws, 84, "%GRR ≤ this: acceptable", 0.10, "MSA p. 35 ('%GRR <= 10% : gauge generally considered to be acceptable')")
    constant_row(ws, 85, "%GRR ≤ this: may be acceptable; above: not acceptable", 0.30,
                 "MSA p. 35 (10 % - 30 %: depends on the application; > 30 %: not acceptable)")
    ws["B84"].number_format = ws["B85"].number_format = "0%"
    label(ws, 86, 1, "Other course thresholds: 'Gauge R&R ≤ 20 % of tolerance' (deck Les 4 p. 61); σ²_measure / σ²_observed "
                     "≤ 0.1 good, 0.1-0.3 marginal, ≥ 0.3 unacceptable (Dummies p. 178).", italic=True)
    label(ws, 87, 1, "Multiplier 6 (99.73 %) as in the MSA formulas; historically 5.15 (99 %), MSA p. 24. "
                     "ndc (number of distinct categories) is not in the course.", italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the Gage R&R calculator."""
    write_header(ws, HEADER)
    _data(ws)
    _per_cell(ws)
    _size(ws)
    _average_range(ws)
    _anova(ws)
    _thresholds(ws)
    ws.column_dimensions["A"].width = 50
    for letter in "BCDEFGHIJKL":
        ws.column_dimensions[letter].width = 13
