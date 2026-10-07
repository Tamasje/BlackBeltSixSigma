"""Mean & proportion sheet ('Gemiddelde & proportie'): confidence intervals and tests for one mean (z, t), two
means (unpaired pooled t, paired t), one proportion (normal approximation, exact binomial, Z-test) and two proportions.

Course (Ottoy, Les 2): CI for μ with t (Confidence Intervals.pdf p. 9-10; CI Further Reading (Dutch) p. 5-8), for
μ1 − μ2 unpaired and paired (p. 9-18), for π (p. 20, normal approximation and exact binomial, cross-checked there with
R binom.test); test recipes: t-test μ, unpaired and paired t-test, Z-test π with n·π0 > 5 (Test Recipes p. 4-10);
z-test with σ known (Testing of Hypotheses Further Reading p. 10); Dummies p. 189-198.
Decision 5: α input prefilled 0.05. Decision 6: two-sided, lower-only and upper-only side by side. Decision 8: pooled
variance with n1 + n2 − 2 (Test Recipes p. 5; CI Further Reading p. 15 prints n1 + n2 − 1).
Exact interval for π = the Clopper-Pearson interval of R binom.test, computed with BETA.INV.

Row plan: settings 8-9; one mean 11-39; unpaired 41-70; paired 72-92; one proportion 94-110; two proportions 112-125.
Raw data: L10:L509 (sample 1), M10:M509 (sample 2), N = L − M (computed, for paired data).
"""
from __future__ import annotations

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    font,
    input_cell,
    input_row,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Gemiddelde & proportie"

HEADER = HeaderBlock(
    tool="Gemiddelden en proporties: betrouwbaarheidsintervallen (confidence intervals, BI) en toetsen (z, t, "
         "ongepaarde en gepaarde t, proportie exact en benaderd, Z-toets)",
    source="source/course/Les 2/20260529_ottoy_Confidence Intervals.pdf p. 4, 9-15; Confidence Intervals - Further "
           "Reading (Dutch).pdf p. 5-20; Testing of Hypotheses.pdf p. 12; Testing of Hypotheses - Further Reading "
           "(Dutch).pdf p. 10-14; Test Recipes - Further Reading (Dutch).pdf p. 4-10",
    convention="Conventiebeslissingen 5, 6, 8: invoer α vooraf ingevuld op 0,05; tweezijdig, alleen ondergrens en "
               "alleen bovengrens naast elkaar; gepoolde variantie (pooled variance) met n1 + n2 − 2.",
    status=Status.VERIFIED,
    status_detail="getest tegen uitgewerkte cursusvoorbeelden (worked examples) S03-WE01, S03-WE02, S03-WE04, "
                  "S03-WE10, S03-WE13, S03-WE14, S03-WE15; ook tegen S08-WE12, S08-WE13 uit Dummies p. 197 "
                  "(extra, Dummies; niet te kennen)",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Betrouwbaarheidsintervallen (BI) en toetsen voor één gemiddelde (σ gekend: z; σ ongekend: t), het "
            "verschil van twee gemiddelden (ongepaarde gepoolde t, gepaarde t), één proportie (normale benadering, "
            "exact binomiaal, Z-toets) en twee proporties.",
    inputs="α; per blok n, gemiddelde, s of σ, μ0 / d0 / π0; of ruwe data geplakt in kolommen L en M (gepaard: L en M "
           "rij per rij). Proporties: n en het aantal successen x.",
    audit="niet nodig (er bestaan uitgewerkte cursusvoorbeelden)",
    disagreements=(
        "CI Further Reading (Dutch) p. 15 drukt de gepoolde variantie met noemer n1 + n2 − 1; het blad gebruikt "
        "n1 + n2 − 2 zoals Test Recipes p. 5 en de t(n1 + n2 − 2) van dezelfde slide (conventiebeslissing 8). Het "
        "uitgewerkte cursusvoorbeeld S03-WE13 (s_p = 6,59) klopt met n1 + n2 − 2.",
        "Confidence Intervals.xlsx 'example t-test' F13 deelt door SQRT(19) met n = 20 (conventiebeslissing 9); "
        "geen testdoel.",
        "Dummies p. 197 (S08-WE13): '0.08 ± 0.076 = [0.004, 0.156]' kapt de halve breedte 0,0765 af; onafgerond "
        "[0,0035; 0,1565] (extra, Dummies; niet te kennen).",
    ),
)

ALPHA = "B9"
DATA_1, DATA_2, DIFF = "L10:L509", "M10:M509", "N10:N509"
NUMBER = "0.000000"
PERCENT = "0.0000%"
DECISION = '=IF(ISNUMBER({p}),IF({p}<' + ALPHA + ',"verwerp H0","H0 niet verwerpen"),"")'

INPUTS = {
    "alpha": ALPHA, "n": "B12", "mean": "B13", "s": "B14", "sigma": "B15", "mu0": "B16",
    "n1": "B42", "mean1": "B43", "s1": "B44", "n2": "B45", "mean2": "B46", "s2": "B47", "d0": "B48",
    "np": "B73", "vbar": "B74", "sv": "B75", "d0p": "B76",
    "pn": "B95", "px": "B96", "pi0": "B97",
    "qn1": "B113", "qx1": "B114", "qn2": "B115", "qx2": "B116",
}
RESULTS = {
    "n_used": "B17", "mean_used": "B18", "s_used": "B19", "se_sigma": "B20", "se_s": "B21", "df": "B22",
    "n1_used": "B49", "mean1_used": "B50", "s1_used": "B51", "n2_used": "B52", "mean2_used": "B53", "s2_used": "B54",
    "sp": "B55", "df_u": "B56", "se_u": "B57", "diff_u": "B58",
    "np_used": "B77", "vbar_used": "B78", "sv_used": "B79", "se_p": "B80", "df_p": "B81",
    "p": "B98", "se_prop": "B99", "npi0_check": "B100",
    "q1": "B117", "q2": "B118", "qdiff": "B119", "qse": "B120",
}
ONE_MEAN_CI = {"two_sided": 25, "lower_only": 26, "upper_only": 27}     # B,C = z; D,E = t
Z_TEST = {"two_sided": 30, "greater": 31, "less": 32}                  # B z, C/D critical z, E/F critical x̄, G p, H decision
T_TEST = {"two_sided": 35, "greater": 36, "less": 37}                  # same layout with t
UNPAIRED_CI = {"two_sided": 61, "lower_only": 62, "upper_only": 63}    # B from, C to, D half-width, E t used
UNPAIRED_TEST = {"two_sided": 66, "greater": 67, "less": 68}
PAIRED_CI = {"two_sided": 84, "lower_only": 85, "upper_only": 86}
PAIRED_TEST = {"two_sided": 89, "greater": 90, "less": 91}
PROPORTION_CI = {"two_sided": 103, "lower_only": 104, "upper_only": 105}  # B,C normal approx; D,E exact
PROPORTION_TEST = {"two_sided": 108, "greater": 109, "less": 110}
TWO_PROPORTION_CI = {"two_sided": 123, "lower_only": 124, "upper_only": 125}  # B from, C to, D half-width


def _data_columns(ws: Worksheet) -> None:
    """Columns L, M: optional raw data; N: their row-by-row difference (for paired data)."""
    label(ws, 8, 12, "Optioneel: plak hieronder ruwe waarden. n, gemiddelde en s worden er dan uit berekend.",
          italic=True)
    for column, text in (("L", "Steekproef 1"), ("M", "Steekproef 2"), ("N", "L − M (gepaard)")):
        ws[f"{column}9"] = text
        ws[f"{column}9"].font = font(bold=True)
    for row in range(10, 510):
        input_cell(ws, f"L{row}")
        input_cell(ws, f"M{row}")
        output_cell(ws, f"N{row}", f'=IF(AND(ISNUMBER(L{row}),ISNUMBER(M{row})),L{row}-M{row},"")', "General")


def _used(ws: Worksheet, row: int, text: str, data: str, typed: str, kind: str) -> None:
    """'Value used' cell: from the pasted data when it holds at least 2 values, else the typed input."""
    stat = {"n": f"COUNT({data})", "mean": f"AVERAGE({data})", "s": f"_xlfn.STDEV.S({data})"}[kind]
    result_row(ws, row, text, f'=IF(COUNT({data})>1,{stat},IF(ISNUMBER({typed}),{typed},""))',
               "0" if kind == "n" else NUMBER,
               f"uit kolom {data[0]} als die 2+ waarden bevat, anders de getypte waarde")


def _interval_rows(ws: Worksheet, rows: dict[str, int], have: str, centre: str, half_two: str, half_one: str,
                   columns: tuple[str, str], percent: bool = False) -> None:
    """Two-sided centre ± half_two, lower-only centre − half_one, upper-only centre + half_one."""
    low, high = columns
    fmt = PERCENT if percent else NUMBER
    output_cell(ws, f"{low}{rows['two_sided']}", f'=IF({have},{centre}-{half_two},"")', fmt)
    output_cell(ws, f"{high}{rows['two_sided']}", f'=IF({have},{centre}+{half_two},"")', fmt)
    output_cell(ws, f"{low}{rows['lower_only']}", f'=IF({have},{centre}-{half_one},"")', fmt)
    output_cell(ws, f"{high}{rows['lower_only']}", f'=IF({have},"+∞","")', fmt)
    output_cell(ws, f"{low}{rows['upper_only']}", f'=IF({have},"−∞","")', fmt)
    output_cell(ws, f"{high}{rows['upper_only']}", f'=IF({have},{centre}+{half_one},"")', fmt)


def _interval_labels(ws: Worksheet, rows: dict[str, int], what: str) -> None:
    """Row labels of an interval table."""
    label(ws, rows["two_sided"], 1, "tweezijdig (two-sided)")
    label(ws, rows["lower_only"], 1, f"alleen ondergrens ({what} minstens …)")
    label(ws, rows["upper_only"], 1, f"alleen bovengrens ({what} hoogstens …)")


def _test_table(ws: Worksheet, rows: dict[str, int], have: str, stat: str, dist: str, df: str | None,
                labels: tuple[str, str, str]) -> None:
    """Statistic, critical value(s), p-value and decision for HA ≠, >, <, for a z (dist='z') or t statistic."""
    if dist == "z":
        crit_two, crit_one = f"_xlfn.NORM.S.INV(1-{ALPHA}/2)", f"_xlfn.NORM.S.INV(1-{ALPHA})"
        p_two, p_greater, p_less = "2*_xlfn.NORM.S.DIST(-ABS({s}),TRUE)", "_xlfn.NORM.S.DIST(-{s},TRUE)", \
            "_xlfn.NORM.S.DIST({s},TRUE)"
    else:
        crit_two, crit_one = f"_xlfn.T.INV(1-{ALPHA}/2,{df})", f"_xlfn.T.INV(1-{ALPHA},{df})"
        p_two, p_greater, p_less = f"_xlfn.T.DIST.2T(ABS({{s}}),{df})", f"_xlfn.T.DIST.RT({{s}},{df})", \
            f"_xlfn.T.DIST({{s}},{df},TRUE)"
    spec = {
        "two_sided": (labels[0], f"-{crit_two}", crit_two, p_two),
        "greater": (labels[1], crit_one, None, p_greater),
        "less": (labels[2], f"-{crit_one}", None, p_less),
    }
    for key, (text, critical, critical_2, p_value) in spec.items():
        r = rows[key]
        label(ws, r, 1, text)
        output_cell(ws, f"B{r}", f'=IF({have},{stat},"")', "0.0000")
        output_cell(ws, f"C{r}", f'=IF({have},{critical},"")', "0.0000")
        output_cell(ws, f"D{r}", f'=IF({have},{critical_2},"")' if critical_2 else '=""', "0.0000")
        output_cell(ws, f"G{r}", f'=IF({have},{p_value.format(s=f"B{r}")},"")', PERCENT)
        output_cell(ws, f"H{r}", DECISION.format(p=f"G{r}"))


def _one_mean(ws: Worksheet) -> None:
    """Sections 1-2: α, and one mean with σ known (z) or unknown (t)."""
    section_title(ws, 8, "1. Instellingen")
    input_row(ws, 9, "Significantieniveau α (betrouwbaarheid = 1 − α)",
              "de cursus geeft geen standaardwaarde; de meeste cursusvoorbeelden gebruiken 5 %; gebruik de waarde uit "
              "de vraag", "0.0000")
    ws[ALPHA] = 0.05

    section_title(ws, 11, "2. Eén gemiddelde μ")
    input_row(ws, 12, "n", "of plak de data in kolom L")
    input_row(ws, 13, "steekproefgemiddelde (sample mean) x̄", "", NUMBER)
    input_row(ws, 14, "s = steekproefstandaardafwijking (STDEV.S)", "voor de t-rijen", NUMBER)
    input_row(ws, 15, "σ = gekende standaardafwijking van de populatie",
              "voor de z-rijen (alleen als de vraag σ geeft)", NUMBER)
    input_row(ws, 16, "μ0 voor de toets", "", NUMBER)
    _used(ws, 17, "n gebruikt", DATA_1, "B12", "n")
    _used(ws, 18, "x̄ gebruikt", DATA_1, "B13", "mean")
    _used(ws, 19, "s gebruikt", DATA_1, "B14", "s")
    result_row(ws, 20, "standaardfout (standard error) met σ: σ / √n",
               '=IF(AND(ISNUMBER(B15),ISNUMBER(B17)),B15/SQRT(B17),"")', NUMBER)
    result_row(ws, 21, "standaardfout met s: s / √n", '=IF(AND(ISNUMBER(B19),ISNUMBER(B17)),B19/SQRT(B17),"")', NUMBER)
    result_row(ws, 22, "vrijheidsgraden (degrees of freedom, df) n − 1", '=IF(ISNUMBER(B17),B17-1,"")', "0")

    column_titles(ws, 24, ["Betrouwbaarheidsinterval (confidence interval, BI) voor μ (1 − α)", "σ gekend (z): van",
                           "tot", "σ ongekend (t): van", "tot", "Bron in de cursus"])
    _interval_labels(ws, ONE_MEAN_CI, "μ")
    have_z = f"AND(ISNUMBER(B18),ISNUMBER(B20),ISNUMBER({ALPHA}))"
    have_t = f"AND(ISNUMBER(B18),ISNUMBER(B21),B22>0,ISNUMBER({ALPHA}))"
    _interval_rows(ws, ONE_MEAN_CI, have_z, "B18", f"_xlfn.NORM.S.INV(1-{ALPHA}/2)*B20",
                   f"_xlfn.NORM.S.INV(1-{ALPHA})*B20", ("B", "C"))
    _interval_rows(ws, ONE_MEAN_CI, have_t, "B18", f"_xlfn.T.INV(1-{ALPHA}/2,B22)*B21",
                   f"_xlfn.T.INV(1-{ALPHA},B22)*B21", ("D", "E"))
    label(ws, 25, 6, "CI Further Reading p. 5-8; Confidence Intervals.pdf p. 9-10", italic=True)
    label(ws, 27, 6, "Confidence Intervals.pdf p. 15 ('one-sided 98 %-CI ]−∞, 9.98[')", italic=True)

    titles = ["{} toets van H0: μ = μ0", "toetsgrootheid (test statistic)", "kritieke waarde (critical value)",
              "2de kritieke waarde", "kritieke x̄", "2de kritieke x̄", "p-waarde (p-value)", "Besluit bij α"]
    column_titles(ws, 29, [titles[0].format("z-toets (z-test, σ gekend):")] + ["z = (x̄ − μ0) / (σ/√n)"] + titles[2:])
    have = f"AND({have_z},ISNUMBER(B16))"
    _test_table(ws, Z_TEST, have, "(B18-B16)/B20", "z", None, ("HA: μ ≠ μ0", "HA: μ > μ0", "HA: μ < μ0"))
    column_titles(ws, 34, [titles[0].format("t-toets (t-test, σ ongekend):")] + ["t = (x̄ − μ0) / (s/√n)"] + titles[2:])
    have = f"AND({have_t},ISNUMBER(B16))"
    _test_table(ws, T_TEST, have, "(B18-B16)/B21", "t", "B22", ("HA: μ ≠ μ0", "HA: μ > μ0", "HA: μ < μ0"))
    for rows, se in ((Z_TEST, "B20"), (T_TEST, "B21")):
        for r in rows.values():  # critical values on the scale of x̄, as Testing of Hypotheses FR p. 10 does
            output_cell(ws, f"E{r}", f'=IF(ISNUMBER(C{r}),$B$16+C{r}*{se},"")', "0.0000")
            output_cell(ws, f"F{r}", f'=IF(ISNUMBER(D{r}),$B$16+D{r}*{se},"")', "0.0000")
    label(ws, 38, 1, "Bronnen: z-toets Testing of Hypotheses - Further Reading (Dutch) p. 10-14; t-toets Test Recipes "
                     "p. 4, Testing of Hypotheses.pdf p. 12.", italic=True)


def _unpaired(ws: Worksheet) -> None:
    """Section 3: two unpaired means with the pooled standard deviation (decision 8)."""
    section_title(ws, 41, "3. Twee gemiddelden, ongepaarde steekproeven (unpaired; σ1 = σ2 verondersteld: "
                          "gepoolde s_p)")
    input_row(ws, 42, "n1", "of plak steekproef 1 in kolom L")
    input_row(ws, 43, "x̄1", "", NUMBER)
    input_row(ws, 44, "s1", "", NUMBER)
    input_row(ws, 45, "n2", "of plak steekproef 2 in kolom M")
    input_row(ws, 46, "x̄2", "", NUMBER)
    input_row(ws, 47, "s2", "", NUMBER)
    input_row(ws, 48, "d0 = μ1 − μ2 onder H0", "leeg laten voor 0", NUMBER)
    for row, text, data, typed, kind in ((49, "n1 gebruikt", DATA_1, "B42", "n"),
                                         (50, "x̄1 gebruikt", DATA_1, "B43", "mean"),
                                         (51, "s1 gebruikt", DATA_1, "B44", "s"),
                                         (52, "n2 gebruikt", DATA_2, "B45", "n"),
                                         (53, "x̄2 gebruikt", DATA_2, "B46", "mean"),
                                         (54, "s2 gebruikt", DATA_2, "B47", "s")):
        _used(ws, row, text, data, typed, kind)
    have = "AND(ISNUMBER(B49),ISNUMBER(B50),ISNUMBER(B51),ISNUMBER(B52),ISNUMBER(B53),ISNUMBER(B54),N(B49)+N(B52)>2)"  # N(): AND does not short-circuit
    result_row(ws, 55, "s_p = √[((n1 − 1)s1² + (n2 − 1)s2²) / (n1 + n2 − 2)]",
               f'=IF({have},SQRT(((B49-1)*B51^2+(B52-1)*B54^2)/(B49+B52-2)),"")', NUMBER,
               "Test Recipes p. 5 (conventiebeslissing 8; CI Further Reading p. 15 drukt n1 + n2 − 1)")
    result_row(ws, 56, "vrijheidsgraden n1 + n2 − 2", f'=IF({have},B49+B52-2,"")', "0")
    result_row(ws, 57, "standaardfout s_p √(1/n1 + 1/n2)", '=IF(ISNUMBER(B55),B55*SQRT(1/B49+1/B52),"")', NUMBER)
    result_row(ws, 58, "x̄1 − x̄2", f'=IF({have},B50-B53,"")', NUMBER)
    _difference_block(ws, 60, UNPAIRED_CI, UNPAIRED_TEST, "B58", "B57", "B56", "B48", "μ1 − μ2",
                      "CI Further Reading p. 15-17 (S03-WE13: −5,6 ± 6,19)")


def _paired(ws: Worksheet) -> None:
    """Section 4: paired samples via the differences v = sample 1 − sample 2."""
    section_title(ws, 72, "4. Twee gemiddelden, gepaarde steekproeven (paired; verschillen v = steekproef 1 − "
                          "steekproef 2)")
    input_row(ws, 73, "n = aantal paren", "of plak de paren in kolommen L en M (zelfde rijen)")
    input_row(ws, 74, "v̄ = gemiddeld verschil", "", NUMBER)
    input_row(ws, 75, "s_v = standaardafwijking van de verschillen", "", NUMBER)
    input_row(ws, 76, "d0 onder H0", "leeg laten voor 0", NUMBER)
    _used(ws, 77, "n gebruikt", DIFF, "B73", "n")
    _used(ws, 78, "v̄ gebruikt", DIFF, "B74", "mean")
    _used(ws, 79, "s_v gebruikt", DIFF, "B75", "s")
    have = "AND(ISNUMBER(B77),ISNUMBER(B79),B77>1)"
    result_row(ws, 80, "standaardfout s_v / √n", f'=IF({have},B79/SQRT(B77),"")', NUMBER)
    result_row(ws, 81, "vrijheidsgraden n − 1", f'=IF({have},B77-1,"")', "0")
    _difference_block(ws, 83, PAIRED_CI, PAIRED_TEST, "B78", "B80", "B81", "B76", "μ1 − μ2",
                      "CI Further Reading p. 16, 18 (S03-WE14: −3,3 ± 1,72); Test Recipes p. 7-8")


def _difference_block(ws: Worksheet, top: int, ci_rows: dict[str, int], test_rows: dict[str, int], centre: str,
                      se: str, df: str, d0: str, what: str, source: str) -> None:
    """CI (from, to, half-width, t used) and t-test table for a difference of means."""
    column_titles(ws, top, [f"BI voor {what} (1 − α)", "van", "tot", "± (halve breedte, half-width)", "t gebruikt",
                            "Bron in de cursus"])
    _interval_labels(ws, ci_rows, what)
    have = f"AND(ISNUMBER({centre}),ISNUMBER({se}),ISNUMBER({df}),{df}>0,ISNUMBER({ALPHA}))"
    t_two, t_one = f"_xlfn.T.INV(1-{ALPHA}/2,{df})", f"_xlfn.T.INV(1-{ALPHA},{df})"
    _interval_rows(ws, ci_rows, have, centre, f"{t_two}*{se}", f"{t_one}*{se}", ("B", "C"))
    for key, t in (("two_sided", t_two), ("lower_only", t_one), ("upper_only", t_one)):
        r = ci_rows[key]
        output_cell(ws, f"D{r}", f'=IF({have},{t}*{se},"")', NUMBER)
        output_cell(ws, f"E{r}", f'=IF({have},{t},"")', "0.0000")
    label(ws, ci_rows["two_sided"], 6, source, italic=True)
    column_titles(ws, top + 5, [f"t-toets van H0: {what} = d0", "t = (verschil − d0) / SE", "kritieke t",
                                "2de kritieke t", "", "", "p-waarde", "Besluit bij α"])
    stat = f"({centre}-IF(ISNUMBER({d0}),{d0},0))/{se}"
    _test_table(ws, test_rows, have, stat, "t", df, (f"HA: {what} ≠ d0", f"HA: {what} > d0", f"HA: {what} < d0"))


def _one_proportion(ws: Worksheet) -> None:
    """Section 5: one proportion π."""
    section_title(ws, 94, "5. Eén proportie π (bv. fractie defect, fraction defective)")
    input_row(ws, 95, "n = steekproefgrootte (sample size)")
    input_row(ws, 96, "x = aantal successen (bv. defecte stuks)")
    input_row(ws, 97, "π0 voor de toets (fractie)", "", "0.0000")
    have = "AND(ISNUMBER(B95),ISNUMBER(B96),B95>0,B96>=0,B96<=B95)"
    result_row(ws, 98, "p = x / n", f'=IF({have},B96/B95,"")', PERCENT)
    result_row(ws, 99, "standaardfout √(p(1 − p)/n)", '=IF(ISNUMBER(B98),SQRT(B98*(1-B98)/B95),"")', NUMBER)
    result_row(ws, 100, "Voorwaarde voor de Z-toets: n · π0 > 5", '=IF(AND(ISNUMBER(B95),ISNUMBER(B97)),IF(B95*B97>5,'
               '"voldaan","NIET voldaan: gebruik het exacte (binomiale) interval"),"")', "General",
               "Test Recipes p. 9-10")
    column_titles(ws, 102, ["BI voor π (1 − α)", "normale benadering: van", "tot", "exact (binomiaal): van", "tot",
                            "Bron in de cursus"])
    _interval_labels(ws, PROPORTION_CI, "π")
    ok = f"AND(ISNUMBER(B98),ISNUMBER({ALPHA}))"
    _interval_rows(ws, PROPORTION_CI, ok, "B98", f"_xlfn.NORM.S.INV(1-{ALPHA}/2)*B99",
                   f"_xlfn.NORM.S.INV(1-{ALPHA})*B99", ("B", "C"), percent=True)
    exact = {  # Clopper-Pearson: the interval of R binom.test (CI Further Reading p. 20)
        "two_sided": (f"IF(B96=0,0,_xlfn.BETA.INV({ALPHA}/2,B96,B95-B96+1))",
                      f"IF(B96=B95,1,_xlfn.BETA.INV(1-{ALPHA}/2,B96+1,B95-B96))"),
        "lower_only": (f"IF(B96=0,0,_xlfn.BETA.INV({ALPHA},B96,B95-B96+1))", '"+∞"'),
        "upper_only": ("0", f"IF(B96=B95,1,_xlfn.BETA.INV(1-{ALPHA},B96+1,B95-B96))"),
    }
    for key, (low, high) in exact.items():
        r = PROPORTION_CI[key]
        output_cell(ws, f"D{r}", f'=IF({ok},{low},"")', PERCENT)
        output_cell(ws, f"E{r}", f'=IF({ok},{high},"")', PERCENT)
    label(ws, 103, 6, "CI Further Reading p. 20 (normale benadering; exact = R binom.test); Confidence Intervals.pdf "
                      "p. 4", italic=True)
    column_titles(ws, 107, ["Z-toets van H0: π = π0", "z = (p − π0) / √(π0(1 − π0)/n)", "kritieke z", "2de kritieke z",
                            "", "", "p-waarde", "Besluit bij α"])
    have_test = f"AND({ok},ISNUMBER(B97),B97>0,B97<1)"
    _test_table(ws, PROPORTION_TEST, have_test, "(B98-B97)/SQRT(B97*(1-B97)/B95)", "z", None,
                ("HA: π ≠ π0", "HA: π > π0", "HA: π < π0"))


def _two_proportions(ws: Worksheet) -> None:
    """Section 6: difference of two proportions (normal approximation)."""
    section_title(ws, 112, "6. Twee proporties: verschil π1 − π2 (normale benadering)")
    input_row(ws, 113, "n1")
    input_row(ws, 114, "x1 = successen in steekproef 1")
    input_row(ws, 115, "n2")
    input_row(ws, 116, "x2 = successen in steekproef 2")
    have = "AND(ISNUMBER(B113),ISNUMBER(B114),ISNUMBER(B115),ISNUMBER(B116),B113>0,B115>0)"
    result_row(ws, 117, "p1 = x1 / n1", f'=IF({have},B114/B113,"")', PERCENT)
    result_row(ws, 118, "p2 = x2 / n2", f'=IF({have},B116/B115,"")', PERCENT)
    result_row(ws, 119, "p1 − p2", f'=IF({have},B117-B118,"")', PERCENT)
    result_row(ws, 120, "standaardfout √(p1(1 − p1)/n1 + p2(1 − p2)/n2)",
               f'=IF({have},SQRT(B117*(1-B117)/B113+B118*(1-B118)/B115),"")', NUMBER,
               "CI Further Reading p. 20")
    column_titles(ws, 122, ["BI voor π1 − π2 (1 − α)", "van", "tot", "± (halve breedte)", "", "Bron in de cursus"])
    _interval_labels(ws, TWO_PROPORTION_CI, "π1 − π2")
    ok = f"AND(ISNUMBER(B119),ISNUMBER(B120),ISNUMBER({ALPHA}))"
    z_two, z_one = f"_xlfn.NORM.S.INV(1-{ALPHA}/2)", f"_xlfn.NORM.S.INV(1-{ALPHA})"
    _interval_rows(ws, TWO_PROPORTION_CI, ok, "B119", f"{z_two}*B120", f"{z_one}*B120", ("B", "C"), percent=True)
    for key, z in (("two_sided", z_two), ("lower_only", z_one), ("upper_only", z_one)):
        output_cell(ws, f"D{TWO_PROPORTION_CI[key]}", f'=IF({ok},{z}*B120,"")', PERCENT)
    label(ws, 123, 6, "CI Further Reading p. 20. Voorbeeld S08-WE13: Dummies p. 197 (extra, Dummies; niet te kennen)",
          italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the mean & proportion calculator."""
    write_header(ws, HEADER)
    _data_columns(ws)
    _one_mean(ws)
    _unpaired(ws)
    _paired(ws)
    _one_proportion(ws)
    _two_proportions(ws)
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="Fractie",
                              error="Typ een fractie, bv. 0,05.")
    ws.add_data_validation(fraction)
    for coordinate in (ALPHA, "B97"):
        fraction.add(coordinate)
    ws.column_dimensions["A"].width = 54
    for letter in "BCDEFG":
        ws.column_dimensions[letter].width = 15
    ws.column_dimensions["H"].width = 17
    ws.column_dimensions["K"].width = 3
    for letter in "LMN":
        ws.column_dimensions[letter].width = 12
