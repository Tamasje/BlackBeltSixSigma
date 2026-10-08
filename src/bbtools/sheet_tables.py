"""The Tables sheet: every transcribed control-chart constant table in full, disagreements highlighted.

Calculators never hard-code a constant; they look it up here through workbook names such as `T18_d2`
(values) and `T18_n` (subgroup sizes). `lookup_formula` builds that lookup.
"""
from __future__ import annotations

from collections import defaultdict

from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.constants import USED_TABLE, AverageRangeTable, ConstantTable, load_average_range_table
from bbtools.printed import decimals_printed, parse_printed, rounding_consistent
from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    BOX,
    FLAG_FILL,
    HeaderBlock,
    Status,
    column_titles,
    font,
    label,
    section_title,
    write_header,
)

SHEET = "Tabellen"

# Tables whose only source is a book that is not examinable (Six Sigma For Dummies): kept for reference, but
# written last under EXTRA_HEADING so they are not mistaken for course tables. Their disagreement flags stay.
EXTRA_TABLES = frozenset({"DUM"})
EXTRA_HEADING = "Extra (niet te kennen voor het examen): Six Sigma For Dummies, Table 10-2"

HEADER = HeaderBlock(
    tool="Tabellen: constanten voor regelkaarten (control charts), capabiliteit en MSA",
    source="source/course/Les 4/___4.1 tabellen SPC.pdf p. 1-2; Control charts - constants.pdf p. 1-2; "
           "Les 5/20260619_ottoy_tabel MSA.pdf p. 1; extra (niet te kennen): Six Sigma For Dummies.pdf p. 250",
    convention="Beslissing 4: de rekenbladen gebruiken Table 18; c4 en d3 uit Table A; A3, E2, B5, B6 uit Six Sigma "
               "Demystified. Elke bron staat er volledig, zoals gedrukt.",
    status=Status.VERIFIED,
    status_detail="transcripties van beide gescande tabellen cel per cel dubbel gecontroleerd; waarden vergeleken "
                  "met hun wiskundige definities in tests/test_constants.py",
)


DOC = SheetDoc(
    sheet=SHEET,
    purpose="Elke constantentabel voor regelkaarten uit de cursus, precies zoals gedrukt, met de opzoekbereiken "
            "(lookup ranges) die de rekenbladen gebruiken. Oranje cellen verschillen meer dan afronding van een "
            "andere gedrukte bron (beweeg erover voor de waarden). Six Sigma For Dummies Table 10-2 staat onderaan "
            "als extra (niet te kennen voor het examen).",
    inputs="geen (naslagblad)",
    audit="niet nodig (transcripties dubbel gecontroleerd; waarden vergeleken met hun definities)",
    disagreements=(),
)


def excel_name(key: str, symbol: str) -> str:
    """Workbook name for one constant column, e.g. ('T18', '1/d2') -> 'T18_inv_d2'."""
    return f"{key}_{symbol.replace('1/', 'inv_').replace('*', 'star')}"


def lookup_formula(symbol: str, n_ref: str) -> str:
    """Excel expression (no leading '=') returning `symbol` for subgroup size `n_ref`, from the used table."""
    key = USED_TABLE[symbol]
    return f"INDEX({excel_name(key, symbol)},MATCH({n_ref},{key}_n,0))"


# where each used constant comes from, for the tables shown on the calculator sheets
SOURCE_NAMES = {"T18": "Table 18 (___4.1 tabellen SPC.pdf p. 2)", "TA": "Table A (___4.1 tabellen SPC.pdf p. 1)",
                "SSD1": "Six Sigma Demystified (Control charts - constants.pdf p. 1)",
                "SSD2": "Six Sigma Demystified (Control charts - constants.pdf p. 2)"}


def used_constants_table(ws: Worksheet, top: int, column: int, symbols: tuple[str, ...], purpose: str) -> int:
    """The constants a calculator uses, for n = 2 … 25, written from (top, column); returns the next free row.

    Every value is a lookup into the Tabellen sheet, the same lookup the calculator's formulas use, so the table on
    the sheet always shows exactly the value the sheet computes with.
    """
    label(ws, top, column, f"Constanten die dit blad gebruikt ({purpose})", bold=True)
    sources = sorted({USED_TABLE[s] for s in symbols}, key=list(SOURCE_NAMES).index)
    label(ws, top + 1, column, "Bron: " + "; ".join(
        f"{', '.join(s for s in symbols if USED_TABLE[s] == key)} uit {SOURCE_NAMES[key]}" for key in sources)
        + ". Alle bronnen naast elkaar: blad Tabellen.", italic=True)
    column_titles(ws, top + 2, ["n", *symbols], first_column=column)
    for offset, n in enumerate(range(2, 26)):
        row = top + 3 + offset
        n_cell = ws.cell(row=row, column=column, value=n)
        n_cell.font, n_cell.border = font(bold=True), BOX
        for k, symbol in enumerate(symbols, start=1):
            cell = ws.cell(row=row, column=column + k)
            cell.value = f'=IFERROR({lookup_formula(symbol, get_column_letter(column) + str(row))},"")'
            cell.font, cell.border, cell.number_format = font(), BOX, "0.0000"
    return top + 3 + 24


def msa_table(ws: Worksheet, top: int, column: int) -> int:
    """The d2* table of tabel MSA.pdf (g = 1 … 20 rows, m = 2 … 20 columns) and its d2 row, from (top, column).

    Lookups into the Tabellen sheet, like the Gage R&R formulas; returns the next free row.
    """
    table = load_average_range_table()
    label(ws, top, column, "Tabel d2* (tabel MSA.pdf): K1 = 1/d2 met m = r (onderste rij, g → ∞); K2 = 1/d2* met m = k, "
                           "K3 = 1/d2* met m = n (rij g = 1)", bold=True)
    column_titles(ws, top + 1, ["g \\ m", *[str(m) for m in table.m]], first_column=column)
    for offset, g in enumerate(table.g):
        row = top + 2 + offset
        g_cell = ws.cell(row=row, column=column, value=g)
        g_cell.font, g_cell.border = font(bold=True), BOX
        for k in range(len(table.m)):
            cell = ws.cell(row=row, column=column + 1 + k)
            cell.value = f"=INDEX(MSA_d2star,{offset + 1},{k + 1})"
            cell.font, cell.border, cell.number_format = font(), BOX, "0.00000"
    row = top + 2 + len(table.g)
    d2_label = ws.cell(row=row, column=column, value="d2 (g → ∞)")
    d2_label.font, d2_label.border = font(bold=True), BOX
    for k in range(len(table.m)):
        cell = ws.cell(row=row, column=column + 1 + k)
        cell.value = f"=INDEX(MSA_d2,1,{k + 1})"
        cell.font, cell.border, cell.number_format = font(), BOX, "0.00000"
    return row + 1


def d2_star_formula(g_ref: str, m_ref: str) -> str:
    """Excel expression for d2* of g subgroups of size m (tabel MSA.pdf)."""
    return f"INDEX(MSA_d2star,MATCH({g_ref},MSA_g,0),MATCH({m_ref},MSA_m,0))"


def msa_d2_formula(m_ref: str) -> str:
    """Excel expression for d2 (g → ∞) of subgroup size m, from the last row of tabel MSA.pdf."""
    return f"INDEX(MSA_d2,1,MATCH({m_ref},MSA_m,0))"


def disagreeing_cells(tables: tuple[ConstantTable, ...]) -> dict[tuple[str, str, str], list[str]]:
    """(table key, symbol, n) -> other sources' printed values, for every value that disagrees beyond rounding."""
    seen: dict[tuple[str, str], list[tuple[str, str]]] = defaultdict(list)
    for table in tables:
        for symbol, column in table.symbol_columns().items():
            for row in table.rows:
                value = row[column].strip()
                if value and _is_number(value):
                    seen[(symbol, row[0].strip())].append((table.source.key, value))
    flagged: dict[tuple[str, str, str], list[str]] = {}
    for (symbol, n), entries in seen.items():
        if len(entries) > 1 and not rounding_consistent([v for _, v in entries]):
            for key, value in entries:
                flagged[(key, symbol, n)] = [f"{other_key}: {other}" for other_key, other in entries if other_key != key]
    return flagged


def _is_number(text: str) -> bool:
    """True if the printed text is a single number."""
    try:
        parse_printed(text)
    except ValueError:
        return False
    return True


def _write_cell(ws: Worksheet, row: int, column: int, printed: str) -> None:
    """Write a printed value: numbers as numbers displayed with the printed decimals, the rest as text."""
    cell = ws.cell(row=row, column=column)
    if printed.strip() and _is_number(printed):
        cell.value = float(parse_printed(printed))
        decimals = decimals_printed(printed)
        cell.number_format = "0" if decimals == 0 else "0." + "0" * decimals
    else:
        cell.value = printed
    cell.font = font()
    cell.border = BOX


def _write_table(ws: Worksheet, wb: Workbook, top: int, table: ConstantTable,
                 flagged: dict[tuple[str, str, str], list[str]]) -> int:
    """Write one table starting at row `top`, define its workbook names, return the next free row."""
    label(ws, top, 1, table.title, bold=True)
    label(ws, top + 1, 1, f"Bron: {table.source_file} p. {table.source_page}. Volgens die pagina overgenomen uit: "
                          f"{table.source.origin}.", italic=True)
    label(ws, top + 2, 1, table.source.role, bold=True)
    column_titles(ws, top + 3, list(table.columns))
    first = top + 4
    for offset, row in enumerate(table.rows):
        for column, printed in enumerate(row, start=1):
            _write_cell(ws, first + offset, column, printed)
    last = first + len(table.rows) - 1
    symbols = {i: s for s, i in table.symbol_columns().items()}
    for offset, row in enumerate(table.rows):
        n = row[0].strip()
        if n == "0":  # Six Sigma Demystified p. 1 prints a stray row '0 | 2.606' (review_items.md)
            ws.cell(row=first + offset, column=1).comment = Comment(
                "Zo gedrukt in de bron (een losse rij na n = 2). Geen subgroepgrootte; nooit gebruikt.", "bbtools")
        for column, symbol in symbols.items():
            others = flagged.get((table.source.key, symbol, n))
            if others:
                cell = ws.cell(row=first + offset, column=column + 1)
                cell.fill = FLAG_FILL
                cell.comment = Comment("Andere gedrukte bronnen: " + "; ".join(others), "bbtools")
    _define(wb, f"{table.source.key}_n", 1, first, last)
    for column, symbol in symbols.items():
        _define(wb, excel_name(table.source.key, symbol), column + 1, first, last)
    return last + 2


def _define(wb: Workbook, name: str, column: int, first: int, last: int) -> None:
    """Workbook-level name for one column range on the Tables sheet."""
    letter = get_column_letter(column)
    wb.defined_names[name] = DefinedName(name, attr_text=f"{SHEET}!${letter}${first}:${letter}${last}")


def _define_area(wb: Workbook, name: str, top_left: tuple[int, int], bottom_right: tuple[int, int]) -> None:
    """Workbook-level name for a rectangular range (row, column) on the Tables sheet."""
    (r1, c1), (r2, c2) = top_left, bottom_right
    area = f"${get_column_letter(c1)}${r1}:${get_column_letter(c2)}${r2}"
    wb.defined_names[name] = DefinedName(name, attr_text=f"{SHEET}!{area}")


def _write_average_range(ws: Worksheet, wb: Workbook, top: int, table: AverageRangeTable) -> int:
    """The MSA d2* table as two grids (d2* and ν, rows g, columns m) plus the d2 and cd rows; returns the next row."""
    label(ws, top, 1, table.title, bold=True)
    label(ws, top + 1, 1, f"Bron: {table.source_file} p. {table.source_page}. Volgens die pagina overgenomen uit: "
                          f"Measurement Systems Analysis Reference Manual (DaimlerChrysler, Ford, GM), 2002.", italic=True)
    label(ws, top + 2, 1, "GEBRUIKT door het blad Gage R&R (K1 = 1/d2 met g → ∞; K2, K3 = 1/d2* met g = 1). Elke "
                          "gedrukte cel 'ν / d2*' is gesplitst in de twee rasters hieronder.", bold=True)
    row = top + 3
    for grid_name, values, fmt_source in (("d2*", table.d2_star, "d2_star"), ("ν (vrijheidsgraden, df)", table.nu, "nu")):
        column_titles(ws, row, [f"{grid_name}: g \\ m"] + [str(m) for m in table.m])
        for column, m in enumerate(table.m, start=2):
            ws.cell(row=row, column=column).value = m  # numeric, so MATCH finds it
        for offset, (g, printed) in enumerate(zip(table.g, values, strict=True), start=1):
            ws.cell(row=row + offset, column=1, value=g).border = BOX
            for column, text in enumerate(printed, start=2):
                _write_cell(ws, row + offset, column, text)
        name = "MSA_d2star" if fmt_source == "d2_star" else "MSA_nu"
        _define_area(wb, name, (row + 1, 2), (row + len(table.g), 1 + len(table.m)))
        if fmt_source == "d2_star":
            _define_area(wb, "MSA_m", (row, 2), (row, 1 + len(table.m)))
            _define(wb, "MSA_g", 1, row + 1, row + len(table.g))
        row += len(table.g) + 2
    for text, values, name in (("d2 (g → ∞)", table.d2, "MSA_d2"), ("cd", table.cd, "MSA_cd")):
        label(ws, row, 1, text, bold=True)
        for column, printed in enumerate(values, start=2):
            _write_cell(ws, row, column, printed)
        _define_area(wb, name, (row, 2), (row, 1 + len(table.m)))
        row += 1
    return row + 1


def build_tables_sheet(wb: Workbook, tables: tuple[ConstantTable, ...]) -> Worksheet:
    """Add the Tables sheet to `wb` and define the lookup names the calculators use.

    Course tables first, then the MSA table, then the book-only tables (EXTRA_TABLES) under their own heading.
    """
    ws = wb.create_sheet(SHEET)
    write_header(ws, HEADER)
    label(ws, 8, 1, "Leeswijzer: elke tabel hieronder staat precies zoals gedrukt in de cursus. De rekenbladen nemen "
                    "d2, A2, D3, D4 ... uit Table 18; c4 en d3 uit Table A; A3, E2, B5, B6 uit Six Sigma Demystified. "
                    "Oranje = die waarde verschilt (meer dan afronding) van een andere bron voor dezelfde n; beweeg "
                    "over de cel om de andere waarden te zien. Onderaan, als extra: Six Sigma For Dummies (niet te "
                    "kennen voor het examen).")
    # flags compare every printed source, the extra book included, so a course value that differs from it stays orange
    flagged = disagreeing_cells(tables)
    row = 10
    for table in tables:
        if table.source.key not in EXTRA_TABLES:
            row = _write_table(ws, wb, row, table, flagged)
    row = _write_average_range(ws, wb, row, load_average_range_table())
    section_title(ws, row, EXTRA_HEADING)
    row += 2
    for table in tables:
        if table.source.key in EXTRA_TABLES:
            row = _write_table(ws, wb, row, table, flagged)
    ws.column_dimensions["A"].width = 16
    for column in range(2, 20):
        ws.column_dimensions[get_column_letter(column)].width = 9
    return ws
