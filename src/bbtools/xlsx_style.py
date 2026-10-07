"""Shared look of every sheet: fonts, the input/output colours and the mandatory header block.

CLAUDE.md: inputs in one fill colour, outputs in another; every sheet opens with a header block giving
tool name, course source (file + pages), convention used and status VERIFIED / UNVERIFIED.
"""
from __future__ import annotations

from dataclasses import dataclass
from enum import Enum

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")    # light yellow: type here
OUTPUT_FILL = PatternFill("solid", fgColor="E2EFDA")   # light green: computed, do not type
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")   # light blue: header block and column titles
FLAG_FILL = PatternFill("solid", fgColor="F8CBAD")     # orange: value disagrees with another source
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_ROWS = 6  # rows 1-6 are the header block on every sheet; content starts at row 8


class Status(Enum):
    """Verification status shown in the header block (CLAUDE.md definition of VERIFIED)."""

    VERIFIED = "VERIFIED"
    UNVERIFIED = "UNVERIFIED"

    @property
    def label(self) -> str:
        """Dutch label with the CLAUDE.md term in brackets, e.g. 'GEVERIFIEERD (VERIFIED)'."""
        return {"VERIFIED": "GEVERIFIEERD (VERIFIED)", "UNVERIFIED": "NIET GEVERIFIEERD (UNVERIFIED)"}[self.value]


@dataclass(frozen=True)
class HeaderBlock:
    """The six header lines every sheet starts with."""

    tool: str
    source: str        # course file(s) + pages
    convention: str    # which course convention the sheet applies, or "shows all versions side by side"
    status: Status
    status_detail: str  # e.g. the worked-example IDs the tests check


def font(bold: bool = False, size: int = 10, color: str = "000000", italic: bool = False) -> Font:
    """The workbook font (Arial) with the given emphasis."""
    return Font(name=FONT_NAME, bold=bold, size=size, color=color, italic=italic)


def write_header(ws: Worksheet, header: HeaderBlock) -> None:
    """Write the header block into rows 1-6, labels in column A and text in column B."""
    ws["A1"] = header.tool
    ws["A1"].font = font(bold=True, size=14)
    lines = [
        ("Bron in de cursus", header.source),
        ("Gebruikte conventie", header.convention),
        ("Status", f"{header.status.label}: {header.status_detail}"),
        ("Legende", "Gele cel = vul hier in (invoer).  Groene cel = resultaat, niet overtypen.  "
                    "Oranje cel (Tabellen) = verschilt van een andere gedrukte bron."),
        ("Controle", "Blijft een groene cel leeg, dan ontbreekt nog een invoer die ze nodig heeft."),
    ]
    for row, (label, text) in enumerate(lines, start=2):
        ws.cell(row=row, column=1, value=label).font = font(bold=True)
        ws.cell(row=row, column=1).fill = HEADER_FILL
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = font(bold=(label == "Status"), color="C00000" if header.status is Status.UNVERIFIED and
                         label == "Status" else "000000")
        cell.alignment = Alignment(wrap_text=False, vertical="top")


def section_title(ws: Worksheet, row: int, text: str) -> None:
    """A bold section title in column A, e.g. '1. Specification and process'."""
    ws.cell(row=row, column=1, value=text).font = font(bold=True, size=11)


def label(ws: Worksheet, row: int, column: int, text: str, bold: bool = False, italic: bool = False) -> None:
    """Plain descriptive text."""
    cell = ws.cell(row=row, column=column, value=text)
    cell.font = font(bold=bold, italic=italic)


def input_cell(ws: Worksheet, coordinate: str, number_format: str = "General") -> None:
    """Mark a cell as an input: yellow, boxed, left empty for the user."""
    cell = ws[coordinate]
    cell.fill = INPUT_FILL
    cell.border = BOX
    cell.font = font(color="0000FF")  # blue text = typed input
    cell.number_format = number_format


def output_cell(ws: Worksheet, coordinate: str, formula: str, number_format: str = "General") -> None:
    """Write a formula into a green, boxed result cell."""
    cell = ws[coordinate]
    cell.value = formula
    cell.fill = OUTPUT_FILL
    cell.border = BOX
    cell.font = font()
    cell.number_format = number_format


def input_row(ws: Worksheet, row: int, text: str, note: str = "", number_format: str = "General") -> str:
    """Label in A, yellow input in B, italic note in C; returns the input's coordinate."""
    label(ws, row, 1, text)
    input_cell(ws, f"B{row}", number_format)
    if note:
        label(ws, row, 3, note, italic=True)
    return f"B{row}"


def result_row(ws: Worksheet, row: int, text: str, formula: str, number_format: str = "General",
               source: str = "") -> str:
    """Label in A, green result in B, italic course source in C; returns the result's coordinate."""
    label(ws, row, 1, text)
    output_cell(ws, f"B{row}", formula, number_format)
    if source:
        label(ws, row, 3, source, italic=True)
    return f"B{row}"


def constant_row(ws: Worksheet, row: int, text: str, value: float, source: str) -> str:
    """A course constant in its own boxed cell (not an input, not a result), with its source."""
    label(ws, row, 1, text)
    cell = ws.cell(row=row, column=2, value=value)
    cell.font = font(bold=True)
    cell.border = BOX
    label(ws, row, 3, source, italic=True)
    return f"B{row}"


def column_titles(ws: Worksheet, row: int, titles: list[str], first_column: int = 1) -> None:
    """Bold, shaded, wrapped column titles for a result table."""
    for offset, title in enumerate(titles):
        cell = ws.cell(row=row, column=first_column + offset, value=title)
        cell.font = font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
