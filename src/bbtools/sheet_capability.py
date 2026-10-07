"""Capability sheet: Cp, Cpk (Pp, Ppk) and % out of specification, one result row per σ estimate.

Convention decision 1 (inventory/conventions.md): the course estimates σ in several ways, so the sheet does
not choose. Each σ the user fills in gets its own result row, labelled with the course page that uses it.
Decision 2: the "6 sigma criterion" is shown in both readings the course uses.
The MR̄ / 1.128 estimate comes only from Six Sigma For Dummies (not examinable): it lives on the Extra sheet.

Formulas (course): Cp = (UL − LL)/6σ (deck p. 34); Cpk = min{(UL − x̄)/3σ, (x̄ − LL)/3σ} (deck p. 35);
% out of spec from the normal distribution (deck p. 18-19, 46; Excel check NORM.VERD on deck p. 46).
All quantities are in the user's measurement unit; percentages are stored as fractions.
"""
from __future__ import annotations

from dataclasses import dataclass

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.sheet_tables import lookup_formula
from bbtools.xlsx_style import (
    BOX,
    HeaderBlock,
    Status,
    column_titles,
    font,
    input_cell,
    label,
    output_cell,
    section_title,
    write_header,
)

SHEET = "Capabiliteit"

HEADER = HeaderBlock(
    tool="Procescapabiliteit (process capability): Cp, Cpk (Pp, Ppk) en % buiten specificatie (out of spec)",
    source="source/course/Les 4/2026 Lean - Six Sigma v13 - Capabiliteit - SPC.pdf p. 18-21, 32-47, 50, 64, 83; "
           "___4.1 tabellen SPC.pdf p. 1-2; __Xbar R kaart data - berekeningen oefening 2.xlsx",
    convention="Beslissingen 1-2: één resultaatrij per σ-schatting die de cursus gebruikt (er wordt niet voor je "
               "gekozen); het '6 sigma'-criterium in beide lezingen van de cursus.",
    status=Status.VERIFIED,
    status_detail="getest tegen de uitgewerkte voorbeelden S06-WE01, S06-WE02, S06-WE04, S06-WE07, S06-WE08, "
                  "S06-WE12, S06-WE13 uit de cursus; enkele gedrukte waarden op deck p. 46 en 50 wijken af van de "
                  "berekening (zie build/README.md)",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Cp, Cpk (of Pp, Ppk), Z-afstanden en % / ppm buiten specificatie uit LSL, USL, gemiddelde en een spreiding; "
            "de twee lezingen van het '6 sigma'-criterium in de cursus; de Cp-niveaus van deck p. 40.",
    inputs="LSL en/of USL, gemiddelde; dan eender welke van: σ gegeven, R̄ (+ n), s̄ (+ n), totale s. Elke spreiding "
           "krijgt haar eigen resultaatrij, met de cursuspagina die ze gebruikt. MR̄ / 1,128 (alleen in Six Sigma For "
           "Dummies) staat op blad Extra (boeken).",
    audit="niet nodig (er bestaan uitgewerkte voorbeelden in de cursus)",
    disagreements=(
        "deck p. 46 (S06-WE01): Cp gedrukt als '1,166' (1,4/1,2 = 1,1667 rondt af op 1,167); gecentreerde Cpk idem.",
        "deck p. 46 notities (S06-WE01): '5 % totaal' en '2,5 % per zijde' voor gemiddelde 71,8, σ 0,2, specificatie "
        "71,4-72,8; de normale verdeling geeft 2,275 % onder LSL en 0,00003 % boven USL (2,275 % totaal).",
        "deck p. 46 notities (S06-WE01, gecentreerd): '0,04 % / 400 ppm' tweezijdig verdubbelt de afgeronde "
        "0,02 % / 200 ppm; berekend 0,0465 % / 465 ppm.",
        "deck p. 50 (S06-WE12, Minitab): '% out of spec' 8,74 en 9,33 tegenover 8,731 en 9,315 berekend uit het "
        "gedrukte gemiddelde 0,14852; Minitab rekende met niet-afgeronde data. Pp, Ppk, Cp en Cpk kloppen.",
    ),
)

# Input cells (tests write here).
INPUTS: dict[str, str] = {
    "lsl": "B9", "usl": "B10", "mean": "B11",
    "sigma_given": "B14", "rbar": "B15", "sbar": "B16", "n": "B17", "s_overall": "B18",
}

# One result row per σ estimate: key, row, label, σ formula, constant formula, constant text, course source.
@dataclass(frozen=True)
class SigmaRow:
    """One line of the results table."""

    key: str
    row: int
    label: str
    sigma: str        # formula for the σ used (column B)
    constant: str     # formula for the table constant used (column Q), '' if none
    constant_note: str
    source: str


SIGMA_ROWS: tuple[SigmaRow, ...] = (
    SigmaRow("given", 23, "σ gegeven in de vraag", "=IF(ISNUMBER(B14),B14,\"\")", "", "",
             "de vraag"),
    SigmaRow("rbar_d2", 24, "R̄ / d2  (korte termijn: Cp, Cpk)", "=IF(AND(ISNUMBER(B15),ISNUMBER(Q24)),B15/Q24,\"\")",
             f"=IF(AND(ISNUMBER(B15),ISNUMBER(B17)),IFERROR({lookup_formula('d2', 'B17')},\"n niet in Tabel 18\"),\"\")",
             "d2 voor n, Tabel 18 (___4.1 tabellen SPC.pdf p. 2)",
             "deck p. 64 (sprekersnotities); oefenwerkboeken oefening 2, 3, 5"),
    SigmaRow("sbar", 25, "s̄ rechtstreeks gebruikt  (zoals deck p. 46)", "=IF(ISNUMBER(B16),B16,\"\")", "", "",
             "deck p. 46 (Cp = (72,8 − 71,4)/(6 · 0,2))"),
    SigmaRow("sbar_c4", 26, "s̄ / c4  (zuiver, unbiased; Tabel A)", "=IF(AND(ISNUMBER(B16),ISNUMBER(Q26)),B16/Q26,\"\")",
             f"=IF(AND(ISNUMBER(B16),ISNUMBER(B17)),IFERROR({lookup_formula('c4', 'B17')},\"n niet in Tabel A\"),\"\")",
             "c4 voor n, Tabel A (___4.1 tabellen SPC.pdf p. 1)",
             "___4.1 tabellen SPC.pdf p. 1 ('An unbiased est. of SD(X) = s̄/c4')"),
    SigmaRow("overall", 27, "totale s (overall; lange termijn: dit zijn Pp, Ppk)", "=IF(ISNUMBER(B18),B18,\"\")", "", "",
             "deck p. 33"),
)
RESULT_ROWS: dict[str, int] = {r.key: r.row for r in SIGMA_ROWS}

# Result columns (tests read here).
RESULT_COLUMNS: dict[str, str] = {
    "sigma": "B", "cp": "C", "cp_level": "D", "cpu": "E", "cpl": "F", "cpk": "G", "cpk_capable": "H",
    "z_lsl": "I", "z_usl": "J", "below_lsl": "K", "above_usl": "L", "out_total": "M", "ppm_total": "N",
}

# Constants of the course, in their own labelled cells so no formula hides a number.
CP_LEVELS = {"just": "B40", "acceptable": "B41", "good": "B42", "six": "B43", "cpk_capable": "B45"}
SIX_SIGMA = {"shift": "B35", "z": "B36"}
CRITERION_ROWS = {"long_term_shift": 32, "short_term_centred": 33}


def _inputs(ws: Worksheet) -> None:
    """Sections 1 and 2: the yellow input cells with their labels and validation."""
    section_title(ws, 8, "1. Specificatie en proces")
    rows = [
        (9, "Onderste specificatiegrens LSL (lower specification limit)", "leeg laten bij een eenzijdige specificatie"),
        (10, "Bovenste specificatiegrens USL (upper specification limit)", "leeg laten bij een eenzijdige specificatie"),
        (11, "Procesgemiddelde x̄, of het gemiddelde van de subgroepgemiddelden (X̿)", ""),
    ]
    for row, text, note in rows:
        label(ws, row, 1, text)
        input_cell(ws, f"B{row}")
        label(ws, row, 3, note, italic=True)
    # A warning line of its own, so it never hides the notes in column C; red text, shown only when wrong.
    ws["A12"] = '=IF(AND(ISNUMBER(B9),ISNUMBER(B10),B10<=B9),"Controle: USL moet boven LSL liggen","")'
    ws["A12"].font = font(bold=True, color="C00000")

    section_title(ws, 13, "2. Spreiding: vul alleen in wat de vraag geeft. Elke ingevulde lijn krijgt haar eigen "
                          "resultaatrij in deel 3.")
    rows = [
        (14, "σ gegeven in de vraag", "bv. 'standaardafwijking van 10 mm'"),
        (15, "R̄ (R-bar) = gemiddelde spreidingsbreedte (range) van de subgroepen", "heeft n hieronder nodig"),
        (16, "s̄ (s-bar) = gemiddelde standaardafwijking van de subgroepen", "heeft n hieronder nodig voor s̄ / c4"),
        (17, "n = subgroepgrootte (voor R̄ en s̄)", "2 tot 25 (Tabel 18); Tabel A ook 30-100"),
        (18, "s = totale standaardafwijking van alle waarden (overall, STDEV.S)", "lange termijn"),
    ]
    for row, text, note in rows:
        label(ws, row, 1, text)
        input_cell(ws, f"B{row}")
        label(ws, row, 3, note, italic=True)
    label(ws, 19, 1, "MR̄ / 1,128 (individuele waarden) komt alleen uit Six Sigma For Dummies: zie blad 'Extra (boeken)'.",
          italic=True)

    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Spreiding moet positief zijn",
                              error="σ, R̄, s̄ en s moeten groter zijn dan 0.")
    subgroup = DataValidation(type="whole", operator="between", formula1="2", formula2="100", allow_blank=True,
                              showErrorMessage=True, errorTitle="Subgroepgrootte", error="n is een geheel getal, 2 tot 100.")
    ws.add_data_validation(positive)
    ws.add_data_validation(subgroup)
    for coordinate in ("B14", "B15", "B16", "B18"):
        positive.add(coordinate)
    subgroup.add("B17")


def _result_row(ws: Worksheet, spec: SigmaRow) -> None:
    """One row of section 3: σ, indices, level texts, Z distances and fractions out of spec."""
    r = spec.row
    label(ws, r, 1, spec.label)
    number = "0.0000"
    # Reversed limits would give nonsense such as 197 % out of spec: blank σ, and with it the whole row.
    guarded = f'=IF(AND(ISNUMBER($B$9),ISNUMBER($B$10),$B$10<=$B$9),"",{spec.sigma[1:]})'
    output_cell(ws, f"B{r}", guarded, number)
    output_cell(ws, f"C{r}", f"=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$9),ISNUMBER($B$10)),($B$10-$B$9)/(6*B{r}),\"\")", "0.000")
    output_cell(ws, f"D{r}", f"=IF(ISNUMBER(C{r}),IF(C{r}>=$B$43,\"6 Sigma-kwaliteitsniveau (6 Sigma quality level)\","
                             f"IF(C{r}>=$B$42,\"goed (good)\",IF(C{r}>=$B$41,\"acceptabel (acceptable)\","
                             f"IF(C{r}>=$B$40,\"net capabel (just capable)\",\"niet capabel (not capable)\")))),\"\")")
    output_cell(ws, f"E{r}", f"=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$10),ISNUMBER($B$11)),($B$10-$B$11)/(3*B{r}),\"\")", "0.000")
    output_cell(ws, f"F{r}", f"=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$9),ISNUMBER($B$11)),($B$11-$B$9)/(3*B{r}),\"\")", "0.000")
    output_cell(ws, f"G{r}", f"=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),MIN(E{r},F{r}),IF(ISNUMBER(E{r}),E{r},"
                             f"IF(ISNUMBER(F{r}),F{r},\"\")))", "0.000")
    output_cell(ws, f"H{r}", f"=IF(ISNUMBER(G{r}),IF(G{r}>$B$45,\"ja\",\"nee\"),\"\")")
    output_cell(ws, f"I{r}", f"=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$9),ISNUMBER($B$11)),($B$11-$B$9)/B{r},\"\")", "0.000")
    output_cell(ws, f"J{r}", f"=IF(AND(ISNUMBER(B{r}),ISNUMBER($B$10),ISNUMBER($B$11)),($B$10-$B$11)/B{r},\"\")", "0.000")
    # NORM.S.DIST(-z) rather than 1 - NORM.S.DIST(z): same value, no loss of digits in the far tail.
    output_cell(ws, f"K{r}", f"=IF(ISNUMBER(I{r}),_xlfn.NORM.S.DIST(-I{r},TRUE),\"\")", "0.0000%")
    output_cell(ws, f"L{r}", f"=IF(ISNUMBER(J{r}),_xlfn.NORM.S.DIST(-J{r},TRUE),\"\")", "0.0000%")
    output_cell(ws, f"M{r}", f"=IF(OR(ISNUMBER(K{r}),ISNUMBER(L{r})),SUM(K{r},L{r}),\"\")", "0.0000%")
    output_cell(ws, f"N{r}", f"=IF(ISNUMBER(M{r}),M{r}*1000000,\"\")", "#,##0.0")
    if spec.constant:
        output_cell(ws, f"Q{r}", spec.constant, "0.0000")
    label(ws, r, 18, spec.constant_note, italic=True)
    label(ws, r, 19, spec.source, italic=True)


def _results(ws: Worksheet) -> None:
    """Section 3: the results table."""
    section_title(ws, 21, "3. Resultaten: één rij per spreidingsschatting (een rij blijft leeg tot haar invoer is "
                          "ingevuld)")
    column_titles(ws, 22, [
        "Spreidingsschatting", "Gebruikte σ", "Cp  (laatste rij: Pp)", "Cp-niveau (deck p. 40)",
        "Cpu = (USL − gemiddelde) / 3σ", "Cpl = (gemiddelde − LSL) / 3σ", "Cpk = min(Cpu, Cpl)  (laatste rij: Ppk)",
        "Cpk > 1,33? (deck p. 41; oefening 2)", "Z: LSL ligt … σ onder het gemiddelde",
        "Z: USL ligt … σ boven het gemiddelde", "% onder LSL", "% boven USL", "% buiten specificatie, totaal",
        "ppm buiten specificatie, totaal", "", "", "Gebruikte constante", "Welke constante", "Bron in de cursus",
    ])
    for spec in SIGMA_ROWS:
        _result_row(ws, spec)


def _criterion(ws: Worksheet) -> None:
    """Section 4: the '6 sigma' criterion in both course readings (decision 2)."""
    section_title(ws, 30, "4. Het '6 sigma'-criterium: het maximum buiten specificatie. De cursus gebruikt twee "
                          "lezingen.")
    column_titles(ws, 31, ["Lezing", "Gebruikte Z", "Max % buiten specificatie", "Max ppm", "Bron in de cursus"])
    label(ws, 32, 1, "Lange termijn, met de 1,5σ-verschuiving (1.5σ shift): één staart voorbij Z = 6 − 1,5")
    output_cell(ws, "B32", f"={SIX_SIGMA['z']}-{SIX_SIGMA['shift']}", "0.0")
    output_cell(ws, "C32", "=_xlfn.NORM.S.DIST(-B32,TRUE)", "0.00000%")
    output_cell(ws, "D32", "=C32*1000000", "0.0")
    label(ws, 32, 5, "deck p. 21 (tabel: 6 -> 3,4 DPMO), 37 (+ notities)", italic=True)
    label(ws, 33, 1, "Korte termijn, gecentreerd proces met Cp = 2: beide staarten voorbij ±6σ")
    output_cell(ws, "B33", f"={SIX_SIGMA['z']}", "0.0")
    output_cell(ws, "C33", "=2*_xlfn.NORM.S.DIST(-B33,TRUE)", "0.0000000%")
    output_cell(ws, "D33", "=C33*1000000", "0.000")
    label(ws, 33, 5, "deck p. 40: 'CP = 2  6 Sigma quality level – 2 defect per billion of opportunities (short term)'",
          italic=True)
    _constant(ws, 35, "Verschuiving die de cursus aanneemt (σ)", 1.5, "deck p. 37")
    _constant(ws, 36, "Six sigma: afstand van het gemiddelde tot de specificatie (σ)", 6,
              "deck p. 40 (CP = 2: 6 Sigma quality level)")


def _levels(ws: Worksheet) -> None:
    """Section 5: the course's interpretation thresholds, referenced by columns D and H."""
    section_title(ws, 38, "5. Interpretatiedrempels gebruikt in kolommen D en H")
    label(ws, 39, 1, "CP < 1: het proces is niet capabel (not capable)", italic=True)
    _constant(ws, 40, "Cp ≥ dit: net capabel (just capable)", 1, "deck p. 40 ('CP = 1 The process is just capable')")
    _constant(ws, 41, "Cp ≥ dit: acceptabel (acceptable)", 1.33,
              "deck p. 40 ('CP >= 1.33 The capability of the process is acceptable')")
    _constant(ws, 42, "Cp ≥ dit: goed (good)", 1.67, "deck p. 40 ('CP >= 1.67 The capability of the process is good')")
    _constant(ws, 43, "Cp ≥ dit: 6 Sigma-kwaliteitsniveau (6 Sigma quality level)", 2,
              "deck p. 40 ('CP = 2 6 Sigma quality level')")
    _constant(ws, 45, "Cpk groter dan dit: goed (good)", 1.33,
              "deck p. 41 (figuur: Cpk 1,33 'Good'); __Xbar R kaart data - berekeningen oefening 2.xlsx "
              "('Cpk: 1,33 = good')")
    section_title(ws, 47, "6. Capabiliteit verbeteren (cursus)")
    label(ws, 48, 1, "Centreer het proces (gemiddelde naar het midden van LSL en USL): Cpk stijgt tot Cp. "
                     "Deck p. 36, 42-45, 46.")
    label(ws, 49, 1, "Verklein de variatie (σ): Cp en Cpk stijgen allebei. Deck p. 22, 36.")


def _constant(ws: Worksheet, row: int, text: str, value: float, source: str) -> None:
    """A course constant in its own boxed cell (not an input, not a result), with its source."""
    label(ws, row, 1, text)
    cell = ws.cell(row=row, column=2, value=value)
    cell.font = font(bold=True)
    cell.border = BOX
    label(ws, row, 3, source, italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the capability calculator."""
    write_header(ws, HEADER)
    _inputs(ws)
    _results(ws)
    _criterion(ws)
    _levels(ws)
    ws.column_dimensions["A"].width = 52
    for letter in "BCDEFGHIJKLMN":
        ws.column_dimensions[letter].width = 13
    ws.column_dimensions["D"].width = 30  # fits the longest Dutch level text with its English term
    for letter in "OP":
        ws.column_dimensions[letter].width = 2
    ws.column_dimensions["Q"].width = 10
    ws.column_dimensions["R"].width = 44
    ws.column_dimensions["S"].width = 60
    ws.row_dimensions[22].height = 42
