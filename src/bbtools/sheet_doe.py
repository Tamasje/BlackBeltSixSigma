"""DOE 2^k sheet: effects and ANOVA of a 2^k factorial design (k = 2 … 5, up to 20 replicates).

Course (De Vuyst, Les 3, DOE.pdf): levels coded −1/+1 (p. 46), in standard order (p. 70, 74); effect = mean response
at +1 minus mean at −1 = contrast / (n 2^(k−1)), SS = contrast² / (n 2^k) (p. 48, 57, 66); s.e.(effect) =
√(σ̂² / (n 2^(k−2))) and effect ± 2 s.e. (p. 67-68); single replicate: pool the higher-order interactions into the
error (p. 72, 77); Dummies p. 233: coefficient = effect / 2, β0 = mean of all runs. R²_adj shown with n − p − 1 (the
course outputs, DOE p. 61) and as printed on Regression p. 56.

Row plan: α 9; how to fill 11-14; k and pooling 16-19; runs 20-52 (replicates in G … Z, per-run summaries AA … AE,
sign columns from AF); error estimate 54-63; effects 65-97; model table 99-105.
"""
from __future__ import annotations

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    font,
    input_block,
    input_row,
    instructions,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "DOE 2^k"

HEADER = HeaderBlock(
    tool="2^k-factoriële proefopzet (factorial design, k = 2 tot 5): effecten, kwadratensommen, F-toetsen, ANOVA van "
         "het model en R²",
    source="source/course/Les 3/20260605_de vuyst_BB_DOE.pdf p. 46-92",
    convention="Invoer α vooraf ingevuld op 0,05 (beslissing 5); factoren gecodeerd −1/+1 in standaardvolgorde (standard "
               "order; DOE p. 46, 70, 74); één herhaling (single replicate): interacties van een gekozen orde en hoger "
               "in de fout poolen (DOE p. 72, 77); R²_adj op beide manieren getoond.",
    status=Status.VERIFIED,
    status_detail="getest tegen uitgewerkte voorbeelden van de cursus S05-WE05 tot S05-WE09, en tegen S08-WE16 (extra, "
                  "Dummies; niet te kennen); gedrukte waarden die niet kloppen staan in build/README.md",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Effecten, coëfficiënten (extra, Dummies; niet te kennen), kwadratensommen (SS), F-toetsen en intervallen "
            "±2 s.e. van een 2^k-factoriële proefopzet met zuivere fout (pure error) uit herhalingen en/of gepoolde "
            "interacties van hogere orde, de ANOVA van het model en R².",
    inputs="α; k, optioneel de orde vanaf waar gepoold wordt; tot 20 herhalingen per run in standaardvolgorde.",
    audit="niet nodig (er bestaan uitgewerkte voorbeelden in de cursus)",
    disagreements=(
        "DOE p. 50 (S05-WE05): '[AB] = 5,78 – 4,92 = 0,857'; het exacte effect is 0,8583 (gemiddelden 5,7767 en "
        "4,9183).",
        "DOE p. 71 (S05-WE08): SS_ABC (en MS) gedrukt als 5,5625; de gegevens op p. 70 geven contrast 9 en SS 81/16 = "
        "5,0625, wat ook de gedrukte F0 2,08, P 0,19 en het totaal 92,9375 impliceren. De gedrukte P van A, "
        "2,54 × 10^-3, is 2,534 × 10^-3 voor F0 18,69 op F(1, 8).",
        "DOE p. 53 (S05-WE06): 'AB = (52 + 20)/2 − (30 + 40)/2 = −1'; die uitdrukking is gelijk aan +1, wat het blad "
        "geeft.",
        "Regression p. 56 drukt R²_adj = 1 − (1 − R²)(n − 1)/(n − k − 2); de eigen outputs van de cursus gebruiken "
        "n − k − 1 (DOE p. 61 Adj R-Squared 0,8666). Het blad toont beide, met label.",
    ),
)

ALPHA = "$B$9"
NUMBER = "0.000000"
DECISION = '=IF(ISNUMBER({p}),IF({p}<' + ALPHA + ',"verwerp H0","H0 niet verwerpen"),"")'

FACTORS = "ABCDE"
REPLICATES = 20
REPLICATE_COLUMNS = [get_column_letter(7 + i) for i in range(REPLICATES)]  # G … Z
RUN_N, RUN_MEAN, RUN_SUM, RUN_SS, RUN_IN = "AA", "AB", "AC", "AD", "AE"   # per-run summaries
RUNS, FIRST_RUN, FIRST_EFFECT = 32, 21, 66
SIGN_FIRST_COLUMN = 32  # column AF: sign column of effect 1 (A); effect m sits in column AF + m − 1
MODEL_ROWS = {"model": 100, "error": 101, "total": 102}  # B SS, C df, D MS, E F, F p, G decision
INPUTS = {"alpha": "B9", "k": "B17", "pool_order": "B18"}
RESULTS = {
    "n_rep": "B54", "complete": "B55", "N": "B56", "beta0": "B57", "ss_pe": "B58", "df_pe": "B59",
    "ss_pool": "B60", "df_pool": "B61", "mse": "B62", "se_effect": "B63",
    "r2_doe": "B103", "r2_adj_doe": "B104", "r2_adj_doe_p56": "B105",
}


def effect_name(mask: int) -> str:
    """Effect label of a factor bitmask: 1 -> 'A', 3 -> 'AB', 7 -> 'ABC' (bit j = factor j)."""
    return "".join(FACTORS[j] for j in range(len(FACTORS)) if mask >> j & 1)


def sign(run: int, mask: int) -> int:
    """Coded level (−1/+1) of effect `mask` in standard-order run `run` (0-based): the product of its factors' levels.

    In standard order factor j is +1 in run r exactly when bit j of r is set (A alternates every run, B every 2 runs).
    """
    value = 1
    for j in range(len(FACTORS)):
        if mask >> j & 1:
            value *= 1 if run >> j & 1 else -1
    return value


def run_row(run: int) -> int:
    """Sheet row of standard-order run `run` (0-based)."""
    return FIRST_RUN + run


def effect_row(mask: int) -> int:
    """Sheet row of effect `mask` (1 … 31) in the effects table."""
    return FIRST_EFFECT + mask - 1


def response_cell(run: int, replicate: int) -> str:
    """Input cell of replicate `replicate` (0-based) of standard-order run `run` (0-based)."""
    return f"{REPLICATE_COLUMNS[replicate]}{run_row(run)}"


def _runs(ws: Worksheet) -> None:
    """Rows 20-52: the 32 standard-order runs, their replicate inputs and per-run summaries."""
    column_titles(ws, 20, ["Run", *FACTORS, *[f"y herhaling {i + 1}" for i in range(REPLICATES)], "n",
                           "gemiddelde van de run", "som van de run", "Σ (y − gemiddelde van de run)²",
                           "in de proefopzet (1/0)"])
    input_block(ws, (run_row(0), run_row(RUNS - 1)), (column_index_from_string(REPLICATE_COLUMNS[0]),
                                                     column_index_from_string(REPLICATE_COLUMNS[-1])))
    for run in range(RUNS):
        r = run_row(run)
        ws[f"A{r}"] = run + 1
        ws[f"A{r}"].font = font()
        output_cell(ws, f"{RUN_IN}{r}", f"=IF(AND(ISNUMBER($B$17),$B$17>=2,$B$17<=5,A{r}<=2^$B$17),1,0)", "0")
        for j, letter in enumerate("BCDEF"):
            output_cell(ws, f"{letter}{r}", f'=IF(AND({RUN_IN}{r}=1,{j + 1}<=$B$17),{1 if run >> j & 1 else -1},"")', "0")
        ys = f"{REPLICATE_COLUMNS[0]}{r}:{REPLICATE_COLUMNS[-1]}{r}"
        output_cell(ws, f"{RUN_N}{r}", f'=IF(AND({RUN_IN}{r}=1,COUNT({ys})>0),COUNT({ys}),"")', "0")
        output_cell(ws, f"{RUN_MEAN}{r}", f'=IF(ISNUMBER({RUN_N}{r}),AVERAGE({ys}),"")', NUMBER)
        output_cell(ws, f"{RUN_SUM}{r}", f"=IF(ISNUMBER({RUN_N}{r}),SUM({ys}),0)", NUMBER)
        output_cell(ws, f"{RUN_SS}{r}", f"=IF(ISNUMBER({RUN_N}{r}),DEVSQ({ys}),0)", NUMBER)
    label(ws, 19, SIGN_FIRST_COLUMN, "Teken (−1/+1) van elk effect in elke run: gebruikt door de contrastformules",
          italic=True)
    for mask in range(1, RUNS):
        column = SIGN_FIRST_COLUMN + mask - 1
        ws.cell(row=20, column=column, value=effect_name(mask)).font = font(bold=True)
        for run in range(RUNS):
            ws.cell(row=run_row(run), column=column, value=sign(run, mask)).font = font(color="808080")


def _error_estimate(ws: Worksheet) -> None:
    """Rows 54-63: n, completeness, β0, pure error, pooled interactions, MS_E and s.e.(effect)."""
    first, last = run_row(0), run_row(RUNS - 1)
    n_col, in_design = f"{RUN_N}{first}:{RUN_N}{last}", f"{RUN_IN}{first}:{RUN_IN}{last}"
    result_row(ws, 54, "n = herhalingen (replicates) per run", f'=IF(COUNT({n_col})>0,MAX({n_col}),"")', "0")
    result_row(ws, 55, "Volledige proefopzet (runs 1 … 2^k hebben elk n responsen)?",
               f'=IF(COUNT({n_col})=0,"",IF(AND(COUNT({n_col})=SUM({in_design}),MIN({n_col})=MAX({n_col})),"ja",'
               f'"NEE: vul runs 1 … 2^k in, elk met hetzelfde aantal herhalingen"))')
    ok = '$B$55="ja"'
    pooling = "AND(ISNUMBER($B$18),$B$18>=2)"
    ss, order = f"F{FIRST_EFFECT}:F{effect_row(RUNS - 1)}", f"B{FIRST_EFFECT}:B{effect_row(RUNS - 1)}"
    result_row(ws, 56, "N = 2^k · n waarnemingen", f'=IF({ok},2^$B$17*$B$54,"")', "0")
    result_row(ws, 57, "totaal gemiddelde (grand mean) = β0", f'=IF({ok},SUM({RUN_SUM}{first}:{RUN_SUM}{last})/$B$56,"")', NUMBER,
               "β0: Dummies p. 233 (extra, Dummies; niet te kennen)")
    result_row(ws, 58, "SS zuivere fout (pure error, tussen de herhalingen)", f'=IF({ok},SUM({RUN_SS}{first}:{RUN_SS}{last}),"")',
               "0.0000")
    result_row(ws, 59, "df zuivere fout = 2^k (n − 1)", f'=IF({ok},2^$B$17*($B$54-1),"")', "0")
    result_row(ws, 60, "SS gepoolde interacties (orde ≥ B18)",
               f'=IF({ok},IF({pooling},SUMIFS({ss},{order},">="&$B$18),0),"")', "0.0000", "DOE p. 72, 77")
    result_row(ws, 61, "df gepoolde interacties",
               f'=IF({ok},IF({pooling},COUNTIFS({order},">="&$B$18,{ss},">=0"),0),"")', "0")
    result_row(ws, 62, "σ̂² = MS_E = (SS zuivere fout + SS gepoold) / (df + df)",
               f'=IF(AND({ok},N($B$59)+N($B$61)>0),($B$58+$B$60)/($B$59+$B$61),"")', NUMBER)
    result_row(ws, 63, "standaardfout (standard error) s.e.(effect) = √(σ̂² / (n 2^(k−2)))",
               '=IF(ISNUMBER($B$62),SQRT($B$62/($B$54*2^($B$17-2))),"")', NUMBER, "DOE p. 67")


def _effect_row(ws: Worksheet, mask: int) -> None:
    """One row of the effects table: contrast, effect, coefficient, SS, F-test and effect ± 2 s.e."""
    r, first, last = effect_row(mask), run_row(0), run_row(RUNS - 1)
    signs = get_column_letter(SIGN_FIRST_COLUMN + mask - 1)
    ws[f"A{r}"] = effect_name(mask)
    ws[f"A{r}"].font = font(bold=True)
    ws[f"B{r}"] = bin(mask).count("1")
    ws[f"B{r}"].font = font()
    in_design = f'AND($B$55="ja",{mask}<2^$B$17)'
    output_cell(ws, f"C{r}", f'=IF({in_design},SUMPRODUCT({signs}{first}:{signs}{last},{RUN_SUM}{first}:{RUN_SUM}{last}),"")', "0.0000")
    output_cell(ws, f"D{r}", f'=IF(ISNUMBER(C{r}),C{r}/($B$54*2^($B$17-1)),"")', "0.0000")
    output_cell(ws, f"E{r}", f'=IF(ISNUMBER(D{r}),D{r}/2,"")', "0.0000")
    output_cell(ws, f"F{r}", f'=IF(ISNUMBER(C{r}),C{r}^2/($B$54*2^$B$17),"")', "0.0000")
    output_cell(ws, f"G{r}", f'=IF(ISNUMBER(C{r}),IF(AND(ISNUMBER($B$18),$B$18>=2,B{r}>=$B$18),'
                             f'"gepoold in de fout","in het model"),"")')
    tested = f'AND(G{r}="in het model",ISNUMBER($B$62),N($B$62)>0)'
    output_cell(ws, f"H{r}", f'=IF({tested},F{r}/$B$62,"")', "0.0000")
    output_cell(ws, f"I{r}", f'=IF(ISNUMBER(H{r}),_xlfn.F.DIST.RT(H{r},1,$B$59+$B$61),"")', "0.000000")
    output_cell(ws, f"J{r}", DECISION.format(p=f"I{r}"))
    output_cell(ws, f"K{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER($B$63)),D{r}-2*$B$63,"")', "0.0000")
    output_cell(ws, f"L{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER($B$63)),D{r}+2*$B$63,"")', "0.0000")


def _effects(ws: Worksheet) -> None:
    """Rows 65-97: the table of all 31 effects (only those of the chosen k fill in)."""
    column_titles(ws, 65, ["Effect", "orde", "contrast", "effect = contrast / (n 2^(k−1))",
                            "coëfficiënt = effect / 2 (extra, Dummies; niet te kennen)", "SS = contrast² / (n 2^k)",
                            "in het model of gepoold (pooled)", "F0 = SS / MS_E", "p-waarde", "Besluit bij α",
                            "effect − 2 s.e.", "effect + 2 s.e."])
    for mask in range(1, RUNS):
        _effect_row(ws, mask)
    label(ws, 97, 1, "Effect = gemiddelde bij +1 − gemiddelde bij −1 (DOE p. 48, 66); effect ± 2 s.e. ≈ 95 %-BI: "
                      "bevat het 0, dan niet significant bij 5 % (DOE p. 68). Coëfficiënt = effect / 2, Dummies p. 233 "
                      "(extra, Dummies; niet te kennen).", italic=True)


def _model_table(ws: Worksheet) -> None:
    """Rows 99-105: ANOVA of the model (effects in model) against the error, R² and both R²_adj."""
    ss = f"F{FIRST_EFFECT}:F{effect_row(RUNS - 1)}"
    ok = '$B$55="ja"'
    column_titles(ws, 99, ["Bron (source)", "SS", "df", "MS", "F0", "p-waarde", "Besluit bij α"])
    rows = {
        100: ("Model (alle effecten in het model)", f"SUM({ss})-$B$60", "2^$B$17-1-$B$61"),
        101: ("Fout (error: zuivere fout + gepoold)", "$B$58+$B$60", "$B$59+$B$61"),
        102: ("Totaal", f"SUM({ss})+$B$58", "$B$56-1"),
    }
    for row, (text, sum_sq, df) in rows.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({ok},{sum_sq},"")', "0.0000")
        output_cell(ws, f"C{row}", f'=IF({ok},{df},"")', "0")
    output_cell(ws, "D100", '=IF(AND(ISNUMBER(B100),N(C100)>0),B100/C100,"")', "0.0000")
    output_cell(ws, "D101", '=IF(ISNUMBER($B$62),$B$62,"")', "0.0000")
    output_cell(ws, "E100", '=IF(AND(ISNUMBER(D100),ISNUMBER(D101),N(D101)>0),D100/D101,"")', "0.0000")
    output_cell(ws, "F100", '=IF(ISNUMBER(E100),_xlfn.F.DIST.RT(E100,C100,C101),"")', "0.000000")
    output_cell(ws, "G100", DECISION.format(p="F100"))
    label(ws, 101, 5, "zoals de Design-Expert-output van de cursus, DOE p. 61", italic=True)
    result_row(ws, 103, "R² = SS_model / SS_totaal", '=IF(AND(ISNUMBER(B102),N(B102)>0),B100/B102,"")', "0.0000",
               "DOE p. 61")
    result_row(ws, 104, "R²_adj = 1 − (1 − R²)(N − 1)/(N − p − 1), p = df van het model",
               '=IF(AND(ISNUMBER(B103),N(C101)>0),1-(1-B103)*C102/C101,"")', "0.0000",
               "zoals de outputs van de cursus (DOE p. 61; Regression p. 22, 57)")
    result_row(ws, 105, "R²_adj zoals gedrukt op Regression p. 56: 1 − (1 − R²)(N − 1)/(N − p − 2)",
               '=IF(AND(ISNUMBER(B103),N(C101)>1),1-(1-B103)*C102/(C101-1),"")', "0.0000",
               "Regression p. 56 (verschilt van de outputs van de cursus)")


def _factorial(ws: Worksheet) -> None:
    """Section 2: 2^k factorial design, effects and ANOVA."""
    section_title(ws, 16, "2^k-factoriële proefopzet (factorial design); standaardvolgorde (standard order): A "
                          "wisselt elke run, B om de 2 runs, C om de 4, …")
    input_row(ws, 17, "k = aantal factoren (2 tot 5)",
              "een fractie 2^(k−p) (fractional factorial): voer ze in als de volledige proefopzet van haar k − p "
              "basisfactoren; elk effect is dan een aliasketen (alias chain) (DOE p. 81-82, 88-90)", "0")
    input_row(ws, 18, "Interacties van deze orde en hoger in de fout poolen (optioneel, 2 tot 5)",
              "één herhaling (single replicate): nodig voor F-toetsen, bv. 3 zoals DOE p. 77", "0")
    label(ws, 19, 1, "Typ elke respons in de rij waarvan de tekens A … E overeenkomen met je tabel (DOE p. 70, 74 "
                     f"gebruiken deze volgorde); één kolom per herhaling (replicate), tot {REPLICATES}.", italic=True)
    _runs(ws)
    _error_estimate(ws)
    _effects(ws)
    _model_table(ws)


def _validation(ws: Worksheet) -> None:
    """k and the pooling order: whole numbers 2 to 5; α a fraction."""
    whole = DataValidation(type="whole", operator="between", formula1="2", formula2="5", allow_blank=True,
                           showErrorMessage=True, errorTitle="k of orde voor poolen",
                           error="Een geheel getal van 2 tot 5.")
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="α", error="α ligt tussen 0 en 1.")
    for rule, cells in ((whole, ("B17", "B18")), (fraction, ("B9",))):
        ws.add_data_validation(rule)
        for cell in cells:
            rule.add(cell)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the 2^k factorial calculator."""
    write_header(ws, HEADER)
    section_title(ws, 8, "Instellingen")
    input_row(ws, 9, "Significantieniveau α (significance level; betrouwbaarheid = 1 − α)",
              "de cursus heeft geen standaardwaarde; de meeste cursusvoorbeelden gebruiken 5 %; gebruik de waarde uit "
              "de vraag", "0.0000")
    ws["B9"] = 0.05
    instructions(ws, 11, 1, "Zo vul je de runs in", (
        "• Typ k in B17. Kolommen B-F tonen dan de tekens (−1/+1) van A … E per run, in standaardvolgorde.",
        f"• Typ elke respons (GETAL) in de rij van de run met dezelfde tekens als in je opgave, één kolom per herhaling: "
        f"herhaling 1 in G, 2 in H, … tot {REPLICATES} (kolom {REPLICATE_COLUMNS[-1]}). Elke run evenveel herhalingen.",
        "• Eén herhaling: typ in B18 vanaf welke orde de interacties in de fout gaan (bv. 3), anders zijn er geen "
        "F-toetsen. Rij 55 zegt of de proefopzet volledig is.",
    ))
    _factorial(ws)
    _validation(ws)
    ws.column_dimensions["A"].width = 58
    for column in range(2, SIGN_FIRST_COLUMN - 1):
        ws.column_dimensions[get_column_letter(column)].width = 12
    for column in range(SIGN_FIRST_COLUMN, SIGN_FIRST_COLUMN + RUNS - 1):
        ws.column_dimensions[get_column_letter(column)].width = 6
