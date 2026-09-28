"""Variance sheet: confidence intervals and tests for σ², σ (χ²) and for a ratio of two variances (F).

Course: (n−1)s²/σ² ~ χ²(n−1) and the 95 % CI (n−1)s²/χ²_0.025 ≤ σ² ≤ (n−1)s²/χ²_0.975 (CI Further Reading
(Dutch) p. 21); χ²-test for σ (Test Recipes p. 11; Testing of Hypotheses p. 13); (s1²/σ1²)/(s2²/σ2²) ~
F(n1−1, n2−1) and the F-test (Test Recipes p. 12-13); exam Q2's hint (s1²/s2²)(σ2²/σ1²) ~ F(n1−1, n2−1).
Decision 5: α is an input prefilled 0.05. Decision 6: every interval and test is shown two-sided, lower-only
and upper-only. Decision 7: F quantiles come straight from F.INV / F.INV.RT; both ratio orientations shown.
Excel names: CHISQ.INV.RT(p; df) is the value with upper-tail area p, CHISQ.INV(p; df) with lower-tail area p.

Row plan: settings 8-9; one sample 11-29 (inputs 12-14, used 15-18, CIs 21-23, χ²-tests 27-29);
two samples 31-51 (inputs 32-35, used 36-41, CIs 44-46, F-tests 49-51); raw data in H10:I509.
"""
from __future__ import annotations

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    font,
    input_cell,
    input_row,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Variance CI & tests"

HEADER = HeaderBlock(
    tool="Variances: CI and test for σ² and σ (χ²), CI and test for a ratio of two variances (F)",
    source="source/course/Les 2/20260529_ottoy_Confidence Intervals - Further Reading (Dutch).pdf p. 21; "
           "Confidence Intervals.pdf p. 16; Testing of Hypotheses.pdf p. 13; Test Recipes - Further Reading "
           "(Dutch).pdf p. 11-14; Six Sigma For Dummies.pdf p. 192-196",
    convention="Decisions 5-7: α input prefilled 0.05; two-sided, lower-only and upper-only side by side; "
               "F from F.INV / F.INV.RT, both σ1²/σ2² and σ2²/σ1².",
    status=Status.VERIFIED,
    status_detail="tested against course worked examples S03-WE03, S03-WE05, S08-WE10 and the Dummies χ² and F "
                  "tables (p. 194, 196); the Dummies F-interval example S08-WE11 disagrees (build/README.md)",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="One sample: CI for σ² and σ, χ²-test of σ = σ0. Two samples: CI for σ1²/σ2² and σ2²/σ1², F-test of "
            "σ1 = σ2 (e.g. 'is machine M1 more precise than M2?'). Every result two-sided, lower-only and upper-only.",
    inputs="α; n and s per sample, or the raw values pasted in columns H (sample 1) and I (sample 2); σ0 for the "
           "χ²-test.",
    audit="stats-auditor PASS (2026-09-28) for the F part: all 28 values (CIs in both orientations, F-tests ≠, >, <) "
          "for one input agree with an independent computation from Test Recipes p. 12-14, CI Further Reading "
          "p. 21 and exam Q2's hint. Note: the course prints no explicit two-sample ratio CI; it follows from the "
          "F pivot of p. 12 inverted as on p. 21.",
    disagreements=(
        "Dummies p. 196 (S08-WE11): CI for σA²/σB² printed [(1/3.633)(4/7.5), 5.999(4/7.5)] = [0.147, 3.199]. "
        "With (sA²/σA²)/(sB²/σB²) ~ F(nA−1, nB−1) (Test Recipes p. 12, exam Q2 hint) the same 5 % tail values give "
        "[0.0889, 1.937]: the book swaps the two F values. Its table values themselves (Table 8-3) are correct.",
        "Dummies p. 192-196 call '95 %' what is ±2σ (95.45 %, 2.275 % per tail) for χ², and a 5 % upper tail for F. "
        "Use α = 0.0455 to reproduce its χ² example (S08-WE10).",
        "Dummies Table 8-3 p. 196: F for n1 = n2 = 2 printed 161.446 (definition 161.448). Table 8-2 p. 194: "
        "99.7 % upper value for n = 5 printed 17.800 (definition 17.8006). All other entries agree.",
    ),
)

ALPHA = "B9"
DATA_1, DATA_2 = "H10:H509", "I10:I509"
INPUTS = {"alpha": ALPHA, "n": "B12", "s": "B13", "sigma0": "B14", "n1": "B32", "s1": "B33", "n2": "B34", "s2": "B35"}
RESULTS = {
    "n_used": "B15", "s_used": "B16", "var_used": "B17", "df": "B18",
    "n1_used": "B36", "s1_used": "B37", "n2_used": "B38", "s2_used": "B39", "f": "B40", "df1": "B41", "df2": "C41",
}
CI_ROWS = {"two_sided": 21, "lower_only": 22, "upper_only": 23}          # B,C = σ² from/to; D,E = σ from/to
CHI2_TEST_ROWS = {"two_sided": 27, "greater": 28, "less": 29}           # B stat, C/D critical, E p, F decision
F_CI_ROWS = {"two_sided": 44, "ratio_at_least": 45, "ratio_at_most": 46}  # B,C = σ1²/σ2²; D,E = σ2²/σ1²
F_TEST_ROWS = {"two_sided": 49, "greater": 50, "less": 51}              # B stat, C/D critical, E p, F decision
NUMBER = "0.000000"
DECISION = '=IF(ISNUMBER(E{r}),IF(E{r}<' + ALPHA + ',"reject H0","do not reject H0"),"")'


def _data_columns(ws: Worksheet) -> None:
    """Columns H and I: optional raw data for sample 1 and sample 2 (500 values each)."""
    label(ws, 8, 8, "Optional: paste raw values below; n and s are then computed from them", italic=True)
    for column, text in (("H", "Sample 1 data"), ("I", "Sample 2 data")):
        ws[f"{column}9"] = text
        ws[f"{column}9"].font = font(bold=True)
        for row in range(10, 510):
            input_cell(ws, f"{column}{row}")


def _used(ws: Worksheet, row: int, text: str, data: str, typed: str, is_n: bool) -> None:
    """A green 'value used' cell: from the pasted data when it holds at least 2 values, else the typed input."""
    if is_n:
        formula, fmt = f'=IF(COUNT({data})>1,COUNT({data}),IF(ISNUMBER({typed}),{typed},""))', "0"
    else:
        formula, fmt = f'=IF(COUNT({data})>1,_xlfn.STDEV.S({data}),IF(ISNUMBER({typed}),{typed},""))', NUMBER
    result_row(ws, row, text, formula, fmt, f"from column {data[0]} if it holds 2+ values, else the typed value")


def _test_rows(ws: Worksheet, rows: dict[int, tuple[str, str, str | None, str]], have: str, stat: str) -> None:
    """Rows of a test table: statistic, critical value(s), p-value and decision at α."""
    for row, (text, critical, critical_2, p_value) in rows.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({have},{stat},"")', "0.0000")
        output_cell(ws, f"C{row}", f'=IF({have},{critical},"")', "0.0000")
        output_cell(ws, f"D{row}", f'=IF({have},{critical_2},"")' if critical_2 else '=""', "0.0000")
        output_cell(ws, f"E{row}", f'=IF({have},{p_value},"")', "0.0000%")
        output_cell(ws, f"F{row}", DECISION.format(r=row))


def _one_sample(ws: Worksheet) -> None:
    """Sections 1-2: α and one sample (CI for σ², σ and the χ²-test)."""
    section_title(ws, 8, "1. Settings")
    input_row(ws, 9, "Significance α (confidence = 1 − α)",
              "course has no default; most course examples use 5 %; use the value the question gives", "0.0000")
    ws[ALPHA] = 0.05

    section_title(ws, 11, "2. One sample: σ² and σ")
    input_row(ws, 12, "n (sample size)", "or paste the data in column H")
    input_row(ws, 13, "s (sample standard deviation, STDEV.S)", "", NUMBER)
    input_row(ws, 14, "σ0 for the test (the claimed σ)", "e.g. 0.01", NUMBER)
    _used(ws, 15, "n used", DATA_1, "B12", is_n=True)
    _used(ws, 16, "s used", DATA_1, "B13", is_n=False)
    result_row(ws, 17, "s² used", '=IF(ISNUMBER(B16),B16^2,"")', "0.0000000000")
    result_row(ws, 18, "degrees of freedom n − 1", '=IF(ISNUMBER(B15),B15-1,"")', "0")

    column_titles(ws, 20, ["Confidence interval (1 − α)", "σ² from", "σ² to", "σ from", "σ to", "Course source"])
    ok = f"AND(ISNUMBER(B17),ISNUMBER(B18),B18>0,ISNUMBER({ALPHA}))"
    q = "B18*B17"  # (n−1) s²
    intervals = {
        21: ("two-sided", f"{q}/_xlfn.CHISQ.INV.RT({ALPHA}/2,B18)", f"{q}/_xlfn.CHISQ.INV({ALPHA}/2,B18)",
             "CI Further Reading (Dutch) p. 21"),
        22: ("lower bound only (σ at least …; goes with HA: σ > σ0)", f"{q}/_xlfn.CHISQ.INV.RT({ALPHA},B18)", None,
             "Confidence Intervals.pdf p. 16 ('one-sided 98 %-CI ]0.0088, +∞[')"),
        23: ("upper bound only (σ at most …; goes with HA: σ < σ0)", None, f"{q}/_xlfn.CHISQ.INV({ALPHA},B18)", ""),
    }
    for row, (text, low, high, source) in intervals.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({ok},{low if low else 0},"")', NUMBER)
        output_cell(ws, f"C{row}", f'=IF({ok},{high},"")' if high else f'=IF({ok},"+∞","")', NUMBER)
        output_cell(ws, f"D{row}", f'=IF(ISNUMBER(B{row}),SQRT(B{row}),"")', NUMBER)
        output_cell(ws, f"E{row}", f'=IF(ISNUMBER(C{row}),SQRT(C{row}),C{row})', NUMBER)
        label(ws, row, 6, source, italic=True)

    column_titles(ws, 26, ["χ²-test of H0: σ = σ0", "Statistic (n−1)s²/σ0²", "Critical value", "2nd critical value",
                           "p-value", "Decision at α"])
    _test_rows(ws, {
        27: ("HA: σ ≠ σ0 (two-sided)", f"_xlfn.CHISQ.INV({ALPHA}/2,B18)", f"_xlfn.CHISQ.INV.RT({ALPHA}/2,B18)",
             "2*MIN(_xlfn.CHISQ.DIST(B27,B18,TRUE),_xlfn.CHISQ.DIST.RT(B27,B18))"),
        28: ("HA: σ > σ0 (e.g. 'more erratic')", f"_xlfn.CHISQ.INV.RT({ALPHA},B18)", None, "_xlfn.CHISQ.DIST.RT(B28,B18)"),
        29: ("HA: σ < σ0", f"_xlfn.CHISQ.INV({ALPHA},B18)", None, "_xlfn.CHISQ.DIST(B29,B18,TRUE)"),
    }, f"AND({ok},ISNUMBER(B14),B14>0)", "B18*B17/B14^2")
    label(ws, 30, 1, "Source: Test Recipes - Further Reading (Dutch) p. 11; Testing of Hypotheses.pdf p. 13", italic=True)


def _two_samples(ws: Worksheet) -> None:
    """Section 3: two samples (CI for the variance ratio in both orientations, and the F-test)."""
    section_title(ws, 31, "3. Two samples: ratio of variances (F)   e.g. exam Q2: is machine M1 more precise than M2?")
    input_row(ws, 32, "n1 (sample 1)", "or paste sample 1 in column H")
    input_row(ws, 33, "s1", "", NUMBER)
    input_row(ws, 34, "n2 (sample 2)", "or paste sample 2 in column I")
    input_row(ws, 35, "s2", "", NUMBER)
    _used(ws, 36, "n1 used", DATA_1, "B32", is_n=True)
    _used(ws, 37, "s1 used", DATA_1, "B33", is_n=False)
    _used(ws, 38, "n2 used", DATA_2, "B34", is_n=True)
    _used(ws, 39, "s2 used", DATA_2, "B35", is_n=False)
    have = f"AND(ISNUMBER(B36),ISNUMBER(B37),ISNUMBER(B38),ISNUMBER(B39),B36>1,B38>1,B39>0,ISNUMBER({ALPHA}))"
    result_row(ws, 40, "F = s1² / s2²", f'=IF({have},B37^2/B39^2,"")', "0.0000",
               "Test Recipes p. 12-13: (s1²/σ1²)/(s2²/σ2²) ~ F(n1 − 1, n2 − 1)")
    label(ws, 41, 1, "degrees of freedom (numerator n1 − 1, denominator n2 − 1)")
    output_cell(ws, "B41", '=IF(ISNUMBER(B36),B36-1,"")', "0")
    output_cell(ws, "C41", '=IF(ISNUMBER(B38),B38-1,"")', "0")

    column_titles(ws, 43, ["Confidence interval (1 − α)", "σ1²/σ2² from", "σ1²/σ2² to", "σ2²/σ1² from", "σ2²/σ1² to",
                           "Course source"])
    f, v1, v2 = "B40", "B41", "C41"
    intervals = {
        44: ("two-sided", f"{f}/_xlfn.F.INV.RT({ALPHA}/2,{v1},{v2})", f"{f}/_xlfn.F.INV({ALPHA}/2,{v1},{v2})",
             "exam Q2 hint; Test Recipes p. 12"),
        45: ("one-sided: σ1²/σ2² at least … (σ2²/σ1² at most …)", f"{f}/_xlfn.F.INV.RT({ALPHA},{v1},{v2})", None,
             "goes with HA: σ1 > σ2"),
        46: ("one-sided: σ1²/σ2² at most … (σ2²/σ1² at least …)", None, f"{f}/_xlfn.F.INV({ALPHA},{v1},{v2})",
             "goes with HA: σ1 < σ2 (e.g. M1 more precise)"),
    }
    for row, (text, low, high, source) in intervals.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({have},{low if low else 0},"")', NUMBER)
        output_cell(ws, f"C{row}", f'=IF({have},{high},"")' if high else f'=IF({have},"+∞","")', NUMBER)
        # σ2²/σ1² is the reciprocal interval: from = 1 / (σ1²/σ2² to), to = 1 / (σ1²/σ2² from)
        output_cell(ws, f"D{row}", f'=IF(ISNUMBER(C{row}),1/C{row},IF(C{row}="+∞",0,""))', NUMBER)
        output_cell(ws, f"E{row}", f'=IF(ISNUMBER(B{row}),IF(B{row}=0,"+∞",1/B{row}),"")', NUMBER)
        label(ws, row, 6, source, italic=True)

    column_titles(ws, 48, ["F-test of H0: σ1 = σ2", "Statistic F = s1²/s2²", "Critical value", "2nd critical value",
                           "p-value", "Decision at α"])
    _test_rows(ws, {
        49: ("HA: σ1 ≠ σ2 (two-sided)", f"_xlfn.F.INV({ALPHA}/2,{v1},{v2})", f"_xlfn.F.INV.RT({ALPHA}/2,{v1},{v2})",
             f"2*MIN(_xlfn.F.DIST(B49,{v1},{v2},TRUE),_xlfn.F.DIST.RT(B49,{v1},{v2}))"),
        50: ("HA: σ1 > σ2", f"_xlfn.F.INV.RT({ALPHA},{v1},{v2})", None, f"_xlfn.F.DIST.RT(B50,{v1},{v2})"),
        51: ("HA: σ1 < σ2", f"_xlfn.F.INV({ALPHA},{v1},{v2})", None, f"_xlfn.F.DIST(B51,{v1},{v2},TRUE)"),
    }, have, f)
    label(ws, 52, 1, "Source: Test Recipes - Further Reading (Dutch) p. 13-14; Dummies Table 8-3 p. 196 "
                     "(5 % upper-tail F values)", italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the variance calculator."""
    write_header(ws, HEADER)
    _data_columns(ws)
    _one_sample(ws)
    _two_samples(ws)
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="α", error="α is a fraction, e.g. 0.05.")
    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Must be positive", error="Must be greater than 0.")
    ws.add_data_validation(fraction)
    ws.add_data_validation(positive)
    fraction.add(ALPHA)
    for coordinate in ("B12", "B13", "B14", "B32", "B33", "B34", "B35"):
        positive.add(coordinate)
    ws.column_dimensions["A"].width = 52
    for letter in "BCDE":
        ws.column_dimensions[letter].width = 17
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 4
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 14
