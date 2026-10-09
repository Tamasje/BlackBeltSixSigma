"""Tests for bbtools.sheet_charts, evaluated by LibreOffice headless.

Expected values: course worked examples from the lecturers' exercise workbooks (S06-WE03, S06-WE05, S06-WE06,
S10-WE04a/b; Excel cached floats at rel=1e-9) with their input data read from the course files named by the
oracle's data_ref (read-only), and the Dummies example S08-WE18 at printed precision. An independent Python
computation for random subgroups. Printed values that disagree are strict xfails naming the source.
"""
from __future__ import annotations

import random
import statistics
from collections.abc import Callable
from typing import Any

import openpyxl
import pytest
import xlrd

from bbtools.constants import ROOT
from bbtools.printed import agrees_at_printed_precision
from bbtools.build_workbook import build_workbook
from bbtools.sheet_charts import (
    CHARTS,
    CHECK,
    FLAGS,
    INPUTS,
    LIMITS,
    PLOT_INDEX,
    PLOT_R,
    PLOT_X,
    POINTS,
    SHEET,
    SUMMARY,
    TYPED,
    X_COLUMNS,
    _plot_columns,
    subgroup_row,
)

pytestmark = pytest.mark.libreoffice

Evaluate = Callable[[str, dict[str, float]], Any]
Printed = Callable[[str, str, str, str], str]
COURSE = ROOT / "source" / "course"


def raw_subgroups(rows: list[list[float]]) -> dict[str, float]:
    """Cells for subgroups given as raw values (up to 25 per subgroup)."""
    return {f"{X_COLUMNS[j]}{subgroup_row(i)}": x for i, row in enumerate(rows) for j, x in enumerate(row)}


def typed_subgroups(means: list[float], ranges: list[float], n: int) -> dict[str, float]:
    """Cells for subgroups given as typed x-bar and R, with the subgroup size n."""
    cells: dict[str, float] = {INPUTS["n_typed"]: n}
    for i, (m, r) in enumerate(zip(means, ranges, strict=True)):
        cells[f"{TYPED['xbar']}{subgroup_row(i)}"], cells[f"{TYPED['r']}{subgroup_row(i)}"] = m, r
    return cells


def limit(ws: Any, chart: str, which: str) -> Any:
    """LCL, CL or UCL of one chart ('xbar_r', 'r', 'xbar_s', 's')."""
    return ws[f"{ {'lcl': 'B', 'cl': 'C', 'ucl': 'D'}[which] }{LIMITS[chart]}"].value


def oefening_2_data() -> list[list[float]]:
    """__Gegevens oefeningen.xlsx 'Gegevens oefening 2'!B5:F24: 20 subgroups of 5 net weights (S06-WE03 data_ref)."""
    ws = openpyxl.load_workbook(COURSE / "Les 4" / "__Gegevens oefeningen.xlsx", data_only=True)["Gegevens oefening 2"]
    return [[ws.cell(r, c).value for c in range(2, 7)] for r in range(5, 25)]


def oefening_3_data() -> tuple[list[float], list[float]]:
    """'Gegevens oefening 3'!C2:C25 (Xbar') and E2:E25 (R') for 24 samples (S06-WE05 data_ref)."""
    ws = openpyxl.load_workbook(COURSE / "Les 4" / "__Gegevens oefeningen.xlsx", data_only=True)["Gegevens oefening 3"]
    return [ws.cell(r, 3).value for r in range(2, 26)], [ws.cell(r, 5).value for r in range(2, 26)]


def rheostat_data(sheet: str) -> list[list[float]]:
    """Rheostat Knob Data.xls sheet A3:E29: 27 subgroups of 5 (S10-WE04a/b data_ref)."""
    book = xlrd.open_workbook(str(COURSE / "Les 5" / "20260619_ottoy_Rheostat Knob Data.xls"))
    s = book.sheet_by_name(sheet)
    return [s.row_values(r)[:5] for r in range(2, 29)]


def test_oefening_2_xbar_r_chart_from_raw_data_s06_we03(oracle: dict[str, Any], evaluate: Evaluate) -> None:
    # arrange -- bleach net weight, 20 subgroups of n = 5
    answers = oracle["S06-WE03"]["stated_answers"]
    ws = evaluate(SHEET, raw_subgroups(oefening_2_data()))
    # act / assert -- Excel cached floats, rel=1e-9 (same arithmetic, same constants A2 0.577, D4 2.115)
    assert ws[SUMMARY["xbarbar"]].value == pytest.approx(float(answers["Xbarbar (X_streep_streep)"]), rel=1e-9)
    assert ws[SUMMARY["rbar"]].value == pytest.approx(float(answers["Rbar (R_streep)"]), rel=1e-9)
    assert limit(ws, "xbar_r", "ucl") == pytest.approx(float(answers["UCL_xbar"]), rel=1e-9)
    assert limit(ws, "xbar_r", "lcl") == pytest.approx(float(answers["LCL_xbar"]), rel=1e-9)
    assert limit(ws, "r", "ucl") == pytest.approx(float(answers["UCL_R"]), rel=1e-9)
    assert limit(ws, "r", "lcl") == float(answers["LCL_R"])
    assert ws[SUMMARY["sigma_r"]].value == pytest.approx(float(answers["sigma_estimate_Rbar_over_d2"]), rel=1e-9)


@pytest.mark.parametrize(("drop", "prefix"), [(None, "all_24_samples__"), (8, "revised_excl_outlier__")])
def test_oefening_3_from_typed_means_and_ranges_s06_we05(drop: int | None, prefix: str, oracle: dict[str, Any],
                                                         evaluate: Evaluate) -> None:
    # arrange -- 24 samples of n = 5 given as x-bar' and R'; the revision drops sample 9 ('uitschieter ! eruit halen')
    answers = oracle["S06-WE05"]["stated_answers"]
    means, ranges = oefening_3_data()
    if drop is not None:
        del means[drop], ranges[drop]
    ws = evaluate(SHEET, typed_subgroups(means, ranges, 5))
    # act / assert -- Excel cached floats, rel=1e-9
    for name, key in (("xbarbar", "Xbarbarbar"), ("rbar", "Rbar")):
        assert ws[SUMMARY[name]].value == pytest.approx(float(answers[prefix + key]), rel=1e-9)
    for chart, which, key in (("xbar_r", "ucl", "UCL_xbar"), ("xbar_r", "lcl", "LCL_xbar"), ("r", "ucl", "UCL_R")):
        assert limit(ws, chart, which) == pytest.approx(float(answers[prefix + key]), rel=1e-9)


@pytest.mark.parametrize("n", [3, 8])
def test_constants_for_another_subgroup_size_s06_we06(n: int, oracle: dict[str, Any], evaluate: Evaluate) -> None:
    # arrange -- oefening 4 re-monitors with n = 3 and n = 8 using the constants of Table 18
    answers = oracle["S06-WE06"]["stated_answers"]
    ws = evaluate(SHEET, typed_subgroups([1.0], [1.0], n))
    # act / assert -- the constants the sheet looks up equal the ones the workbook uses
    assert ws[f"F{LIMITS['xbar_r']}"].value == float(answers[f"n={n}__A2"])
    assert ws[f"F{LIMITS['r']}"].value == float(answers[f"n={n}__D3"])
    assert ws[f"G{LIMITS['r']}"].value == float(answers[f"n={n}__D4"])
    assert ws[SUMMARY["d2"]].value == float(answers[f"n={n}__d2"])


@pytest.mark.parametrize(("example_id", "sheet"), [("S10-WE04a", "originele data"), ("S10-WE04b", "afgeronde data")])
def test_rheostat_xbar_limits_s10_we04(example_id: str, sheet: str, oracle: dict[str, Any], evaluate: Evaluate) -> None:
    # arrange -- Les 5: 27 subgroups of 5, original (1/1000 inch) and rounded (1/100 inch) data
    answers = oracle[example_id]["stated_answers"]
    ws = evaluate(SHEET, raw_subgroups(rheostat_data(sheet)))
    # act / assert -- Excel cached floats, rel=1e-9
    assert limit(ws, "xbar_r", "cl") == pytest.approx(float(answers["CL_Xbar"]), rel=1e-9)
    assert limit(ws, "xbar_r", "lcl") == pytest.approx(float(answers["LCL_Xbar"]), rel=1e-9)
    assert limit(ws, "xbar_r", "ucl") == pytest.approx(float(answers["UCL_Xbar"]), rel=1e-9)
    assert limit(ws, "r", "cl") == pytest.approx(float(answers["CL_R (=Rbar)"]), rel=1e-9)


@pytest.mark.xfail(reason="Rheostat Knob Data.xls (Les 5) uses D4 = 2.114 (Six Sigma Demystified); the sheet uses "
                          "Table 18's 2.115 like the Les 4 workbooks (decision 4)")
@pytest.mark.parametrize(("example_id", "sheet"), [("S10-WE04a", "originele data"), ("S10-WE04b", "afgeronde data")])
def test_rheostat_r_chart_upper_limit_s10_we04(example_id: str, sheet: str, oracle: dict[str, Any], evaluate: Evaluate) -> None:
    # arrange
    ws = evaluate(SHEET, raw_subgroups(rheostat_data(sheet)))
    # act / assert
    assert limit(ws, "r", "ucl") == pytest.approx(float(oracle[example_id]["stated_answers"]["UCL_R"]), rel=1e-9)


def test_dummies_xbar_r_chart_s08_we18(oracle: dict[str, Any], evaluate: Evaluate) -> None:
    # arrange -- Dummies p. 252 (Figure 10-9): X-bar 84.5, R-bar 5.75. The subgroup size is not printed;
    #            n = 5 (A2 0.577) is the only size whose A2 reproduces the printed limits.
    answers = oracle["S08-WE18"]["stated_answers"]
    ws = evaluate(SHEET, typed_subgroups([answers["Xbar"]], [answers["Rbar"]], 5))
    # act / assert -- printed precision
    assert agrees_at_printed_precision(limit(ws, "xbar_r", "ucl"), str(answers["UCL_Xbar"]))
    assert agrees_at_printed_precision(limit(ws, "xbar_r", "lcl"), str(answers["LCL_Xbar"]))
    assert agrees_at_printed_precision(limit(ws, "r", "ucl"), str(answers["UCL_R"]))


@pytest.mark.parametrize("seed", [1, 2])
def test_xbar_r_and_xbar_s_charts_match_an_independent_computation(seed: int, evaluate: Evaluate) -> None:
    # arrange -- constants as printed in Table 18 / Table A / Six Sigma Demystified for the chosen n
    rng = random.Random(seed)
    n = rng.choice([2, 3, 4, 5, 6, 8, 10])
    subgroups = [[rng.gauss(50, 3) for _ in range(n)] for _ in range(rng.randint(5, 30))]
    ws = evaluate(SHEET, raw_subgroups(subgroups))
    a2, d3, d4 = (ws[f"F{LIMITS['xbar_r']}"].value, ws[f"F{LIMITS['r']}"].value, ws[f"G{LIMITS['r']}"].value)
    xbb = statistics.mean(statistics.mean(g) for g in subgroups)
    rbar = statistics.mean(max(g) - min(g) for g in subgroups)
    sbar = statistics.mean(statistics.stdev(g) for g in subgroups)
    a3, b3, b4 = ws[f"F{LIMITS['xbar_s']}"].value, ws[f"F{LIMITS['s']}"].value, ws[f"G{LIMITS['s']}"].value
    # act / assert -- same arithmetic in Python; rel=1e-9 covers summation order
    expected = {
        (None, "xbar_r", "ucl"): xbb + a2 * rbar, (None, "xbar_r", "lcl"): xbb - a2 * rbar,
        (None, "r", "ucl"): d4 * rbar, (None, "r", "lcl"): d3 * rbar,
        (None, "xbar_s", "ucl"): xbb + a3 * sbar, (None, "s", "ucl"): b4 * sbar, (None, "s", "lcl"): b3 * sbar,
    }
    for (_, chart, which), target in expected.items():
        assert limit(ws, chart, which) == pytest.approx(target, rel=1e-9, abs=1e-12), (chart, which)
    for i, g in enumerate(subgroups):  # every flag says what the limits say
        m, flag = statistics.mean(g), ws[f"{FLAGS['xbar_r']}{subgroup_row(i)}"].value
        high, low = limit(ws, "xbar_r", "ucl"), limit(ws, "xbar_r", "lcl")
        assert (flag or "") == ("boven UCL" if m > high else "onder LCL" if m < low else "")


def test_many_large_subgroups_and_the_row_check(evaluate: Evaluate) -> None:
    # arrange -- 200 subgroups of 25 values (beyond the old 50 × 10 table), one row with raw values and a typed x̄
    rng = random.Random(3)
    groups = [[rng.gauss(10, 1) for _ in range(25)] for _ in range(200)]
    cells = raw_subgroups(groups) | {f"{TYPED['xbar']}{subgroup_row(0)}": 99.0}
    # act
    ws = evaluate(SHEET, cells)
    # assert -- summary from all 200 subgroups; the typed x̄ next to raw values is ignored and reported
    assert ws[SUMMARY["k"]].value == 200 and ws[SUMMARY["n"]].value == 25
    assert ws[SUMMARY["xbarbar"]].value == pytest.approx(statistics.mean(statistics.mean(g) for g in groups), rel=1e-12)
    assert ws[SUMMARY["rbar"]].value == pytest.approx(statistics.mean(max(g) - min(g) for g in groups), rel=1e-12)
    assert ws[SUMMARY["sbar"]].value == pytest.approx(statistics.mean(statistics.stdev(g) for g in groups), rel=1e-12)
    assert ws[f"{CHECK}{subgroup_row(0)}"].value.startswith("ruwe waarden én getypt")
    assert ws[f"{CHECK}{subgroup_row(1)}"].value is None


def test_plot_helper_columns_hold_points_limits_and_the_points_outside(evaluate: Evaluate) -> None:
    # arrange -- 20 subgroups of 5 values, subgroup 7 shifted far above the rest; the charts start at subgroup 3
    rng = random.Random(8)
    rows = [[rng.gauss(10, 0.3) for _ in range(5)] for _ in range(20)]
    rows[6] = [x + 5 for x in rows[6]]
    # act
    ws = evaluate(SHEET, {**raw_subgroups(rows), INPUTS["first_plotted"]: 3})
    # assert -- row j of the helper block is subgroup 3 + j; the x̄ column is the subgroup mean, the limits repeat the table
    first = subgroup_row(0)
    assert ws[f"{PLOT_INDEX}{first}"].value == 3 and ws[f"{PLOT_INDEX}{first + POINTS - 1}"].value == 3 + POINTS - 1
    assert ws[f"{PLOT_X}{first}"].value == pytest.approx(statistics.mean(rows[2]), rel=1e-9)
    assert ws[f"{PLOT_R}{first}"].value == pytest.approx(max(rows[2]) - min(rows[2]), rel=1e-9)
    columns = _plot_columns("xbar_r")
    assert ws[f"{columns['ucl']}{first}"].value == pytest.approx(limit(ws, "xbar_r", "ucl"), rel=1e-12)
    assert ws[f"{columns['lcl']}{first + 5}"].value == pytest.approx(limit(ws, "xbar_r", "lcl"), rel=1e-12)
    flagged = {r for r in range(first, first + POINTS) if isinstance(ws[f"{columns['out']}{r}"].value, float)}
    in_table = {first + j for j in range(POINTS) if ws[f"{FLAGS['xbar_r']}{subgroup_row(2 + j)}"].value}
    assert flagged == in_table and first + 4 in flagged   # the same subgroups the table flags, shifted subgroup 7 among them
    # beyond the 20 subgroups there is nothing to draw: the cell holds #N/A on purpose (a gap in the chart)
    assert ws[f"{PLOT_X}{first + 25}"].value == "#N/A" and ws[f"{PLOT_INDEX}{first + 25}"].value == 28


def test_the_sheet_carries_four_charts_with_five_series_each() -> None:
    # arrange / act
    ws = build_workbook(only=(SHEET,))[SHEET]
    # assert
    assert len(ws._charts) == len(CHARTS) == 4
    assert all(len(chart.series) == 5 for chart in ws._charts)
