"""Tests for bbtools.sheet_tables: the Tables sheet shows every source as printed and names the lookup ranges."""
from __future__ import annotations

from typing import Any

import pytest
from openpyxl import Workbook

from bbtools.constants import TABLE_SOURCES, USED_TABLE, load_all, load_average_range_table
from bbtools.printed import decimals_printed, parse_printed
from bbtools.sheet_tables import (
    EXTRA_HEADING,
    EXTRA_TABLES,
    SHEET,
    build_tables_sheet,
    disagreeing_cells,
    dutch_title,
    excel_name,
    lookup_formula,
)
from bbtools.xlsx_style import FLAG_FILL


def build() -> Workbook:
    """A workbook holding only the Tables sheet."""
    wb = Workbook()
    build_tables_sheet(wb, load_all())
    return wb


def test_excel_name_makes_every_symbol_a_valid_name() -> None:
    # act / assert
    assert excel_name("T18", "d2") == "T18_d2"
    assert excel_name("T18", "1/d2") == "T18_inv_d2"


def test_lookup_formula_uses_the_table_of_decision_4() -> None:
    # act / assert
    assert lookup_formula("d2", "B17") == "INDEX(T18_d2,MATCH(B17,T18_n,0))"
    assert lookup_formula("c4", "B17") == "INDEX(TA_c4,MATCH(B17,TA_n,0))"
    assert lookup_formula("E2", "2") == "INDEX(SSD2_E2,MATCH(2,SSD2_n,0))"


def test_every_printed_cell_appears_with_its_printed_decimals() -> None:
    # arrange
    wb = build()
    ws = wb[SHEET]
    by_position = {(c.row, c.column): c for row in ws.iter_rows() for c in row}
    # act -- walk each table by its n-column name, which points at its first data row
    for table in load_all():
        first = int(wb.defined_names[f"{table.source.key}_n"].attr_text.split("$")[2].split(":")[0])
        for offset, row in enumerate(table.rows):
            for column, text in enumerate(row, start=1):
                cell = by_position.get((first + offset, column))
                # assert -- numbers stored as numbers and shown with the printed number of decimals
                if text.strip() and _is_number(text):
                    assert cell.value == float(parse_printed(text)), (table.source.key, row[0], column)
                    decimals = decimals_printed(text)
                    assert cell.number_format == ("0" if decimals == 0 else "0." + "0" * decimals)
                else:
                    assert (cell.value or "") == text


def test_every_symbol_column_has_a_workbook_name() -> None:
    # arrange
    wb = build()
    # act
    expected = {f"{t.source.key}_n" for t in load_all()} | {
        excel_name(t.source.key, s) for t in load_all() for s in t.symbol_columns()}
    # assert
    assert expected <= set(wb.defined_names)
    assert len(TABLE_SOURCES) == 5


def test_disagreements_beyond_rounding_are_flagged_and_rounding_is_not() -> None:
    # arrange
    flagged = disagreeing_cells(load_all())
    # act / assert -- D4(5): Table 18 prints 2.115, Six Sigma Demystified 2.114 -> flagged
    assert ("T18", "D4", "5") in flagged and ("SSD2", "D4", "5") in flagged
    # d3(10): Table A 0.7971 vs Demystified 0.797 -> same value at 3 decimals, not flagged
    assert ("TA", "d3", "10") not in flagged


def test_flagged_cells_are_orange_on_the_sheet() -> None:
    # arrange
    wb = build()
    ws = wb[SHEET]
    first = int(wb.defined_names["T18_D4"].attr_text.split("$")[2].split(":")[0])
    column = wb.defined_names["T18_D4"].attr_text.split("$")[1]
    # act -- row for n = 5 is the 4th data row (n = 2, 3, 4, 5)
    cell = ws[f"{column}{first + 3}"]
    # assert
    assert cell.value == 2.115
    assert cell.fill.fgColor.rgb.endswith(FLAG_FILL.fgColor.rgb[-6:])


def test_book_only_table_comes_last_under_the_extra_heading() -> None:
    # arrange -- Six Sigma For Dummies is not examinable: its table goes below every course table
    wb = build()
    ws = wb[SHEET]

    def first_row(name: str) -> int:
        """First data row of a workbook name on the Tables sheet."""
        return int(wb.defined_names[name].attr_text.split("$")[2].split(":")[0])

    heading = next(c.row for c in ws["A"] if c.value == EXTRA_HEADING)
    # act / assert -- heading after every course table and the MSA table, right above the Dummies table
    assert EXTRA_TABLES == {"DUM"}
    course = [first_row(f"{t.source.key}_n") for t in load_all() if t.source.key not in EXTRA_TABLES]
    assert max(course) < first_row("MSA_g") < heading < first_row("DUM_n")
    assert ws.cell(row=heading + 2, column=1).value == dutch_title(next(t.title for t in load_all() if t.source.key == "DUM"))
    # its disagreement colouring stays: D4(5) differs from Table 18
    assert ("DUM", "D4", "5") in disagreeing_cells(load_all())
    column = wb.defined_names["DUM_D4"].attr_text.split("$")[1]
    n_column = wb.defined_names["DUM_n"].attr_text.split("$")[1]
    row = next(r for r in range(first_row("DUM_n"), ws.max_row + 1) if ws[f"{n_column}{r}"].value == 5)
    assert ws[f"{column}{row}"].fill.fgColor.rgb.endswith(FLAG_FILL.fgColor.rgb[-6:])


def _is_number(text: str) -> bool:
    """True if the printed text is one number."""
    try:
        parse_printed(text)
    except ValueError:
        return False
    return True


def test_msa_table_grids_hold_every_printed_value() -> None:
    # arrange -- tabel MSA.pdf: each printed cell 'nu / d2*' is split into two grids
    wb = build()
    ws = wb[SHEET]
    table = load_average_range_table()

    def origin(name: str) -> tuple[int, int]:
        """(row, column) of the top-left cell of a workbook name."""
        ref = wb.defined_names[name].attr_text.split("!")[1].split(":")[0].replace("$", "")
        column = "".join(ch for ch in ref if ch.isalpha())
        return int(ref[len(column):]), ws[f"{column}1"].column

    # act / assert -- every value numeric, equal to the printed string, with the printed decimals
    for name, grid in (("MSA_d2star", table.d2_star), ("MSA_nu", table.nu)):
        row0, col0 = origin(name)
        for i, printed_row in enumerate(grid):
            for j, text in enumerate(printed_row):
                cell = ws.cell(row=row0 + i, column=col0 + j)
                assert cell.value == float(parse_printed(text)), (name, table.g[i], table.m[j])
    row0, col0 = origin("MSA_d2")
    assert [ws.cell(row=row0, column=col0 + j).value for j in range(len(table.m))] == [float(v) for v in table.d2]
    assert table.m == tuple(range(2, 21)) and table.g == tuple(range(1, 21))


@pytest.mark.libreoffice
@pytest.mark.parametrize(("sheet", "top", "column", "symbols"), [
    ("Regelkaarten", 22, 10, ("A2", "D3", "D4", "A3", "B3", "B4", "d2", "c4")),
    ("Capabiliteit", 51, 1, ("d2", "c4")),
])
def test_constants_shown_on_a_sheet_are_the_used_printed_values(sheet: str, top: int, column: int,
                                                                symbols: tuple[str, ...], evaluate: Any) -> None:
    # arrange -- the printed value of each symbol in the table decision 4 assigns to it
    printed = {}
    for table in load_all():
        for symbol, index in table.symbol_columns().items():
            if USED_TABLE.get(symbol) == table.source.key:
                printed |= {(symbol, row[0].strip()): row[index] for row in table.rows}
    # act
    ws = evaluate(sheet, {})
    # assert -- every n from 2 to 25 that the source prints
    for offset, n in enumerate(range(2, 26)):
        for k, symbol in enumerate(symbols, start=1):
            value = ws.cell(row=top + 3 + offset, column=column + k).value
            text = printed.get((symbol, str(n)), "").strip()
            if text and text not in ("-", "—"):
                assert value == pytest.approx(float(parse_printed(text))), (sheet, symbol, n)
