"""Every number that study/parts/06_acceptance_sampling.html marks "zelf berekend", every exercise answer
(data-answer) of Deel 06, and a re-check of the printed course results that the text quotes.

Sources (all read-only):
- source/course/Les 2/20260529_ottoy_Acceptance Sampling.pdf (AS), slide numbers in the comments;
- source/course/Les 2/20260529_ottoy_Acceptance Sampling - Further Reading.pdf (FR);
- source/course/Les 2/20260529_ottoy_Testing of Hypotheses.pdf (TH) and its .xlsx;
- source/course/Les 2/20260529_ottoy_Acceptance Sampling.xlsm (lot data and cached cell values).
Formulas are the course's own (page cited at each block). Run from anywhere:
    python3 study/parts/06_numbers.py
The script prints OK / MISMATCH for every printed course value it re-checks; a MISMATCH is a known
disagreement that the guide text reports. It always exits 0 when it runs to the end.
"""
from __future__ import annotations

import warnings
from decimal import ROUND_HALF_UP, Decimal
from math import comb, floor, log, sqrt
from pathlib import Path

import numpy as np
import openpyxl
from scipy import optimize, stats

ROOT = Path(__file__).resolve().parents[2]
LES2 = ROOT / "source" / "course" / "Les 2"
XLSM = LES2 / "20260529_ottoy_Acceptance Sampling.xlsm"
TH_XLSX = LES2 / "20260529_ottoy_Testing of Hypotheses.xlsx"
CI_XLSX = LES2 / "20260529_ottoy_Confidence Intervals.xlsx"

binom, hyper, norm = stats.binom, stats.hypergeom, stats.norm


def show(label: str, value: float | str, note: str = "") -> None:
    """Print one computed value."""
    text = f"{value:.6g}" if isinstance(value, (float, np.floating)) else str(value)
    print(f"  {label:<60} {text:>16}  {note}")


def answer(ex_id: str, label: str, value: float) -> None:
    """Print one exercise answer (the data-answer of the HTML)."""
    print(f"  ANSWER {ex_id:<9} {label:<50} {value:.6g}")


def half_up(value: float, decimals: int) -> Decimal:
    """Round half-up to a number of decimals (conventions.md decision 10)."""
    # 15 significant digits first, so that binary noise (0.0001945 stored as 0.00019449999...) does not flip
    # a half-up rounding
    return Decimal(f"{float(value):.15g}").quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)


def check(label: str, computed: float, printed: str, note: str = "") -> None:
    """Compare a computed value with a printed course value at the printed precision."""
    decimals = len(printed.split(".")[1]) if "." in printed else 0
    ok = half_up(computed, decimals) == Decimal(printed)
    print(f"  {'OK' if ok else 'MISMATCH':<8} {label:<56} computed {computed:.6g}  printed {printed}  {note}")


def check_close(label: str, computed: float, printed: float, rel_tol: float, note: str = "") -> None:
    """Compare two Monte Carlo results: they can only agree within a relative tolerance, never exactly."""
    ok = abs(computed - printed) <= rel_tol * abs(printed)
    print(f"  {'OK' if ok else 'MISMATCH':<8} {label:<56} computed {computed:.6g}  printed {printed:g}  "
          f"(within {rel_tol:.0%}) {note}")


def cells(path: Path, sheet: str, refs: list[str]) -> dict[str, object]:
    """Cached values of some cells of a course workbook (read-only, as saved by Excel)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    return {ref: ws[ref].value for ref in refs}


def oc_hyper(c: int, N: int, p: float, n: int) -> float:
    """OC(p) = H(c; N, Mp, n) with Mp = [Np] (FR p. 4)."""
    return float(hyper.cdf(c, N, int(floor(N * p + 1e-9)), n))


# ----------------------------------------------------------------------------------------------- blocks
def lot_defectives() -> None:
    """AS p. 10 and 12, xlsm 'defective probabilities': number of defectives i in a lot of N = 10000."""
    print("\n1. Defectives in a lot of N = 10000 (AS p. 10, 12; xlsm 'defective probabilities')")
    n_lot = 10000
    pi3 = 2 * norm.cdf(-3)          # +-3 sigma inside the specs (Cpk = 1)
    pi4 = 2 * norm.cdf(-4)          # +-4 sigma (Cpk = 4/3, the slide's "1.33")
    show("pi at +-3 sigma (fraction)", pi3)
    show("pi at +-4 sigma (ppm)", pi4 * 1e6)
    check("AS p. 10: 0.27 % defective at 6 sigma spec width", pi3 * 100, "0.27")
    check("AS p. 10: 63 ppm at 8 sigma spec width", pi4 * 1e6, "63")
    check("AS p. 10: 99.9937 % conform at 8 sigma", (1 - pi4) * 100, "99.9937")
    check("AS p. 12: Cpk 1.33 -> pi = 0.0063 %", pi4 * 100, "0.0063")
    e3, e4 = n_lot * pi3, n_lot * pi4
    check("AS p. 12: E[i] = 27 at Cpk = 1", e3, "27")
    check("AS p. 12: E[i] = 0.6 at Cpk = 1.33", e4, "0.6")
    show("sd of i at Cpk = 1 = sqrt(N pi (1-pi))", sqrt(n_lot * pi3 * (1 - pi3)))
    show("P(12 <= i <= 44) at Cpk = 1", binom.cdf(44, n_lot, pi3) - binom.cdf(11, n_lot, pi3), "zelf berekend")
    check("AS p. 12 notes: P(i = 100) = 2.1e-27 (x 1e27)", binom.pmf(100, n_lot, pi3) * 1e27, "2.1")
    p0_4 = binom.pmf(0, n_lot, pi4)
    p_le3 = binom.cdf(3, n_lot, pi4)
    show("P(i = 0) at Cpk = 1.33", p0_4, "zelf berekend")
    show("P(i <= 3) at Cpk = 1.33", p_le3, "zelf berekend")
    show("P(i > 3) at Cpk = 1.33", 1 - p_le3, "zelf berekend")
    xl = cells(XLSM, "defective probabilities", ["C3", "C4", "J1", "K1", "H3"])
    check("xlsm C3 (3 sigma fraction defective)", pi3, f"{xl['C3']:.10f}")
    check("xlsm J1 (E[i], 3 sigma)", e3, f"{xl['J1']:.6f}")
    check("xlsm K1 (E[i], 4 sigma)", e4, f"{xl['K1']:.6f}")
    check("xlsm H3 (P[i = 0], 4 sigma)", p0_4, f"{xl['H3']:.8f}")
    answer("ex-06-2", "pi at Cpk = 1 in %", pi3 * 100)
    answer("ex-06-2", "E[i] at Cpk = 1", e3)
    answer("ex-06-2", "E[i] at Cpk = 1.33 (4 sigma)", e4)
    answer("ex-06-2", "P(i = 0) at Cpk = 1.33", p0_4)


def simple_random_sample() -> None:
    """AS p. 16 (+ notes), xlsm 'sampling distribution': SRS of n = 100 from N = 10000, pi = 2 %."""
    print("\n2. Simple random sample (AS p. 16; xlsm 'sampling distribution')")
    n_lot, d_lot, n = 10000, 200, 100
    total = comb(n_lot, n)
    same = comb(d_lot, 2) * comb(n_lot - d_lot, n - 2)
    check("AS p. 16: C(10000,100) = 6.5e241 (/1e241)", total / 1e241, "6.5")
    check("AS p. 16 notes: samples with d = 2 = 1.8e241 (/1e241)", same / 1e241, "1.8")
    check("AS p. 16 notes: their share 27.5 %", same / total * 100, "27.5")
    check("AS p. 16 notes: all-defective samples 9.05e58 (/1e58)", comb(d_lot, n) / 1e58, "9.05")
    check("AS p. 16 notes: their share 1.4e-183 (x 1e183)", comb(d_lot, n) / total * 1e183, "1.4")
    p0_h = hyper.pmf(0, n_lot, d_lot, n)
    p0_b = binom.pmf(0, n, 0.02)
    check("AS p. 20 notes: SRS P[d = 0] = 13 % (hypergeometric)", p0_h * 100, "13")
    show("P[d = 0] binomial", p0_b, "zelf berekend")
    show("P[d = 2] hypergeometric", hyper.pmf(2, n_lot, d_lot, n), "zelf berekend")
    var_b = 0.02 * 0.98 / n
    var_h = float(hyper.var(n_lot, d_lot, n)) / n ** 2
    show("sigma^2[P] binomial = pi(1-pi)/n", var_b)
    show("sigma^2[P] hypergeometric (exact)", var_h, "zelf berekend")
    show("sigma[P] n = 100", sqrt(var_b), "zelf berekend")
    show("sigma[P] n = 500", sqrt(0.02 * 0.98 / 500), "zelf berekend")
    xl = cells(XLSM, "sampling distribution", ["D13", "E13", "D15", "G45"])
    check("xlsm D13 (hypergeometric P[d = 0])", p0_h, f"{xl['D13']:.10f}")
    check("xlsm D15 (hypergeometric P[d = 2])", hyper.pmf(2, n_lot, d_lot, n), f"{xl['D15']:.10f}")
    check("xlsm G45 (variance of P, SRS)", var_b, f"{xl['G45']:.9f}")
    answer("ex-06-3", "sigma^2[P] (binomial)", var_b)
    answer("ex-06-3", "sigma[P] n = 100", sqrt(var_b))
    answer("ex-06-3", "P[d = 0] binomial", p0_b)
    answer("ex-06-3", "sigma[P] n = 500", sqrt(0.02 * 0.98 / 500))


def stratification() -> None:
    """AS p. 17-19: proportional and optimal pre-stratification."""
    print("\n3. Stratification (AS p. 17-19)")
    w_a, w_b, p_a, p_b, n = 0.6, 0.4, 0.03, 0.005, 100
    pi = w_a * p_a + w_b * p_b
    s2, s2a, s2b = pi * (1 - pi), p_a * (1 - p_a), p_b * (1 - p_b)
    v_s = (w_a * s2a + w_b * s2b) / n
    v_p = v_s + w_a * w_b * (p_a - p_b) ** 2 / n
    check("AS p. 17: pi = 2 %", pi * 100, "2")
    check("AS p. 18: sigma^2 = 0.0196", s2, "0.0196")
    check("AS p. 18: sigma_A^2 = 0.0291", s2a, "0.0291")
    check("AS p. 18: sigma_B^2 = 0.004975", s2b, "0.004975")
    check("AS p. 18: sigma^2[Ps] = 0.000195", v_s, "0.000195")
    check("AS p. 18: sigma^2[P] = 0.000196", v_p, "0.000196")
    share = w_a * sqrt(s2a) / (w_a * sqrt(s2a) + w_b * sqrt(s2b))
    check("AS p. 18 notes: optimal share of A = 78 %", share * 100, "78")
    xl = cells(XLSM, "sampling distribution", ["Y48"])
    check("xlsm Y48 (variance of Ps)", v_s, f"{xl['Y48']:.10f}")
    answer("ex-06-4", "sigma_A^2", s2a)
    answer("ex-06-4", "sigma^2[Ps]", v_s)
    answer("ex-06-4", "sigma^2[P] (SRS)", v_p)
    answer("ex-06-4", "optimal n_A (of 100)", share * n)
    # AS p. 19 notes: two normal strata of equal size
    mu_a, mu_b, s2s = 10.0, 15.0, 4.0
    mu = (mu_a + mu_b) / 2
    sig2 = s2s + (mu_a - mu_b) ** 2 / 4
    v_xs = s2s / n
    v_x = v_xs + (mu_a - mu_b) ** 2 / (4 * n)
    check("AS p. 19: population mean 12.5", mu, "12.5")
    show("sigma^2 of the population", sig2, "zelf berekend")
    check("xlsm 'normal distribution' AU: variance 0.1025", v_x, "0.1025")
    check("xlsm 'normal distribution' AV: variance 0.04", v_xs, "0.04")
    answer("ex-06-5", "population variance sigma^2", sig2)
    answer("ex-06-5", "sigma^2[Xbar] (SRS)", v_x)
    answer("ex-06-5", "sigma^2[Xbar_S] (stratified)", v_xs)


def lot_boxes() -> np.ndarray:
    """Defectives per box of the course lot, shaped (10 palettes, 100 boxes) (xlsm 'lot fraction defectives').

    Column B is the box number (1-1000). The cluster procedure on the same sheet puts boxes 1-100 on palette 1,
    101-200 on palette 2, ... (cells Y9, Z9, ...), so palettes are formed from the box number, not from column C.
    """
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(XLSM, read_only=True, data_only=True)
    rows = list(wb["lot fraction defectives"].iter_rows(min_row=5, max_row=10004, max_col=5, values_only=True))
    box = np.array([r[1] for r in rows], int)
    defective = np.array([r[4] for r in rows], float)
    return np.bincount(box, weights=defective)[1:].reshape(10, 100)


def cluster() -> None:
    """AS p. 20: cluster sample of 2 palettes x 5 boxes; extreme case; long-run equivalence."""
    print("\n4. Cluster sampling (AS p. 20; xlsm 'lot fraction defectives', 'sampling distribution')")
    n_samples = comb(10, 2) * comb(100, 5) ** 2
    check("AS p. 20: number of cluster samples 2.6e17 (/1e17)", n_samples / 1e17, "2.6")
    p0_c = (comb(9, 2) * comb(100, 5) ** 2 + 9 * comb(100, 5) * comb(80, 5)) / (comb(10, 2) * comb(100, 5) ** 2)
    check("AS p. 20 notes: extreme case P[d_C = 0] = 86 %", p0_c * 100, "86")
    answer("ex-06-6", "P[d_C = 0] extreme case", p0_c)
    answer("ex-06-6", "P[d = 0] SRS (hypergeometric)", hyper.pmf(0, 10000, 200, 100))
    boxes = lot_boxes()
    show("defectives per palette (course lot)", str(boxes.sum(axis=1).astype(int).tolist()))
    show("total defectives", float(boxes.sum()))
    rng = np.random.default_rng(20260529)          # fixed seed: reproducible
    runs = 200_000
    pals = np.argsort(rng.random((runs, 10)), axis=1)[:, :2]
    picks = np.argsort(rng.random((runs, 2, 100)), axis=2)[:, :, :5]
    d = boxes[pals[:, :, None], picks].sum(axis=(1, 2))
    show("own Monte Carlo (200000 runs): mean of P_C", d.mean() / 100, "zelf berekend")
    show("own Monte Carlo: sigma^2[P_C]", d.var() / 100 ** 2, "zelf berekend")
    show("own Monte Carlo: P[d_C = 0]", float((d == 0).mean()), "zelf berekend")
    xl = cells(XLSM, "sampling distribution", ["AL33", "AL34"])
    check_close("AS p. 20 / xlsm AL34: sigma^2[P_C] = 0.000267 (course MC, 10000 runs)", d.var() / 1e4,
                0.000267, 0.05, "Monte Carlo against Monte Carlo")
    show("xlsm AL33 / AL34 (course MC)", f"{xl['AL33']} / {xl['AL34']}")
    # long-run equivalence, AS p. 20 notes
    coef = 2 * (comb(10, 3) + 10 * comb(10, 2))
    check("AS p. 20 notes: C(20,3) = 1140", comb(20, 3), "1140")
    check("AS p. 20 notes: 2(C(10,3) + 10 C(10,2)) = 1140", coef, "1140")
    answer("ex-06-7", "P(3 defectives) at pi = 2 %", 1140 * 0.02 ** 3 * 0.98 ** 17)


def frame_bias() -> None:
    """AS p. 22: SRS of 100 from a frame with 1 % defectives, population 4 %."""
    print("\n5. Non-response bias (AS p. 22)")
    n = 100
    show("E[d] from the frame", n * 0.01)
    show("E[d] if drawn from the whole population", n * 0.04)
    check("AS p. 22 notes: under-estimate by 3 % points", (0.04 - 0.01) * 100, "3")
    show("P[d = 0] from the frame", binom.pmf(0, n, 0.01), "zelf berekend")
    show("sigma[P] from the frame", sqrt(0.01 * 0.99 / n), "zelf berekend")
    answer("ex-06-8", "expected estimate P from the frame in %", 1.0)
    answer("ex-06-8", "bias in % points", 1.0 - 4.0)
    answer("ex-06-8", "P[d = 0] from the frame", binom.pmf(0, n, 0.01))


def oc_examples() -> None:
    """TH p. 5-8, 11 and TH.xlsx: plan (100,4) with AQL 2 %, plan (130,5), hypergeometric and n = 5000."""
    print("\n6. OC curves of single plans (TH p. 5-8; TH.xlsx; FR p. 4)")
    oc2, oc8 = binom.cdf(4, 100, 0.02), binom.cdf(4, 100, 0.08)
    check("TH p. 5: P[d <= 4 | 2 %] = 95 %", oc2 * 100, "95")
    check("TH p. 6: alpha = 5 %", (1 - oc2) * 100, "5")
    check("TH p. 5: P[d <= 4 | 8 %] = 10 %", oc8 * 100, "10", "p. 7 prints 9 % for the same value")
    check("TH p. 7: beta at 8 % = 9 %", oc8 * 100, "9")
    check("TH p. 5 notes: P[d = 4 | 2 %] almost 10 %", binom.pmf(4, 100, 0.02) * 100, "9")
    check("TH p. 5 notes: P[d = 4 | 8 %] = 5 %", binom.pmf(4, 100, 0.08) * 100, "5")
    a_c5 = 1 - binom.cdf(5, 100, 0.02)
    check("TH p. 6: c = 5 -> alpha = 1.5 %", a_c5 * 100, "1.5")
    a130, b130 = 1 - binom.cdf(5, 130, 0.02), binom.cdf(5, 130, 0.08)
    check("TH p. 8: (130,5) alpha = 5 %", a130 * 100, "5")
    check("TH p. 8: (130,5) beta at 8 % = 5 %", b130 * 100, "5")
    xl = cells(TH_XLSX, "OC-curve (binomial)", ["D10", "E10"])
    check("TH.xlsx 'OC-curve (binomial)' D10", oc2, f"{xl['D10']:.10f}")
    check("TH.xlsx 'OC-curve (binomial)' E10", 1 - a130, f"{xl['E10']:.10f}")
    h100 = oc_hyper(4, 10000, 0.02, 100)
    h5000 = oc_hyper(111, 10000, 0.02, 5000)
    h5000_25 = oc_hyper(111, 10000, 0.025, 5000)
    xh = cells(TH_XLSX, "OC-curve (hypergeometric)", ["D11", "E11", "D12"])
    check("TH.xlsx hypergeometric (100,4) at 2 %", h100, f"{xh['E11']:.10f}")
    check("TH.xlsx hypergeometric (5000,111) at 2 %", h5000, f"{xh['D11']:.10f}")
    check("TH.xlsx hypergeometric (5000,111) at 2.5 %", h5000_25, f"{xh['D12']:.10f}")
    show("(100,4) binomial OC at 2.5 %", binom.cdf(4, 100, 0.025), "zelf berekend")
    # CI.xlsx 'OC-curve (n,c)': helper lines drawn at three fractions (cells G7, G10, G13)
    marks = cells(CI_XLSX, "OC-curve (n,c)", ["G7", "G10", "G13"])
    for ref, printed in [("G7", "0.90"), ("G10", "0.10"), ("G13", "0.99")]:
        check(f"CI.xlsx {ref}: OC of (100,4) at pi = {marks[ref]}", binom.cdf(4, 100, float(marks[ref])), printed)
    answer("ex-06-11", "OC(2 %) of (100,4)", oc2)
    answer("ex-06-11", "alpha in %", (1 - oc2) * 100)
    answer("ex-06-11", "beta at 8 % in %", oc8 * 100)
    answer("ex-06-11", "alpha with c = 5 in %", a_c5 * 100)
    answer("ex-06-12", "alpha of (130,5) in %", a130 * 100)
    answer("ex-06-12", "beta at 8 % of (130,5) in %", b130 * 100)
    answer("ex-06-13", "OC(2 %) hypergeometric (100,4), N = 10000", h100)
    answer("ex-06-13", "OC(2.5 %) of (5000,111)", h5000_25)
    answer("ex-06-13", "OC(2.5 %) of (100,4) binomial", binom.cdf(4, 100, 0.025))


def peach() -> None:
    """FR p. 4: method of Peach example (0.5 %, 95 %) and (3.5 %, 5 %) -> (164,2); chart plan (160,2)."""
    print("\n7. Peach example (FR p. 4)")
    check("FR p. 4: R0 = p2/p1 = 7", 0.035 / 0.005, "7")
    for n, c in [(164, 2), (160, 2)]:
        show(f"({n},{c}) OC(0.5 %) binomial", binom.cdf(c, n, 0.005), "zelf berekend")
        show(f"({n},{c}) OC(3.5 %) binomial", binom.cdf(c, n, 0.035), "zelf berekend")
    check("FR p. 4: (164,2) gives beta = 5 % at 3.5 %", binom.cdf(2, 164, 0.035) * 100, "5",
          "approximate method; printed value disagrees")
    p_at_5 = optimize.brentq(lambda p: binom.cdf(2, 164, p) - 0.05, 0.001, 0.2)
    show("p where OC of (164,2) = 5 %", p_at_5, "zelf berekend")
    answer("ex-06-14", "R0", 7.0)
    answer("ex-06-14", "OC(0.5 %) of (164,2)", binom.cdf(2, 164, 0.005))
    answer("ex-06-14", "OC(3.5 %) of (164,2) in %", binom.cdf(2, 164, 0.035) * 100)


def double_plan() -> None:
    """AS p. 25, FR p. 5: (175,8) versus (90,2,7)+(90,8); OC, Pi and ASN (binomial)."""
    print("\n8. Double plan (AS p. 25; FR p. 5)")
    n1, c1, c2, n2, c3 = 90, 2, 7, 90, 8

    def double(p: float) -> tuple[float, float, float]:
        """OC, probability of a decision after the first sample (Pi), and ASN of the double plan."""
        oc = binom.cdf(c1, n1, p) + sum(binom.pmf(j, n1, p) * binom.cdf(c3 - j, n2, p) for j in range(c1 + 1, c2))
        pi_first = binom.cdf(c1, n1, p) + 1 - binom.cdf(c2 - 1, n1, p)
        return oc, pi_first, n1 * pi_first + (n1 + n2) * (1 - pi_first)

    for p in [0.01, 0.02, 0.03, 0.04, 0.05, 0.06, 0.08, 0.10]:
        oc, pi_first, asn = double(p)
        show(f"p = {p:.2f}: OC single (175,8) / OC double / ASN",
             f"{binom.cdf(8, 175, p):.4f} / {oc:.4f} / {asn:.1f}", "zelf berekend")
    grid = np.linspace(0.001, 0.2, 2000)
    asns = [double(p)[2] for p in grid]
    show("maximum ASN of the double plan and its p", f"{max(asns):.1f} at {grid[int(np.argmax(asns))]:.3f}",
         "zelf berekend")
    oc3, pi3, asn3 = double(0.03)
    acc1, rej1 = binom.cdf(c1, n1, 0.03), 1 - binom.cdf(c2 - 1, n1, 0.03)
    answer("ex-06-16", "P(accept after sample 1) at 3 %", acc1)
    answer("ex-06-16", "P(reject after sample 1) at 3 %", rej1)
    answer("ex-06-16", "Pi(3 %) = decision after sample 1", pi3)
    answer("ex-06-16", "ASN(3 %)", asn3)
    answer("ex-06-16", "OC(3 %) double", oc3)
    answer("ex-06-16", "OC(3 %) single (175,8)", binom.cdf(8, 175, 0.03))


def sprt_constants(p0: float, pt: float, a: float, b: float) -> tuple[float, float, float]:
    """h1, h2 and s of Wald's SPRT (FR p. 7)."""
    g = log(pt * (1 - p0) / (p0 * (1 - pt)))
    return log((1 - a) / b) / g, log((1 - b) / a) / g, log((1 - p0) / (1 - pt)) / g


def sprt_asn(p: float, p0: float, pt: float, a: float, b: float) -> tuple[float, float]:
    """Wald's approximate ASN(p) and OC(p) of the SPRT (FR p. 8), via the parameter tau."""
    h1, h2, s = sprt_constants(p0, pt, a, b)
    if abs(p - s) < 1e-12:
        return h1 * h2 / (s * (1 - s)), float("nan")
    big_a, big_b = (1 - b) / a, b / (1 - a)
    r, rr = (1 - pt) / (1 - p0), pt / p0
    if p == 0:
        oc = 1.0
    else:
        tau = optimize.brentq(lambda t: (1 - r ** t) / (rr ** t - r ** t) - p if abs(t) > 1e-12 else s - p, -60, 60)
        oc = (big_a ** tau - 1) / (big_a ** tau - big_b ** tau)
    num = oc * log(big_b) + (1 - oc) * log(big_a)
    den = p * log(rr) + (1 - p) * log(r)
    return num / den, oc


def sequential() -> None:
    """FR p. 7-8: SPRT examples."""
    print("\n9. Sequential sampling, SPRT (FR p. 7-8)")
    h1, h2, s = sprt_constants(0.01, 0.05, 0.05, 0.05)
    check("FR p. 7: h1 = 1.78", h1, "1.78")
    check("FR p. 7: h2 = 1.78", h2, "1.78")
    check("FR p. 7: s = 0.025", s, "0.025")
    acc, rej = floor(h1 / s) + 1, floor(h2 / (1 - s)) + 1
    show("shortest path to accept n = [h1/s] + 1", acc, "zelf berekend")
    show("shortest path to reject n = [h2/(1-s)] + 1", rej, "zelf berekend")
    asn_s = h1 * h2 / (s * (1 - s))
    show("ASN at p = s", asn_s, "zelf berekend")
    answer("ex-06-17", "h1 (= h2)", h1)
    answer("ex-06-17", "s", s)
    answer("ex-06-17", "shortest path to accept", acc)
    answer("ex-06-17", "shortest path to reject", rej)
    answer("ex-06-17", "ASN at p = s", asn_s)
    h1b, h2b, sb = sprt_constants(0.005, 0.035, 0.05, 0.05)
    show("(0.5 %, 3.5 %) h1 = h2 / s", f"{h1b:.4f} / {sb:.5f}", "zelf berekend")
    for p in [0.0, 0.005, 0.01, 0.012, sb, 0.02, 0.035, 0.05]:
        asn, oc = sprt_asn(p, 0.005, 0.035, 0.05, 0.05)
        oc_text = "-" if np.isnan(oc) else f"{oc:.3f}"
        show(f"(0.5 %, 3.5 %) ASN({p:.4f}) / OC", f"{asn:.1f} / {oc_text}", "zelf berekend")
    grid = np.linspace(0.0005, 0.1, 400)
    vals = [sprt_asn(p, 0.005, 0.035, 0.05, 0.05)[0] for p in grid]
    show("(0.5 %, 3.5 %) maximum ASN and its p", f"{max(vals):.1f} at {grid[int(np.argmax(vals))]:.4f}",
         "zelf berekend; FR p. 8 compares with n = 164")
    answer("ex-06-18", "ASN(0)", sprt_asn(0.0, 0.005, 0.035, 0.05, 0.05)[0])
    answer("ex-06-18", "ASN at p = s", sprt_asn(sb, 0.005, 0.035, 0.05, 0.05)[0])


def variables() -> None:
    """AS p. 27 (notes) and xlsm 'variable sampling plan'; FR p. 9 (k, n and the OC approximation)."""
    print("\n10. Sampling plans for variables (AS p. 26-28; FR p. 9; xlsm 'variable sampling plan')")
    z = norm.ppf
    mu, sigma, n, alpha, p0 = 10.0, 2.0, 10, 0.05, 0.02
    xi = z(p0, mu, sigma)
    z1p, za = z(1 - p0), z(alpha)
    t = (z1p + za * sqrt(1 / n + z1p ** 2 / (2 * n) - za ** 2 / (2 * n ** 2))) / (1 - za ** 2 / (2 * n))
    xl = cells(XLSM, "variable sampling plan", ["C14", "I4", "I5", "I6", "F4", "F5", "F6", "C17"])
    check("xlsm C14 xi = X_p0", xi, f"{xl['C14']:.10f}")
    check("xlsm I6 t(1-alpha, p0, n)", t, f"{xl['I6']:.10f}")
    xbar, s = float(xl["F4"]), float(xl["F5"])
    q = (xbar - xi) / s
    check("xlsm F6 Q", q, f"{xl['F6']:.10f}")

    def oc_var(p: float, k: float, size: float) -> float:
        """OC(p) ~ 1 - Phi((z_p + k) sqrt(n) / sqrt(1 + k^2/2)) (FR p. 9)."""
        return float(1 - norm.cdf((z(p) + k) * sqrt(size) / sqrt(1 + k * k / 2)))

    oc_k2, oc_kt = oc_var(p0, 2.0, n), oc_var(p0, t, n)
    show("OC(p0) with k = 2, n = 10 (FR p. 9 approximation)", oc_k2, "zelf berekend; xlsm C17 = 0.61 from 100 runs")
    show("OC(p0) with k = t (should be 1 - alpha)", oc_kt, "zelf berekend")
    show("sd of a 100-run Monte Carlo fraction near 0.6 = sqrt(0.61*0.39/100)", sqrt(0.61 * 0.39 / 100),
         "zelf berekend (sigma^2[P] = pi(1-pi)/n, AS p. 16)")
    answer("ex-06-19", "xi", xi)
    answer("ex-06-19", "t(1-alpha, p0, n)", t)
    answer("ex-06-19", "Q", q)
    answer("ex-06-19", "OC(p0) with k = 2 (approximation)", oc_k2)
    # k and n from AQL and LQL (FR p. 9)
    for p_0, p_t, a, b, ex in [(0.005, 0.035, 0.05, 0.05, "ex-06-20"), (0.005, 0.035, 0.05, 0.10, "ex-06-20b"),
                               (0.02, 0.04, 0.05, 0.05, "ex-06-24")]:
        k = (z(p_t) * z(a) + z(p_0) * z(b)) / (z(1 - a) + z(1 - b))
        size = (z(1 - a) + z(1 - b)) ** 2 * (1 + k * k / 2) / (z(p_t) - z(p_0)) ** 2
        show(f"({p_0}, {p_t}, alpha {a}, beta {b}) k / n", f"{k:.4f} / {size:.2f}")
        show("  OC at AQL and LQL with n rounded up", f"{oc_var(p_0, k, np.ceil(size)):.4f} / "
             f"{oc_var(p_t, k, np.ceil(size)):.4f}", "zelf berekend")
        if ex == "ex-06-20":
            check("FR p. 9: k = 2.1939", k, "2.1939")
            check("FR p. 9: n = 64 (rounded up)", float(np.ceil(size)), "64")
        answer(ex.replace("b", ""), f"k ({p_0}, {p_t}, beta {b})", k)
        answer(ex.replace("b", ""), f"n unrounded ({p_0}, {p_t}, beta {b})", size)
        answer(ex.replace("b", ""), f"n rounded up ({p_0}, {p_t}, beta {b})", float(np.ceil(size)))


def rectifying() -> None:
    """FR p. 10: AOQ, AOQL, ATI for (250,5) and N = 1000."""
    print("\n11. Rectifying inspection (FR p. 10)")
    N, n, c = 1000, 250, 5
    rows = []
    for m in range(0, 101):
        p = m / N
        oc = float(hyper.cdf(c, N, m, n))
        rs = float(hyper.cdf(c - 1, N - 1, m - 1, n - 1)) / oc if m > 0 else 0.0
        exact = p * oc * (1 - n / N * rs)
        direct = sum((p - i / N) * hyper.pmf(i, N, m, n) for i in range(c + 1))
        rows.append((p, oc, rs, exact, direct, p * oc, p * binom.cdf(c, n, p), n * oc + N * (1 - oc)))
    r = np.array(rows)
    assert np.allclose(r[:, 3], r[:, 4]), "exact AOQ formula must equal the direct sum of FR p. 10 notes"
    i_ex, i_ap = int(np.argmax(r[:, 3])), int(np.argmax(r[:, 5]))
    show("AOQL exact (slide formula) and its p", f"{r[i_ex, 3]:.5f} at {r[i_ex, 0]:.3f}", "zelf berekend")
    show("AOQL approximation p*OC(p) (hypergeometric OC) and its p", f"{r[i_ap, 5]:.5f} at {r[i_ap, 0]:.3f}",
         "zelf berekend")
    i_b = int(np.argmax(r[:, 6]))
    show("AOQL approximation p*OC(p) (binomial OC) and its p", f"{r[i_b, 6]:.5f} at {r[i_b, 0]:.3f}", "zelf berekend")
    check("FR p. 10: AOQL approximately 1.3 % (approximation)", r[i_ap, 5] * 100, "1.3")
    check("FR p. 10: AOQL approximately 1.3 % (exact formula)", r[i_ex, 3] * 100, "1.3", "printed value disagrees")
    check("FR p. 10: at about 1.8 % (location of the maximum)", r[i_ap, 0] * 100, "1.8", "printed value disagrees")
    show("AOQ approximation at p = 1.8 %", r[18, 5], "zelf berekend (curve is flat near its top)")
    for m in (10, 17, 20, 30):
        p, oc, rs, exact, _, approx, _, ati = r[m]
        show(f"p = {p:.3f}: OC / Rs* / AOQ exact / AOQ approx / ATI",
             f"{oc:.4f} / {rs:.4f} / {exact:.5f} / {approx:.5f} / {ati:.1f}", "zelf berekend")
    p, oc, rs, exact, _, approx, _, ati = r[20]
    answer("ex-06-21", "OC(2 %) hypergeometric", oc)
    answer("ex-06-21", "AOQ(2 %) approximation in %", approx * 100)
    answer("ex-06-21", "AOQ(2 %) exact in %", exact * 100)
    answer("ex-06-21", "ATI(2 %)", ati)
    answer("ex-06-21", "AOQL exact in %", r[i_ex, 3] * 100)
    answer("ex-06-21", "AOQL approximation in %", r[i_ap, 5] * 100)


def skip_lot() -> None:
    """FR p. 11-12: qualification probability for skip-lot inspection; AQL of (80,2) with N = 1000."""
    print("\n12. Skip-lot and ISO 2859 (FR p. 11-12)")
    for p, printed in [(0.01, "2.4"), (0.001, "85")]:
        pq = binom.cdf(3, 640, p) * binom.cdf(0, 80, p) ** 2
        check(f"FR p. 11: Pq at p = {p}", pq * 100, printed)
        answer("ex-06-22", f"Pq at p = {p} in %", pq * 100)
    oc_h, oc_b = oc_hyper(2, 1000, 0.01, 80), binom.cdf(2, 80, 0.01)
    show("(80,2), N = 1000: OC(1 %) hypergeometric / binomial", f"{oc_h:.4f} / {oc_b:.4f}", "zelf berekend")
    aql_b = optimize.brentq(lambda p: binom.cdf(2, 80, p) - 0.95, 1e-4, 0.1)
    show("(80,2): p with OC = 95 % (binomial)", aql_b, "zelf berekend")
    answer("ex-06-22", "OC(1 %) of (80,2), N = 1000, hypergeometric", oc_h)


def further_exercises() -> None:
    """FR p. 17-18: plan (2 %, 95 %) and (4 %, 5 %); inventive inspection; lot size."""
    print("\n13. Further Reading exercises (FR p. 17-18)")
    found = None
    for c in range(0, 60):
        n = c + 1
        while binom.cdf(c, n, 0.04) > 0.05:
            n += 1
        if binom.cdf(c, n, 0.02) >= 0.95:
            found = (n, c)
            break
    assert found is not None
    n, c = found
    show("smallest binomial single plan with OC(2 %) >= 95 % and OC(4 %) <= 5 %", f"({n},{c})", "zelf berekend")
    show("its OC(2 %) / OC(4 %)", f"{binom.cdf(c, n, 0.02):.4f} / {binom.cdf(c, n, 0.04):.4f}", "zelf berekend")
    answer("ex-06-24", "n of the smallest plan", n)
    answer("ex-06-24", "c of the smallest plan", c)
    answer("ex-06-24", "OC(2 %) of that plan", binom.cdf(c, n, 0.02))
    answer("ex-06-24", "OC(4 %) of that plan", binom.cdf(c, n, 0.04))
    # inventive inspection
    lql_56 = 1 - 0.1 ** (1 / 56)
    lql_95 = optimize.brentq(lambda p: binom.cdf(1, 95, p) - 0.10, 1e-4, 0.5)
    combined = (1 - lql_56) ** 56 + 56 * lql_56 * (1 - lql_56) ** 55 * (1 - lql_56) ** 39
    show("LQL of (56,0) / of (95,1)", f"{lql_56:.5f} / {lql_95:.5f}", "zelf berekend")
    show("acceptance probability of the inspector's procedure at that LQL", combined, "zelf berekend")
    for p in (0.01, 0.02, 0.03, 0.05):
        comb_p = (1 - p) ** 56 + 56 * p * (1 - p) ** 94
        show(f"p = {p}: (56,0) / (95,1) / inspector", f"{(1 - p) ** 56:.4f} / {binom.cdf(1, 95, p):.4f} / "
             f"{comb_p:.4f}", "zelf berekend")
    answer("ex-06-15", "LQL of (56,0) in %", lql_56 * 100)
    answer("ex-06-15", "LQL of (95,1) in %", lql_95 * 100)
    answer("ex-06-15", "P(accept) inspector's procedure at LQL in %", combined * 100)
    # lot size
    for lot in (200, 500, 1000, 10000):
        show(f"(80,2) OC(2 %) with N = {lot} (hypergeometric)", oc_hyper(2, lot, 0.02, 80), "zelf berekend")
    show("(80,2) OC(2 %) binomial", binom.cdf(2, 80, 0.02), "zelf berekend")
    answer("ex-06-25", "OC(2 %) of (80,2), N = 200", oc_hyper(2, 200, 0.02, 80))
    answer("ex-06-25", "OC(2 %) of (80,2), N = 10000", oc_hyper(2, 10000, 0.02, 80))


def intermediate_steps() -> None:
    """Intermediate values that the worked solutions of the HTML show (all from the course data above)."""
    print("\n14. Intermediate values shown in the solutions")
    z = norm.ppf
    show("sqrt(5): error ratio n = 100 versus n = 500", sqrt(5), "zelf berekend")
    show("sigma_A / sigma_B (tiles)", f"{sqrt(0.0291):.4f} / {sqrt(0.004975):.4f}", "zelf berekend")
    show("sigma of Xbar_S / of Xbar (normal strata)", f"{sqrt(0.04):.4f} / {sqrt(0.1025):.4f}", "zelf berekend")
    show("cluster: P(bad palette not chosen) = 36/45", comb(9, 2) / comb(10, 2), "zelf berekend")
    show("cluster: P(bad palette chosen) = 9/45", 9 / comb(10, 2), "zelf berekend")
    show("cluster: C(80,5)/C(100,5)", comb(80, 5) / comb(100, 5), "zelf berekend")
    xh = cells(TH_XLSX, "OC-curve (hypergeometric)", ["E12", "F12"])
    check("TH.xlsx hypergeometric (100,4) at 2.5 %", oc_hyper(4, 10000, 0.025, 100), f"{xh['E12']:.10f}")
    check("TH.xlsx binomial (100,4) at 2.5 %", binom.cdf(4, 100, 0.025), f"{xh['F12']:.10f}")
    p = 1 - 0.1 ** (1 / 56)
    show("inspector at LQL: (1-p)^56 / second-chance term",
         f"{(1 - p) ** 56:.4f} / {56 * p * (1 - p) ** 94:.4f}", "zelf berekend")
    show("double plan: 175 - ASN(3 %)", 175 - 134.118, "zelf berekend (ASN from block 8)")
    g1 = log(0.05 * 0.99 / (0.01 * 0.95))
    show("SPRT 1 %/5 %: g, ln 19, ln(0.99/0.95)", f"{g1:.4f} / {log(19):.4f} / {log(0.99 / 0.95):.5f}", "zelf berekend")
    show("SPRT 1 %/5 %: h1/s and h2/(1-s)", f"{sprt_constants(.01, .05, .05, .05)[0] / sprt_constants(.01, .05, .05, .05)[2]:.2f}"
         f" / {sprt_constants(.01, .05, .05, .05)[1] / (1 - sprt_constants(.01, .05, .05, .05)[2]):.3f}", "zelf berekend")
    g2 = log(0.035 * 0.995 / (0.005 * 0.965))
    show("SPRT 0.5 %/3.5 %: g, ln(0.05/0.95), ln(0.965/0.995)",
         f"{g2:.4f} / {log(0.05 / 0.95):.4f} / {log(0.965 / 0.995):.5f}", "zelf berekend")
    show("z quantiles 0.98 / 0.05 / 0.035 / 0.005 / 0.10 / 0.04 / 0.02",
         " / ".join(f"{z(q):.4f}" for q in (0.98, 0.05, 0.035, 0.005, 0.10, 0.04, 0.02)))
    n_t, za, z1p = 10, z(0.05), z(0.98)
    show("t: 1/n + z1p^2/(2n) - za^2/(2n^2), its root",
         f"{1 / n_t + z1p ** 2 / (2 * n_t) - za ** 2 / (2 * n_t ** 2):.5f} / "
         f"{sqrt(1 / n_t + z1p ** 2 / (2 * n_t) - za ** 2 / (2 * n_t ** 2)):.4f}", "zelf berekend")
    show("OC argument (z_0.02 + 2) sqrt(10) / sqrt(3)", (z(0.02) + 2) * sqrt(10) / sqrt(3), "zelf berekend")
    show("skip-lot: B(3;640,0.01) / B(0;80,0.01) / B(3;640,0.001) / B(0;80,0.001)",
         f"{binom.cdf(3, 640, 0.01):.4f} / {binom.cdf(0, 80, 0.01):.4f} / {binom.cdf(3, 640, 0.001):.4f} / "
         f"{binom.cdf(0, 80, 0.001):.4f}", "zelf berekend")
    last = 781
    while binom.cdf(22, last + 1, 0.02) >= 0.95:
        last += 1
    show("c = 22: largest n that still has OC(2 %) >= 95 %", last, "zelf berekend")


def same_ratio() -> None:
    """Pitfall box of 06.9: plan (1000,40) has the same c/n = 4 % as the course plan (100,4) (TH p. 5-8),
    yet a different OC-curve: AQL and LQL both move towards c/n (binomial, as in TH.xlsx)."""
    print("\n15. Pitfall 06.9: (100,4) versus (1000,40), same c/n (binomial)")
    for n, c in [(100, 4), (1000, 40)]:
        # AQL = p with OC = 1 - alpha = 95 %, LQL = p with OC = beta = 10 % (FR p. 4)
        aql = optimize.brentq(lambda p: binom.cdf(c, n, p) - 0.95, 1e-6, 0.5)
        lql = optimize.brentq(lambda p: binom.cdf(c, n, p) - 0.10, 1e-6, 0.5)
        show(f"({n},{c}) OC at 2 % / 5 % / 8 %",
             " / ".join(f"{binom.cdf(c, n, p):.4f}" for p in (0.02, 0.05, 0.08)), "zelf berekend")
        show(f"({n},{c}) p with OC = 95 % (AQL) / OC = 10 % (LQL)", f"{aql:.4f} / {lql:.4f}", "zelf berekend")


def main() -> None:
    """Run every block in the order of the text."""
    print("Deel 06 - numbers (zelf berekend), exercise answers and checks of printed course values")
    lot_defectives()
    simple_random_sample()
    stratification()
    cluster()
    frame_bias()
    oc_examples()
    peach()
    double_plan()
    sequential()
    variables()
    rectifying()
    skip_lot()
    further_exercises()
    intermediate_steps()
    same_ratio()
    print("\nDone.")


if __name__ == "__main__":
    main()
