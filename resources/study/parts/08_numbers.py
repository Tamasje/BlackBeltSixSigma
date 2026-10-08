"""Deel 08 (ANOVA en proefopzet) of the study guide: every number the fragment 08_anova_doe.html marks "zelf
berekend", every exercise answer (data-answer), and a re-check of every printed course result the fragment quotes.

Data come from inventory/constants (each row carries its course file and page), from the course workbooks
source/course/Les 3/20260605_devuyst_DOE_demo.xlsx and source/course/Les 5/20260619_ottoy_GRR - ANOVA - avegage and
range - 2.xlsx (read-only, openpyxl), or, for data that exist only in a figure, from the slide cited next to them.
Lines "OK" / "MISMATCH" compare a computation with a printed course value (half-up rounding to the printed
precision, convention decision 10); "MISMATCH (erratum)" marks a known slide error the guide text explains.
Run from the project root:  python3 study/parts/08_numbers.py
"""
from __future__ import annotations

import csv
from decimal import ROUND_HALF_UP, Decimal
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.formula.api as smf
from scipy import stats
from statsmodels.stats.anova import anova_lm

ROOT = Path(__file__).resolve().parents[2]
CONSTANTS = ROOT / "inventory" / "constants"
DEMO = ROOT / "source" / "course" / "Les 3" / "20260605_devuyst_DOE_demo.xlsx"
GRR = ROOT / "source" / "course" / "Les 5" / "20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx"
UNEXPECTED: list[str] = []


def rows(stem: str) -> list[dict[str, str]]:
    """Rows of a course table in inventory/constants."""
    with (CONSTANTS / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def sheet(path: Path, name: str):
    """A worksheet with cached values, opened read-only."""
    return openpyxl.load_workbook(path, data_only=True, read_only=True)[name]


def show(label: str, value: object) -> None:
    """Print one derived number (or a short text)."""
    text = f"{value:.6g}" if isinstance(value, (float, np.floating)) else str(value)
    print(f"  {label:<58} {text}")


def check(label: str, computed: float, printed: str, erratum: bool = False) -> None:
    """Compare a computed value with a printed one after half-up rounding to the printed precision."""
    places = len(printed.split(".")[1]) if "." in printed else 0
    rounded = Decimal(repr(float(computed))).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    ok = rounded == Decimal(printed)
    status = "OK" if ok else ("MISMATCH (erratum)" if erratum else "MISMATCH")
    if not ok and not erratum:
        UNEXPECTED.append(label)
    print(f"  {status:<19} {label:<52} printed {printed:>12}  computed {float(computed):.6g}")


def check_sci(label: str, computed: float, printed: str, erratum: bool = False) -> None:
    """Compare a small p-value with a printed one in scientific notation (same number of mantissa digits)."""
    digits = len(printed.lower().split("e")[0].replace("-", "").replace(".", "")) - 1
    ok = f"{computed:.{digits}e}".lower() == f"{float(printed):.{digits}e}".lower()
    status = "OK" if ok else ("MISMATCH (erratum)" if erratum else "MISMATCH")
    if not ok and not erratum:
        UNEXPECTED.append(label)
    print(f"  {status:<19} {label:<52} printed {printed:>12}  computed {computed:.6g}")


def one_way(title: str, groups: list[list[float]], alpha: float) -> dict[str, float]:
    """One-way ANOVA (DOE p. 6-8): prints and returns the table's numbers."""
    print(title)
    y = np.concatenate([np.asarray(g, float) for g in groups])
    a, n_total = len(groups), len(y)
    means = [float(np.mean(g)) for g in groups]
    ss_t = float(np.sum((y - y.mean()) ** 2))
    ss_tr = float(sum(len(g) * (m - y.mean()) ** 2 for g, m in zip(groups, means)))
    ss_e = ss_t - ss_tr
    df_tr, df_e = a - 1, n_total - a
    f0 = (ss_tr / df_tr) / (ss_e / df_e)
    out = {"ss_t": ss_t, "ss_tr": ss_tr, "ss_e": ss_e, "ms_tr": ss_tr / df_tr, "ms_e": ss_e / df_e, "f0": f0,
           "p": float(stats.f.sf(f0, df_tr, df_e)), "fcrit": float(stats.f.isf(alpha, df_tr, df_e)),
           "pooled": float(np.sqrt(ss_e / df_e))}
    show("group totals / means", f"{[float(np.sum(g)) for g in groups]} / {[round(m, 4) for m in means]}")
    show("grand total / grand mean", f"{y.sum():g} / {y.mean():.4f}")
    show("sum of squared values / correction (total)^2 / N", f"{np.sum(y ** 2):g} / {y.sum() ** 2 / n_total:.4f}")
    for key in ("ss_t", "ss_tr", "ss_e", "ms_tr", "ms_e", "f0", "p", "fcrit", "pooled"):
        show(key, out[key])
    show("df treatments / error / total", f"{df_tr} / {df_e} / {n_total - 1}")
    return out


def signs(k: int) -> np.ndarray:
    """Coded design matrix of a 2^k in standard order (A changes fastest), shape (2^k, k)."""
    return np.array([[1 if r >> j & 1 else -1 for j in range(k)] for r in range(2 ** k)])


def effects(k: int, runs: list[list[float]]) -> dict[str, tuple[float, float, float]]:
    """Contrast, effect and SS of every effect of a balanced 2^k in standard order (DOE p. 57, 66)."""
    x, letters, n = signs(k), "ABCDE"[:k], len(runs[0])
    totals = np.array([sum(r) for r in runs])
    out = {}
    for size in range(1, k + 1):
        for combo in combinations(range(k), size):
            contrast = float(np.prod(x[:, combo], axis=1) @ totals)
            out["".join(letters[j] for j in combo)] = (contrast, contrast / (n * 2 ** (k - 1)), contrast ** 2 / (n * 2 ** k))
    return out


def pure_error(runs: list[list[float]]) -> float:
    """Sum of squared deviations of replicates from their run mean."""
    return float(sum(np.sum((np.asarray(r, float) - np.mean(r)) ** 2) for r in runs))


def word_product(a: str, b: str) -> str:
    """Product of two effect words with A·A = I (DOE p. 81): the letters that occur in exactly one of them."""
    return "".join(sorted(set(a.replace("I", "")) ^ set(b.replace("I", "")))) or "I"


def defining_relation(generator_words: list[str]) -> list[str]:
    """All words of the complete defining relation from the generator words (DOE p. 91)."""
    words = set()
    for size in range(1, len(generator_words) + 1):
        for combo in combinations(generator_words, size):
            word = "I"
            for g in combo:
                word = word_product(word, g)
            words.add(word)
    return sorted(words, key=lambda w: (len(w), w))


def anova_part() -> None:
    """8.1-8.8: one-way ANOVA examples and exercises."""
    paper = [[float(r[f"Obs{i}"]) for i in range(1, 7)]
             for r in rows("S05_table_4_4_tensile_strength_of_paper_psi") if r["Obs1"]]
    res = one_way("8.5 Paper tensile strength (DOE p. 4, 9-10), alpha 0.01", paper, 0.01)
    printed = {r["Source"]: r for r in rows("S05_table_4_7_minitab_anova_output_for_paper_tensile_strength_ex")}
    check("SS_T (p. 9)", res["ss_t"], "512.96")
    check("SS_Treatments (p. 9)", res["ss_tr"], "382.79")
    check("SS_E (p. 9)", res["ss_e"], "130.17")
    check("MS Factor (p. 10)", res["ms_tr"], printed["Factor"]["MS"])
    check("MS Error (p. 10)", res["ms_e"], printed["Error"]["MS"])
    check("F (p. 9-10)", res["f0"], printed["Factor"]["F"])
    check("F_0.01,3,20 (p. 9)", res["fcrit"], "4.94")
    check_sci("P (p. 9: 3.59 x 10^-6)", res["p"], "3.59e-6")
    check("Pooled StDev (p. 10)", res["pooled"], "2.551")
    for r in rows("S05_table_4_7_continued_level_summary_individual_95_cis_for_mean"):
        group = paper[[5, 10, 15, 20].index(int(r["Level"]))]
        check(f"level {r['Level']}: mean (p. 10)", np.mean(group), r["Mean"])
        check(f"level {r['Level']}: StDev (p. 10)", np.std(group, ddof=1), r["StDev"])

    fibre_csv = [[float(r[f"Obs{i}"]) for i in range(1, 6)]
                 for r in rows("S05_table_3_1_data_in_lb_in_2_from_the_tensile_strength_experime") if r["Obs1"]]
    ws = sheet(DEMO, "1ANOVA")
    fibre = [[float(v) for v in r[1:6]] for r in ws.iter_rows(min_row=4, max_row=8, values_only=True)]
    show("DOE_demo '1ANOVA' equals the fibre data (Montgomery Table 3-1)", fibre == fibre_csv)
    res = one_way("8.6 Synthetic fibre = DOE_demo '1ANOVA' (DOE p. 13-19), alpha 0.05", fibre, 0.05)
    out = {r["Source of Variation"]: r for r in rows("S05_anova_computer_output_excel_for_synthetic_fiber_cotton_weigh")}
    check("Between Groups SS", res["ss_tr"], out["Between Groups"]["SS"])
    check("Within Groups SS", res["ss_e"], out["Within Groups"]["SS"])
    check("Total SS", res["ss_t"], out["Total"]["SS"])
    check("Between MS", res["ms_tr"], out["Between Groups"]["MS"])
    check("Within MS", res["ms_e"], out["Within Groups"]["MS"])
    check("F", res["f0"], out["Between Groups"]["F"])
    check_sci("P-value", res["p"], out["Between Groups"]["P-value"])
    check("F crit (alpha 0.05)", res["fcrit"], out["Between Groups"]["F crit"])
    for r, group in zip(rows("S05_anova_computer_output_excel_group_summary_for_synthetic_fibe"), fibre):
        check(f"{r['Groups']} Average", np.mean(group), r["Average"])
        check(f"{r['Groups']} Variance", np.var(group, ddof=1), r["Variance"])
    check("scale factor sqrt(MS_E / n) (p. 19)", np.sqrt(res["ms_e"] / 5), "1.27")
    show("F_0.01,4,20 (p. 16 figure marks it, no value printed)", f"{stats.f.isf(0.01, 4, 20):.3f}")

    machine_rows = rows("S05_table_10_1_sample_outputs_of_three_machines")
    machines = [[float(r[c]) for r in machine_rows if r[c].strip().lstrip("-").replace(".", "").isdigit()]
                for c in ("Machine 1", "Machine 2", "Machine 3")]
    res = one_way("8.8 Three machines (DOE p. 20; no course solution), alpha 0.05", machines, 0.05)
    means = [np.mean(m) for m in machines]
    for name, value, printed_value in zip(("X-bar 1", "X-bar 2", "X-bar 3"), means, ("49", "56", "51")):
        check(f"{name} (p. 20)", value, printed_value)
    check("grand mean (p. 20)", np.mean(means), "52")
    check("sum (X-bar_i - X-bar)^2 (p. 20)", sum((m - np.mean(means)) ** 2 for m in means), "26")


def two_way_part() -> None:
    """8.9: two-way ANOVA. Method checked on the course's Excel output (GRR workbook), then the DOE demo sheets."""
    print("8.9 Two-way ANOVA with interaction: check on GRR workbook '2way anova' (MSA Les 5)")
    # 'measurements' rows 4-13: each part (P1..P5) has two rows (trials); columns D, E, F are appraisers A, B, C.
    grid = [r for r in sheet(GRR, "measurements").iter_rows(min_row=4, max_row=13, values_only=True)]
    grr = pd.DataFrame([(f"P{i // 2 + 1}", op, float(v)) for i, r in enumerate(grid) for op, v in zip("ABC", r[3:6])],
                       columns=["part", "operator", "y"])
    refit = anova_lm(smf.ols("y ~ C(part) * C(operator)", grr).fit())
    printed = {}
    for r in sheet(GRR, "2way anova").iter_rows(min_row=43, max_row=48, values_only=True):
        if r[1]:
            printed[r[1]] = r[2:8]
    show("GRR '2way anova' cached SS: Sample / Columns / Interaction / Within / Total",
         " / ".join(f"{printed[k][0]:.4f}" for k in ("Sample", "Columns", "Interaction", "Within", "Total")))
    identity = printed["Total"][0] - printed["Sample"][0] - printed["Columns"][0] - printed["Within"][0]
    check("SS_Interaction = SS_T - SS_parts - SS_operators - SS_within", identity, f"{printed['Interaction'][0]:.4f}")
    for k, term in (("Sample", "C(part)"), ("Columns", "C(operator)"), ("Interaction", "C(part):C(operator)"),
                    ("Within", "Residual")):
        check(f"statsmodels SS {k} = workbook", refit.loc[term, "sum_sq"], f"{printed[k][0]:.4f}")
        if k != "Within":
            check(f"statsmodels F {k} = workbook", refit.loc[term, "F"], f"{printed[k][3]:.4f}")
            check(f"statsmodels P {k} = workbook", refit.loc[term, "PR(>F)"], f"{printed[k][4]:.6f}")

    ws = sheet(DEMO, "2ANOVA")
    table = [r for r in ws.iter_rows(min_row=4, max_row=9, values_only=True)]
    machines = list(table[0][2:6])
    long = pd.DataFrame([(r[1], m, float(v)) for r in table[1:] for m, v in zip(machines, r[2:6])],
                        columns=["operator", "machine", "y"])
    print("8.9 DOE_demo '2ANOVA': 4 machines x 5 operators, one value per cell (no interaction term possible)")
    show("machine means", long.groupby("machine", sort=False).y.mean().to_dict())
    show("operator means", long.groupby("operator", sort=False).y.mean().to_dict())
    show("grand mean", long.y.mean())
    model = anova_lm(smf.ols("y ~ C(machine) + C(operator)", long).fit())
    for term, label in (("C(machine)", "machines"), ("C(operator)", "operators"), ("Residual", "error")):
        row = model.loc[term]
        show(f"{label}: df / SS / MS / F / p", f"{row.df:g} / {row.sum_sq:.4f} / {row.mean_sq:.4f} / "
             f"{row.F:.4f} / {row['PR(>F)']:.4f}" if term != "Residual" else
             f"{row.df:g} / {row.sum_sq:.4f} / {row.mean_sq:.4f}")
    show("F crit 0.05: machines (3, 12) / operators (4, 12)",
         f"{stats.f.isf(0.05, 3, 12):.3f} / {stats.f.isf(0.05, 4, 12):.3f}")
    show("SS machines by hand = 5 * sum (machine mean - grand mean)^2",
         5 * float(((long.groupby("machine").y.mean() - long.y.mean()) ** 2).sum()))
    show("SS operators by hand = 4 * sum (operator mean - grand mean)^2",
         4 * float(((long.groupby("operator").y.mean() - long.y.mean()) ** 2).sum()))
    show("SS total", float(((long.y - long.y.mean()) ** 2).sum()))

    ws = sheet(DEMO, "2ANOVArep")
    grid = [r for r in ws.iter_rows(min_row=4, max_row=14, values_only=True)]
    sun = list(grid[0][2:6])
    records = []
    for i, r in enumerate(grid[1:]):
        water = "Daily" if i < 5 else "Weekly"
        records += [(water, s, float(v)) for s, v in zip(sun, r[2:6])]
    rep = pd.DataFrame(records, columns=["water", "sun", "y"])
    print("8.9 DOE_demo '2ANOVArep': watering (2) x sunlight (4), n = 5")
    show("cell means", rep.groupby(["water", "sun"], sort=False).y.mean().round(2).to_dict())
    show("watering means", rep.groupby("water", sort=False).y.mean().round(3).to_dict())
    show("sunlight means", rep.groupby("sun", sort=False).y.mean().round(3).to_dict())
    model = anova_lm(smf.ols("y ~ C(water) * C(sun)", rep).fit())
    for term, label in (("C(water)", "watering"), ("C(sun)", "sunlight"), ("C(water):C(sun)", "interaction")):
        row = model.loc[term]
        show(f"{label}: df / SS / MS / F / p", f"{row.df:g} / {row.sum_sq:.5f} / {row.mean_sq:.5f} / {row.F:.4f} / {row['PR(>F)']:.3g}")
    row = model.loc["Residual"]
    show("within: df / SS / MS", f"{row.df:g} / {row.sum_sq:.4f} / {row.mean_sq:.6f}")
    show("total SS", f"{float(((rep.y - rep.y.mean()) ** 2).sum()):.5f}")
    show("F crit 0.05: (1, 32) / (3, 32)", f"{stats.f.isf(0.05, 1, 32):.3f} / {stats.f.isf(0.05, 3, 32):.3f}")


def design_part() -> None:
    """8.11-8.13: golf (DOE p. 27, 42) and the two yield experiments (DOE p. 35, 37)."""
    print("8.11 Golf (DOE p. 27 figure: driver O/R, ball B/T, two scores per corner)")
    golf = {"OB": [88, 90], "RB": [93, 91], "OT": [88, 91], "RT": [92, 94]}
    eff = effects(2, [golf["OB"], golf["RB"], golf["OT"], golf["RT"]])
    for name, label in (("A", "driver effect R - O"), ("B", "ball effect T - B"), ("AB", "interaction (diagonals)")):
        show(f"{label}: contrast / effect", f"{eff[name][0]:g} / {eff[name][1]:.4f}")
    show("side means R / O / T / B", f"{np.mean(golf['RB'] + golf['RT']):.2f} / {np.mean(golf['OB'] + golf['OT']):.2f} / "
         f"{np.mean(golf['OT'] + golf['RT']):.2f} / {np.mean(golf['OB'] + golf['RB']):.2f}")
    show("diagonal means OB+RT / RB+OT", f"{np.mean(golf['OB'] + golf['RT']):.2f} / {np.mean(golf['RB'] + golf['OT']):.2f}")

    print("8.13 Yield: experiment 1 (DOE p. 35) and experiment 2 (DOE p. 37)")
    conc1, time1, y1 = [0, 4, 8, 12], [10, 15, 20, 25], [17, 19, 21, 23]
    check("correlation concentration-time, experiment 1 (p. 38)", np.corrcoef(conc1, time1)[0, 1], "1")
    conc2, time2 = [0, 12, 0, 12], [10, 10, 25, 25]
    check("correlation concentration-time, experiment 2 (p. 38)", np.corrcoef(conc2, time2)[0, 1], "0")
    eff = effects(2, [[17], [16], [24], [23]])  # standard order (1)=(0,10), a=(12,10), b=(0,25), ab=(12,25)
    for name, label in (("A", "concentration"), ("B", "time"), ("AB", "interaction")):
        show(f"experiment 2 effect of {label}", f"{eff[name][1]:.2f}")
    show("experiment 2 means: time 25 / time 10 / conc 12 / conc 0",
         f"{(24 + 23) / 2} / {(17 + 16) / 2} / {(16 + 23) / 2} / {(17 + 24) / 2}")
    show("experiment 1: yield rises per run (17, 19, 21, 23)", np.diff(y1).tolist())


def two_k_part() -> None:
    """8.15-8.20: effects, contrasts, ANOVA of 2^k designs."""
    print("8.15 Coded 2^2 example (DOE p. 47-50), standard order (1), a, b, ab")
    coded = [[5.94, 5.50, 6.10], [4.16, 6.72, 2.72], [5.50, 5.25, 5.16], [6.25, 5.90, 4.97]]
    y = np.array(coded)
    a_sign, b_sign = np.array([-1, 1, -1, 1]), np.array([-1, -1, 1, 1])
    for label, mask, printed in (("mean at A=+1 (p. 48)", a_sign > 0, "5.12"), ("mean at A=-1 (p. 48)", a_sign < 0, "5.575"),
                                 ("mean at B=+1 (p. 49)", b_sign > 0, "5.505"), ("mean at B=-1 (p. 49)", b_sign < 0, "5.19"),
                                 ("mean at AB=+1 (p. 50)", a_sign * b_sign > 0, "5.78"),
                                 ("mean at AB=-1 (p. 50)", a_sign * b_sign < 0, "4.92")):
        check(label, y[mask].mean(), printed)
    eff = effects(2, coded)
    check("[A] (p. 48)", eff["A"][1], "-0.455")
    check("[B] (p. 49)", eff["B"][1], "0.315")
    check("[AB] (p. 50)", eff["AB"][1], "0.857", erratum=True)
    show("[AB] exact", f"{eff['AB'][1]:.4f}")
    show("[AB] from the rounded means 5.78 - 4.92", f"{5.78 - 4.92:.2f}")
    show("sums at AB = +1 / AB = -1", f"{y[a_sign * b_sign > 0].sum():.2f} / {y[a_sign * b_sign < 0].sum():.2f}")

    print("8.16 Corner figures (DOE p. 53-54); standard order (1), a, b, ab")
    for page, corners, printed in ((53, [20, 40, 30, 52], ("21", "11", "-1")), (54, [20, 50, 40, 12], ("1", "-9", "-29"))):
        eff = effects(2, [[c] for c in corners])
        for name, value in zip(("A", "B", "AB"), printed):
            check(f"p. {page}: {name}", eff[name][1], value, erratum=(page == 53 and name == "AB"))

    print("8.18 Chemical process (DOE p. 59-61)")
    chem = {(r["A"], r["B"]): [float(r[f"Rep {i}"]) for i in ("I", "II", "III")]
            for r in rows("S05_chemical_process_example_treatment_combinations_and_replicat")}
    runs = [chem[(a, b)] for b in "-+" for a in "-+"]
    eff = effects(2, runs)
    for name, printed in (("A", "8.33"), ("B", "-5.00"), ("AB", "1.67")):
        show(f"contrast {name} / contrast^2", f"{eff[name][0]:g} / {eff[name][0] ** 2:g}")
        check(f"effect {name} (p. 60)", eff[name][1], printed)
    table = {r["Source"]: r for r in rows("S05_anova_for_selected_factorial_model_router_chemical_process_e")}
    sspe = pure_error(runs)
    mse = sspe / 8
    for label, run in zip(("(1)", "a", "b", "ab"), runs):
        dev = [(v - np.mean(run)) ** 2 for v in run]
        show(f"pure error run {label}: mean / squared deviations", f"{np.mean(run):.2f} / {[round(float(d), 2) for d in dev]}")
    show("F crit 0.05 (1, 8)", f"{stats.f.isf(0.05, 1, 8):.3f}")
    model_ss = sum(v[2] for v in eff.values())
    check("Pure Error SS (p. 61)", sspe, table["Pure Error"]["Sum of Squares"])
    check("Pure Error MS (p. 61)", mse, table["Pure Error"]["Mean Square"])
    check("Model SS (p. 61)", model_ss, table["Model"]["Sum of Squares"])
    check("Model F (p. 61)", model_ss / 3 / mse, table["Model"]["F Value"])
    check("Model Prob > F (p. 61)", stats.f.sf(model_ss / 3 / mse, 3, 8), table["Model"]["Prob > F"])
    for name in ("A", "B", "AB"):
        f0 = eff[name][2] / mse
        check(f"SS {name} (p. 61)", eff[name][2], table[name]["Sum of Squares"])
        check(f"F {name} (p. 61)", f0, table[name]["F Value"])
        if name != "A":
            check(f"Prob > F {name} (p. 61)", stats.f.sf(f0, 1, 8), table[name]["Prob > F"])
        else:
            show("Prob > F A (p. 61: < 0.0001)", f"{stats.f.sf(f0, 1, 8):.6f}")
    total = model_ss + sspe
    check("Cor Total (p. 61)", total, table["Cor Total"]["Sum of Squares"])
    check("R-Squared (p. 61)", model_ss / total, "0.9030")
    check("Adj R-Squared, (N-1)/(N-p-1) (p. 61)", 1 - (1 - model_ss / total) * 11 / 8, "0.8666")
    show("Adj R-Squared with the REG p. 56 print (N-1)/(N-p-2)", f"{1 - (1 - model_ss / total) * 11 / 7:.4f}")
    check("Std. Dev. = sqrt(MS_E) (p. 61)", np.sqrt(mse), "1.98")
    check("Mean (p. 61)", np.mean(runs), "27.50")

    print("8.19 Surface finish (DOE p. 69-71), n = 2")
    surface = [[float(v) for v in r["Surface Finish obs1,obs2"].split(",")]
               for r in rows("S05_surface_finish_quality_example_data_2_3_design_n_2")]
    printed_totals = [r["Total"] for r in rows("S05_surface_finish_quality_example_data_2_3_design_n_2")]
    for run, total_printed in zip(surface, printed_totals):
        check(f"run total {run}", sum(run), total_printed)
    eff = effects(3, surface)
    sspe = pure_error(surface)
    mse = sspe / 8
    table = {r["Source of Variation"]: r for r in rows("S05_anova_table_surface_finish_quality_example_2_3_n_2")}
    check("Error SS (p. 71)", sspe, table["Error"]["Sum of Squares"])
    check("Error MS (p. 71)", mse, table["Error"]["Mean Square"])
    check("Total SS (p. 71)", sum(v[2] for v in eff.values()) + sspe, table["Total"]["Sum of Squares"])
    se = float(np.sqrt(mse / (2 * 2 ** (3 - 2))))
    for name, (contrast, effect, ss) in eff.items():
        f0 = ss / mse
        check(f"SS {name} (p. 71)", ss, table[name]["Sum of Squares"], erratum=(name == "ABC"))
        check(f"F0 {name} (p. 71)", f0, table[name]["F0"])
        p = stats.f.sf(f0, 1, 8)
        if name == "A":
            check_sci("P A (p. 71: 2.54 x 10^-3)", p, "2.54e-3", erratum=True)
            show("P A exact", f"{p:.4e}")
        else:
            check(f"P {name} (p. 71)", p, table[name]["P-value"])
        show(f"{name}: contrast / effect / effect +- 2 s.e.",
             f"{contrast:g} / {effect:.4f} / [{effect - 2 * se:.3f}, {effect + 2 * se:.3f}]")
    show("s.e.(effect) = sqrt(2.4375 / (2 * 2))", f"{se:.4f}")
    show("2 s.e.", f"{2 * se:.4f}")
    show("t quantile 0.975, df 8 (versus the rule's 2)", f"{stats.t.ppf(0.975, 8):.3f}")
    show("grand mean", f"{np.mean(surface):.4f}")

    print("8.20 Etch rate (DOE p. 73-77), single replicate 2^4")
    etch = [[float(r["Etch Rate (Angstrom/min)"])] for r in rows("S05_etch_rate_example_data_single_replicate_2_4_design")]
    eff = effects(4, etch)
    for r in rows("S05_etch_rate_example_estimated_effects_2_4_single_replicate"):
        check(f"effect {r['Effect']} (p. 75)", eff[r["Effect"]][1], r["Estimate"])
    for name in ("ABC", "ABD", "ACD", "BCD", "ABCD", "AD"):
        show(f"SS {name} = 4 * effect^2", f"{eff[name][2]:.4f}")
    pooled = sum(v[2] for name, v in eff.items() if len(name) >= 3)
    show("SS pooled 3- and 4-factor interactions (df 5)", f"{pooled:.4f}")
    show("MS_E pooled", f"{pooled / 5:.4f}")
    for r in rows("S05_nitride_etch_example_r_anova_output_pooled_model_a_b_c_d_2_t"):
        term = r["Term"].replace(":", "")
        if term == "Residuals":
            check("Residuals SS (p. 77)", pooled, r["Sum Sq"])
            check("Residuals MS (p. 77)", pooled / 5, r["Mean Sq"])
            continue
        f0 = eff[term][2] / (pooled / 5)
        check(f"SS {r['Term']} (p. 76-77)", eff[term][2], r["Sum Sq"])
        check(f"F {r['Term']} (p. 77)", f0, r["F value"])
        p_printed = r["Pr(>F)"].split()[0]
        if "e" in p_printed:
            check_sci(f"P {r['Term']} (p. 77)", stats.f.sf(f0, 1, 5), p_printed)
        else:
            check(f"P {r['Term']} (p. 77)", stats.f.sf(f0, 1, 5), p_printed)
    show("F crit 0.05 (1, 5)", f"{stats.f.isf(0.05, 1, 5):.3f}")


def fraction_part() -> None:
    """8.21-8.23: half fraction, aliases, resolution, generators (DOE p. 78-93)."""
    print("8.21 Fractions (DOE p. 78-93)")
    check("k = 11: runs of the full 2^k (p. 87)", 2 ** 11, "2048")
    check("k = 11: runs of the half fraction (p. 87)", 2 ** 10, "1024")
    check("k = 11 in 32 runs: fraction 1/2^6 = 1/... (p. 87)", 2 ** 11 // 32, "64")
    x = signs(3)
    abc = np.prod(x, axis=1)
    labels = ["(1)", "a", "b", "ab", "c", "ac", "bc", "abc"]
    show("runs with ABC = + (p. 80: a, b, c, abc)", [lab for lab, s in zip(labels, abc) if s > 0])
    for effect in ("A", "B", "C", "AB", "AC", "BC"):
        show(f"I = ABC: alias of {effect} (p. 82 gives A, B, AB)", word_product(effect, "ABC"))
    half = abc > 0
    show("in the half fraction column A equals column BC", bool(np.all(x[half, 0] == x[half, 1] * x[half, 2])))

    show("2^(6-2): number of runs", 2 ** (6 - 2))
    relation = defining_relation(["ABCE", "BCDF"])
    show("2^(6-2), E = ABC, F = BCD: complete defining relation (p. 91)", " = ".join(["I"] + relation))
    check("2^(6-2): resolution = shortest word (p. 91: IV)", min(len(w) for w in relation), "4")
    base = signs(4)
    e_col, f_col = base[:, 0] * base[:, 1] * base[:, 2], base[:, 1] * base[:, 2] * base[:, 3]
    for name, column, printed in (("E = ABC", e_col, "-++-+--+-++-+--+"), ("F = BCD", f_col, "--++++--++----++")):
        text = "".join("+" if v > 0 else "-" for v in column)
        status = "OK" if text == printed else "MISMATCH"
        if text != printed:
            UNEXPECTED.append(name)
        print(f"  {status:<19} {'p. 90 column ' + name + ', runs 1-16':<52} printed {printed}  computed {text}")
    for effect in ("A", "B", "E", "AB"):
        show(f"2^(6-2): aliases of {effect}", ", ".join(word_product(effect, w) for w in relation))
    show("2^(5-1), I = ABCDE: shortest word", min(len(w) for w in defining_relation(["ABCDE"])))
    for effect in ("A", "AB"):
        show(f"2^(5-1): alias of {effect}", word_product(effect, "ABCDE"))


def demo_full() -> None:
    """8.24: DOE_demo '2^5 full rep' (n = 3): effects, F-tests, significant effects."""
    ws = sheet(DEMO, "2^5 full rep")
    data = np.array([r[:6] for r in ws.iter_rows(min_row=3, max_row=98, values_only=True)], float)
    x, y = data[:, :5], data[:, 5]
    print(f"8.24 DOE_demo '2^5 full rep': {len(y)} observations")
    runs: dict[tuple, list[float]] = {}
    for key, value in zip(map(tuple, x), y):
        runs.setdefault(key, []).append(value)
    sspe, df_pe = pure_error(list(runs.values())), len(y) - len(runs)
    mse = sspe / df_pe
    se = np.sqrt(mse / (3 * 2 ** (5 - 2)))
    show("runs / replicates per run", f"{len(runs)} / {sorted({len(v) for v in runs.values()})}")
    show("SS pure error / df / MS_E", f"{sspe:.4f} / {df_pe} / {mse:.6f}")
    show("s.e.(effect) = sqrt(MS_E / (3 * 2^3)) / 2 s.e.", f"{se:.4f} / {2 * se:.4f}")
    show("t quantile 0.975, df 64", f"{stats.t.ppf(0.975, df_pe):.3f}")
    show("F crit 0.05 (1, 64)", f"{stats.f.isf(0.05, 1, df_pe):.3f}")
    show("grand mean / SS total", f"{y.mean():.4f} / {np.sum((y - y.mean()) ** 2):.4f}")
    table = []
    for size in range(1, 6):
        for combo in combinations(range(5), size):
            column = np.prod(x[:, combo], axis=1)
            effect = y[column > 0].mean() - y[column < 0].mean()
            ss = (column @ y) ** 2 / len(y)
            table.append(("".join("ABCDE"[j] for j in combo), effect, ss, ss / mse, stats.f.sf(ss / mse, 1, df_pe)))
    show("effects with |effect| > 2 s.e.", ", ".join(name for name, effect, *_ in table if abs(effect) > 2 * se))
    for name, effect, ss, f0, p in table:
        flag = "  significant at 0.05" if p < 0.05 else ""
        show(f"{name}: effect / SS / F0 / p", f"{effect:.4f} / {ss:.4f} / {f0:.3f} / {p:.4g}{flag}")
    show("number of effects with p < 0.05", sum(1 for *_, p in table if p < 0.05))
    coded_example = [runs[(a, b, -1.0, -1.0, -1.0)] for b in (-1.0, 1.0) for a in (-1.0, 1.0)]
    show("runs with C = D = E = -1 are the DOE p. 47 example", [[float(v) for v in r] for r in coded_example])


def demo_fractions() -> None:
    """8.24: DOE_demo '2^5-1 fract rep' and '2^5-2 fract rep'."""
    ws = sheet(DEMO, "2^5-1 fract rep")
    data = np.array([r[:6] for r in ws.iter_rows(min_row=3, max_row=50, values_only=True)], float)
    x, y = data[:, :5], data[:, 5]
    a, b, c, d, e = x.T
    print(f"8.24 DOE_demo '2^5-1 fract rep': {len(y)} observations")
    show("every run has ABCDE = +1", bool(np.all(np.prod(x, axis=1) == 1)))
    show("column BCDE identical to column A", bool(np.all(b * c * d * e == a)))
    design = np.column_stack([np.ones(len(y)), a, b, c, d, e, b * c * d * e])
    show("model Y = A+B+C+D+E+BCDE: columns / rank of the design matrix", f"{design.shape[1]} / {np.linalg.matrix_rank(design)}")
    full_ws = sheet(DEMO, "2^5 full rep")
    full: dict[tuple, list[float]] = {}
    for r in full_ws.iter_rows(min_row=3, max_row=98, values_only=True):
        full.setdefault(tuple(float(v) for v in r[:5]), []).append(float(r[5]))
    runs: dict[tuple, list[float]] = {}
    for key, value in zip(map(tuple, x), y):
        runs.setdefault(key, []).append(value)
    show("the 16 runs are the ABCDE = +1 half of '2^5 full rep'", all(sorted(full[k]) == sorted(v) for k, v in runs.items()))
    sspe, df_pe = pure_error(list(runs.values())), len(y) - len(runs)
    mse = sspe / df_pe
    show("SS pure error / df / MS_E", f"{sspe:.4f} / {df_pe} / {mse:.4f}")
    show("F crit 0.05 (1, 32)", f"{stats.f.isf(0.05, 1, df_pe):.3f}")
    for size in (1, 2):
        for combo in combinations(range(5), size):
            column = np.prod(x[:, combo], axis=1)
            effect = y[column > 0].mean() - y[column < 0].mean()
            ss = (column @ y) ** 2 / len(y)
            name = "".join("ABCDE"[j] for j in combo)
            p = stats.f.sf(ss / mse, 1, df_pe)
            flag = "  significant at 0.05" if p < 0.05 else ""
            show(f"{name} (alias {word_product(name, 'ABCDE')}): effect / F0 / p", f"{effect:.4f} / {ss / mse:.3f} / {p:.4g}{flag}")
    coef = np.linalg.lstsq(np.column_stack([np.ones(len(y)), a, b, c, d, e]), y, rcond=None)[0]
    show("model A..E: regression coefficient of A (= effect / 2)", f"{coef[1]:.4f}")

    ws2 = sheet(DEMO, "2^5-2 fract rep")
    values = [float(r[5]) for r in ws2.iter_rows(min_row=4, max_row=27, values_only=True)]
    print("8.24 DOE_demo '2^5-2 fract rep'")
    show("its 24 values equal rows 1-24 of '2^5-1 fract rep'", values == list(y[:24]))
    show("in those rows of '2^5-1': E values / D = -ABC", f"{sorted({float(v) for v in e[:24]})} / {bool(np.all(d[:24] == -(a * b * c)[:24]))}")
    options = ["AB", "AC", "BC", "ABC"]
    for gd in options:
        for ge in options:
            relation = defining_relation([gd + "D", ge + "E"])
            show(f"D = {gd}, E = {ge}: defining relation / resolution",
                 f"I = {' = '.join(relation)} / {min(len(w) for w in relation)}")
    relation = defining_relation(["ABD", "ACE"])
    for effect in ("A", "B", "C", "D", "E"):
        show(f"D = AB, E = AC: aliases of {effect}", ", ".join(word_product(effect, w) for w in relation))


def pitfall_numbers() -> None:
    """Numbers quoted in the "Valkuilen en strikvragen" boxes that are not printed elsewhere in this script."""
    print("Valkuilen: numbers used in the pitfall boxes")
    # 8.4: swapping the numerator and denominator df gives a very different critical value
    show("F crit 0.01 with df (3, 20), as on DOE p. 9", f"{stats.f.isf(0.01, 3, 20):.2f}")
    show("F crit 0.01 with the df swapped (20, 3)", f"{stats.f.isf(0.01, 20, 3):.2f}")
    # 8.5: the pooled StDev is sqrt(MS_E), not the plain mean of the group standard deviations
    paper = [[float(r[f"Obs{i}"]) for i in range(1, 7)]
             for r in rows("S05_table_4_4_tensile_strength_of_paper_psi") if r["Obs1"]]
    sds = [float(np.std(g, ddof=1)) for g in paper]
    show("paper: group standard deviations", [round(s, 3) for s in sds])
    show("paper: plain mean of the four standard deviations", f"{np.mean(sds):.3f}")
    show("paper: sqrt(mean of the four variances) = pooled StDev", f"{np.sqrt(np.mean(np.square(sds))):.3f}")


def main() -> None:
    """Print the numbers of Deel 08 in the order of its units."""
    anova_part()
    two_way_part()
    design_part()
    two_k_part()
    fraction_part()
    demo_full()
    demo_fractions()
    pitfall_numbers()
    print("unexpected mismatches:", UNEXPECTED or "none")


if __name__ == "__main__":
    main()
