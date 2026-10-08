"""Conditional probability sheet ('Voorwaardelijke kansen'): a cross table typed in or counted from raw data, the
joint, marginal and conditional distributions, any P(E | G) chosen from dropdowns, and (in)dependence.

Course (Naert, Les 1, 20260522_naert_big data.pdf): marginal distribution = sum over the other variable (p. 18);
independence P(X, Y) = P(X)·P(Y) ⇔ P(Y | X) = P(Y) (p. 19, 22); conditional distribution P(Y | X) = P(X, Y)/P(X),
"filter the rows" (p. 20); the worked exercise of p. 24-25 decides "not independent" because P(K | L = lijn 1) ≠ P(K).
The χ²-test for a cross table that is a sample: Test Recipes - Further Reading (Dutch) p. 18-20 (e_kl > 5; Yates for
2 × 2). Decision 5: α prefilled 0.05. P(not G) = 1 − P(G) is general probability, not printed by the course.

Row plan: settings 8-12; how to fill 14-18; the question 20-26 and its answer 28-34; data used 36-38; independence and
χ² 40-50; tables 52-252; the data last (typed table from row 256, raw data in W and X from row 257 down), so it can
grow. Helper columns from W (rows 29-78) and Z-AC (raw rows): not for the user.
"""
from __future__ import annotations

from collections.abc import Callable

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    font,
    input_block,
    input_row,
    instructions,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Voorwaardelijke kansen"

HEADER = HeaderBlock(
    tool="Voorwaardelijke kansen (conditional probability): kruistabel of ruwe data, gezamenlijke, marginale en "
         "voorwaardelijke verdelingen, P(gevraagd | gegeven), (on)afhankelijkheid",
    source="source/course/Les 1/20260522_naert_big data.pdf p. 17-25; 20261005_naert_web_lecture 1 notes.pdf p. 11-12; "
           "Les 2/20260529_ottoy_Test Recipes - Further Reading (Dutch).pdf p. 18-20 (χ²)",
    convention="Onafhankelijk in de tabel zelf als P(Y | X) = P(Y) voor elke rij (zoals Data p. 25); is de tabel een "
               "steekproef, dan de χ²-toets (α vooraf ingevuld op 0,05, beslissing 5). P(niet G) = 1 − P(G) is "
               "algemene kansrekening.",
    status=Status.VERIFIED,
    status_detail="getest tegen het uitgewerkte cursusvoorbeeld S02-WE01 (Data p. 24-25), als kruistabel en als ruwe "
                  "data; χ² tegen scipy",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Een kruistabel (getypt of geteld uit ruwe data) met de gezamenlijke kansen, de marginale verdelingen P(X) "
            "en P(Y), P(Y | X), P(X | Y), P(X)·P(Y); een gekozen kans P(gevraagd | gegeven) met teller en noemer; het "
            "oordeel onafhankelijk/afhankelijk en de χ²-toets.",
    inputs="α; optioneel de namen van X en Y; een kruistabel tot 20 × 20 met namen (rij 256 en kolom A vanaf rij 257) "
           "of ruwe data: per waarneming de categorie van X (kolom W) en van Y (kolom X), tot 10 000 regels; de vraag "
           "met de keuzelijsten van sectie 3.",
    audit="niet nodig (er bestaat een uitgewerkt cursusvoorbeeld)",
    disagreements=(),
)

CATEGORIES = 20                                        # rows and columns of the cross table
RAW_ROWS = 10_000                                      # raw observations
COLUMNS = [get_column_letter(2 + j) for j in range(CATEGORIES)]  # B … U
TOTAL = "V"                                            # row totals (and P(X) in the P(X | Y) table)
NUMBER = "0.0000"
ALPHA = "$B$9"

# the question (inputs in B) and the helper list of the variable names
QUESTION = {"e_var": "B21", "e_is": "B22", "e_cat": "B23", "g_var": "B24", "g_is": "B25", "g_cat": "B26"}
NAMES = {"x": "$D$10", "y": "$D$11", "none": "$D$12"}
NOTHING = "(niets gegeven)"
ANSWER_ROW = 29
ANSWERS = {"conditional": 29, "p_e": 30, "p_g": 31, "p_eg": 32, "reverse": 33, "product": 34}  # B value, C/D n
STATUS = {"source": "B37", "warning": "B38"}
INDEPENDENCE = {"max_difference": "B41", "verdict": "B42", "chi2": "B43", "df": "B44", "critical": "B45",
                "p": "B46", "decision": "B47", "small": "B48", "yates": "B49", "p_yates": "B50"}

# tables: title row, first category row (20 rows), and for some a margin row below
TABLES = {"counts": 53, "joint": 76, "given_x": 99, "given_y": 122, "product": 144, "difference": 166,
          "expected": 188, "contribution": 210, "yates": 232}


def first_row(table: str) -> int:
    """First category row of a table (its titles are one row higher)."""
    return TABLES[table] + 1


def last_row(table: str) -> int:
    """Last category row of a table."""
    return first_row(table) + CATEGORIES - 1


COUNTS = f"$B${first_row('counts')}:$U${last_row('counts')}"
ROW_TOTALS = f"${TOTAL}${first_row('counts')}:${TOTAL}${last_row('counts')}"
MARGIN_ROW = last_row("counts") + 1                    # column totals; V of this row = N
COL_TOTALS = f"$B${MARGIN_ROW}:$U${MARGIN_ROW}"
N = f"${TOTAL}${MARGIN_ROW}"
LABEL_ROW = TABLES["counts"]                           # the category names of Y, B … U

# the data block at the bottom
TYPED_HEADER = 256
TYPED_FIRST = TYPED_HEADER + 1
TYPED_LAST = TYPED_FIRST + CATEGORIES - 1
RAW_FIRST, RAW_LAST = TYPED_FIRST, TYPED_FIRST + RAW_ROWS - 1
RAW_X, RAW_Y = f"$W${RAW_FIRST}:$W${RAW_LAST}", f"$X${RAW_FIRST}:$X${RAW_LAST}"
RAW = "$X$37"                                          # 1 when raw data are present

# helpers: counters and numbers of new categories in the raw data (Z, AA for X; AB, AC for Y)
X_COUNTER, X_NUMBER, Y_COUNTER, Y_NUMBER = "Z", "AA", "AB", "AC"
# helpers next to the counts table: dropdown lists, row masks, row sums under the column masks
E_LIST, G_LIST, E_ROWS, G_ROWS, EG_ROWS, E_SUMS, G_SUMS, EG_SUMS = "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE"
MASK_FIRST_COLUMN = 24                                 # column X: column masks in X … AQ, rows 76-78
COLUMN_MASK_ROWS = {"e": 76, "g": 77, "eg": 78}
HELP_COUNTS = {"n_e": "$X$29", "n_g": "$X$30", "n_eg": "$X$31"}


def typed_cell(i: int, j: int) -> str:
    """Input cell of the typed count in row i, column j (0-based) of the cross table."""
    return f"{COLUMNS[j]}{TYPED_FIRST + i}"


def raw_cells(i: int) -> tuple[str, str]:
    """Input cells (category of X, category of Y) of raw observation i (0-based)."""
    return f"W{RAW_FIRST + i}", f"X{RAW_FIRST + i}"


def table_cell(table: str, i: int, j: int) -> str:
    """Cell of category row i and column j (0-based) of a result table."""
    return f"{COLUMNS[j]}{first_row(table) + i}"


def _settings(ws: Worksheet) -> None:
    """Rows 8-12: α and the names of X and Y (the names used fill D10:D12, the source of the dropdowns)."""
    section_title(ws, 8, "1. Instellingen")
    input_row(ws, 9, "Significantieniveau α voor de χ²-toets", "de cursus heeft geen standaardwaarde; gebruik de waarde "
                                                              "uit de vraag", "0.00")
    ws["B9"] = 0.05
    input_row(ws, 10, "Naam van X, de variabele in de rijen (tekst, bv. L)")
    input_row(ws, 11, "Naam van Y, de variabele in de kolommen (tekst, bv. K)")
    label(ws, 10, 3, "gebruikt:")
    output_cell(ws, "D10", '=IF(TRIM(B10)="","X",TRIM(B10))')
    output_cell(ws, "D11", '=IF(TRIM(B11)="","Y",TRIM(B11))')
    ws["D12"] = NOTHING
    ws["D12"].font = font(color="808080")


def _instructions(ws: Worksheet) -> None:
    """Rows 14-18: how to enter the data, with a link to the data block."""
    instructions(ws, 14, 1, f"2. Zo vul je de data in (onderaan, vanaf rij {TYPED_HEADER}); kies daarna in sectie 3 "
                            "wat gevraagd en wat gegeven is", (
        f"• Kruistabel (aantallen): rij {TYPED_HEADER} = de namen van de kolommen (categorieën van Y, TEKST), kolom A = "
        "de namen van de rijen (categorieën van X, TEKST), daarin de aantallen (GETALLEN). Tot 20 × 20.",
        "• Een rij of kolom 'Totaal' mag mee: ze telt niet (de totalen worden herberekend).",
        f"• Of ruwe data: één waarneming per regel, de categorie van X in kolom W en die van Y in kolom X (tekst of "
        f"getal), vanaf rij {RAW_FIRST}, tot {RAW_ROWS} regels. Staan er ruwe data, dan telt de getypte tabel niet.",
        "• Plakken uit een ander bestand: Plakken speciaal → Waarden. Rij 37-38 zegt welke data gebruikt worden en wat "
        "er mis is.",
    ))
    link = ws.cell(row=14, column=12, value=f"→ naar de data (rij {TYPED_HEADER})")
    link.hyperlink = f"#'{SHEET}'!A{TYPED_HEADER}"
    link.font = font(color="0563C1")


def _question(ws: Worksheet) -> None:
    """Rows 20-26: what is asked (E) and what is given (G), each a variable, 'is' or 'is niet', and a category."""
    section_title(ws, 20, "3. Jouw vraag: kies uit de lijsten (eerst de variabele, dan de categorie)")
    rows = (("e_var", "Gevraagd (E): variabele", "X of Y"),
            ("e_is", "Gevraagd: is / is niet", "'is niet' = de complementaire gebeurtenis"),
            ("e_cat", "Gevraagd: categorie", "de lijst volgt de gekozen variabele"),
            ("g_var", "Gegeven (G): variabele, of (niets gegeven)", "'als', 'gegeven dat', 'van de …'"),
            ("g_is", "Gegeven: is / is niet", ""),
            ("g_cat", "Gegeven: categorie", ""))
    for row, (key, text, note) in enumerate(rows, start=21):
        input_row(ws, row, text, note)
    ws[QUESTION["e_is"]] = ws[QUESTION["g_is"]] = "is"
    ws[QUESTION["g_var"]] = NOTHING
    lists = {QUESTION["e_var"]: f"{NAMES['x']}:{NAMES['y']}", QUESTION["g_var"]: f"{NAMES['x']}:{NAMES['none']}",
             QUESTION["e_is"]: '"is,is niet"', QUESTION["g_is"]: '"is,is niet"',
             QUESTION["e_cat"]: f"${E_LIST}${first_row('counts')}:${E_LIST}${last_row('counts')}",
             QUESTION["g_cat"]: f"${G_LIST}${first_row('counts')}:${G_LIST}${last_row('counts')}"}
    for cell, source in lists.items():
        rule = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True,
                              errorTitle="Kies uit de lijst", error="Kies een waarde uit de keuzelijst.")
        ws.add_data_validation(rule)
        rule.add(cell)


def _event_text(var: str, is_: str, cat: str) -> str:
    """Excel text 'L = Lijn 2' or 'L ≠ Lijn 2' of an event."""
    return f'{var}&IF({is_}="is niet"," ≠ "," = ")&{cat}'


def _answer(ws: Worksheet) -> None:
    """Rows 28-34: P(E | G) with numerator and denominator, and the related probabilities."""
    q = {key: f"${cell[0]}${cell[1:]}" for key, cell in QUESTION.items()}
    have_e = f'AND({q["e_var"]}<>"",{q["e_cat"]}<>"")'
    have_g = f'AND({q["g_var"]}<>"",{q["g_var"]}<>{NAMES["none"]},{q["g_cat"]}<>"")'
    e = _event_text(q["e_var"], q["e_is"], q["e_cat"])
    g = _event_text(q["g_var"], q["g_is"], q["g_cat"])
    n_e, n_g, n_eg = HELP_COUNTS["n_e"], HELP_COUNTS["n_g"], HELP_COUNTS["n_eg"]
    column_titles(ws, 28, ["Kans", "waarde", "teller (aantal)", "noemer (aantal)", "uitleg"])
    ok = f"AND({have_e},N({N})>0)"
    rows = {
        "conditional": (f'=IF({have_e},IF({have_g},"P("&{e}&" | "&{g}&") = n(E en G) / n(G)","P("&{e}&") = n(E) / N '
                        f'(niets gegeven)"),"kies in sectie 3 wat gevraagd is")',
                        f'=IF({ok},IF({have_g},IF({n_g}>0,{n_eg}/{n_g},""),{n_e}/{N}),"")',
                        f'=IF({ok},IF({have_g},{n_eg},{n_e}),"")', f'=IF({ok},IF({have_g},{n_g},{N}),"")',
                        "Data p. 20, 25: houd alleen de waarnemingen waar G geldt"),
        "p_e": (f'=IF({have_e},"P("&{e}&") = n(E) / N","")', f'=IF({ok},{n_e}/{N},"")', f'=IF({ok},{n_e},"")',
                f'=IF({ok},{N},"")', "marginale kans"),
        "p_g": (f'=IF({have_g},"P("&{g}&") = n(G) / N","")', f'=IF(AND({ok},{have_g}),{n_g}/{N},"")',
                f'=IF(AND({ok},{have_g}),{n_g},"")', f'=IF(AND({ok},{have_g}),{N},"")', "marginale kans"),
        "p_eg": (f'=IF({have_g},"P("&{e}&" en "&{g}&") = n(E en G) / N","")', f'=IF(AND({ok},{have_g}),{n_eg}/{N},"")',
                 f'=IF(AND({ok},{have_g}),{n_eg},"")', f'=IF(AND({ok},{have_g}),{N},"")', "gezamenlijke kans"),
        "reverse": (f'=IF({have_g},"omgekeerd: P("&{g}&" | "&{e}&") = n(E en G) / n(E)","")',
                    f'=IF(AND({ok},{have_g},{n_e}>0),{n_eg}/{n_e},"")', f'=IF(AND({ok},{have_g}),{n_eg},"")',
                    f'=IF(AND({ok},{have_g}),{n_e},"")', "andere noemer: niet dezelfde kans"),
        "product": (f'=IF({have_g},"P(E)·P(G): gelijk aan P(E en G) als E en G onafhankelijk zijn","")',
                    f'=IF(AND(ISNUMBER(B30),ISNUMBER(B31)),B30*B31,"")', '=""', '=""',
                    '=IF(AND(ISNUMBER(B32),ISNUMBER(B34)),IF(ABS(B32-B34)<=1E-12,"E en G onafhankelijk",'
                    '"E en G afhankelijk (P(E en G) ≠ P(E)·P(G))"),"")'),
    }
    for key, (text, value, numerator, denominator, note) in rows.items():
        r = ANSWERS[key]
        output_cell(ws, f"A{r}", text)
        ws[f"A{r}"].font = font(bold=key == "conditional")
        output_cell(ws, f"B{r}", value, NUMBER)
        ws[f"B{r}"].font = font(bold=key == "conditional")
        output_cell(ws, f"C{r}", numerator, "0")
        output_cell(ws, f"D{r}", denominator, "0")
        if note.startswith("="):
            output_cell(ws, f"E{r}", note)
        else:
            label(ws, r, 5, note, italic=True)
    # helper counts n(E), n(G), n(E en G): row sums under the column masks, weighted by the row masks
    first, last = first_row("counts"), last_row("counts")
    for cell, (sums, rows_) in zip(HELP_COUNTS.values(), ((E_SUMS, E_ROWS), (G_SUMS, G_ROWS), (EG_SUMS, EG_ROWS))):
        ws[cell.replace("$", "")] = f"=SUMPRODUCT(${sums}${first}:${sums}${last},${rows_}${first}:${rows_}${last})"
        ws[cell.replace("$", "")].font = font(color="808080")
    label(ws, 28, 23, "hulp: n(E), n(G), n(E en G)", italic=True)


def _status(ws: Worksheet) -> None:
    """Rows 36-38: which data are used, and what is wrong with them."""
    section_title(ws, 36, "Gebruikte data")
    ws[RAW.replace("$", "")] = f"=IF(COUNTA({RAW_X},{RAW_Y})>0,1,0)"
    ws[RAW.replace("$", "")].font = font(color="808080")
    label(ws, 37, 23, "hulp: ruwe data?", italic=True)
    typed = f"$B${TYPED_FIRST}:$U${TYPED_LAST}"
    result_row(ws, 37, "Gebruikt", f'=IF({RAW}=1,"ruwe data (kolommen W en X): "&COUNTIFS({RAW_X},"<>",{RAW_Y},"<>")&'
                                    f'" waarnemingen",IF(COUNT({typed})>0,"de getypte kruistabel","nog geen data"))')
    one_only = f'SUMPRODUCT(({RAW_X}<>"")*({RAW_Y}=""))+SUMPRODUCT(({RAW_X}="")*({RAW_Y}<>""))'
    many = (f"OR(MAX(${X_COUNTER}${RAW_FIRST}:${X_COUNTER}${RAW_LAST})>{CATEGORIES},"
            f"MAX(${Y_COUNTER}${RAW_FIRST}:${Y_COUNTER}${RAW_LAST})>{CATEGORIES})")
    result_row(ws, 38, "Controle",
               f'=IF({RAW}=1,IF({one_only}>0,"LET OP: "&{one_only}&" regel(s) met maar één categorie tellen niet. ","")'
               f'&IF({many},"LET OP: meer dan {CATEGORIES} categorieën: alleen de eerste {CATEGORIES} tellen. ","")'
               f'&IF(COUNT({typed})>0,"De getypte tabel telt niet (er staan ruwe data).",""),'
               f'IF(COUNTA({typed})>COUNT({typed}),"LET OP: cellen met tekst tussen de aantallen tellen niet mee.",""))')


def _labels_and_counts(ws: Worksheet) -> None:
    """The used cross table: category names (typed, or found in the raw data), counts and totals."""
    top, first = TABLES["counts"], first_row("counts")
    section_title(ws, 52, "4. Tabellen (berekend)")
    output_cell(ws, f"A{top}", f'="aantallen n: "&{NAMES["x"]}&" \\ "&{NAMES["y"]}')
    ws[f"{TOTAL}{top}"] = "totaal"
    ws[f"{TOTAL}{top}"].font = font(bold=True)
    for j, letter in enumerate(COLUMNS):
        typed_name = f"TRIM({letter}${TYPED_HEADER})"
        typed = (f'IF({typed_name}="",IF(COUNT({letter}${TYPED_FIRST}:{letter}${TYPED_LAST})>0,"kolom {j + 1}",""),'
                 f'IF(OR({typed_name}="totaal",{typed_name}="total"),"",{typed_name}))')
        raw = f'IFERROR(INDEX({RAW_Y},MATCH({j + 1},${Y_NUMBER}${RAW_FIRST}:${Y_NUMBER}${RAW_LAST},0)),"")'
        output_cell(ws, f"{letter}{top}", f"=IF({RAW}=1,{raw},{typed})")
        ws[f"{letter}{top}"].font = font(bold=True)
    for i in range(CATEGORIES):
        r, typed_row = first + i, TYPED_FIRST + i
        typed_name = f"TRIM($A${typed_row})"
        typed = (f'IF({typed_name}="",IF(COUNT($B${typed_row}:$U${typed_row})>0,"rij {i + 1}",""),'
                 f'IF(OR({typed_name}="totaal",{typed_name}="total"),"",{typed_name}))')
        raw = f'IFERROR(INDEX({RAW_X},MATCH({i + 1},${X_NUMBER}${RAW_FIRST}:${X_NUMBER}${RAW_LAST},0)),"")'
        output_cell(ws, f"A{r}", f"=IF({RAW}=1,{raw},{typed})")
        ws[f"A{r}"].font = font(bold=True)
        for j, letter in enumerate(COLUMNS):
            output_cell(ws, f"{letter}{r}", f'=IF(OR($A{r}="",{letter}${top}=""),"",IF({RAW}=1,COUNTIFS({RAW_X},$A{r},'
                                            f'{RAW_Y},{letter}${top}),IF(ISNUMBER({letter}{typed_row}),'
                                            f'{letter}{typed_row},0)))', "0")
        output_cell(ws, f"{TOTAL}{r}", f'=IF($A{r}="","",SUM(B{r}:U{r}))', "0")
    label(ws, MARGIN_ROW, 1, "totaal", bold=True)
    for letter in COLUMNS:
        output_cell(ws, f"{letter}{MARGIN_ROW}", f'=IF({letter}${top}="","",SUM({letter}{first}:{letter}{last_row("counts")}))',
                    "0")
    output_cell(ws, N.replace("$", ""), f"=SUM({ROW_TOTALS})", "0")


def _grid(ws: Worksheet, table: str, title: str, cell: Callable[[int, str], str],
          right: tuple[str, Callable[[int], str]] | None = None,
          bottom: tuple[str, Callable[[str], str]] | None = None) -> None:
    """A 20 × 20 result table under `title`, the category names around it, `cell(i, letter)` in every cell.

    `right` = (title, formula of category row i) fills column V; `bottom` = (label, formula of column letter) fills
    the row below the table.
    """
    top, first = TABLES[table], first_row(table)
    output_cell(ws, f"A{top}", f'="{title}"')
    ws[f"A{top}"].font = font(bold=True)
    for letter in COLUMNS:
        output_cell(ws, f"{letter}{top}", f"={letter}${LABEL_ROW}")
        ws[f"{letter}{top}"].font = font(bold=True)
    if right:
        ws[f"{TOTAL}{top}"] = right[0]
        ws[f"{TOTAL}{top}"].font = font(bold=True)
    for i in range(CATEGORIES):
        r = first + i
        output_cell(ws, f"A{r}", f"=$A{first_row('counts') + i}")
        ws[f"A{r}"].font = font(bold=True)
        for letter in COLUMNS:
            output_cell(ws, f"{letter}{r}", "=" + cell(i, letter), NUMBER)
        if right:
            output_cell(ws, f"{TOTAL}{r}", "=" + right[1](i), NUMBER)
    if bottom:
        r = last_row(table) + 1
        label(ws, r, 1, bottom[0], bold=True)
        for letter in COLUMNS:
            output_cell(ws, f"{letter}{r}", "=" + bottom[1](letter), NUMBER)


def _count(i: int, letter: str) -> str:
    """Used count in category row i (0-based), column `letter`."""
    return f"{letter}{first_row('counts') + i}"


def _row_total(i: int) -> str:
    """Row total of category row i (0-based)."""
    return f"${TOTAL}{first_row('counts') + i}"


def _col_total(letter: str) -> str:
    """Column total of column `letter`."""
    return f"{letter}${MARGIN_ROW}"


def _tables(ws: Worksheet) -> None:
    """Joint, conditional, product, difference and χ² tables."""
    def share(total: str) -> str:
        """A total divided by N."""
        return f'IF(AND(ISNUMBER({total}),N({N})>0),{total}/{N},"")'

    p_x = ("P(X)", lambda i: share(_row_total(i)))
    p_y = ("P(Y)", lambda letter: share(_col_total(letter)))
    _grid(ws, "joint", "gezamenlijke kansen P(X, Y) = n / N; rechts P(X), onderaan P(Y) (de marginale verdelingen)",
          lambda i, c: f'IF(AND(ISNUMBER({_count(i, c)}),N({N})>0),{_count(i, c)}/{N},"")', p_x, p_y)
    _grid(ws, "given_x", "P(Y | X) = n / rijtotaal (elke rij is een verdeling); onderaan P(Y) om te vergelijken",
          lambda i, c: f'IF(AND(ISNUMBER({_count(i, c)}),N({_row_total(i)})>0),{_count(i, c)}/{_row_total(i)},"")',
          None, p_y)
    _grid(ws, "given_y", "P(X | Y) = n / kolomtotaal (elke kolom is een verdeling); rechts P(X) om te vergelijken",
          lambda i, c: f'IF(AND(ISNUMBER({_count(i, c)}),N({_col_total(c)})>0),{_count(i, c)}/{_col_total(c)},"")', p_x)
    _grid(ws, "product", "P(X)·P(Y): zo zou P(X, Y) zijn als X en Y onafhankelijk waren",
          lambda i, c: (f'IF(AND(ISNUMBER({_count(i, c)}),N({N})>0),({_row_total(i)}/{N})*({_col_total(c)}/{N}),"")'))

    def difference(i: int, c: str) -> str:
        """P(Y | X) − P(Y) for one cell, from the P(Y | X) table and its P(Y) row."""
        given, marginal = f"{c}{first_row('given_x') + i}", f"{c}${last_row('given_x') + 1}"
        return f'IF(AND(ISNUMBER({given}),ISNUMBER({marginal})),{given}-{marginal},"")'

    _grid(ws, "difference", "P(Y | X) − P(Y): overal 0 = onafhankelijk (Data p. 22, 25)", difference)
    _grid(ws, "expected", "verwachte aantallen e = rijtotaal · kolomtotaal / N (χ²-toets, Test Recipes p. 18-20)",
          lambda i, c: (f'IF(AND(ISNUMBER({_count(i, c)}),N({_row_total(i)})>0,N({_col_total(c)})>0),'
                        f'{_row_total(i)}*{_col_total(c)}/{N},"")'))

    def expected(i: int, c: str) -> str:
        """Expected count of one cell."""
        return f"{c}{first_row('expected') + i}"

    _grid(ws, "contribution", "bijdrage aan χ²: (n − e)² / e",
          lambda i, c: f'IF(ISNUMBER({expected(i, c)}),({_count(i, c)}-{expected(i, c)})^2/{expected(i, c)},"")')
    _grid(ws, "yates", "bijdrage met Yates (alleen 2 × 2): (|n − e| − 0,5)² / e",
          lambda i, c: (f'IF(ISNUMBER({expected(i, c)}),MAX(0,ABS({_count(i, c)}-{expected(i, c)})-0.5)^2/'
                        f'{expected(i, c)},"")'))


def _independence(ws: Worksheet) -> None:
    """Rows 40-50: is X independent of Y in the table, and the χ²-test when the table is a sample."""
    section_title(ws, 40, "5. Zijn X en Y onafhankelijk?")
    difference = f"B{first_row('difference')}:U{last_row('difference')}"
    rows_used, cols_used = f'COUNTIF({ROW_TOTALS},">0")', f'COUNTIF({COL_TOTALS},">0")'
    result_row(ws, 41, "grootste |P(Y | X) − P(Y)| in de tabel",
               f'=IF(COUNT({difference})>0,MAX(MAX({difference}),-MIN({difference})),"")', NUMBER,
               "tabel 'P(Y | X) − P(Y)' hieronder")
    result_row(ws, 42, "in de tabel zelf (zoals Data p. 25)",
               f'=IF(ISNUMBER({INDEPENDENCE["max_difference"]}),IF({INDEPENDENCE["max_difference"]}<=1E-12,"onafhankelijk: P(Y | X) = '
               f'P(Y) voor elke rij","afhankelijk: P(Y | X) verschilt van P(Y) voor minstens één rij"),"")')
    ok = f"AND({rows_used}>=2,{cols_used}>=2)"
    contributions = f"B{first_row('contribution')}:U{last_row('contribution')}"
    yates = f"B{first_row('yates')}:U{last_row('yates')}"
    expected = f"B{first_row('expected')}:U{last_row('expected')}"
    result_row(ws, 43, "Is de tabel een steekproef? χ² = Σ (n − e)² / e", f'=IF({ok},SUM({contributions}),"")', NUMBER,
               "Test Recipes p. 18-20")
    result_row(ws, 44, "vrijheidsgraden (r − 1)(c − 1)", f'=IF({ok},({rows_used}-1)*({cols_used}-1),"")', "0")
    result_row(ws, 45, "kritieke waarde CHISQ.INV.RT(α; df)", f'=IF(ISNUMBER(B44),_xlfn.CHISQ.INV.RT({ALPHA},B44),"")',
               NUMBER)
    result_row(ws, 46, "p-waarde", '=IF(ISNUMBER(B43),_xlfn.CHISQ.DIST.RT(B43,B44),"")', "0.000000")
    result_row(ws, 47, "besluit bij α", f'=IF(ISNUMBER(B46),IF(B46<{ALPHA},"verwerp H0: afhankelijk",'
                                         '"H0 (onafhankelijk) niet verwerpen"),"")')
    result_row(ws, 48, "cellen met e ≤ 5 (voorwaarde e > 5)", f'=IF({ok},COUNTIFS({expected},">0",{expected},"<=5"),"")',
               "0", "Test Recipes p. 19: is dit meer dan 0, dan is de toets niet betrouwbaar")
    two = f"AND({rows_used}=2,{cols_used}=2)"
    result_row(ws, 49, "2 × 2: χ² met Yates-correctie", f'=IF({two},SUM({yates}),"")', NUMBER, "Test Recipes p. 19-20")
    result_row(ws, 50, "2 × 2: p-waarde met Yates", '=IF(ISNUMBER(B49),_xlfn.CHISQ.DIST.RT(B49,1),"")', "0.000000")


def _helpers(ws: Worksheet) -> None:
    """Dropdown lists of categories, row and column masks of E and G, and row sums under the column masks."""
    q = {key: f"${cell[0]}${cell[1:]}" for key, cell in QUESTION.items()}
    first = first_row("counts")
    label(ws, TABLES["counts"], column_index_from_string(E_LIST), "hulp: keuzelijsten en maskers (niet invullen)", italic=True)
    have_g = f'AND({q["g_var"]}<>"",{q["g_var"]}<>{NAMES["none"]},{q["g_cat"]}<>"")'

    def match(name: str, var: str, is_: str, cat: str) -> str:
        """1 when the category `name` satisfies 'var is (niet) cat', else 0."""
        return f'IF({is_}="is niet",IF({name}<>{cat},1,0),IF({name}={cat},1,0))'

    for i in range(CATEGORIES):
        r = first + i
        name, col_name = f"$A{r}", f"INDEX($B${LABEL_ROW}:$U${LABEL_ROW},{i + 1})"
        ws[f"{E_LIST}{r}"] = f'=IF({q["e_var"]}={NAMES["x"]},{name},IF({q["e_var"]}={NAMES["y"]},{col_name},""))'
        ws[f"{G_LIST}{r}"] = f'=IF({q["g_var"]}={NAMES["x"]},{name},IF({q["g_var"]}={NAMES["y"]},{col_name},""))'
        ws[f"{E_ROWS}{r}"] = (f'=IF({name}="",0,IF({q["e_var"]}={NAMES["x"]},'
                              f'{match(name, NAMES["x"], q["e_is"], q["e_cat"])},1))')
        ws[f"{G_ROWS}{r}"] = (f'=IF({name}="",0,IF(AND({have_g},{q["g_var"]}={NAMES["x"]}),'
                              f'{match(name, NAMES["x"], q["g_is"], q["g_cat"])},1))')
        ws[f"{EG_ROWS}{r}"] = f"={E_ROWS}{r}*{G_ROWS}{r}"
        for sums, key in ((E_SUMS, "e"), (G_SUMS, "g"), (EG_SUMS, "eg")):
            row = COLUMN_MASK_ROWS[key]
            masks = f"${get_column_letter(MASK_FIRST_COLUMN)}${row}:${get_column_letter(MASK_FIRST_COLUMN + CATEGORIES - 1)}${row}"
            ws[f"{sums}{r}"] = f"=SUMPRODUCT($B{r}:$U{r},{masks})"
    for j, letter in enumerate(COLUMNS):
        name = f"{letter}${LABEL_ROW}"
        mask = get_column_letter(MASK_FIRST_COLUMN + j)
        ws[f"{mask}{COLUMN_MASK_ROWS['e']}"] = (f'=IF({name}="",0,IF({q["e_var"]}={NAMES["y"]},'
                                                f'{match(name, NAMES["y"], q["e_is"], q["e_cat"])},1))')
        ws[f"{mask}{COLUMN_MASK_ROWS['g']}"] = (f'=IF({name}="",0,IF(AND({have_g},{q["g_var"]}={NAMES["y"]}),'
                                                f'{match(name, NAMES["y"], q["g_is"], q["g_cat"])},1))')
        ws[f"{mask}{COLUMN_MASK_ROWS['eg']}"] = f"={mask}{COLUMN_MASK_ROWS['e']}*{mask}{COLUMN_MASK_ROWS['g']}"
    for column in range(column_index_from_string(E_LIST), MASK_FIRST_COLUMN + CATEGORIES):
        for row in list(range(first, first + CATEGORIES)) + list(COLUMN_MASK_ROWS.values()):
            cell = ws.cell(row=row, column=column)
            if cell.value is not None:
                cell.font = font(color="808080")


def _data(ws: Worksheet) -> None:
    """The data block: the typed cross table (row 256 and below) and the raw data in W and X with their helpers."""
    section_title(ws, TYPED_HEADER - 2, "6. Data (zie sectie 2): typ of plak de kruistabel hier, of de ruwe data in "
                                        "kolommen W en X")
    label(ws, TYPED_HEADER - 1, 1, "Kruistabel: namen van de kolommen in deze rij →, namen van de rijen in kolom A ↓",
          italic=True)
    label(ws, TYPED_HEADER - 1, 23, "Ruwe data: één waarneming per regel ↓", italic=True)
    label(ws, TYPED_HEADER, 1, "rijnaam (tekst) \\ kolomnaam (tekst) →", bold=True)
    input_block(ws, (TYPED_HEADER, TYPED_HEADER), (2, 1 + CATEGORIES))
    input_block(ws, (TYPED_FIRST, TYPED_LAST), (1, 1))
    input_block(ws, (TYPED_FIRST, TYPED_LAST), (2, 1 + CATEGORIES), "0")
    for column, text in ((23, "categorie van X"), (24, "categorie van Y")):
        ws.cell(row=TYPED_HEADER, column=column, value=text).font = font(bold=True)
    input_block(ws, (RAW_FIRST, RAW_LAST), (23, 24))
    label(ws, TYPED_HEADER - 1, column_index_from_string(X_COUNTER), "hulp: nieuwe categorieën tellen (niet invullen)", italic=True)
    for letter in (X_COUNTER, Y_COUNTER):
        ws[f"{letter}{TYPED_HEADER}"] = 0
    for r in range(RAW_FIRST, RAW_LAST + 1):
        for raw, counter, number in (("W", X_COUNTER, X_NUMBER), ("X", Y_COUNTER, Y_NUMBER)):
            # running count of the distinct categories so far; a new category gets its number on its first row
            ws[f"{counter}{r}"] = (f'={counter}{r - 1}+IF({raw}{r}="",0,IF(ISNA(MATCH({raw}{r},{raw}${TYPED_HEADER}:'
                                   f'{raw}{r - 1},0)),1,0))')
            ws[f"{number}{r}"] = f'=IF({counter}{r}>{counter}{r - 1},{counter}{r},"")'


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the conditional probability calculator."""
    write_header(ws, HEADER)
    _settings(ws)
    _instructions(ws)
    _question(ws)
    _answer(ws)
    _status(ws)
    _independence(ws)
    _labels_and_counts(ws)
    _tables(ws)
    _helpers(ws)
    _data(ws)
    ws.column_dimensions["A"].width = 62
    for letter in [*COLUMNS, TOTAL]:
        ws.column_dimensions[letter].width = 11
    ws.column_dimensions["B"].width = 18   # the dropdowns of the question
    for letter in ("W", "X"):
        ws.column_dimensions[letter].width = 16
