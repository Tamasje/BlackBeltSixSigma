"""One-way ANOVA sheet: up to 50 groups of any size, the ANOVA table and the F-test.

Course (De Vuyst, Les 3, DOE.pdf): SS_Total = ΣΣ (Yij − Ȳ..)², SS_Error = ΣΣ (Yij − Ȳi.)², SS_Treatment =
Σ n_i (Ȳi. − Ȳ..)² (p. 6), F = MS_Treatment / MS_Error ~ F(a − 1, N − a) (p. 7-8, Table 4.6). Decision 5: α is an
input prefilled 0.05.

Row plan: α 9; how to fill the groups 11-16; per group n, mean, s and the SS terms 18-23; ANOVA table 25-31; the
groups last (names in row 34, values from row 35 down, open-ended), so they can grow without moving anything.
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    DATA_LAST_ROW,
    HeaderBlock,
    Status,
    column_titles,
    input_block,
    input_row,
    instructions,
    label,
    open_input_column,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "ANOVA (eenweg)"

HEADER = HeaderBlock(
    tool="Eenwegs-ANOVA (one-way ANOVA): zijn de gemiddelden van a groepen (factorniveaus) gelijk?",
    source="source/course/Les 3/20260605_de vuyst_BB_DOE.pdf p. 3-15",
    convention="Invoer α vooraf ingevuld op 0,05 (beslissing 5); groepen mogen verschillend groot zijn.",
    status=Status.VERIFIED,
    status_detail="getest tegen de uitgewerkte voorbeelden van de cursus S05-WE01 en S05-WE02",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Eenwegs-ANOVA-tabel (SS, df, MS, F0, p-waarde, kritieke F) met per groep n, gemiddelde en s.",
    inputs="α; tot 50 groepen (één kolom per groep, B tot AY), elk met zoveel waarden als nodig, vanaf rij 35; "
           "optioneel een naam per groep in rij 34.",
    audit="niet nodig (er bestaan uitgewerkte voorbeelden in de cursus)",
    disagreements=(),
)

ALPHA = "$B$9"
NUMBER = "0.000000"
DECISION = '=IF(ISNUMBER({p}),IF({p}<' + ALPHA + ',"verwerp H0","H0 niet verwerpen"),"")'

GROUPS = 50
GROUP_COLUMNS = [get_column_letter(2 + g) for g in range(GROUPS)]  # B … AY
NAME_ROW, FIRST_GROUP_ROW = 34, 35
GROUP_STATS = {"n": 18, "mean": 19, "sd": 20}  # rows 21 and 22 hold the two SS terms per group
CHECK_ROW = 23
ANOVA_ROWS = {"treatments": 26, "error": 27, "total": 28}  # B SS, C df, D MS, E F, F p, G F crit, H decision
INPUTS = {"alpha": "B9"}
RESULTS = {"grand_mean": "B29", "pooled_sd": "B30", "groups": "B31"}
FIRST, LAST = GROUP_COLUMNS[0], GROUP_COLUMNS[-1]
EVERYTHING = f"${FIRST}${FIRST_GROUP_ROW}:${LAST}${DATA_LAST_ROW}"


def group_cell(group: int, i: int) -> str:
    """Input cell of value i (0-based) of group `group` (0-based)."""
    return f"{GROUP_COLUMNS[group]}{FIRST_GROUP_ROW + i}"


def _values(letter: str) -> str:
    """All values of one group column."""
    return f"{letter}${FIRST_GROUP_ROW}:{letter}${DATA_LAST_ROW}"


def _group_statistics(ws: Worksheet) -> None:
    """Rows 18-23: per group n, mean, s, the two terms of SS_Treatment and SS_Error, and a check for text."""
    for row, text in ((18, "n_i"), (19, "gemiddelde Ȳ_i"), (20, "standaardafwijking s_i"), (21, "n_i (Ȳ_i − Ȳ)²"),
                      (22, "Σ (Y_ij − Ȳ_i)²"), (CHECK_ROW, "controle: tekst in de waarden telt niet mee")):
        label(ws, row, 1, text)
    for letter in GROUP_COLUMNS:
        column = _values(letter)
        output_cell(ws, f"{letter}18", f'=IF(COUNT({column})>0,COUNT({column}),"")', "0")
        output_cell(ws, f"{letter}19", f'=IF(COUNT({column})>0,AVERAGE({column}),"")', NUMBER)
        output_cell(ws, f"{letter}20", f'=IF(COUNT({column})>1,_xlfn.STDEV.S({column}),"")', NUMBER)
        output_cell(ws, f"{letter}21", f'=IF(ISNUMBER({letter}19),{letter}18*({letter}19-$B$29)^2,"")', NUMBER)
        output_cell(ws, f"{letter}22", f'=IF(COUNT({column})>0,DEVSQ({column}),"")', NUMBER)
        output_cell(ws, f"{letter}{CHECK_ROW}", f'=IF(COUNTA({column})>COUNT({column}),"TEKST!","")')


def _table(ws: Worksheet) -> None:
    """Rows 25-31: the ANOVA table, grand mean, pooled s and the number of groups."""
    n_row = f"B18:{LAST}18"
    ok = f"AND(ISNUMBER($B$31),$B$31>=2,SUM({n_row})>$B$31)"  # at least 2 groups and error df > 0
    column_titles(ws, 25, ["Bron (source)", "SS", "df", "MS", "F0", "p-waarde", "kritieke F bij α", "Besluit bij α"])
    rows = {
        26: ("Behandelingen (treatments, tussen de groepen)", f"SUM(B21:{LAST}21)", "$B$31-1"),
        27: ("Fout (error, binnen de groepen)", f"SUM(B22:{LAST}22)", f"SUM({n_row})-$B$31"),
        28: ("Totaal", f"DEVSQ({EVERYTHING})", f"SUM({n_row})-1"),
    }
    for row, (text, ss, df) in rows.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({ok},{ss},"")', "0.0000")
        output_cell(ws, f"C{row}", f'=IF({ok},{df},"")', "0")
    output_cell(ws, "D26", '=IF(ISNUMBER(B26),B26/C26,"")', "0.0000")
    output_cell(ws, "D27", '=IF(ISNUMBER(B27),B27/C27,"")', "0.0000")
    output_cell(ws, "E26", '=IF(AND(ISNUMBER(D27),N(D27)>0),D26/D27,"")', "0.0000")
    output_cell(ws, "F26", '=IF(ISNUMBER(E26),_xlfn.F.DIST.RT(E26,C26,C27),"")', "0.000000")
    output_cell(ws, "G26", f'=IF(AND(ISNUMBER(C26),ISNUMBER(C27),N(C26)>0,N(C27)>0,ISNUMBER({ALPHA})),_xlfn.F.INV.RT({ALPHA},C26,C27),"")', "0.0000")
    output_cell(ws, "H26", DECISION.format(p="F26"))
    label(ws, 27, 5, "H0: alle groepsgemiddelden gelijk (DOE p. 7-8, Table 4.6)", italic=True)
    result_row(ws, 29, "totaal gemiddelde Ȳ (grand mean)", f'=IF(COUNT({EVERYTHING})>0,AVERAGE({EVERYTHING}),"")',
               NUMBER)
    result_row(ws, 30, "gepoolde standaardafwijking √MS_E", '=IF(ISNUMBER(D27),SQRT(D27),"")', NUMBER)
    result_row(ws, 31, "a = aantal groepen", f'=IF(COUNT({n_row})>0,COUNT({n_row}),"")', "0")


def _groups(ws: Worksheet) -> None:
    """The data block: a name per group (row 34) and its values below (row 35 down, open-ended)."""
    section_title(ws, 32 + 1, f"De groepen: één kolom per groep (B tot {LAST}); de waarden eronder vanaf rij "
                              f"{FIRST_GROUP_ROW}")
    label(ws, NAME_ROW, 1, "naam van de groep (tekst, mag leeg) →", bold=True)
    label(ws, FIRST_GROUP_ROW, 1, "waarden (alleen getallen), onder elkaar ↓", bold=True)
    input_block(ws, (NAME_ROW, NAME_ROW), (2, 1 + GROUPS))
    for letter in GROUP_COLUMNS:
        open_input_column(ws, letter, FIRST_GROUP_ROW)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the one-way ANOVA calculator."""
    write_header(ws, HEADER)
    section_title(ws, 8, "Instellingen")
    input_row(ws, 9, "Significantieniveau α (significance level; betrouwbaarheid = 1 − α)",
              "de cursus heeft geen standaardwaarde; de meeste cursusvoorbeelden gebruiken 5 %; gebruik de waarde uit "
              "de vraag", "0.0000")
    ws["B9"] = 0.05
    instructions(ws, 11, 1, f"Zo vul je de groepen in (onderaan, vanaf rij {FIRST_GROUP_ROW})", (
        "• Eén kolom per groep (factorniveau, factor level): groep 1 in kolom B, groep 2 in C, … Rij 34: de naam van "
        "de groep (TEKST, mag leeg).",
        f"• Vanaf rij {FIRST_GROUP_ROW}: alleen GETALLEN, één waarneming per cel, onder elkaar; zoveel als je wilt. "
        "Groepen mogen verschillend groot zijn.",
        "• Het blad telt zelf de groepen en de waarden (rij 18). 'TEKST!' in rij 23: die kolom heeft een cel met tekst "
        "die niet meetelt.",
        "• Plakken uit een ander bestand: Plakken speciaal → Waarden. Leegmaken: selecteer de kolommen vanaf rij 34 "
        "en druk Delete.",
    ))
    column_titles(ws, 17, ["1. Per groep (berekend)"] + [f"groep {g + 1}" for g in range(GROUPS)])
    _group_statistics(ws)
    _table(ws)
    _groups(ws)
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="α", error="α ligt tussen 0 en 1.")
    ws.add_data_validation(fraction)
    fraction.add("B9")
    ws.column_dimensions["A"].width = 58
    for letter in GROUP_COLUMNS:
        ws.column_dimensions[letter].width = 14
