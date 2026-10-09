"""Simple linear regression sheet: y = b0 + b1 x with ANOVA, R², t- and F-tests, CIs of β0 and β1, the CI of the
mean response and the prediction interval at x0.

Course (De Vuyst, Les 3, Regression.pdf): b1 = S_xy / S_xx (p. 17-19), SS_T = SS_R + SS_E, R² = SS_R / SS_T
(p. 24-27), MS_E = SS_E / (n − 2) (p. 28), Var(b1) = σ²/S_xx, Var(b0) = σ²(1/n + x̄²/S_xx) (p. 29), t-tests (p. 30-31),
F-test (p. 33), CIs (p. 34), CI of the mean response (p. 35-36) and prediction interval (p. 38). R²_adj: p. 56 prints
1 − (1 − R²)(n − 1)/(n − k − 2); the course's own outputs (Regression p. 22: 87.1 %) and σ̂ = √(SSE/(n − k − 1))
(p. 57) use n − k − 1. Both are shown. Decisions 5 and 6: α prefilled 0.05; one- and two-sided side by side.

Row plan: settings 8-12; the verdict "is the regression significant?" 13; how to enter the pairs 14-18; completeness 20; fit 21-42; t-tests 44-46; intervals 48-52;
the pairs last (x in B, y in C, from row 58 down, open-ended), so they can grow without moving anything.
"""
from __future__ import annotations

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    DATA_LAST_ROW,
    HeaderBlock,
    Status,
    column_titles,
    data_check,
    input_row,
    instructions,
    label,
    open_input_column,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Regressie"

HEADER = HeaderBlock(
    tool="Enkelvoudige lineaire regressie (simple linear regression): schatting, ANOVA, R², toetsen, BI en PI",
    source="source/course/Les 3/20260605_de vuyst_BB_Regression.pdf p. 16-38, 56-57",
    convention="Invoer α vooraf ingevuld op 0,05 (beslissing 5); BI's en t-toetsen eenzijdig en tweezijdig naast "
               "elkaar (beslissing 6); R²_adj op beide manieren getoond.",
    status=Status.VERIFIED,
    status_detail="getest tegen de uitgewerkte voorbeelden van de cursus S05-WE17 en S05-WE18",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Enkelvoudige lineaire regressie met ANOVA, R², t-toetsen, BI's van β0 en β1, en het BI van de gemiddelde "
            "respons en het predictie-interval bij x0.",
    inputs="α; paren (x, y) in kolommen B en C vanaf rij 58, zoveel als nodig; x0 en de H0-waarden van β1 en β0.",
    audit="niet nodig (er bestaan uitgewerkte voorbeelden in de cursus)",
    disagreements=(
        "Regression p. 56 drukt R²_adj = 1 − (1 − R²)(n − 1)/(n − k − 2); de eigen outputs van de cursus gebruiken "
        "n − k − 1 (Regression p. 22 Minitab R-Sq(adj) 87,1 %; σ̂ op Regression p. 57). Het blad toont beide, met "
        "label.",
        "Regression p. 21 (S05-WE17): de gefitte rechte 'ŷ = 74.20 + 14.97x' van Figure 11-4 verschilt van de "
        "kleinste-kwadratenwaarden (least squares) 74,283 en 14,947 van de Minitab-output op p. 22, die het blad "
        "reproduceert.",
    ),
)

ALPHA = "$B$9"
NUMBER = "0.000000"
DECISION = '=IF(ISNUMBER({p}),IF({p}<' + ALPHA + ',"verwerp H0","H0 niet verwerpen"),"")'

FIRST_PAIR = 58
X_DATA = f"$B${FIRST_PAIR}:$B${DATA_LAST_ROW}"
Y_DATA = f"$C${FIRST_PAIR}:$C${DATA_LAST_ROW}"
CHECKS = {"x": f"B{FIRST_PAIR - 3}", "y": f"B{FIRST_PAIR - 2}"}
T_TESTS = {"b1": 45, "b0": 46}  # B estimate, C H0 value, D t, E-G p (≠, >, <), H-J decisions
INTERVALS = {"b1": 49, "b0": 50, "mean": 51, "prediction": 52}  # B centre, C s.e., D-E two-sided, F lower, G upper
INPUTS = {"alpha": "B9", "x0": "B10", "beta1_0": "B11", "beta0_0": "B12"}
RESULTS = {
    "pairs": "B20", "n": "B21", "xbar": "B22", "ybar": "B23", "sxx": "B24", "b1": "B25", "b0": "B26",
    "sst": "B27", "ssr": "B28", "sse": "B29", "mse_reg": "B30", "s": "B31", "r2": "B32", "r2_adj": "B33",
    "r2_adj_p56": "B34", "F": "B35", "p_F": "B36", "se_b1": "B38", "se_b0": "B39", "fit": "B40",
    "se_fit": "B41", "se_pred": "B42",
}


def pair_cells(i: int) -> tuple[str, str]:
    """Input cells (x, y) of pair i (0-based)."""
    return f"B{FIRST_PAIR + i}", f"C{FIRST_PAIR + i}"


def _regression_inputs(ws: Worksheet) -> None:
    """Rows 8-20: α, x0, the H0 values, how to enter the pairs and their completeness; the pairs block at the bottom."""
    section_title(ws, 8, "Instellingen")
    input_row(ws, 9, "Significantieniveau α (significance level; betrouwbaarheid = 1 − α)",
              "de cursus heeft geen standaardwaarde; de meeste cursusvoorbeelden gebruiken 5 %; gebruik de waarde uit "
              "de vraag", "0.0000")
    ws["B9"] = 0.05
    input_row(ws, 10, "x0 voor de gefitte waarde, het BI van de gemiddelde respons en het predictie-interval", "",
              NUMBER)
    input_row(ws, 11, "β1,0 = helling (slope) onder H0", "0 = toets op significantie van de regressie "
                                                          "(Regression p. 30, 33)", NUMBER)
    input_row(ws, 12, "β0,0 = intercept onder H0", "Regression p. 31", NUMBER)
    ws["B11"] = 0
    ws["B12"] = 0
    instructions(ws, 14, 1, f"Zo vul je de paren (x, y) in (onderaan, vanaf rij {FIRST_PAIR})", (
        "• Kolom B = x (de verklarende variabele), kolom C = y (de respons); één paar per rij, alleen GETALLEN.",
        "• Zoveel rijen als je wilt; plak twee kolommen uit Excel (Plakken speciaal → Waarden), x eerst.",
        f"• Rij 20 zegt of elke rij x én y heeft; rijen {FIRST_PAIR - 3}-{FIRST_PAIR - 2} tellen de getallen en melden "
        "tekst (die telt niet mee).",
        "• Leegmaken: selecteer B en C vanaf de eerste paar-rij en druk Delete.",
    ))
    count = f"COUNT({X_DATA})"
    both = f"SUMPRODUCT(ISNUMBER({X_DATA})*ISNUMBER({Y_DATA}))"
    result_row(ws, 20, "Paren volledig (elke rij heeft x én y, minstens 3 rijen)?",
               f'=IF({count}+COUNT({Y_DATA})=0,"",IF(AND({count}=COUNT({Y_DATA}),{both}={count},{count}>2),"ja",'
               f'"NEE: elke rij heeft x én y nodig, minstens 3 rijen"))')
    section_title(ws, FIRST_PAIR - 4, f"De paren (x, y): x in kolom B, y in kolom C, vanaf rij {FIRST_PAIR}")
    label(ws, FIRST_PAIR - 3, 1, "controle x (kolom B)")
    data_check(ws, CHECKS["x"], X_DATA)
    label(ws, FIRST_PAIR - 2, 1, "controle y (kolom C)")
    data_check(ws, CHECKS["y"], Y_DATA)
    column_titles(ws, FIRST_PAIR - 1, ["", "x", "y"])
    for letter in "BC":
        open_input_column(ws, letter, FIRST_PAIR)


def _verdict(ws: Worksheet) -> None:
    """Row 13: the answer to "is the regression significant?" in words, from the F-test (H0: β1 = 0)."""
    label(ws, 13, 1, "Is de regressie significant? (F-toets, H0: β1 = 0)", bold=True)
    p = "$B$36"
    shown = f'IF({p}<0.0001,"< 0,0001","= "&FIXED({p},4))'
    yes = (f'"JA, de regressie is significant: p-waarde "&{shown}&" < α = "&FIXED({ALPHA},3)&". H0 (β1 = 0) wordt '
           f'verworpen: er is een lineair verband tussen x en y (R² = "&FIXED($B$32,3)&")."')
    no = (f'"NEE, de regressie is niet significant: p-waarde "&{shown}&" ≥ α = "&FIXED({ALPHA},3)&". H0 (β1 = 0) wordt '
          f'niet verworpen: geen lineair verband aangetoond."')
    output_cell(ws, "B13", f'=IF(AND(ISNUMBER({p}),ISNUMBER({ALPHA})),IF({p}<{ALPHA},{yes},{no}),"")')


def _regression_fit(ws: Worksheet) -> None:
    """Rows 21-42: estimates, ANOVA, R², F-test and standard errors."""
    ok = '$B$20="ja"'
    rows = [
        (21, "n = aantal paren", f'=IF({ok},COUNT({X_DATA}),"")', "0", ""),
        (22, "x̄", f'=IF({ok},AVERAGE({X_DATA}),"")', NUMBER, ""),
        (23, "ȳ", f'=IF({ok},AVERAGE({Y_DATA}),"")', NUMBER, ""),
        (24, "S_xx = Σ (x − x̄)²", f'=IF({ok},DEVSQ({X_DATA}),"")', NUMBER, ""),
        (25, "b1 = S_xy / S_xx (helling, slope)", f'=IF(AND({ok},N(B24)>0),SLOPE({Y_DATA},{X_DATA}),"")', NUMBER,
         "Regression p. 17-19"),
        (26, "b0 = ȳ − b1 x̄ (intercept)", f'=IF(ISNUMBER(B25),INTERCEPT({Y_DATA},{X_DATA}),"")', NUMBER, ""),
        (27, "SS_T = Σ (y − ȳ)², totale kwadratensom (sum of squares), df n − 1",
         f'=IF(ISNUMBER(B25),DEVSQ({Y_DATA}),"")', NUMBER, "Regression p. 24-27"),
        (28, "SS_R = Σ (ŷ − ȳ)² = b1² S_xx, df 1", '=IF(ISNUMBER(B25),B25^2*B24,"")', NUMBER, ""),
        (29, "SS_E = Σ (y − ŷ)² = SS_T − SS_R, df n − 2", '=IF(ISNUMBER(B25),B27-B28,"")', NUMBER, ""),
        (30, "σ̂² = MS_E = SS_E / (n − 2)", '=IF(ISNUMBER(B29),B29/(B21-2),"")', NUMBER, "Regression p. 28"),
        (31, "σ̂ = √MS_E (Minitab 'S')", '=IF(ISNUMBER(B30),SQRT(B30),"")', NUMBER, ""),
        (32, "R² = SS_R / SS_T", '=IF(AND(ISNUMBER(B28),N(B27)>0),B28/B27,"")', "0.0000", "Regression p. 27"),
        (33, "R²_adj = 1 − (1 − R²)(n − 1)/(n − 2)", '=IF(ISNUMBER(B32),1-(1-B32)*(B21-1)/(B21-2),"")', "0.0000",
         "zoals de Minitab-output van de cursus (Regression p. 22)"),
        (34, "R²_adj zoals gedrukt op Regression p. 56: 1 − (1 − R²)(n − 1)/(n − 3)",
         '=IF(AND(ISNUMBER(B32),N(B21)>3),1-(1-B32)*(B21-1)/(B21-3),"")', "0.0000",
         "Regression p. 56 met k = 1 (verschilt van de outputs van de cursus)"),
        (35, "F0 = MS_R / MS_E, df (1, n − 2)", '=IF(AND(ISNUMBER(B30),N(B30)>0),B28/B30,"")', "0.0000",
         "Regression p. 33"),
        (36, "p-waarde van F0 (H0: β1 = 0)", '=IF(ISNUMBER(B35),_xlfn.F.DIST.RT(B35,1,B21-2),"")', "0.000000",
         ""),
        (37, "kritieke t, df n − 2: tweezijdig t(1 − α/2); eenzijdig t(1 − α) in kolom D",
         f'=IF(AND(N(B21)>2,ISNUMBER({ALPHA})),_xlfn.T.INV(1-{ALPHA}/2,B21-2),"")', "0.0000", ""),
        (38, "s.e.(b1) = √(MS_E / S_xx)", '=IF(ISNUMBER(B30),SQRT(B30/B24),"")', NUMBER, "Regression p. 29-30"),
        (39, "s.e.(b0) = √(MS_E (1/n + x̄²/S_xx))", '=IF(ISNUMBER(B30),SQRT(B30*(1/B21+B22^2/B24)),"")', NUMBER,
         "Regression p. 29, 31"),
        (40, "ŷ0 = b0 + b1 x0", '=IF(AND(ISNUMBER(B25),ISNUMBER($B$10)),B26+B25*$B$10,"")', NUMBER,
         "Regression p. 35"),
        (41, "s.e.(ŷ0) = √(MS_E (1/n + (x0 − x̄)²/S_xx))",
         '=IF(ISNUMBER(B40),SQRT(B30*(1/B21+($B$10-B22)^2/B24)),"")', NUMBER, "Regression p. 36"),
        (42, "s.e.(e0) = √(MS_E (1 + 1/n + (x0 − x̄)²/S_xx))",
         '=IF(ISNUMBER(B40),SQRT(B30*(1+1/B21+($B$10-B22)^2/B24)),"")', NUMBER, "Regression p. 38"),
    ]
    for row, text, formula, fmt, source in rows:
        result_row(ws, row, text, formula, fmt, source)
    output_cell(ws, "D36", DECISION.format(p="B36"))
    output_cell(ws, "D37", f'=IF(AND(N(B21)>2,ISNUMBER({ALPHA})),_xlfn.T.INV(1-{ALPHA},B21-2),"")', "0.0000")


def _regression_tests(ws: Worksheet) -> None:
    """Rows 44-46: t-tests of H0: β1 = β1,0 and H0: β0 = β0,0, all three alternatives (decision 6)."""
    column_titles(ws, 44, ["t-toets (t-test), df n − 2 (Regression p. 30-31)", "schatting", "H0-waarde", "t0",
                            "p (HA: ≠)", "p (HA: >)", "p (HA: <)", "Besluit (≠)", "Besluit (>)", "Besluit (<)"])
    df = "$B$21-2"
    for key, (text, estimate, h0, se) in {"b1": ("helling β1", "B25", "$B$11", "B38"),
                                          "b0": ("intercept β0", "B26", "$B$12", "B39")}.items():
        r = T_TESTS[key]
        label(ws, r, 1, text)
        output_cell(ws, f"B{r}", f'=IF(ISNUMBER({estimate}),{estimate},"")', NUMBER)
        output_cell(ws, f"C{r}", f"=N({h0})", NUMBER)
        output_cell(ws, f"D{r}", f'=IF(AND(ISNUMBER({se}),N({se})>0),(B{r}-C{r})/{se},"")', "0.0000")
        output_cell(ws, f"E{r}", f'=IF(ISNUMBER(D{r}),_xlfn.T.DIST.2T(ABS(D{r}),{df}),"")', "0.000000")
        output_cell(ws, f"F{r}", f'=IF(ISNUMBER(D{r}),_xlfn.T.DIST.RT(D{r},{df}),"")', "0.000000")
        output_cell(ws, f"G{r}", f'=IF(ISNUMBER(D{r}),_xlfn.T.DIST(D{r},{df},TRUE),"")', "0.000000")
        for p, decision in zip("EFG", "HIJ"):
            output_cell(ws, f"{decision}{r}", DECISION.format(p=f"{p}{r}"))


def _regression_intervals(ws: Worksheet) -> None:
    """Rows 48-52: CIs of β1, β0, the mean response at x0 and the prediction interval, one- and two-sided."""
    column_titles(ws, 48, ["Betrouwbaarheidsinterval (confidence interval, BI) 1 − α, t met n − 2 df", "centrum",
                            "s.e.", "tweezijdig: van", "tot", "alleen ondergrens (minstens …)",
                            "alleen bovengrens (hoogstens …)", "Bron in de cursus"])
    two, one = f"_xlfn.T.INV(1-{ALPHA}/2,$B$21-2)", f"_xlfn.T.INV(1-{ALPHA},$B$21-2)"
    spec = {
        "b1": ("BI voor de helling β1", "B25", "B38", "Regression p. 34"),
        "b0": ("BI voor het intercept β0", "B26", "B39", "Regression p. 34"),
        "mean": ("BI voor de gemiddelde respons (mean response) bij x0", "B40", "B41", "Regression p. 36"),
        "prediction": ("Predictie-interval (prediction interval, PI) voor een nieuwe y bij x0", "B40", "B42",
                       "Regression p. 38"),
    }
    for key, (text, centre, se, source) in spec.items():
        r = INTERVALS[key]
        label(ws, r, 1, text)
        output_cell(ws, f"B{r}", f'=IF(AND(ISNUMBER({centre}),ISNUMBER({se})),{centre},"")', NUMBER)
        output_cell(ws, f"C{r}", f'=IF(ISNUMBER(B{r}),{se},"")', NUMBER)
        output_cell(ws, f"D{r}", f'=IF(ISNUMBER(B{r}),B{r}-{two}*C{r},"")', NUMBER)
        output_cell(ws, f"E{r}", f'=IF(ISNUMBER(B{r}),B{r}+{two}*C{r},"")', NUMBER)
        output_cell(ws, f"F{r}", f'=IF(ISNUMBER(B{r}),B{r}-{one}*C{r},"")', NUMBER)
        output_cell(ws, f"G{r}", f'=IF(ISNUMBER(B{r}),B{r}+{one}*C{r},"")', NUMBER)
        label(ws, r, 8, source, italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the simple linear regression calculator."""
    write_header(ws, HEADER)
    _regression_inputs(ws)
    _verdict(ws)
    _regression_fit(ws)
    _regression_tests(ws)
    _regression_intervals(ws)
    label(ws, 53, 1, "Meervoudige regressie (multiple regression, Regression p. 46-61) staat niet op dit blad.",
          italic=True)
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="α", error="α ligt tussen 0 en 1.")
    ws.add_data_validation(fraction)
    fraction.add("B9")
    ws.column_dimensions["A"].width = 62
    for letter in "BCDEFGHIJ":
        ws.column_dimensions[letter].width = 14
