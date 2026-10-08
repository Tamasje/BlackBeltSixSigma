"""Gage R&R sheet: average and range method and ANOVA method for a crossed study (k operators × n parts × r trials).

Course (Ottoy, Les 5, Black Belt in Six Sigma - Measurement System Analysis.pdf):
- Average and range (p. 34-35): EV = R̿ · K1, K1 = 1/d2 (subgroup size r); AV = √((X̄_DIFF · K2)² − EV²/(n r)),
  K2 = 1/d2* (one subgroup of size k); PV = √((R_p · K3)² − EV²/(k r)), K3 = 1/d2* (one subgroup of size n);
  GRR² = EV² + AV², TV² = GRR² + PV², %GRR = GRR / TV. Constants from tabel MSA.pdf: d2 from its g → ∞ row,
  d2* from its g = 1 row (as the course workbook does: d2 1.12838, d2*(3) 1.91155, d2*(5) 2.48124).
- ANOVA (p. 36-37): two-way model without interaction; EV = √MSe, AV = √((MSa − EV²)/(n r)),
  PV = √((MSp − EV²)/(k r)), error df = nk(r − 1) + (k − 1)(n − 1).
- %GRR against tolerance 6·σ_GRR / TOL (p. 24); thresholds 10 % / 30 % (p. 35).
Decision 11: multiplier 6; %GRR against total variation and tolerance; the ANOVA with interaction is shown too, with
its p-value, and no pooling α is chosen; ndc omitted (not in the course).

Row plan: settings 8-10; how to enter the data 12-17; study size 19-25; average & range 27-41; ANOVA 43-55;
ANOVA with interaction 57-64; thresholds 66-70; per operator and part from 72; the data block last, so it can
grow: k (up to 20 operators), n (up to 50 parts) and r are read from it.
"""
from __future__ import annotations

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.sheet_tables import d2_star_formula, msa_d2_formula, msa_table
from bbtools.xlsx_style import (
    HeaderBlock,
    Status,
    column_titles,
    constant_row,
    font,
    input_block,
    input_row,
    label,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Gage R&R"

HEADER = HeaderBlock(
    tool="Gage R&R: gemiddelde-en-spreidingsbreedtemethode (average and range method) en ANOVA-methode "
         "(EV, AV, PV, GRR, TV, %GRR)",
    source="source/course/Les 5/20260619_ottoy_Black Belt in Six Sigma - Measurement System Analysis.pdf p. 18, 24, 34-38; "
           "20260619_ottoy_tabel MSA.pdf; 20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx",
    convention="Beslissing 11: vermenigvuldigingsfactor (multiplier) 6; %GRR t.o.v. totale variatie en tolerantie; "
               "ANOVA van de cursus = model zonder interactie (interactie ook getoond); constanten uit tabel MSA.pdf; "
               "ndc niet in de cursus.",
    status=Status.VERIFIED,
    status_detail="getest tegen de uitgewerkte cursusvoorbeelden (worked examples) S10-WE02 (ANOVA) en S10-WE03 "
                  "(gemiddelde-en-spreidingsbreedtemethode) op de eigen studiedata van de cursus",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Herhaalbaarheid (repeatability, EV), reproduceerbaarheid (reproducibility, AV), variatie tussen de delen "
            "(part variation, PV), GRR, TV en %GRR met de gemiddelde-en-spreidingsbreedtemethode (average and range "
            "method) en met de ANOVA-methode, met het oordeel 10 % / 30 % van de cursus; de ANOVA met de interactieterm.",
    inputs="de metingen van een volledige gekruiste studie (crossed study), onderaan het blad: kolom A de naam van de "
           "operator, één rij per herhaling (trial), één kolom per deel (part); het blad leidt k (tot 20), n (tot 50) "
           "en r af uit de data (tot 1000 rijen); optioneel de tolerantie USL − LSL en α voor de F-toetsen.",
    audit="niet nodig (er bestaan uitgewerkte cursusvoorbeelden)",
    disagreements=(
        "GRR-werkboek, blad '2way anova', K45 typt '=412.5+296.667' (de interactie-SS 296,6667 afgerond), dus is zijn "
        "EV² 30,8333478 in plaats van 30,8333333; zijn AV, PV, TV en %GRR verschuiven in het 7e cijfer. Het blad rekent "
        "vanuit de data; de tests vergelijken met een tolerantie die precies dit toelaat.",
    ),
)

MAX_OPERATORS, MAX_PARTS, DATA_ROWS = 20, 50, 1000   # tabel MSA.pdf has d2* up to m = 20; parts beyond 20 only lack K3
FIRST_PART_COLUMN = 2                                  # column B = part 1
LAST_PART_COLUMN = FIRST_PART_COLUMN + MAX_PARTS - 1   # column AY = part 50
OPERATOR_MEAN, OPERATOR_TRIALS, OPERATOR_START, OPERATOR_BLOCK = "AZ", "BA", "BB", "BC"  # section 8 columns
NAME_HELPER, INDEX_HELPER = "BA", "BB"                 # data rows: operator of the row and its number
RANGE_FIRST = 74                                       # section 8: ranges of operator 1 … 20
MEAN_FIRST = RANGE_FIRST + MAX_OPERATORS + 1           # means of operator 1 … 20
PART_MEANS = MEAN_FIRST + MAX_OPERATORS                # mean per part
PART_COUNTS = PART_MEANS + 1                           # number of values per part
FIRST_DATA_ROW = PART_COUNTS + 5
LAST_DATA_ROW = FIRST_DATA_ROW + DATA_ROWS - 1
NUMBER = "0.000000"
PERCENT = "0.00%"
TOL, ALPHA, OK = "$B$9", "$B$10", '$B$25="ja"'
INPUTS = {"tolerance": "B9", "alpha": "B10"}
RESULTS = {
    "k": "B20", "n": "B21", "r": "B22", "count": "B23", "blocks": "B24", "complete": "B25",
    "rbarbar": "B28", "xdiff": "B29", "rp": "B30", "d2": "B31", "d2star_k": "B32", "d2star_n": "B33",
    "ev_ar": "B34", "av_ar": "B35", "pv_ar": "B36", "grr_ar": "B37", "tv_ar": "B38", "pct_ar": "B39", "pct_tol_ar": "B40",
    "ev_anova": "B49", "av_anova": "B50", "pv_anova": "B51", "grr_anova": "B52", "tv_anova": "B53",
    "pct_anova": "B54", "pct_tol_anova": "B55",
}
VERDICTS = {"pct_ar": "D39", "pct_tol_ar": "D40", "pct_anova": "D54", "pct_tol_anova": "D55"}
ANOVA_ROWS = {"parts": 45, "operators": 46, "error": 47, "total": 48}
INTERACTION_ROWS = {"parts": 59, "operators": 60, "interaction": 61, "within": 62, "total": 63}
THRESHOLDS = {"acceptable": "B67", "marginal": "B68"}


def part_letter(part: int) -> str:
    """Column of part `part` (0-based)."""
    return get_column_letter(FIRST_PART_COLUMN + part)


DATA = f"$B${FIRST_DATA_ROW}:${part_letter(MAX_PARTS - 1)}${LAST_DATA_ROW}"
NAMES = f"${NAME_HELPER}${FIRST_DATA_ROW}:${NAME_HELPER}${LAST_DATA_ROW}"
NUMBERS = f"${INDEX_HELPER}${FIRST_DATA_ROW}:${INDEX_HELPER}${LAST_DATA_ROW}"


def data_row(operator: int, trial: int, trials: int) -> int:
    """Row of one trial of one operator (0-based) when every operator has `trials` rows, in operator blocks."""
    return FIRST_DATA_ROW + trials * operator + trial


def data_cell(operator: int, trial: int, part: int, trials: int) -> str:
    """Input cell of one measurement (0-based operator, trial and part) in a study with `trials` trials."""
    return f"{part_letter(part)}{data_row(operator, trial, trials)}"


def operator_cell(operator: int, trials: int) -> str:
    """Name cell (column A) on the first row of an operator's block; the rows below may stay empty."""
    return f"A{data_row(operator, 0, trials)}"


def _instructions(ws: Worksheet) -> None:
    """Section 2: how to enter the study, next to a link to the data block."""
    section_title(ws, 12, f"2. Zo vul je de studiedata in (blok onderaan, vanaf rij {FIRST_DATA_ROW})")
    link = ws.cell(row=12, column=8, value=f"→ naar de data (rij {FIRST_DATA_ROW})")
    link.hyperlink = f"#'{SHEET}'!A{FIRST_DATA_ROW}"
    link.font = font(color="0563C1")
    for offset, text in enumerate((
        "• Eén rij per meting van één operator in één herhaling (trial); één kolom per deel (part): deel 1 in kolom B, "
        f"deel 2 in C, … tot deel {MAX_PARTS} in kolom {part_letter(MAX_PARTS - 1)}.",
        "• Kolom A = naam van de operator (TEKST, bv. A, B, Jan). Typ ze op de eerste rij van elke operator; de rijen "
        "eronder zonder naam horen bij dezelfde operator. Een naam op elke rij mag ook.",
        "• Kolommen B en verder: alleen GETALLEN. Houd de rijen van één operator bij elkaar (eerst alle herhalingen van "
        "operator A, dan B, …). Geen lege rij binnen een operator.",
        f"• Het blad telt zelf het aantal operators k (tot {MAX_OPERATORS}), delen n (tot {MAX_PARTS}) en herhalingen r "
        f"(tot {DATA_ROWS} rijen samen). Sectie 3 zegt of de studie volledig is.",
        "• Plak je uit een ander bestand: Plakken speciaal → Waarden. Leegmaken: selecteer het blok en druk Delete.",
    )):
        label(ws, 13 + offset, 1, text, italic=offset == 4)


def _data(ws: Worksheet) -> None:
    """Section 9 (bottom of the sheet, so it can grow): the study data and two helper columns per row."""
    section_title(ws, FIRST_DATA_ROW - 3, "9. Studiedata (zie sectie 2): kolom A = operator (tekst), kolommen B … = "
                                          "delen (getallen)")
    label(ws, FIRST_DATA_ROW - 2, int(column_index_from_string(NAME_HELPER)),
          "hulpkolommen: operator van de rij en zijn volgnummer (niet invullen)", italic=True)
    column_titles(ws, FIRST_DATA_ROW - 1, ["Operator (tekst)"] + [f"Deel {j + 1}" for j in range(MAX_PARTS)])
    input_block(ws, (FIRST_DATA_ROW, LAST_DATA_ROW), (1, 1))
    input_block(ws, (FIRST_DATA_ROW, LAST_DATA_ROW), (FIRST_PART_COLUMN, FIRST_PART_COLUMN + MAX_PARTS - 1))
    last_part = part_letter(MAX_PARTS - 1)
    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        above = f"{NAME_HELPER}{r - 1}"
        # the operator of a row: its own name, else the operator of the row above (a block names only its first row)
        ws[f"{NAME_HELPER}{r}"] = (f'=IF(COUNT(B{r}:{last_part}{r})=0,"",IF(TRIM(A{r})<>"",TRIM(A{r}),'
                                   f'IF({above}="","(geen naam)",{above})))')
        # 1, 2, 3 … on the first row of each new operator
        ws[f"{INDEX_HELPER}{r}"] = (f'=IF({NAME_HELPER}{r}="","",IF(ISNA(MATCH({NAME_HELPER}{r},'
                                    f'{NAME_HELPER}${FIRST_DATA_ROW - 1}:{above},0)),'
                                    f'MAX({INDEX_HELPER}${FIRST_DATA_ROW - 1}:{INDEX_HELPER}{r - 1})+1,""))')
        ws[f"{NAME_HELPER}{r}"].font = ws[f"{INDEX_HELPER}{r}"].font = font(color="808080")


def _per_cell(ws: Worksheet) -> None:
    """Section 8: per operator and part the range and mean over the trials, operator means and part means."""
    section_title(ws, RANGE_FIRST - 2, "8. Per operator en deel (berekend uit sectie 9)")
    titles = ["Spreidingsbreedte (range) over de herhalingen"] + [f"Deel {j + 1}" for j in range(MAX_PARTS)]
    column_titles(ws, RANGE_FIRST - 1, titles + ["", "herhalingen r", "eerste rij", "rijen bij elkaar?"])
    column_titles(ws, MEAN_FIRST - 1, ["Gemiddelde over de herhalingen"] + titles[1:] + ["gemiddelde operator"])
    for o in range(MAX_OPERATORS):
        r, m = RANGE_FIRST + o, MEAN_FIRST + o
        name, trials, start = f"$A{r}", f"${OPERATOR_TRIALS}{r}", f"${OPERATOR_START}{r}"
        output_cell(ws, f"A{r}", f'=IFERROR(INDEX({NAMES},MATCH({o + 1},{NUMBERS},0)),"")')
        output_cell(ws, f"A{m}", f"=A{r}")
        output_cell(ws, f"{OPERATOR_TRIALS}{r}", f'=IF({name}="","",COUNTIF({NAMES},{name}))', "0")
        output_cell(ws, f"{OPERATOR_START}{r}", f'=IF({name}="","",MATCH({name},{NAMES},0))', "0")
        output_cell(ws, f"{OPERATOR_BLOCK}{r}", f'=IF({name}="","",IF(COUNTIF(OFFSET(${NAME_HELPER}${FIRST_DATA_ROW},'
                                               f'{start}-1,0,{trials},1),{name})={trials},"ja","NEE"))')
        for part in range(MAX_PARTS):
            letter = part_letter(part)
            block = f"OFFSET({letter}${FIRST_DATA_ROW},{start}-1,0,{trials},1)"
            used = f'OR({name}="",{letter}${PART_COUNTS}=0)'
            output_cell(ws, f"{letter}{r}", f'=IF({used},"",IF(COUNT({block})>1,MAX({block})-MIN({block}),""))', NUMBER)
            output_cell(ws, f"{letter}{m}", f'=IF({used},"",IF(COUNT({block})>0,AVERAGE({block}),""))', NUMBER)
        everything = f"OFFSET($B${FIRST_DATA_ROW},{start}-1,0,{trials},{MAX_PARTS})"
        output_cell(ws, f"{OPERATOR_MEAN}{m}", f'=IF({name}="","",AVERAGE({everything}))', NUMBER)
    label(ws, PART_MEANS, 1, "Gemiddelde per deel (alle operators en herhalingen)")
    label(ws, PART_COUNTS, 1, "aantal waarden per deel")
    for part in range(MAX_PARTS):
        letter = part_letter(part)
        values = f"{letter}${FIRST_DATA_ROW}:{letter}${LAST_DATA_ROW}"
        output_cell(ws, f"{letter}{PART_MEANS}", f'=IF(COUNT({values})>0,AVERAGE({values}),"")', NUMBER)
        output_cell(ws, f"{letter}{PART_COUNTS}", f"=COUNT({values})", "0")


def _size(ws: Worksheet) -> None:
    """Section 3: k, n, r and the completeness checks."""
    section_title(ws, 19, "3. Omvang van de studie (afgeleid uit de data)")
    last = RANGE_FIRST + MAX_OPERATORS - 1
    trials = f"{OPERATOR_TRIALS}{RANGE_FIRST}:{OPERATOR_TRIALS}{last}"
    result_row(ws, 20, "k = aantal operators", f'=IF(COUNT({NUMBERS})>0,MAX({NUMBERS}),0)', "0")
    result_row(ws, 21, "n = aantal delen", f'=COUNTIF(B{PART_COUNTS}:{part_letter(MAX_PARTS - 1)}{PART_COUNTS},">0")',
               "0")
    result_row(ws, 22, "r = aantal herhalingen (rijen per operator)",
               f'=IF(B20=0,0,IF(MIN({trials})=MAX({trials}),{OPERATOR_TRIALS}{RANGE_FIRST},"NEE: niet elke operator '
               f'heeft evenveel rijen"))', "0")
    result_row(ws, 23, "aantal metingen", f"=COUNT({DATA})", "0")
    result_row(ws, 24, "Rijen van elke operator bij elkaar?",
               f'=IF(B20=0,"",IF(COUNTIF({OPERATOR_BLOCK}{RANGE_FIRST}:{OPERATOR_BLOCK}{last},"NEE")=0,"ja",'
               f'"NEE: zet alle rijen van één operator onder elkaar"))')
    result_row(ws, 25, "Volledige studie (k ≥ 2, n ≥ 2, r ≥ 2, k·n·r waarden)?",
               f'=IF(B23=0,"",IF(AND(ISNUMBER(B22),B24="ja"),IF(AND(B20>=2,B21>=2,N(B22)>=2,B23=B20*B21*N(B22),'
               f'B20<={MAX_OPERATORS}),"ja","NEE: elke operator moet elk deel in elke herhaling meten"),'
               f'"NEE: zie r en de rijen per operator hierboven"))')


def _verdict(value: str) -> str:
    """Course verdict for a %GRR (MSA p. 35), using the thresholds of section 7."""
    low, high = (f"${c[0]}${c[1:]}" for c in (THRESHOLDS["acceptable"], THRESHOLDS["marginal"]))
    return (f'=IF(ISNUMBER({value}),IF({value}<={low},"aanvaardbaar (acceptable, ≤ 10 %)",IF({value}<={high},'
            f'"mogelijk aanvaardbaar (may be acceptable, 10 - 30 %)","niet aanvaardbaar (not acceptable, > 30 %)")),"")')


def _results(ws: Worksheet, first: int, ev: str, av: str, pv: str, source: str) -> None:
    """EV, AV, PV, GRR, TV, %GRR (TV and tolerance) and the two verdicts, from row `first`."""
    r = first
    result_row(ws, r, "EV = herhaalbaarheid (repeatability)", f'=IF({OK},{ev},"")', NUMBER, source)
    result_row(ws, r + 1, "AV = reproduceerbaarheid (reproducibility)", f'=IF({OK},{av},"")', NUMBER)
    result_row(ws, r + 2, "PV = variatie tussen de delen (part variation)", f'=IF({OK},{pv},"")', NUMBER)
    result_row(ws, r + 3, "GRR = √(EV² + AV²)", f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(B{r + 1})),SQRT(B{r}^2+B{r + 1}^2),"")',
               NUMBER, "MSA p. 18, 35")
    result_row(ws, r + 4, "TV = totale variatie (total variation) = √(GRR² + PV²)",
               f'=IF(AND(ISNUMBER(B{r + 3}),ISNUMBER(B{r + 2})),SQRT(B{r + 3}^2+B{r + 2}^2),"")', NUMBER)
    result_row(ws, r + 5, "%GRR t.o.v. totale variatie = GRR / TV", f'=IF(ISNUMBER(B{r + 4}),B{r + 3}/B{r + 4},"")',
               PERCENT, "MSA p. 35")
    result_row(ws, r + 6, "%GRR t.o.v. tolerantie = 6 · GRR / TOL",
               f'=IF(AND(ISNUMBER(B{r + 3}),ISNUMBER({TOL}),{TOL}>0),6*B{r + 3}/{TOL},"")', PERCENT, "MSA p. 24")
    # verdicts in column D: column C holds the course source of each row
    output_cell(ws, f"D{r + 5}", _verdict(f"B{r + 5}"))
    output_cell(ws, f"D{r + 6}", _verdict(f"B{r + 6}"))


# section 8 ranges: ranges, operator means, part means
RANGES = f"B{RANGE_FIRST}:{part_letter(MAX_PARTS - 1)}{RANGE_FIRST + MAX_OPERATORS - 1}"
CELL_MEANS = f"B{MEAN_FIRST}:{part_letter(MAX_PARTS - 1)}{MEAN_FIRST + MAX_OPERATORS - 1}"
OPERATOR_MEANS = f"{OPERATOR_MEAN}{MEAN_FIRST}:{OPERATOR_MEAN}{MEAN_FIRST + MAX_OPERATORS - 1}"
PART_MEAN_ROW = f"B{PART_MEANS}:{part_letter(MAX_PARTS - 1)}{PART_MEANS}"


def _average_range(ws: Worksheet) -> None:
    """Section 4: average and range method."""
    section_title(ws, 27, "4. Gemiddelde-en-spreidingsbreedtemethode (average and range method, MSA p. 34-35)")
    result_row(ws, 28, "R̿ = gemiddelde van alle spreidingsbreedtes", f'=IF({OK},AVERAGE({RANGES}),"")', NUMBER)
    result_row(ws, 29, "X̄_DIFF = grootste − kleinste operatorgemiddelde",
               f'=IF({OK},MAX({OPERATOR_MEANS})-MIN({OPERATOR_MEANS}),"")', NUMBER)
    result_row(ws, 30, "R_p = grootste − kleinste deelgemiddelde", f'=IF({OK},MAX({PART_MEAN_ROW})-MIN({PART_MEAN_ROW}),"")',
               NUMBER)
    result_row(ws, 31, "d2 voor m = r (g → ∞): K1 = 1/d2",
               f'=IF({OK},IFERROR({msa_d2_formula("$B$22")},"r niet in de tabel"),"")', "0.00000",
               "tabel MSA.pdf, laatste rij")
    result_row(ws, 32, "d2* voor m = k, g = 1: K2 = 1/d2*",
               f'=IF({OK},IFERROR({d2_star_formula("1", "$B$20")},"k niet in de tabel"),"")', "0.00000",
               "tabel MSA.pdf, g = 1")
    result_row(ws, 33, "d2* voor m = n, g = 1: K3 = 1/d2*",
               f'=IF({OK},IFERROR({d2_star_formula("1", "$B$21")},"n niet in de tabel"),"")', "0.00000",
               "tabel MSA.pdf, g = 1 (tot m = 20)")
    # a constant outside tabel MSA.pdf (e.g. more than 20 parts for K3) leaves that component empty, not an error
    _results(ws, 34, 'IF(ISNUMBER(B31),B28/B31,"")',
             'IF(AND(ISNUMBER(B31),ISNUMBER(B32)),SQRT(MAX(0,(B29/B32)^2-(B28/B31)^2/(B21*B22))),"")',
             'IF(AND(ISNUMBER(B31),ISNUMBER(B33)),SQRT(MAX(0,(B30/B33)^2-(B28/B31)^2/(B20*B22))),"")', "MSA p. 35")


def _anova_table(ws: Worksheet, rows: dict[str, int], entries: dict[str, tuple[str, str, str, str]], denominator: str) -> None:
    """Rows of an ANOVA table: SS, df, MS and for the factors F, p-value and F crit against `denominator`'s MS."""
    for key, (text, ss, df, _) in entries.items():
        r = rows[key]
        label(ws, r, 1, text)
        output_cell(ws, f"B{r}", f'=IF({OK},{ss},"")', "0.0000")
        output_cell(ws, f"C{r}", f'=IF({OK},{df},"")', "0")
        if key != "total":
            output_cell(ws, f"D{r}", f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),N(C{r})>0),B{r}/C{r},"")', "0.0000")
    denominator_row = rows[denominator]
    for key, (_, _, _, is_factor) in entries.items():
        if is_factor:
            r = rows[key]
            ms_e = f"$D${denominator_row}"
            output_cell(ws, f"E{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER({ms_e}),N({ms_e})>0),D{r}/{ms_e},"")', "0.0000")
            output_cell(ws, f"F{r}", f'=IF(ISNUMBER(E{r}),_xlfn.F.DIST.RT(E{r},C{r},$C${denominator_row}),"")', "0.000000")
            output_cell(ws, f"G{r}", f'=IF(AND(ISNUMBER(C{r}),ISNUMBER($C${denominator_row}),N(C{r})>0,N($C${denominator_row})>0,'
                                    f'ISNUMBER({ALPHA})),_xlfn.F.INV.RT({ALPHA},C{r},$C${denominator_row}),"")', "0.0000")


def _anova(ws: Worksheet) -> None:
    """Sections 5-6: ANOVA without interaction (the course's method) and with interaction (to judge it)."""
    ss_parts, ss_operators = f"B20*B22*DEVSQ({PART_MEAN_ROW})", f"B21*B22*DEVSQ({OPERATOR_MEANS})"
    ss_total, ss_cells = f"DEVSQ({DATA})", f"B22*DEVSQ({CELL_MEANS})"
    section_title(ws, 43, "5. ANOVA-methode (MSA p. 36-37): tweewegmodel zonder interactie, zoals de cursus")
    titles = ["Variatiebron (source)", "SS", "df", "MS", "F", "p-waarde", "F krit. bij α"]
    column_titles(ws, 44, titles)
    _anova_table(ws, ANOVA_ROWS, {
        "parts": ("Delen (parts)", ss_parts, "B21-1", "factor"),
        "operators": ("Operators", ss_operators, "B20-1", "factor"),
        "error": ("Fout (error): herhaalbaarheid + interactie", f"{ss_total}-{ss_parts}-{ss_operators}",
                  "B21*B20*(B22-1)+(B20-1)*(B21-1)", ""),
        "total": ("Totaal", ss_total, "B21*B20*B22-1", ""),
    }, "error")
    _results(ws, 49, "SQRT(D47)", "SQRT(MAX(0,(D46-D47)/(B21*B22)))", "SQRT(MAX(0,(D45-D47)/(B20*B22)))", "MSA p. 36")

    section_title(ws, 57, "6. ANOVA met de interactie operator × deel (zoals Excel 'Anova: two-factor with replication')")
    column_titles(ws, 58, titles)
    _anova_table(ws, INTERACTION_ROWS, {
        "parts": ("Delen (parts)", ss_parts, "B21-1", "factor"),
        "operators": ("Operators", ss_operators, "B20-1", "factor"),
        "interaction": ("Interactie operator × deel", f"{ss_cells}-{ss_parts}-{ss_operators}", "(B21-1)*(B20-1)",
                        "factor"),
        "within": ("Binnen (within): herhaalbaarheid", f"{ss_total}-{ss_cells}", "B21*B20*(B22-1)", ""),
        "total": ("Totaal", ss_total, "B21*B20*B22-1", ""),
    }, "within")
    label(ws, 64, 1, "Een grote p-waarde van de interactie steunt het poolen ervan in de fout, zoals sectie 5 doet (de "
                     "'afgeleide ANOVA-tabel zonder interactie' van het cursuswerkboek). De cursus geeft EV, AV, PV alleen "
                     "voor sectie 5.", italic=True)


def _thresholds(ws: Worksheet) -> None:
    """Section 7: the course's acceptance rule and the other thresholds it mentions."""
    section_title(ws, 66, "7. Aanvaarding van het meetsysteem")
    constant_row(ws, 67, "%GRR ≤ dit: aanvaardbaar", 0.10,
                 "MSA p. 35 ('%GRR <= 10% : gauge generally considered to be acceptable')")
    constant_row(ws, 68, "%GRR ≤ dit: mogelijk aanvaardbaar; erboven: niet aanvaardbaar", 0.30,
                 "MSA p. 35 (10 % - 30 %: hangt af van de toepassing; > 30 %: niet aanvaardbaar)")
    ws["B67"].number_format = ws["B68"].number_format = "0%"
    label(ws, 69, 1, "Andere drempels: 'Gauge R&R ≤ 20 % of tolerance' (deck Les 4 p. 61); σ²_measure / σ²_observed "
                     "≤ 0,1 goed, 0,1-0,3 marginaal, ≥ 0,3 onaanvaardbaar: Dummies p. 178 (extra, Dummies; niet te "
                     "kennen).", italic=True)
    label(ws, 70, 1, "Vermenigvuldigingsfactor 6 (99,73 %) zoals in de MSA-formules; vroeger 5,15 (99 %), MSA p. 24. "
                     "ndc (number of distinct categories, aantal onderscheiden categorieën) staat niet in de cursus.",
          italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the Gage R&R calculator."""
    write_header(ws, HEADER)
    section_title(ws, 8, "1. Instellingen")
    input_row(ws, 9, "Tolerantie (tolerance) TOL = USL − LSL (optioneel)", "voor %GRR t.o.v. de tolerantie")
    input_row(ws, 10, "α voor de F-toetsen", "de cursus heeft geen standaardwaarde; de meeste cursusvoorbeelden "
                                             "gebruiken 5 %", "0.00")
    ws[ALPHA.replace("$", "")] = 0.05
    _instructions(ws)
    _size(ws)
    _average_range(ws)
    _anova(ws)
    _thresholds(ws)
    _per_cell(ws)
    _data(ws)
    msa_table(ws, 19, 9)
    ws.column_dimensions["A"].width = 50
    for column in range(FIRST_PART_COLUMN, FIRST_PART_COLUMN + MAX_PARTS):
        ws.column_dimensions[get_column_letter(column)].width = 11
    for letter in (OPERATOR_MEAN, OPERATOR_TRIALS, OPERATOR_START, OPERATOR_BLOCK):
        ws.column_dimensions[letter].width = 14
