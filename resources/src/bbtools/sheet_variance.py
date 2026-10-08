"""Variance sheet ('Varianties BI & toetsen'): confidence intervals and tests for σ², σ (χ²) and for a ratio of two
variances (F).

Course: (n−1)s²/σ² ~ χ²(n−1) and the 95 % CI (n−1)s²/χ²_0.025 ≤ σ² ≤ (n−1)s²/χ²_0.975 (CI Further Reading
(Dutch) p. 21); χ²-test for σ (Test Recipes p. 11; Testing of Hypotheses p. 13); (s1²/σ1²)/(s2²/σ2²) ~
F(n1−1, n2−1) and the F-test (Test Recipes p. 12-13); exam Q2's hint (s1²/s2²)(σ2²/σ1²) ~ F(n1−1, n2−1).
Decision 5: α is an input prefilled 0.05. Decision 6: every interval and test is shown two-sided, lower-only
and upper-only. Decision 7: F quantiles come straight from F.INV / F.INV.RT; both ratio orientations shown.
Excel names: CHISQ.INV.RT(p; df) is the value with upper-tail area p, CHISQ.INV(p; df) with lower-tail area p.

Row plan: settings 8-9; one sample 11-29 (inputs 12-14, used 15-18, CIs 21-23, χ²-tests 27-29);
two samples 31-51 (inputs 32-35, used 36-41, CIs 44-46, F-tests 49-51); raw data in H and I from row 10 down
(open-ended), how to fill them and their checks in column K.
"""
from __future__ import annotations

from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from bbtools.readme import SheetDoc
from bbtools.xlsx_style import (
    DATA_LAST_ROW,
    HeaderBlock,
    Status,
    column_titles,
    data_check,
    font,
    input_row,
    instructions,
    label,
    open_input_column,
    output_cell,
    result_row,
    section_title,
    write_header,
)

SHEET = "Varianties BI & toetsen"

HEADER = HeaderBlock(
    tool="Varianties: betrouwbaarheidsinterval (confidence interval, BI) en toets voor σ² en σ (χ²), BI en toets "
         "voor een verhouding van twee varianties (F)",
    source="source/course/Les 2/20260529_ottoy_Confidence Intervals - Further Reading (Dutch).pdf p. 21; "
           "Confidence Intervals.pdf p. 16; Testing of Hypotheses.pdf p. 13; Test Recipes - Further Reading "
           "(Dutch).pdf p. 11-14",
    convention="Conventiebeslissingen 5-7: invoer α vooraf ingevuld op 0,05; tweezijdig, alleen ondergrens en alleen "
               "bovengrens naast elkaar; F uit F.INV / F.INV.RT, zowel σ1²/σ2² als σ2²/σ1².",
    status=Status.VERIFIED,
    status_detail="getest tegen uitgewerkte cursusvoorbeelden (worked examples) S03-WE03, S03-WE05; ook tegen "
                  "S08-WE10 en de χ²- en F-tabellen van Dummies (p. 194, 196) (extra, Dummies; niet te kennen); het "
                  "F-intervalvoorbeeld S08-WE11 van Dummies klopt niet (resources/build/README.md)",
)

DOC = SheetDoc(
    sheet=SHEET,
    purpose="Eén steekproef: betrouwbaarheidsinterval (BI) voor σ² en σ, χ²-toets van σ = σ0. Twee steekproeven: "
            "BI voor σ1²/σ2² en σ2²/σ1², F-toets van σ1 = σ2 (bv. examenvraag Q2: werkt machine M1 nauwkeuriger dan "
            "M2?). Elk resultaat tweezijdig, alleen ondergrens en alleen bovengrens.",
    inputs="α; n en s per steekproef, of de ruwe waarden geplakt in kolommen H (steekproef 1) en I (steekproef 2); "
           "σ0 voor de χ²-toets.",
    audit="stats-auditor PASS (2026-09-28) voor het F-deel: alle 28 waarden (BI's in beide richtingen, F-toetsen "
          "≠, >, <) voor één invoer kloppen met een onafhankelijke berekening uit Test Recipes p. 12-14, CI Further "
          "Reading p. 21 en de hint van examenvraag Q2. Opmerking: de cursus drukt geen expliciet BI voor de "
          "verhouding van twee varianties; het volgt uit de F-pivot van p. 12, omgekeerd zoals op p. 21.",
    disagreements=(
        "Dummies p. 196 (S08-WE11): BI voor σA²/σB² gedrukt als '[(1/3.633)(4/7.5), 5.999(4/7.5)] = [0.147, 3.199]'. "
        "Met (sA²/σA²)/(sB²/σB²) ~ F(nA−1, nB−1) (Test Recipes p. 12, hint van examenvraag Q2) geven dezelfde "
        "staartwaarden van 5 % [0,0889; 1,938]: het boek verwisselt de twee F-waarden. De tabelwaarden zelf "
        "(Table 8-3) kloppen wel (extra, Dummies; niet te kennen).",
        "Dummies p. 192-196 noemt '95 %' wat ±2σ is (95,45 %, 2,275 % per staart) voor χ², en een rechterstaart van "
        "5 % voor F. Gebruik α = 0,0455 om zijn χ²-voorbeeld (S08-WE10) na te rekenen (extra, Dummies; niet te "
        "kennen).",
        "Dummies Table 8-3 p. 196: F voor n1 = n2 = 2 gedrukt als '161.446' (definitie 161,448). Table 8-2 p. 194: "
        "bovenwaarde bij 99,7 % voor n = 5 gedrukt als '17.800' (definitie 17,8006). Alle andere waarden kloppen "
        "(extra, Dummies; niet te kennen).",
    ),
)

ALPHA = "B9"
DATA_1, DATA_2 = f"H10:H{DATA_LAST_ROW}", f"I10:I{DATA_LAST_ROW}"
CHECKS = {"data_1": "K16", "data_2": "K17"}
INPUTS = {"alpha": ALPHA, "n": "B12", "s": "B13", "sigma0": "B14", "n1": "B32", "s1": "B33", "n2": "B34", "s2": "B35"}
RESULTS = {
    "n_used": "B15", "s_used": "B16", "var_used": "B17", "df": "B18",
    "n1_used": "B36", "s1_used": "B37", "n2_used": "B38", "s2_used": "B39", "f": "B40", "df1": "B41", "df2": "C41",
}
CI_ROWS = {"two_sided": 21, "lower_only": 22, "upper_only": 23}          # B,C = σ² from/to; D,E = σ from/to
CHI2_TEST_ROWS = {"two_sided": 27, "greater": 28, "less": 29}           # B stat, C/D critical, E p, F decision
F_CI_ROWS = {"two_sided": 44, "ratio_at_least": 45, "ratio_at_most": 46}  # B,C = σ1²/σ2²; D,E = σ2²/σ1²
F_TEST_ROWS = {"two_sided": 49, "greater": 50, "less": 51}              # B stat, C/D critical, E p, F decision
NUMBER = "0.000000"
DECISION = '=IF(ISNUMBER(E{r}),IF(E{r}<' + ALPHA + ',"verwerp H0","H0 niet verwerpen"),"")'


def _data_columns(ws: Worksheet) -> None:
    """Columns H and I: optional raw data for sample 1 and sample 2, open-ended; how to fill them in column K."""
    label(ws, 8, 8, "Optioneel: ruwe waarden (uitleg →)", italic=True)
    for column, text in (("H", "Data steekproef 1"), ("I", "Data steekproef 2")):
        ws[f"{column}9"] = text
        ws[f"{column}9"].font = font(bold=True)
        open_input_column(ws, column, 10)
    row = instructions(ws, 8, 11, "Zo gebruik je kolommen H en I (ruwe data)", (
        "• Eén GETAL per cel, onder elkaar, vanaf rij 10; zoveel rijen als je wilt. Geen kop of tekst: de kop staat al "
        "in rij 9.",
        "• Kolom H = steekproef 1: telt voor sectie 2 (één steekproef) én voor steekproef 1 van sectie 3.",
        "• Kolom I = steekproef 2 (sectie 3, F-toets).",
        "• Met 2 of meer getallen in een kolom rekent het blad n en s daaruit; de getypte n en s tellen dan niet.",
        "• Typ je liever n en s: laat H en I leeg (selecteer de kolom vanaf rij 10 en druk Delete).",
        "• Plakken uit een ander bestand: Plakken speciaal → Waarden.",
    ))
    label(ws, row, 11, "Controle van de geplakte data:", bold=True)
    data_check(ws, CHECKS["data_1"], DATA_1)
    data_check(ws, CHECKS["data_2"], DATA_2)
    label(ws, 16, 10, "H:")
    label(ws, 17, 10, "I:")


def _used(ws: Worksheet, row: int, text: str, data: str, typed: str, is_n: bool) -> None:
    """A green 'value used' cell: from the pasted data when it holds at least 2 values, else the typed input."""
    if is_n:
        formula, fmt = f'=IF(COUNT({data})>1,COUNT({data}),IF(ISNUMBER({typed}),{typed},""))', "0"
    else:
        formula, fmt = f'=IF(COUNT({data})>1,_xlfn.STDEV.S({data}),IF(ISNUMBER({typed}),{typed},""))', NUMBER
    result_row(ws, row, text, formula, fmt, f"uit kolom {data[0]} als die 2+ waarden bevat, anders de getypte waarde")


def _test_rows(ws: Worksheet, rows: dict[int, tuple[str, str, str | None, str]], have: str, stat: str,
               crit_have: str) -> None:
    """Rows of a test table: statistic, critical value(s), p-value and decision at α. The critical values need only
    α and the degrees of freedom (`crit_have`), so they show before s is entered."""
    for row, (text, critical, critical_2, p_value) in rows.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({have},{stat},"")', "0.0000")
        output_cell(ws, f"C{row}", f'=IF({crit_have},{critical},"")', "0.0000")
        output_cell(ws, f"D{row}", f'=IF({crit_have},{critical_2},"")' if critical_2 else '=""', "0.0000")
        output_cell(ws, f"E{row}", f'=IF({have},{p_value},"")', "0.0000%")
        output_cell(ws, f"F{row}", DECISION.format(r=row))


def _one_sample(ws: Worksheet) -> None:
    """Sections 1-2: α and one sample (CI for σ², σ and the χ²-test)."""
    section_title(ws, 8, "1. Instellingen")
    input_row(ws, 9, "Significantieniveau α (betrouwbaarheid = 1 − α)",
              "de cursus geeft geen standaardwaarde; de meeste cursusvoorbeelden gebruiken 5 %; gebruik de waarde uit "
              "de vraag", "0.0000")
    ws[ALPHA] = 0.05

    section_title(ws, 11, "2. Eén steekproef: σ² en σ")
    input_row(ws, 12, "n (steekproefgrootte, sample size)", "of plak de data in kolom H")
    input_row(ws, 13, "s (steekproefstandaardafwijking, STDEV.S)", "", NUMBER)
    input_row(ws, 14, "σ0 voor de toets (de beweerde σ)", "bv. 0,01", NUMBER)
    _used(ws, 15, "n gebruikt", DATA_1, "B12", is_n=True)
    _used(ws, 16, "s gebruikt", DATA_1, "B13", is_n=False)
    result_row(ws, 17, "s² gebruikt", '=IF(ISNUMBER(B16),B16^2,"")', "0.0000000000")
    result_row(ws, 18, "vrijheidsgraden (degrees of freedom, df) n − 1", '=IF(ISNUMBER(B15),B15-1,"")', "0")

    column_titles(ws, 20, ["Betrouwbaarheidsinterval (confidence interval, BI) (1 − α)", "σ² van", "σ² tot", "σ van",
                           "σ tot", "Bron in de cursus"])
    ok = f"AND(ISNUMBER(B17),ISNUMBER(B18),B18>0,ISNUMBER({ALPHA}))"
    q = "B18*B17"  # (n−1) s²
    intervals = {
        21: ("tweezijdig (two-sided)", f"{q}/_xlfn.CHISQ.INV.RT({ALPHA}/2,B18)", f"{q}/_xlfn.CHISQ.INV({ALPHA}/2,B18)",
             "CI Further Reading (Dutch) p. 21"),
        22: ("alleen ondergrens (σ minstens …; hoort bij HA: σ > σ0)", f"{q}/_xlfn.CHISQ.INV.RT({ALPHA},B18)", None,
             "Confidence Intervals.pdf p. 16 ('one-sided 98 %-CI ]0.0088, +∞[')"),
        23: ("alleen bovengrens (σ hoogstens …; hoort bij HA: σ < σ0)", None, f"{q}/_xlfn.CHISQ.INV({ALPHA},B18)", ""),
    }
    for row, (text, low, high, source) in intervals.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({ok},{low if low else 0},"")', NUMBER)
        output_cell(ws, f"C{row}", f'=IF({ok},{high},"")' if high else f'=IF({ok},"+∞","")', NUMBER)
        output_cell(ws, f"D{row}", f'=IF(ISNUMBER(B{row}),SQRT(B{row}),"")', NUMBER)
        output_cell(ws, f"E{row}", f'=IF(ISNUMBER(C{row}),SQRT(C{row}),C{row})', NUMBER)
        label(ws, row, 6, source, italic=True)

    column_titles(ws, 26, ["χ²-toets van H0: σ = σ0", "Toetsgrootheid (test statistic) (n−1)s²/σ0²",
                           "Kritieke waarde (critical value)", "2de kritieke waarde", "p-waarde (p-value)",
                           "Besluit bij α"])
    _test_rows(ws, {
        27: ("HA: σ ≠ σ0 (tweezijdig)", f"_xlfn.CHISQ.INV({ALPHA}/2,B18)", f"_xlfn.CHISQ.INV.RT({ALPHA}/2,B18)",
             "2*MIN(_xlfn.CHISQ.DIST(B27,B18,TRUE),_xlfn.CHISQ.DIST.RT(B27,B18))"),
        28: ("HA: σ > σ0 (bv. 'more erratic')", f"_xlfn.CHISQ.INV.RT({ALPHA},B18)", None, "_xlfn.CHISQ.DIST.RT(B28,B18)"),
        29: ("HA: σ < σ0", f"_xlfn.CHISQ.INV({ALPHA},B18)", None, "_xlfn.CHISQ.DIST(B29,B18,TRUE)"),
    }, f"AND({ok},ISNUMBER(B14),B14>0)", "B18*B17/B14^2", f"AND(ISNUMBER(B18),N(B18)>0,ISNUMBER({ALPHA}))")
    label(ws, 30, 1, "Bron: Test Recipes - Further Reading (Dutch) p. 11; Testing of Hypotheses.pdf p. 13", italic=True)


def _two_samples(ws: Worksheet) -> None:
    """Section 3: two samples (CI for the variance ratio in both orientations, and the F-test)."""
    section_title(ws, 31, "3. Twee steekproeven: verhouding van varianties (F)   bv. examenvraag Q2: werkt machine M1 "
                          "nauwkeuriger dan M2?")
    input_row(ws, 32, "n1 (steekproef 1)", "of plak steekproef 1 in kolom H")
    input_row(ws, 33, "s1", "", NUMBER)
    input_row(ws, 34, "n2 (steekproef 2)", "of plak steekproef 2 in kolom I")
    input_row(ws, 35, "s2", "", NUMBER)
    _used(ws, 36, "n1 gebruikt", DATA_1, "B32", is_n=True)
    _used(ws, 37, "s1 gebruikt", DATA_1, "B33", is_n=False)
    _used(ws, 38, "n2 gebruikt", DATA_2, "B34", is_n=True)
    _used(ws, 39, "s2 gebruikt", DATA_2, "B35", is_n=False)
    have = f"AND(ISNUMBER(B36),ISNUMBER(B37),ISNUMBER(B38),ISNUMBER(B39),B36>1,B38>1,B39>0,ISNUMBER({ALPHA}))"
    result_row(ws, 40, "F = s1² / s2²", f'=IF({have},B37^2/B39^2,"")', "0.0000",
               "Test Recipes p. 12-13: (s1²/σ1²)/(s2²/σ2²) ~ F(n1 − 1, n2 − 1)")
    label(ws, 41, 1, "vrijheidsgraden (teller n1 − 1, noemer n2 − 1)")
    output_cell(ws, "B41", '=IF(ISNUMBER(B36),B36-1,"")', "0")
    output_cell(ws, "C41", '=IF(ISNUMBER(B38),B38-1,"")', "0")

    column_titles(ws, 43, ["Betrouwbaarheidsinterval (1 − α)", "σ1²/σ2² van", "σ1²/σ2² tot", "σ2²/σ1² van",
                           "σ2²/σ1² tot", "Bron in de cursus"])
    f, v1, v2 = "B40", "B41", "C41"
    intervals = {
        44: ("tweezijdig (two-sided)", f"{f}/_xlfn.F.INV.RT({ALPHA}/2,{v1},{v2})",
             f"{f}/_xlfn.F.INV({ALPHA}/2,{v1},{v2})", "hint van examenvraag Q2; Test Recipes p. 12"),
        45: ("eenzijdig: σ1²/σ2² minstens … (σ2²/σ1² hoogstens …)", f"{f}/_xlfn.F.INV.RT({ALPHA},{v1},{v2})", None,
             "hoort bij HA: σ1 > σ2"),
        46: ("eenzijdig: σ1²/σ2² hoogstens … (σ2²/σ1² minstens …)", None, f"{f}/_xlfn.F.INV({ALPHA},{v1},{v2})",
             "hoort bij HA: σ1 < σ2 (bv. M1 nauwkeuriger)"),
    }
    for row, (text, low, high, source) in intervals.items():
        label(ws, row, 1, text)
        output_cell(ws, f"B{row}", f'=IF({have},{low if low else 0},"")', NUMBER)
        output_cell(ws, f"C{row}", f'=IF({have},{high},"")' if high else f'=IF({have},"+∞","")', NUMBER)
        # σ2²/σ1² is the reciprocal interval: from = 1 / (σ1²/σ2² to), to = 1 / (σ1²/σ2² from)
        output_cell(ws, f"D{row}", f'=IF(ISNUMBER(C{row}),1/C{row},IF(C{row}="+∞",0,""))', NUMBER)
        output_cell(ws, f"E{row}", f'=IF(ISNUMBER(B{row}),IF(B{row}=0,"+∞",1/B{row}),"")', NUMBER)
        label(ws, row, 6, source, italic=True)

    column_titles(ws, 48, ["F-toets van H0: σ1 = σ2", "Toetsgrootheid F = s1²/s2²", "Kritieke waarde",
                           "2de kritieke waarde", "p-waarde", "Besluit bij α"])
    _test_rows(ws, {
        49: ("HA: σ1 ≠ σ2 (tweezijdig)", f"_xlfn.F.INV({ALPHA}/2,{v1},{v2})", f"_xlfn.F.INV.RT({ALPHA}/2,{v1},{v2})",
             f"2*MIN(_xlfn.F.DIST(B49,{v1},{v2},TRUE),_xlfn.F.DIST.RT(B49,{v1},{v2}))"),
        50: ("HA: σ1 > σ2", f"_xlfn.F.INV.RT({ALPHA},{v1},{v2})", None, f"_xlfn.F.DIST.RT(B50,{v1},{v2})"),
        51: ("HA: σ1 < σ2", f"_xlfn.F.INV({ALPHA},{v1},{v2})", None, f"_xlfn.F.DIST(B51,{v1},{v2},TRUE)"),
    }, have, f, f"AND(ISNUMBER({v1}),ISNUMBER({v2}),N({v1})>0,N({v2})>0,ISNUMBER({ALPHA}))")
    label(ws, 52, 1, "Bron: Test Recipes - Further Reading (Dutch) p. 13-14", italic=True)


def build_sheet(ws: Worksheet) -> None:
    """Fill an empty worksheet with the variance calculator."""
    write_header(ws, HEADER)
    _data_columns(ws)
    _one_sample(ws)
    _two_samples(ws)
    fraction = DataValidation(type="decimal", operator="between", formula1="0.0000001", formula2="0.9999999",
                              allow_blank=True, showErrorMessage=True, errorTitle="α",
                              error="α is een fractie, bv. 0,05.")
    positive = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                              showErrorMessage=True, errorTitle="Moet positief zijn", error="Moet groter zijn dan 0.")
    ws.add_data_validation(fraction)
    ws.add_data_validation(positive)
    fraction.add(ALPHA)
    for coordinate in ("B12", "B13", "B14", "B32", "B33", "B34", "B35"):
        positive.add(coordinate)
    ws.column_dimensions["A"].width = 52
    for letter in "BCDE":
        ws.column_dimensions[letter].width = 17
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 4
    ws.column_dimensions["H"].width = 18
    ws.column_dimensions["I"].width = 18
    ws.column_dimensions["J"].width = 4
