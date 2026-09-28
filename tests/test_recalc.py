"""Tests for bbtools.recalc: LibreOffice computes formulas, and errors are found and located."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from bbtools import recalc as recalc_module
from bbtools.recalc import RecalcError, recalc


@pytest.mark.libreoffice
def test_recalc_computes_formulas_and_locates_errors(tmp_path: Path) -> None:
    # arrange
    path = tmp_path / "small.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "T"
    ws["A1"], ws["A2"], ws["A3"] = "=1+1", "=1/0", "=_xlfn.NORM.S.DIST(0,TRUE)"
    wb.save(path)
    # act
    report = recalc(path)
    values = load_workbook(path, data_only=True)["T"]
    # assert -- NORM.S.DIST(0) = 0.5 exactly
    assert report.total_formulas == 3
    assert report.errors == {"#DIV/0!": ["T!A2"]}
    assert (values["A1"].value, values["A3"].value) == (2, 0.5)


def test_recalc_raises_when_libreoffice_is_missing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # arrange -- patch at the use site
    monkeypatch.setattr(recalc_module.shutil, "which", lambda _name: None)
    path = tmp_path / "x.xlsx"
    Workbook().save(path)
    # act / assert
    with pytest.raises(RecalcError, match="soffice not found"):
        recalc(path)
