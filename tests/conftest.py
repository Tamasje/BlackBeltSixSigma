"""Shared fixtures: the locked oracle and a helper that evaluates a calculator sheet through LibreOffice."""
from __future__ import annotations

import json
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from bbtools.build_workbook import build_workbook
from bbtools.constants import ROOT
from bbtools.recalc import recalc

ORACLE = ROOT / "inventory" / "worked_examples.json"


@pytest.fixture(scope="session")
def oracle() -> dict[str, dict[str, Any]]:
    """Course worked examples by ID (inventory/worked_examples.json, locked by tag oracle-approved)."""
    examples = json.loads(ORACLE.read_text(encoding="utf-8"))["examples"]
    return {example["id"]: example for example in examples}


@pytest.fixture(scope="session")
def printed(oracle: dict[str, dict[str, Any]]) -> Callable[[str, str, str, str], str]:
    """Return a printed number after asserting the oracle really prints it.

    printed("S06-WE08", "stated_answers", "Cp_as_printed", "0,83") checks that the oracle string for that key
    contains '0,83' and returns '0,83'. Tests name the exact printed value instead of parsing prose.
    """
    def lookup(example_id: str, part: str, key: str, value: str) -> str:
        text = str(oracle[example_id][part][key])
        assert value in text, f"{example_id} {part}[{key!r}] = {text!r} does not contain {value!r}"
        return value
    return lookup


@pytest.fixture
def evaluate(tmp_path: Path) -> Callable[[str, dict[str, float]], Any]:
    """Build the toolkit, type `inputs` into sheet `sheet`, recalculate with LibreOffice, return that sheet's values.

    `inputs` maps cell coordinates to values. The returned worksheet holds cached values (data_only=True);
    cells whose formula returned "" read back as None.
    """
    def run(sheet: str, inputs: dict[str, float]) -> Any:
        wb = build_workbook()
        ws = wb[sheet]
        for coordinate, value in inputs.items():
            ws[coordinate] = value
        path = tmp_path / f"{sheet}_{len(list(tmp_path.iterdir()))}.xlsx"
        wb.save(path)
        report = recalc(path)
        assert report.total_errors == 0, f"formula errors after recalculation: {report.errors}"
        return load_workbook(path, data_only=True)[sheet]
    return run
