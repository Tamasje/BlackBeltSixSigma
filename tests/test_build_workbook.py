"""Tests for bbtools.build_workbook: the delivered workbook obeys the Excel rules of CLAUDE.md."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bbtools.build_workbook import build_workbook, save, verify
from bbtools.xlsx_style import INPUT_FILL, OUTPUT_FILL

# Spilling / dynamic-array functions are banned until the exam machine's Excel version is confirmed (CLAUDE.md).
FORBIDDEN = ("LET", "LAMBDA", "XLOOKUP", "XMATCH", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "RANDARRAY")
# Post-2007 names contain a dot (NORM.S.DIST, T.INV.2T, STDEV.S, ...); Excel needs the _xlfn. prefix for them.
DOTTED_FUNCTION = re.compile(r"(?<![A-Za-z0-9_.])([A-Z][A-Z0-9]*(?:\.[A-Z0-9]+)+)\(")


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """The toolkit written to a temporary path, as build/bb_toolkit.xlsx would be."""
    return save(build_workbook(), tmp_path_factory.mktemp("build") / "bb_toolkit.xlsx")


def formulas(path: Path) -> list[tuple[str, str, str]]:
    """(sheet, cell, formula) for every formula in the workbook."""
    wb = load_workbook(path)
    return [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]


@pytest.mark.libreoffice
def test_libreoffice_recalculation_reports_zero_formula_errors(built: Path) -> None:
    # act
    report = verify(built)
    # assert
    assert report.total_formulas > 0
    assert report.errors == {}


def test_no_dynamic_array_functions(built: Path) -> None:
    # act
    offending = [(s, c, f) for s, c, f in formulas(built)
                 if re.search(r"(?<![A-Za-z0-9_.])(" + "|".join(FORBIDDEN) + r")\(", f.upper())]
    # assert
    assert offending == []


def test_every_post_2007_function_carries_the_xlfn_prefix(built: Path) -> None:
    # act -- a dotted name right after '_xlfn.' is fine; anything else shows #NAME? in Excel
    bare = [(s, c, f) for s, c, f in formulas(built)
            for name in DOTTED_FUNCTION.findall(f.replace("_xlfn.", "_xlfn_")) if not name.startswith("_")]
    # assert
    assert bare == []


def test_no_macros_and_no_external_links(built: Path) -> None:
    # act
    names = zipfile.ZipFile(built).namelist()
    # assert
    assert not any("vbaProject" in n for n in names)
    assert not any(n.startswith("xl/externalLinks/") for n in names)


def test_excel_recalculates_everything_on_open(built: Path) -> None:
    # act / assert -- the delivered file has no cached values, so Excel must compute on load
    assert load_workbook(built).calculation.fullCalcOnLoad is True


def test_every_sheet_has_a_complete_header_block(built: Path) -> None:
    # arrange
    wb = load_workbook(built)
    for ws in wb.worksheets:
        # act
        labels = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        texts = [ws.cell(row=r, column=2).value for r in range(2, 5)]
        # assert -- tool name, course source, convention, status (CLAUDE.md)
        assert ws["A1"].value, ws.title
        assert labels == ["Course source", "Convention used", "Status"], ws.title
        assert all(texts), ws.title
        assert texts[2].startswith(("VERIFIED", "UNVERIFIED")), ws.title


def test_calculator_sheets_use_distinct_input_and_output_colours(built: Path) -> None:
    # arrange
    wb = load_workbook(built)
    for ws in wb.worksheets:
        if ws.title == "Tables":
            continue
        # act
        fills = {c.fill.fgColor.rgb[-6:] for row in ws.iter_rows() for c in row if c.fill.fill_type == "solid"}
        # assert
        assert INPUT_FILL.fgColor.rgb[-6:] in fills, ws.title
        assert OUTPUT_FILL.fgColor.rgb[-6:] in fills, ws.title
