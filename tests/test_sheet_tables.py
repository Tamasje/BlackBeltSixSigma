"""Tests for bbtools.sheet_tables: the Tables sheet shows every source as printed and names the lookup ranges."""
from __future__ import annotations

from openpyxl import Workbook

from bbtools.constants import TABLE_SOURCES, load_all, load_average_range_table
from bbtools.printed import decimals_printed, parse_printed
from bbtools.sheet_tables import SHEET, build_tables_sheet, disagreeing_cells, excel_name, lookup_formula
from bbtools.xlsx_style import FLAG_FILL


def build() -> Workbook:
    """A workbook holding only the Tables sheet."""
    wb = Workbook()
    build_tables_sheet(wb, load_all())
    return wb


def test_excel_name_makes_every_symbol_a_valid_name() -> None:
    # act / assert
    assert excel_name("T18", "d2") == "T18_d2"
    assert excel_name("T18", "1/d2") == "T18_inv_d2"


def test_lookup_formula_uses_the_table_of_decision_4() -> None:
    # act / assert
    assert lookup_formula("d2", "B17") == "INDEX(T18_d2,MATCH(B17,T18_n,0))"
    assert lookup_formula("c4", "B17") == "INDEX(TA_c4,MATCH(B17,TA_n,0))"
    assert lookup_formula("E2", "2") == "INDEX(SSD2_E2,MATCH(2,SSD2_n,0))"


def test_every_printed_cell_appears_with_its_printed_decimals() -> None:
    # arrange
    wb = build()
    ws = wb[SHEET]
    by_position = {(c.row, c.column): c for row in ws.iter_rows() for c in row}
    # act -- walk each table by its n-column name, which points at its first data row
    for table in load_all():
        first = int(wb.defined_names[f"{table.source.key}_n"].attr_text.split("$")[2].split(":")[0])
        for offset, row in enumerate(table.rows):
            for column, text in enumerate(row, start=1):
                cell = by_position.get((first + offset, column))
                # assert -- numbers stored as numbers and shown with the printed number of decimals
                if text.strip() and _is_number(text):
                    assert cell.value == float(parse_printed(text)), (table.source.key, row[0], column)
                    decimals = decimals_printed(text)
                    assert cell.number_format == ("0" if decimals == 0 else "0." + "0" * decimals)
                else:
                    assert (cell.value or "") == text


def test_every_symbol_column_has_a_workbook_name() -> None:
    # arrange
    wb = build()
    # act
    expected = {f"{t.source.key}_n" for t in load_all()} | {
        excel_name(t.source.key, s) for t in load_all() for s in t.symbol_columns()}
    # assert
    assert expected <= set(wb.defined_names)
    assert len(TABLE_SOURCES) == 5


def test_disagreements_beyond_rounding_are_flagged_and_rounding_is_not() -> None:
    # arrange
    flagged = disagreeing_cells(load_all())
    # act / assert -- D4(5): Table 18 prints 2.115, Six Sigma Demystified 2.114 -> flagged
    assert ("T18", "D4", "5") in flagged and ("SSD2", "D4", "5") in flagged
    # d3(10): Table A 0.7971 vs Demystified 0.797 -> same value at 3 decimals, not flagged
    assert ("TA", "d3", "10") not in flagged


def test_flagged_cells_are_orange_on_the_sheet() -> None:
    # arrange
    wb = build()
    ws = wb[SHEET]
    first = int(wb.defined_names["T18_D4"].attr_text.split("$")[2].split(":")[0])
    column = wb.defined_names["T18_D4"].attr_text.split("$")[1]
    # act -- row for n = 5 is the 4th data row (n = 2, 3, 4, 5)
    cell = ws[f"{column}{first + 3}"]
    # assert
    assert cell.value == 2.115
    assert cell.fill.fgColor.rgb.endswith(FLAG_FILL.fgColor.rgb[-6:])


def _is_number(text: str) -> bool:
    """True if the printed text is one number."""
    try:
        parse_printed(text)
    except ValueError:
        return False
    return True


def test_msa_table_grids_hold_every_printed_value() -> None:
    # arrange -- tabel MSA.pdf: each printed cell 'nu / d2*' is split into two grids
    wb = build()
    ws = wb[SHEET]
    table = load_average_range_table()

    def origin(name: str) -> tuple[int, int]:
        """(row, column) of the top-left cell of a workbook name."""
        ref = wb.defined_names[name].attr_text.split("!")[1].split(":")[0].replace("$", "")
        column = "".join(ch for ch in ref if ch.isalpha())
        return int(ref[len(column):]), ws[f"{column}1"].column

    # act / assert -- every value numeric, equal to the printed string, with the printed decimals
    for name, grid in (("MSA_d2star", table.d2_star), ("MSA_nu", table.nu)):
        row0, col0 = origin(name)
        for i, printed_row in enumerate(grid):
            for j, text in enumerate(printed_row):
                cell = ws.cell(row=row0 + i, column=col0 + j)
                assert cell.value == float(parse_printed(text)), (name, table.g[i], table.m[j])
    row0, col0 = origin("MSA_d2")
    assert [ws.cell(row=row0, column=col0 + j).value for j in range(len(table.m))] == [float(v) for v in table.d2]
    assert table.m == tuple(range(2, 21)) and table.g == tuple(range(1, 21))
