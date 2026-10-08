"""Tests for the study guide's calculators (study/assets/stats.js and calc.js), run in node.

Expected values: course worked examples from the oracle (S03-WE13, S03-WE20, S03-WE21, S04-WE16, S04-WE18, S05-WE01,
S05-WE07, S05-WE17, S06-WE03, S10-WE02, S10-WE03) on the course's own data, the course sigma tables, and independent
scipy / statsmodels computations for random inputs. The calculators port the workbook's formulas, so the same
course pages and conventions apply (build/README.md).
"""
from __future__ import annotations

import csv
import functools
import importlib.util
import itertools
import json
import math
import random
import re
import shutil
import subprocess
import sys
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
import pytest
import statsmodels.api as sm
import statsmodels.formula.api as smf
from scipy import stats

from bbtools.constants import ROOT, USED_TABLE, load_all, load_average_range_table
from bbtools.printed import agrees_at_printed_precision

ASSETS = ROOT / "study" / "assets"
CONSTANTS_DIR = ROOT / "inventory" / "constants"
COURSE = ROOT / "source" / "course"
GRR_WORKBOOK = COURSE / "Les 5" / "20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx"
# The GRR workbook types the interaction SS rounded ('=412.5+296.667'); 1e-5 allows exactly that (test_sheet_grr.py).
TYPED_ROUNDING = 1e-5


@functools.lru_cache(maxsize=None)
def build_module() -> Any:
    """study/build_study.py as a module (it is a script, not part of the bbtools package)."""
    spec = importlib.util.spec_from_file_location("build_study", ROOT / "study" / "build_study.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module  # its dataclasses look their module up while being created
    spec.loader.exec_module(module)
    return module


def finite(value: Any) -> Any:
    """Undo the node side's encoding of ±Infinity and NaN as strings."""
    if isinstance(value, dict):
        return {k: finite(v) for k, v in value.items()}
    if isinstance(value, list):
        return [finite(v) for v in value]
    if value in ("Infinity", "-Infinity", "NaN"):
        return float(value.replace("Infinity", "inf"))
    return value


def run_js(calls: list[tuple[str, list[Any]]]) -> list[Any]:
    """Evaluate Calc.<name>(*args) or Stats.<name>(*args) in node, in order. "@K" / "@M" stand for the constant tables."""
    node = shutil.which("node")
    assert node, "node is needed to test the study guide's calculators"
    script = f"""
global.Stats = require({json.dumps(str(ASSETS / "stats.js"))});
const Calc = require({json.dumps(str(ASSETS / "calc.js"))});
const input = JSON.parse(require('fs').readFileSync(0, 'utf8'));
const T = Calc.tables(input.constants);
const out = input.calls.map(([name, args]) => {{
  const [lib, fn] = name.split('.');
  const real = args.map(a => a === '@K' ? T.K : a === '@M' ? T.M : a);
  return (lib === 'Stats' ? Stats : Calc)[fn](...real);
}});
process.stdout.write(JSON.stringify(out, (k, v) => typeof v === 'number' && !isFinite(v) ? String(v) : v));
"""
    payload = json.dumps({"constants": build_module().constants_data(), "calls": calls})
    result = subprocess.run([node, "-e", script], input=payload, capture_output=True, text=True, check=True)
    return finite(json.loads(result.stdout))


def course_rows(stem: str) -> list[dict[str, str]]:
    """Rows of a course table extracted to inventory/constants/<stem>.csv."""
    with (CONSTANTS_DIR / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


# ---------- distributions ----------

def test_stats_functions_match_scipy() -> None:
    # arrange -- random arguments over the ranges the calculators meet, tails included
    rng = random.Random(7)
    cases = []
    for _ in range(40):
        p = rng.choice([1e-7, 0.001, 0.005, 0.025, 0.05, 0.1, 0.5, 0.9, 0.95, 0.975, 0.995, rng.random()])
        df, d2 = rng.choice([1, 2, 4, 9, 14, 19, 30, 120]), rng.choice([1, 3, 9, 14, 40])
        z, t, x, f = rng.uniform(-7, 7), rng.uniform(-6, 6), rng.uniform(0.01, 50), rng.uniform(0.01, 12)
        n, k, q = rng.randint(1, 400), 0, rng.random()
        k = rng.randint(0, n)
        big_n = rng.randint(50, 3000)
        defect, sample = rng.randint(0, big_n), rng.randint(1, min(big_n, 250))
        lam = rng.uniform(0.2, 150)
        cases += [
            ("Stats.normCdf", [z], stats.norm.cdf(z)), ("Stats.normInv", [p], stats.norm.ppf(p)),
            ("Stats.tCdf", [t, df], stats.t.cdf(t, df)), ("Stats.tInv", [p, df], stats.t.ppf(p, df)),
            ("Stats.chi2Cdf", [x, df], stats.chi2.cdf(x, df)), ("Stats.chi2Sf", [x, df], stats.chi2.sf(x, df)),
            ("Stats.chi2Inv", [p, df], stats.chi2.ppf(p, df)), ("Stats.chi2Isf", [p, df], stats.chi2.isf(p, df)),
            ("Stats.fCdf", [f, df, d2], stats.f.cdf(f, df, d2)), ("Stats.fSf", [f, df, d2], stats.f.sf(f, df, d2)),
            ("Stats.fInv", [p, df, d2], stats.f.ppf(p, df, d2)), ("Stats.fIsf", [p, df, d2], stats.f.isf(p, df, d2)),
            ("Stats.binomCdf", [k, n, q], stats.binom.cdf(k, n, q)), ("Stats.binomPmf", [k, n, q], stats.binom.pmf(k, n, q)),
            ("Stats.hypergeomCdf", [min(k, sample), big_n, defect, sample], stats.hypergeom.cdf(min(k, sample), big_n, defect, sample)),
            ("Stats.poissonCdf", [k % 200, lam], stats.poisson.cdf(k % 200, lam)),
            ("Stats.betaInv", [p, 1 + k % 30, 1 + n % 40], stats.beta.ppf(p, 1 + k % 30, 1 + n % 40)),
        ]
    # act
    got = run_js([(name, args) for name, args, _ in cases])
    # assert -- double precision libraries; 1e-7 relative leaves room for bisection and continued fractions only
    for (name, args, expected), value in zip(cases, got):
        assert value == pytest.approx(float(expected), rel=1e-7, abs=1e-15), (name, args)


# ---------- normal distribution ----------

def test_normal_solver_recovers_every_unknown() -> None:
    # arrange -- µ, σ, x random; each run leaves one or two of them out
    rng = random.Random(3)
    calls, expected = [], []
    for _ in range(25):
        mu, sigma, x = rng.uniform(-50, 900), rng.uniform(0.01, 80), 0.0
        x = mu + rng.uniform(-3.5, 3.5) * sigma
        pl = float(stats.norm.cdf(x, mu, sigma))
        pr = float(stats.norm.sf(x, mu, sigma))
        calls += [("Calc.normalSolve", [{"mu": mu, "sigma": sigma, "x": x}]),
                  ("Calc.normalSolve", [{"mu": mu, "x": x, "pl": pl}]),
                  ("Calc.normalSolve", [{"sigma": sigma, "x": x, "pr": pr}]),
                  ("Calc.normalSolve", [{"mu": mu, "sigma": sigma, "pl": pl}])]
        expected += [("pl", pl), ("sigma", sigma), ("mu", mu), ("x", x)]
    # act
    got = run_js(calls)
    # assert
    for (key, value), result in zip(expected, got):
        assert result[key] == pytest.approx(value, rel=1e-7, abs=1e-9), key


def test_sigma_from_a_tail_fraction_exam_q6() -> None:
    # arrange -- example exam Q6: 2 in 40 below 720 with mean 820 (study/parts/13_numbers.py)
    expected_sigma = (720 - 820) / stats.norm.ppf(2 / 40)
    # act
    result = run_js([("Calc.normalSolve", [{"mu": 820, "x": 720, "pl": 2 / 40}])])[0]
    # assert
    assert result["sigma"] == pytest.approx(expected_sigma, rel=1e-10)
    assert result["solved"] == ["z", "sigma", "pr"]


def test_contradictory_inputs_are_reported() -> None:
    # arrange -- z = 1 cannot go with µ = 0, σ = 1, x = 2
    # act
    result = run_js([("Calc.normalSolve", [{"mu": 0, "sigma": 1, "x": 2, "z": 1}])])[0]
    # assert
    assert result["notes"] and result["z"] == 2


# ---------- sigma level ----------

@pytest.mark.parametrize("stem, level_column, dpmo_column", [
    ("S06_voc_vs_vop_sigma_capability_defects_per_million_opportunitie", "Sigma Capability", "Defects per Million Opportunities"),
    ("S07_table_1_2_the_sigma_scale", "Sigma", "Defects per Million"),
    ("S08_table_6_3_sigma_score_table_z_dpmo", "Z", "DPMO"),
])
def test_dpmo_with_shift_reproduces_the_course_sigma_tables(stem: str, level_column: str, dpmo_column: str) -> None:
    # arrange -- every course table pairs a sigma level with the one-tail DPMO beyond level − 1.5 (decision 3)
    rows = course_rows(stem)
    # act
    got = run_js([("Calc.dpmoFromSigma", [float(r[level_column])]) for r in rows])
    # assert -- at printed precision; deck p. 21 prints 2σ as 308,537 (308,537.5 rounds to 308,538: build/README.md)
    for row, result in zip(rows, got):
        printed = row[dpmo_column]
        if printed == "308,537":
            assert round(result["shifted"]) == 308538
            continue
        assert agrees_at_printed_precision(result["shifted"], printed, thousands=True), (stem, row[level_column])


def test_sigma_level_from_dpmo_matches_scipy() -> None:
    # arrange
    values = [3.4, 233, 6210, 66807, 308538, 500000]
    # act
    got = run_js([("Calc.sigmaFromDpmo", [v]) for v in values])
    # assert
    for dpmo, result in zip(values, got):
        assert result["z"] == pytest.approx(stats.norm.isf(dpmo / 1e6), rel=1e-9)
        assert result["level"] == pytest.approx(stats.norm.isf(dpmo / 1e6) + 1.5, rel=1e-9)


# ---------- means, proportions, variances ----------

def test_one_mean_ci_and_tests_match_scipy() -> None:
    # arrange
    rng = np.random.default_rng(11)
    xs = rng.normal(10, 2, 15)
    n, mean, s, sigma, mu0, alpha = len(xs), xs.mean(), xs.std(ddof=1), 1.8, 9.2, 0.05
    # act
    r = run_js([("Calc.oneMean", [n, mean, s, sigma, mu0, alpha])])[0]
    # assert
    t_lo, t_hi = stats.t.interval(1 - alpha, n - 1, loc=mean, scale=s / math.sqrt(n))
    assert r["ciT"]["two"] == pytest.approx([t_lo, t_hi], rel=1e-9)
    assert r["ciT"]["lower"][0] == pytest.approx(mean - stats.t.ppf(1 - alpha, n - 1) * s / math.sqrt(n), rel=1e-9)
    z_lo, z_hi = stats.norm.interval(1 - alpha, loc=mean, scale=sigma / math.sqrt(n))
    assert r["ciZ"]["two"] == pytest.approx([z_lo, z_hi], rel=1e-9)
    for key, alternative in (("ne", "two-sided"), ("gt", "greater"), ("lt", "less")):
        expected = stats.ttest_1samp(xs, mu0, alternative=alternative)
        assert r["testT"]["stat"] == pytest.approx(expected.statistic, rel=1e-9)
        assert r["testT"][key]["p"] == pytest.approx(expected.pvalue, rel=1e-7)


def test_unpaired_shoe_soles_s03_we13(oracle: dict[str, Any]) -> None:
    # arrange -- CI Further Reading p. 16-17: LD1 and LD2 of the unpaired experiment
    given, stated = oracle["S03-WE13"]["given"], oracle["S03-WE13"]["stated_answers"]
    a, b = np.array(given["LD1"], float), np.array(given["LD2"], float)
    # act
    r = run_js([("Calc.twoMeansPooled", [len(a), a.mean(), a.std(ddof=1), len(b), b.mean(), b.std(ddof=1), 0, 0.05])])[0]
    # assert -- printed '−5.6 ± 6.19', s_p 6.59; and scipy's pooled t-test
    assert agrees_at_printed_precision(r["sp"], stated["s_p"])
    assert agrees_at_printed_precision(r["diff"], "-5.6")
    assert agrees_at_printed_precision(r["ci"]["half2"], "6.19")
    assert r["test"]["ne"]["p"] == pytest.approx(stats.ttest_ind(a, b, equal_var=True).pvalue, rel=1e-7)


def test_paired_t_matches_scipy() -> None:
    # arrange
    rng = np.random.default_rng(5)
    a = rng.normal(30, 5, 12)
    b = a + rng.normal(1, 2, 12)
    d = a - b
    # act
    r = run_js([("Calc.paired", [len(d), d.mean(), d.std(ddof=1), 0, 0.05])])[0]
    # assert
    for key, alternative in (("ne", "two-sided"), ("gt", "greater"), ("lt", "less")):
        assert r["test"][key]["p"] == pytest.approx(stats.ttest_rel(a, b, alternative=alternative).pvalue, rel=1e-7)


@pytest.mark.parametrize("n, x", [(100, 4), (50, 0), (30, 30), (200, 37)])
def test_exact_proportion_interval_matches_clopper_pearson(n: int, x: int) -> None:
    # arrange -- CI Further Reading p. 20: exact interval = R binom.test (Clopper-Pearson)
    expected = stats.binomtest(x, n).proportion_ci(0.95, method="exact")
    # act
    r = run_js([("Calc.oneProportion", [n, x, None, 0.05])])[0]
    # assert
    assert r["exact"]["two"] == pytest.approx([expected.low, expected.high], rel=1e-8, abs=1e-12)


def test_variance_intervals_and_tests_match_scipy() -> None:
    # arrange
    n1, s1, n2, s2, sigma0, alpha = 20, 0.0125, 15, 0.009, 0.01, 0.05
    # act
    one, two = run_js([("Calc.oneVariance", [n1, s1, sigma0, alpha]), ("Calc.twoVariances", [n1, s1, n2, s2, alpha])])
    # assert -- χ² pivot (n − 1)s²/σ² and F pivot (s1²/σ1²)/(s2²/σ2²) ~ F(n1 − 1, n2 − 1)
    q = (n1 - 1) * s1 ** 2
    assert one["ci"]["two"] == pytest.approx([q / stats.chi2.isf(alpha / 2, n1 - 1), q / stats.chi2.ppf(alpha / 2, n1 - 1)], rel=1e-9)
    assert one["test"]["gt"]["p"] == pytest.approx(stats.chi2.sf(q / sigma0 ** 2, n1 - 1), rel=1e-9)
    f = s1 ** 2 / s2 ** 2
    assert two["ci"]["two"] == pytest.approx([f / stats.f.isf(alpha / 2, n1 - 1, n2 - 1), f / stats.f.ppf(alpha / 2, n1 - 1, n2 - 1)], rel=1e-9)
    assert two["ciInv"]["two"] == pytest.approx([1 / two["ci"]["two"][1], 1 / two["ci"]["two"][0]], rel=1e-12)
    assert two["test"]["gt"]["p"] == pytest.approx(stats.f.sf(f, n1 - 1, n2 - 1), rel=1e-9)


def test_discrete_and_continuous_distributions_match_scipy() -> None:
    # arrange
    calls = [("Calc.binomial", [100, 0.05, 5]), ("Calc.hypergeometric", [1000, 50, 100, 5]), ("Calc.poisson", [2.959, 7]),
             ("Calc.exponential", [175.3, 0.01]), ("Calc.uniform", [1, 2, 1.25])]
    # act
    binom, hyper, pois, expo, unif = run_js(calls)
    # assert
    assert (binom["eq"], binom["le"], binom["ge"]) == pytest.approx(
        (stats.binom.pmf(5, 100, 0.05), stats.binom.cdf(5, 100, 0.05), stats.binom.sf(4, 100, 0.05)), rel=1e-9)
    assert (hyper["mean"], hyper["variance"], hyper["le"]) == pytest.approx(
        (stats.hypergeom.mean(1000, 50, 100), stats.hypergeom.var(1000, 50, 100), stats.hypergeom.cdf(5, 1000, 50, 100)), rel=1e-9)
    assert pois["ge"] == pytest.approx(stats.poisson.sf(6, 2.959), rel=1e-9)
    assert expo["le"] == pytest.approx(stats.expon.cdf(0.01, scale=1 / 175.3), rel=1e-12)
    assert (unif["mean"], unif["variance"], unif["le"]) == pytest.approx((1.5, 1 / 12, 0.25))


def test_contingency_table_probabilities() -> None:
    # arrange
    table = np.array([[120, 30], [80, 270]], float)
    # act
    r = run_js([("Calc.contingency", [table.tolist()])])[0]
    # assert
    total = table.sum()
    assert np.allclose(r["joint"], table / total)
    assert np.allclose(r["colGivenRow"], table / table.sum(axis=1, keepdims=True))
    assert np.allclose(r["rowGivenCol"], table / table.sum(axis=0, keepdims=True))
    assert np.allclose(r["product"], np.outer(table.sum(axis=1), table.sum(axis=0)) / total ** 2)


# ---------- acceptance sampling ----------

def test_oc_values_of_the_course_workbook_s03_we20_s03_we21(oracle: dict[str, Any]) -> None:
    # arrange -- Testing of Hypotheses.xlsx OC sheets, Excel cached floats
    binomial, hyper = oracle["S03-WE20"]["stated_answers"], oracle["S03-WE21"]["stated_answers"]
    # act
    got = run_js([("Calc.samplingPoint", [100, 4, None, 0.02]), ("Calc.samplingPoint", [130, 5, None, 0.02]),
                  ("Calc.samplingPoint", [100, 4, 10000, 0.02]), ("Calc.samplingPoint", [5000, 111, 10000, 0.02])])
    # assert
    assert got[0]["ocBin"] == pytest.approx(float(binomial["P[d<=4] at pi=0.02 (n=100,c=4)"]), rel=1e-9)
    assert got[1]["ocBin"] == pytest.approx(float(binomial["P[d<=5] at pi=0.02 (n=130,c=5)"]), rel=1e-9)
    assert got[2]["ocHyp"] == pytest.approx(float(hyper["P[d<=4] at pi=0.02 (N=10000,n=100,c=4) hypergeometric"]), rel=1e-9)
    assert got[3]["ocHyp"] == pytest.approx(float(hyper["P[d<=111] at pi=0.02 (N=10000,n=5000,c=111)"]), rel=1e-9)


def test_aoql_of_plan_250_5_s04_we18(oracle: dict[str, Any]) -> None:
    # arrange -- Further Reading p. 10: 'approximately 1.3 %' is the approximation p·OC(p) (build/README.md)
    stated = oracle["S04-WE18"]["stated_answers"]
    # act
    r = run_js([("Calc.aoql", [250, 5, 1000])])[0]
    # assert -- the maximum lies at 1.7 % (the slide reads 'about 1.8 %' off its chart); exact formula gives 1.03 %
    assert agrees_at_printed_precision(r["approx"] * 100, "1.3") and stated["AOQL"] == "approximately 1.3%"
    assert r["pApprox"] == pytest.approx(0.017) and r["pExact"] == pytest.approx(0.017)
    assert agrees_at_printed_precision(r["exact"] * 100, "1.03")


def test_variables_plan_s04_we16(oracle: dict[str, Any]) -> None:
    # arrange -- Further Reading p. 9: AQL 0.5 %, LQL 3.5 %, α = β = 5 %
    stated = oracle["S04-WE16"]["stated_answers"]
    # act
    r = run_js([("Calc.variablesPlan", [0.005, 0.035, 0.05, 0.05])])[0]
    # assert
    assert agrees_at_printed_precision(r["k"], stated["k"])
    assert r["nUp"] == int(stated["n"])


# ---------- regression, ANOVA, factorial ----------

def test_regression_oxygen_purity_s05_we17(oracle: dict[str, Any]) -> None:
    # arrange -- Regression p. 2 Table 11-1 and the Minitab output of p. 22
    rows = course_rows("S05_table_11_1_oxygen_and_hydrocarbon_levels")
    xs, ys = [float(r["Hydrocarbon Level x (%)"]) for r in rows], [float(r["Purity y (%)"]) for r in rows]
    stated = oracle["S05-WE17"]["stated_answers"]
    # act
    r = run_js([("Calc.regression", [xs, ys, 1.0, 0, 0, 0.05])])[0]
    # assert -- printed values, then statsmodels for the intervals
    for key, name in (("b0", "beta0hat_minitab"), ("b1", "beta1hat_minitab"), ("seB0", "SE_beta0"), ("seB1", "SE_beta1"),
                      ("sigma", "S")):
        assert agrees_at_printed_precision(r[key], stated[name]), key
    assert agrees_at_printed_precision(r["F"], "128.86") and agrees_at_printed_precision(r["r2"] * 100, "87.7")
    fit = sm.OLS(ys, sm.add_constant(xs)).fit()
    frame = fit.get_prediction(np.array([[1.0, 1.0]])).summary_frame(alpha=0.05)
    assert r["ciMean"]["two"] == pytest.approx([frame["mean_ci_lower"][0], frame["mean_ci_upper"][0]], rel=1e-9)
    assert r["pi"]["two"] == pytest.approx([frame["obs_ci_lower"][0], frame["obs_ci_upper"][0]], rel=1e-9)
    assert r["ciB1"]["two"] == pytest.approx(list(fit.conf_int(0.05)[1]), rel=1e-9)


def test_one_way_anova_paper_strength_s05_we01(oracle: dict[str, Any]) -> None:
    # arrange -- DOE p. 4 Table 4.4, α = 0.01
    groups = [[float(r[f"Obs{i}"]) for i in range(1, 7)]
              for r in course_rows("S05_table_4_4_tensile_strength_of_paper_psi") if r["Obs1"]]
    stated = oracle["S05-WE01"]["stated_answers"]
    # act
    r = run_js([("Calc.anova1", [groups, 0.01])])[0]
    # assert
    for key, name in (("sst", "SS_T"), ("sstr", "SS_Treatments"), ("sse", "SS_E"), ("msTr", "MS_Treatments"), ("msE", "MS_E"),
                      ("F", "F0"), ("Fcrit", "F_crit(0.01,3,20)")):
        assert agrees_at_printed_precision(r[key], stated[name]), key
    assert r["p"] == pytest.approx(stats.f_oneway(*groups).pvalue, rel=1e-8)


def test_factorial_chemical_process_s05_we07(oracle: dict[str, Any]) -> None:
    # arrange -- DOE p. 59-61: 2^2, 3 replicates, runs in standard order (1), a, b, ab
    rows = {(r["A"], r["B"]): [float(r[f"Rep {i}"]) for i in ("I", "II", "III")]
            for r in course_rows("S05_chemical_process_example_treatment_combinations_and_replicat")}
    runs = [rows[(a, b)] for b in "-+" for a in "-+"]
    stated = oracle["S05-WE07"]["stated_answers"]
    # act
    r = run_js([("Calc.factorial", [2, runs, None, 0.05])])[0]
    # assert
    effects = {e["name"]: e for e in r["effects"]}
    for name in ("A", "B", "AB"):
        assert agrees_at_printed_precision(effects[name]["effect"], stated[name]), name
    assert agrees_at_printed_precision(effects["A"]["F"], "53.19") and agrees_at_printed_precision(r["Fmodel"], "24.82")
    assert agrees_at_printed_precision(r["sspe"], "31.33") and agrees_at_printed_precision(r["r2adj"], stated["Adj_R-Squared"])


def test_factorial_with_pooling_matches_statsmodels() -> None:
    # arrange -- 2^4 single replicate, three- and four-factor interactions pooled into the error (DOE p. 72, 77)
    rng = np.random.default_rng(2)
    k = 4
    design = pd.DataFrame([[1 if run >> j & 1 else -1 for j in range(k)] for run in range(2 ** k)], columns=list("ABCD"))
    design["y"] = 50 + 4 * design.A - 3 * design.C + 2 * design.A * design.C + rng.normal(0, 1, 2 ** k)
    # act
    r = run_js([("Calc.factorial", [k, [[y] for y in design.y], 3, 0.05])])[0]
    # assert -- effect = 2 × coded coefficient; the model with all main effects and two-factor interactions
    fit = smf.ols("y ~ (A + B + C + D) ** 2", design).fit()
    effects = {e["name"]: e for e in r["effects"]}
    for term, coefficient in fit.params.items():
        if term == "Intercept":
            continue
        name = "".join(sorted(term.replace(":", "")))
        assert effects[name]["effect"] == pytest.approx(2 * coefficient, rel=1e-9)
        assert effects[name]["p"] == pytest.approx(fit.pvalues[term], rel=1e-6)
    assert r["dfE"] == fit.df_resid and r["mse"] == pytest.approx(fit.mse_resid, rel=1e-9)


def aliases_by_brute_force(generators: dict[str, str]) -> tuple[set[str], int]:
    """Defining relation (words) and resolution of a fractional design, by multiplying generator words out."""
    words = [frozenset(left + right) for left, right in generators.items()]
    relation = set()
    for size in range(1, len(words) + 1):
        for subset in itertools.combinations(words, size):
            product: frozenset[str] = frozenset()
            for word in subset:
                product = product ^ word
            relation.add("".join(sorted(product)))
    return relation, min(len(w) for w in relation)


@pytest.mark.parametrize("text, generators", [
    ("D=ABC", {"D": "ABC"}), ("D=AB; E=AC", {"D": "AB", "E": "AC"}), ("E=ABC, F=BCD", {"E": "ABC", "F": "BCD"}),
    ("D=AB; E=AC; F=BC; G=ABC", {"D": "AB", "E": "AC", "F": "BC", "G": "ABC"}),
])
def test_alias_structure_matches_word_multiplication(text: str, generators: dict[str, str]) -> None:
    # arrange
    relation, resolution = aliases_by_brute_force(generators)
    # act
    r = run_js([("Calc.aliases", [text])])[0]
    # assert
    assert set(r["relation"]) == relation and r["resolution"] == resolution
    assert r["runs"] == 2 ** (len(set("".join(generators.values())) | set(generators)) - len(generators))


# ---------- capability and control charts ----------

def table_constant(symbol: str, n: int) -> float:
    """A printed constant from the table convention decision 4 assigns to it."""
    table = next(t for t in load_all() if t.source.key == USED_TABLE[symbol])
    return float(table.value(symbol, n))


def test_capability_rows_use_the_decision_4_constants() -> None:
    # arrange -- example exam Q3 numbers plus R̄, s̄ and an overall s (MR̄/1.128 is book-only, Deel 14)
    lsl, usl, mean, n = 1400.0, 1460.0, 1440.0, 5
    spreads = {"given": 10.0, "rbar_d2": 23.0 / table_constant("d2", n), "sbar": 9.4, "sbar_c4": 9.4 / table_constant("c4", n),
               "overall": 11.5}
    # act
    rows = run_js([("Calc.capability", [{"lsl": lsl, "usl": usl, "mean": mean, "sigma": 10.0, "rbar": 23.0, "sbar": 9.4, "n": n,
                                         "overall": 11.5}, "@K"])])[0]
    # assert
    assert [r["key"] for r in rows] == list(spreads)
    for row in rows:
        sigma = spreads[row["key"]]
        cpu, cpl = (usl - mean) / (3 * sigma), (mean - lsl) / (3 * sigma)
        out = stats.norm.sf((usl - mean) / sigma) + stats.norm.cdf((lsl - mean) / sigma)
        assert (row["sigma"], row["cp"], row["cpk"], row["out"]) == pytest.approx(((sigma, (usl - lsl) / (6 * sigma), min(cpu, cpl), out)), rel=1e-9)
    assert rows[0]["level"] == "just capable" and rows[0]["capable"] == "nee"


def oefening_2_data() -> list[list[float]]:
    """__Gegevens oefeningen.xlsx 'Gegevens oefening 2'!B5:F24: 20 subgroups of 5 net weights (S06-WE03 data_ref)."""
    ws = openpyxl.load_workbook(COURSE / "Les 4" / "__Gegevens oefeningen.xlsx", data_only=True)["Gegevens oefening 2"]
    return [[ws.cell(r, c).value for c in range(2, 7)] for r in range(5, 25)]


def test_xbar_r_chart_oefening_2_s06_we03(oracle: dict[str, Any]) -> None:
    # arrange -- SPC oefening 2 workbook (Excel cached floats, A2 0.577, D3 0, D4 2.115, d2 2.326)
    stated = oracle["S06-WE03"]["stated_answers"]
    # act
    r = run_js([("Calc.subgroupChart", [oefening_2_data(), "@K"])])[0]
    # assert
    assert r["xbarbar"] == pytest.approx(float(stated["Xbarbar (X_streep_streep)"]), rel=1e-12)
    assert r["rbar"] == pytest.approx(float(stated["Rbar (R_streep)"]), rel=1e-12)
    assert r["xR"] == pytest.approx([float(stated["LCL_xbar"]), r["xbarbar"], float(stated["UCL_xbar"])], rel=1e-12)
    assert r["R"] == pytest.approx([0, r["rbar"], float(stated["UCL_R"])], rel=1e-12)
    assert r["sigmaR"] == pytest.approx(float(stated["sigma_estimate_Rbar_over_d2"]), rel=1e-12)


def test_individuals_and_attribute_charts() -> None:
    # arrange
    xs = [10.0, 12.0, 11.0, 13.0, 12.0, 30.0]
    pairs = [[100, 5], [100, 3], [80, 14]]  # 14/80 = 0.175 lies above its UCL 0.169
    # act
    imr, pchart = run_js([("Calc.individuals", [xs, "@K"]), ("Calc.attributeChart", ["p", pairs])])
    # assert -- Dummies p. 249 and 254 with E2 (n = 2) from Six Sigma Demystified and D4 (n = 2) from Table 18
    mrbar = np.mean(np.abs(np.diff(xs)))
    assert imr["X"] == pytest.approx([np.mean(xs) - table_constant("E2", 2) * mrbar, np.mean(xs), np.mean(xs) + table_constant("E2", 2) * mrbar])
    assert imr["MR"][2] == pytest.approx(table_constant("D4", 2) * mrbar)
    assert imr["points"][-1]["fx"] == "boven UCL"
    pbar = 22 / 280
    assert pchart["rows"][2]["ucl"] == pytest.approx(pbar + 3 * math.sqrt(pbar * (1 - pbar) / 80))
    assert pchart["rows"][2]["flag"] == "boven UCL"


# ---------- Gage R&R ----------

def course_grr_rows() -> list[list[float]]:
    """measurements!D4:F13 (S10-WE02 data_ref) as k·r lines: operator A trial 1, A trial 2, B trial 1, …, 5 parts each."""
    ws = openpyxl.load_workbook(GRR_WORKBOOK, data_only=True)["measurements"]
    value = {(operator, (row - 4) % 2, (row - 4) // 2): ws.cell(row, 4 + operator).value
             for row in range(4, 14) for operator in range(3)}
    return [[value[(operator, trial, part)] for part in range(5)] for operator in range(3) for trial in range(2)]


def test_gage_rr_average_and_range_s10_we03(oracle: dict[str, Any]) -> None:
    # arrange -- 'ranges' sheet of the course workbook: d2 1.12838, d2* 1.91155 and 2.48124 (tabel MSA.pdf)
    answers = oracle["S10-WE03"]["stated_answers"]
    # act
    r = run_js([("Calc.grr", [course_grr_rows(), 3, 2, None, 0.05, "@M"])])[0]
    # assert -- Excel cached floats, same formulas and constants
    assert (r["d2"], r["d2k"], r["d2n"]) == (1.12838, 1.91155, 2.48124)
    for key, name in (("ev", "EV"), ("av", "AV"), ("pv", "PV"), ("tv", "TV"), ("pctTv", "%GRR (formula sqrt((EV^2+AV^2)/TV^2))")):
        assert r["ar"][key] == pytest.approx(float(answers[name]), rel=1e-9), key


def test_gage_rr_anova_s10_we02(oracle: dict[str, Any]) -> None:
    # arrange -- '2way anova' sheet, model without interaction
    answers = oracle["S10-WE02"]["stated_answers"]
    # act
    r = run_js([("Calc.grr", [course_grr_rows(), 3, 2, None, 0.05, "@M"])])[0]
    # assert
    for key, name in (("ev", "EV"), ("av", "AV"), ("pv", "PV"), ("tv", "TV"), ("pctTv", "%GRR (formula sqrt((EV^2+AV^2)/TV^2))")):
        assert r["anova"][key] == pytest.approx(float(answers[name]), rel=TYPED_ROUNDING), key
    assert r["anova"]["verdict"].startswith("niet aanvaardbaar")


def test_gage_rr_anova_tables_match_statsmodels() -> None:
    # arrange -- a random complete study, 3 operators × 2 trials × 6 parts
    rng = np.random.default_rng(9)
    rows = [[round(float(v), 3) for v in 20 + np.arange(6) + rng.normal(0, 0.5, 6) + 0.3 * op] for op in range(3) for _ in range(2)]
    frame = pd.DataFrame([{"op": i // 2, "part": j, "y": rows[i][j]} for i in range(6) for j in range(6)])
    # act
    r = run_js([("Calc.grr", [rows, 3, 2, 1.5, 0.05, "@M"])])[0]
    # assert
    additive = sm.stats.anova_lm(smf.ols("y ~ C(op) + C(part)", frame).fit())
    full = sm.stats.anova_lm(smf.ols("y ~ C(op) * C(part)", frame).fit())
    assert r["anova"]["msErr"] == pytest.approx(additive["mean_sq"]["Residual"], rel=1e-9)
    assert r["anova"]["pOp"] == pytest.approx(additive["PR(>F)"]["C(op)"], rel=1e-7)
    assert r["interaction"]["pInt"] == pytest.approx(full["PR(>F)"]["C(op):C(part)"], rel=1e-7)
    assert r["anova"]["pctTol"] == pytest.approx(6 * r["anova"]["grr"] / 1.5)


def test_embedded_constants_are_the_printed_tables() -> None:
    # arrange
    msa = load_average_range_table()
    # act -- the JSON the page embeds for the calculators
    data = build_module().constants_data()
    # assert -- exactly the printed numbers, from the table decision 4 assigns to each symbol
    assert data["msa"]["d2"] == [float(v) for v in msa.d2]
    assert data["msa"]["d2star"] == [[float(v) for v in row] for row in msa.d2_star]
    for symbol in ("A2", "D3", "D4", "d2", "c4", "A3", "B3", "B4", "E2"):
        for n in (2, 5, 10):
            assert data["chart"][symbol][str(n)] == table_constant(symbol, n), (symbol, n)


# ---------- confusion matrix ----------

def test_confusion_metrics_exam_q5() -> None:
    # arrange -- example exam Q5, model A (study/parts/13_numbers.py): train and test matrices
    train, test = [[480, 20], [15, 485]], [[180, 120], [110, 190]]
    # act
    a_train, a_test = run_js([("Calc.confusion", [train]), ("Calc.confusion", [test])])
    # assert
    assert a_train["accuracy"] == pytest.approx(965 / 1000) and a_test["accuracy"] == pytest.approx(370 / 600)
    assert a_test["c1"]["recall"] == pytest.approx(180 / 300) and a_test["c1"]["precision"] == pytest.approx(180 / 290)
    assert a_test["c2"]["f1"] == pytest.approx(2 * (190 / 300) * (190 / 310) / (190 / 300 + 190 / 310))


# ---------- Les 2 (Ottoy): sample size, tolerance, β/power, χ², rank tests, sampling methods and plans ----------

def test_sample_size_for_a_proportion_and_a_mean() -> None:
    # arrange -- CI p. 7, 10: worst case p = 0.5, FULL width 5 %; CI FR p. 3: σ = 5, full width 6.2 (ex-04-1, 04_numbers)
    # act
    prop, mean = run_js([("Calc.sampleSize", [0.05, 0.05, None, None]), ("Calc.sampleSize", [0.05, 6.2, None, 5])])
    # assert -- (2·z·√(p(1 − p))/W)² = 1536.58 -> 1537 (the slide's "about 1350" is the finite lot, guide d04-nauwkeurigheid)
    assert prop["nPropUp"] == 1537 and prop["nProp"] == pytest.approx((stats.norm.ppf(0.975) / 0.05) ** 2, rel=1e-12)
    assert mean["nMeanUp"] == 10


def test_tolerance_intervals_ci_fr_p22_23(oracle: dict[str, Any]) -> None:
    # arrange -- CI FR p. 22: n 20, Ȳ 5.732, σ 0.2, α 5 %, β 10 % -> LTL 5.40 (S03-WE16); p. 23: σ unknown, distribution-free
    stated = oracle["S03-WE16"]["stated_answers"]
    # act
    known, unknown, free = run_js([("Calc.tolerance", [20, 5.732, 0.2, 0.05, 0.10, True]),
                                   ("Calc.tolerance", [10, 5.8, 0.15, 0.05, 0.10, False]),
                                   ("Calc.toleranceFree", [0.05, 0.10, 72])])
    # assert -- LTL as printed; UTL = Ȳ + kσ (the printed 5.55 is the erratum CI FR p. 22, not a target); 04_numbers values
    assert agrees_at_printed_precision(known["ltl"], stated["situatie 2: LTL"])
    assert known["k1"] == pytest.approx(stats.norm.ppf(0.95) / math.sqrt(20) + stats.norm.ppf(0.90), rel=1e-12)
    assert unknown["k1"] == pytest.approx(2.26307, abs=5e-6) and unknown["ltl"] == pytest.approx(5.46054, abs=5e-6)
    assert free["nMin"] == 72 and free["confidence"] == pytest.approx(1 - 2 * 0.95 ** 72 + 0.9 ** 72, rel=1e-12)


def test_beta_and_power_of_the_z_test_th_fr_p7_9(oracle: dict[str, Any]) -> None:
    # arrange -- TH FR p. 7: µ0 1200, σ 300, n 100, α 5 %, true 1270 -> β 25 %; p. 9: α 1 %, true 1300 -> 16 % (n 100), 1 % (n 195)
    stated = oracle["S03-WE11"]["stated_answers"]
    # act
    a, b, c = run_js([("Calc.powerMean", [1200, 1270, 300, 100, 0.05, None]), ("Calc.powerMean", [1200, 1300, 300, 100, 0.01, 0.01]),
                      ("Calc.powerMean", [1200, 1300, 300, 195, 0.01, None])])
    # assert
    assert agrees_at_printed_precision(a["gt"]["beta"] * 100, "25")
    assert agrees_at_printed_precision(b["gt"]["beta"] * 100, stated["beta at n=100"].rstrip("%"))
    assert agrees_at_printed_precision(c["gt"]["beta"] * 100, stated["beta at n=195"].rstrip("%"))
    assert b["nOneSidedUp"] == 195 and agrees_at_printed_precision(b["gt"]["crit"][0], "1270")


def test_beta_of_the_z_test_for_a_proportion_ex_05_12() -> None:
    # arrange -- TH FR p. 22: π0 2 %, n 400, α 5 %, true 3 %; n for β < 10 % (05_numbers.proportion_test)
    # act
    r = run_js([("Calc.powerProportion", [0.02, 0.03, 400, 0.05, 0.10])])[0]
    # assert
    se0, se1 = math.sqrt(0.02 * 0.98 / 400), math.sqrt(0.03 * 0.97 / 400)
    assert r["gt"]["beta"] == pytest.approx(stats.norm.cdf((0.02 + stats.norm.ppf(0.95) * se0 - 0.03) / se1), rel=1e-10)
    assert r["nSearch"] == 2016


def test_chi_square_frequency_tests_match_scipy() -> None:
    # arrange -- TR p. 15-20 give no worked example: scipy on chosen counts
    observed, table, two = [18, 22, 30, 30], [[20, 30, 50], [30, 30, 40]], [[12, 8], [5, 15]]
    # act
    fit, tab, small = run_js([("Calc.chi2Fit", [observed, [0.25] * 4, 0, 0.05]), ("Calc.chi2Table", [table, 0.05]),
                              ("Calc.chi2Table", [two, 0.05])])
    # assert
    assert fit["p"] == pytest.approx(stats.chisquare(observed).pvalue, rel=1e-10)
    expected = stats.chi2_contingency(table, correction=False)
    assert tab["chi2"] == pytest.approx(expected.statistic, rel=1e-12) and tab["df"] == expected.dof
    assert small["yates"] == pytest.approx(stats.chi2_contingency(two, correction=True).statistic, rel=1e-12)


def test_rank_tests_with_the_course_formulas() -> None:
    # arrange -- shoe soles CI FR p. 17-18 (S03-WE13/14); TR p. 21-26 formulas without tie correction (05_numbers.recipes_t)
    ld1, ld2 = [29, 30, 34, 32, 32, 25, 25, 27, 36, 17], [42, 29, 36, 52, 28, 32, 30, 33, 28, 33]
    paired2 = [28, 30, 38, 35, 34, 30, 30, 30, 42, 23]
    ws = openpyxl.load_workbook(COURSE / "Les 2" / "20260529_ottoy_Testing of Hypotheses.xlsx", data_only=True)["example t-test"]
    series = [ws.cell(r, 2).value for r in range(3, 23)]
    # act
    w, t, runs = run_js([("Calc.rankSum", [ld1, ld2]), ("Calc.signedRank", [[a - b for a, b in zip(ld1, paired2)]]), ("Calc.runsTest", [series])])
    # assert
    assert w["W"] == 84.5 and w["sd"] == pytest.approx(math.sqrt(10 * 10 * 21 / 12)) and w["z"] == pytest.approx((84.5 - 105) / math.sqrt(175))
    assert (t["T"], t["n"]) == (1, 9) and t["z"] == pytest.approx((1 - 22.5) / math.sqrt(9 * 10 * 19 / 24))
    assert runs["R"] == 12 and runs["z"] == pytest.approx((12 - 11) / math.sqrt(19 / 4))


def test_stratified_sampling_variances_s04_we03_we04(oracle: dict[str, Any]) -> None:
    # arrange -- AS p. 17-19: W_A 0.6, π_A 3 %, π_B 0.5 %, n 100; means 10 and 15, σ_S² 4
    s3, s4 = oracle["S04-WE03"]["stated_answers"], oracle["S04-WE04"]["stated_answers"]
    # act
    p, m = run_js([("Calc.samplingVariance", [0.03, 0.005, 0.6, 100]), ("Calc.samplingMeans", [10, 15, 4, 100])])
    # assert
    # printed 0.000195 is the half-up rounding of the exact 0.0001945 (= (0.6·0.0291 + 0.4·0.004975)/100), which a float
    # stores as 0.000194499…; compare with the exact value instead
    assert s3["sigma^2[P_s] (pre-stratified estimator)"] == "0.000195" and p["strat"] == pytest.approx(0.0001945, rel=1e-12)
    assert agrees_at_printed_precision(p["srs"], s3["sigma^2[P] (SRS estimator)"])
    assert agrees_at_printed_precision(p["nAopt"], "78")
    assert m["srs"] == pytest.approx(float(s4["variance parameter used for SRS Xbar in the xlsm chart (=sigma^2[Xbar])"]))


def test_plan_design_inverse_oc_double_plan_sprt(oracle: dict[str, Any]) -> None:
    # arrange -- AS FR p. 4 (Peach (164, 2)), p. 5 (double plan), p. 7 (SPRT, S04-WE14), p. 17 (plan for 2 %/4 %)
    sprt = oracle["S04-WE14"]["stated_answers"]
    # act
    p5, search, double, sp = run_js([("Calc.inverseOC", [164, 2, 0.05]), ("Calc.planSearch", [0.02, 0.04, 0.05, 0.05]),
                                     ("Calc.doublePlan", [90, 2, 7, 90, 8, 0.03, None]), ("Calc.sprt", [0.01, 0.05, 0.05, 0.05, None, None])])
    # assert
    assert stats.binom.cdf(2, 164, p5) == pytest.approx(0.05, rel=1e-9)
    assert (search["n"], search["c"]) == (781, 22) and stats.binom.cdf(22, 780, 0.04) > 0.05
    acc1 = stats.binom.cdf(2, 90, 0.03)
    oc = acc1 + sum(stats.binom.pmf(j, 90, 0.03) * stats.binom.cdf(8 - j, 90, 0.03) for j in range(3, 7))
    assert double["oc"] == pytest.approx(oc, rel=1e-10) and double["acc1"] == pytest.approx(acc1, rel=1e-12)
    for key in ("h1", "h2", "s"):
        assert agrees_at_printed_precision(sp[key], sprt[key]), key


def test_variables_plan_with_given_n_s04_we17() -> None:
    # arrange -- AS.xlsm 'variable sampling plan' (cached cells): ξ C14, t I6, X̄ F4, s F5, Q F6
    ws = openpyxl.load_workbook(COURSE / "Les 2" / "20260529_ottoy_Acceptance Sampling.xlsm", data_only=True)["variable sampling plan"]
    # act
    r = run_js([("Calc.variablesGivenN", [0.02, 0.05, 10, ws["C14"].value, ws["F4"].value, ws["F5"].value, 2])])[0]
    # assert
    assert r["t"] == pytest.approx(ws["I6"].value, rel=1e-12) and r["Q"] == pytest.approx(ws["F6"].value, rel=1e-12)
    assert r["accept"] == "verwerpen" and ws["I9"].value == "NO"


def test_skip_lot_qualification_as_fr_p11() -> None:
    # arrange -- AS FR p. 11 notes: P_q 2.4 % at p = 1 %, 85 % at p = 0.1 %
    # act
    a, b = run_js([("Calc.skipLot", [0.01, 80, 8, 3, 2]), ("Calc.skipLot", [0.001, 80, 8, 3, 2])])
    # assert
    assert agrees_at_printed_precision(a["pq"] * 100, "2.4") and agrees_at_printed_precision(b["pq"] * 100, "85")


# ---------- descriptives, regression extras, multiple regression, two-way ANOVA ----------

def test_descriptives_village_incomes_and_anscombe() -> None:
    # arrange -- VV p. 100 (village incomes, printed summary) and VV p. 134 (Anscombe set 1: r 0.816)
    raw, summary = course_rows("S01_village_income_exercise_ruwe_data_per_geslacht_m_f_per_dorp"), \
        {r["Metric"]: r for r in course_rows("S01_village_income_exercise_samenvattende_statistieken_per_dorp")}
    villages = {v: [float(r[v]) for r in raw if r[v]] for v in ("Abora", "Bladir", "Curo")}
    ans = course_rows("S01_anscombe_s_quartet_dataset_1_ruwe_data")
    # act
    got = run_js([("Calc.descriptives", [xs]) for xs in villages.values()] +
                 [("Calc.correlation", [[float(r["X"]) for r in ans], [float(r["Y"]) for r in ans]])])
    # assert
    for (village, _), d in zip(villages.items(), got[:3]):
        assert agrees_at_printed_precision(d["mean"], summary["average income"][village])
        assert agrees_at_printed_precision(d["median"], summary["median income"][village])
    assert agrees_at_printed_precision(got[3]["r"], "0.816")


def test_regression_from_summaries_and_sums_of_squares() -> None:
    # arrange -- REG p. 18, 22, 52 (oxygen sums from 07_numbers.oxygen; oil ANOVA as printed)
    oil = {r["Variable"] if "Variable" in r else "": r for r in course_rows("S05_oil_consumption_example_excel_coefficients_table")}
    # act
    sums, oxygen, oilss = run_js([("Calc.regSums", [20, 23.92, 1843.21, 29.2892, 2214.6566, None]),
                                  ("Calc.regFromSS", [152.13, 21.25, 20, 1, 0.05]), ("Calc.regFromSS", [228014.6263, 8120.603016, 15, 2, 0.05])])
    # assert -- printed Minitab b1 14.947, b0 74.283, F 128.86, R-Sq 87.7 %; Excel 0.96561037, 0.95987877, 26.0137832, 168.471203
    assert agrees_at_printed_precision(sums["b1"], "14.947") and agrees_at_printed_precision(sums["b0"], "74.283")
    assert agrees_at_printed_precision(oxygen["F"], "128.86") and agrees_at_printed_precision(oxygen["r2"] * 100, "87.7")
    for key, printed in (("r2", "0.96561037"), ("r2adj", "0.95987877"), ("s", "26.0137832"), ("F", "168.471203")):
        assert agrees_at_printed_precision(oilss[key], printed), key
    assert oil  # the coefficient table is checked in the multiple-regression test


def test_regression_intervals_from_summaries_reg_p22() -> None:
    # arrange -- oxygen data (REG p. 20), Minitab at x0 = 1.00: CI (88.486, 89.975), PI (86.830, 91.632)
    rows = course_rows("S05_table_11_1_oxygen_and_hydrocarbon_levels")
    xs, ys = np.array([float(r["Hydrocarbon Level x (%)"]) for r in rows]), np.array([float(r["Purity y (%)"]) for r in rows])
    sxx = float(((xs - xs.mean()) ** 2).sum())
    b1 = float(((xs - xs.mean()) * (ys - ys.mean())).sum()) / sxx
    b0 = float(ys.mean() - b1 * xs.mean())
    mse = float(((ys - b0 - b1 * xs) ** 2).sum()) / 18
    # act
    r = run_js([("Calc.regTests", [b1, b0, mse, sxx, 20, float(xs.mean()), 1.0, 0.05, None, None])])[0]
    # assert
    for value, printed in zip(r["ciMean"]["two"] + r["pi"]["two"], ("88.486", "89.975", "86.830", "91.632")):
        assert agrees_at_printed_precision(value, printed)


def test_multiple_regression_oil_and_acetylene_reg_p52_p62() -> None:
    # arrange -- REG p. 51-52 (oil, Excel) and p. 58, 60, 62 (acetylene, centred second-order model, Minitab)
    oil = course_rows("S05_oil_consumption_example_data_for_15_houses")
    oil_rows = [[float(r["Oil (Gal)"]), float(r["Temp"]), float(r["Insulation"])] for r in oil]
    printed = course_rows("S05_oil_consumption_example_excel_coefficients_table")
    ace = course_rows("S05_table_6_9_the_acetylene_data")
    t = np.array([float(r["Temp T"]) for r in ace]) - 1212.5
    rr = np.array([float(r["Ratio R"]) for r in ace]) - 12.444
    ace_rows = [[float(r["Yield Y"]), a, b, a * b, a * a, b * b] for r, a, b in zip(ace, t, rr)]
    # act
    o, a = run_js([("Calc.multipleRegression", [oil_rows, None, 0.05]), ("Calc.multipleRegression", [ace_rows, None, 0.05])])
    # assert -- the Excel coefficient table as printed, and the Minitab second-order coefficients
    for coef, row in zip(o["coef"], printed):
        for key, column in (("b", "Coefficients"), ("se", "Standard Error"), ("t", "t Stat"), ("lo", "Lower 95%"), ("hi", "Upper 95%")):
            assert agrees_at_printed_precision(coef[key], row[column]), (row["Variable"], column)
    for coef, value in zip(a["coef"], ("36.4339", "0.130476", "0.48005", "-0.0073346", "0.00017820", "-0.02367")):
        assert agrees_at_printed_precision(coef["b"], value)
    assert agrees_at_printed_precision(a["s"], "1.066") and agrees_at_printed_precision(a["F"], "371.49")


def test_partial_f_acetylene_matches_statsmodels() -> None:
    # arrange -- REG p. 61: SS_E(RM) of the first-order model (p. 59), SS_E(FM) of the second-order model (p. 62)
    ace = course_rows("S05_table_6_9_the_acetylene_data")
    frame = pd.DataFrame({"y": [float(r["Yield Y"]) for r in ace], "t": [float(r["Temp T"]) - 1212.5 for r in ace],
                          "r": [float(r["Ratio R"]) - 12.444 for r in ace]})
    reduced, full = smf.ols("y ~ t + r", frame).fit(), smf.ols("y ~ t + r + t:r + I(t**2) + I(r**2)", frame).fit()
    # act
    r = run_js([("Calc.partialF", [reduced.ssr, full.ssr, 3, full.df_resid, 0.05])])[0]
    # assert
    f_value, p_value, _ = full.compare_f_test(reduced)
    assert r["F"] == pytest.approx(f_value, rel=1e-9) and r["p"] == pytest.approx(p_value, rel=1e-7)


def test_two_way_anova_grr_workbook_output(oracle: dict[str, Any]) -> None:
    # arrange -- GRR workbook '2way anova': Excel "Anova: Two-Factor With Replication" on measurements!D4:F13
    #            (rows = parts, 2 trials each = "Rows per sample" 2; columns = appraisers A, B, C)
    ws = openpyxl.load_workbook(GRR_WORKBOOK, data_only=True)["measurements"]
    lines = [[ws.cell(row, col).value for col in (4, 5, 6)] for row in range(4, 14)]
    table = oracle["S10-WE02"]["given"]["anova_table_with_interaction"]
    # act
    r = run_js([("Calc.anova2", [lines, 2, 0.05])])[0]
    # assert -- Sample, Columns, Interaction, Within against the course's cached output
    names = {"Sample": 0, "Columns": 1, "Interaction": 2, "Within": 3}
    for source, values in table.items():
        if source in names:
            row = r["table"][names[source]]
            assert row["ss"] == pytest.approx(float(values["SS"]), rel=1e-9), source
            if "P-value" in values:
                assert row["p"] == pytest.approx(float(values["P-value"]), rel=1e-6), source


def test_two_way_anova_without_replication_matches_statsmodels() -> None:
    # arrange -- DOE demo '2ANOVA' layout (r = 1)
    rng = np.random.default_rng(4)
    lines = (50 + rng.normal(0, 3, (5, 4)) + np.arange(4)).round(2).tolist()
    frame = pd.DataFrame([{"a": i, "b": j, "y": lines[i][j]} for i in range(5) for j in range(4)])
    # act
    r = run_js([("Calc.anova2", [lines, 1, 0.05])])[0]
    # assert
    expected = sm.stats.anova_lm(smf.ols("y ~ C(a) + C(b)", frame).fit())
    assert [row["p"] for row in r["table"][:2]] == pytest.approx(list(expected["PR(>F)"][:2]), rel=1e-7)


# ---------- SPC extras, MSA extras, Bayes, k-class confusion ----------

def test_limits_from_standard_values_spc_p64() -> None:
    # arrange -- SPC p. 64: µ 1.5, σ 0.15, n 5 -> UCL 1.7013, LCL 1.2987 (Table 7.2: µ ± A·σ, A(5) = 1.342)
    # act
    r = run_js([("Calc.standardLimits", [1.5, 0.15, 5, "@K"])])[0]
    # assert
    assert agrees_at_printed_precision(r["xA"][2], "1.7013") and agrees_at_printed_precision(r["xA"][0], "1.2987")
    assert r["R"] == pytest.approx([0, table_constant("d2", 5) * 0.15, table_constant("D2", 5) * 0.15])


def test_run_rules_on_the_spc_p72_chart() -> None:
    # arrange -- ex-10-4: z-values read from SPC p. 72 (10_numbers.py); rules of SPC p. 68 (UNVERIFIED: no course answer)
    z = [-1.05, 0.44, -1.53, -2.49, -2.01, -1.32, -0.44, 0.95, 1.92, 1.47, -0.59, -1.53, -1.08, -1.95, -0.56, 0.41, -1.07,
         2.46, -0.23, -0.59, -1.32, -1.63, -2.38, -2.77, -1.55]
    # act
    hits = run_js([("Calc.runRules", [z, 0, 1])])[0]["hits"]
    # assert
    assert hits["1"] == [] and hits["2"] == [5, 6, 24, 25] and hits["3"] == [5, 6, 7, 24, 25] and hits["4"] == []
    assert hits["5"] == [9, 23, 24]


def test_run_rules_literal_reading_counts_both_sides() -> None:
    # arrange -- SPC p. 68 rules 2 and 3 do not say "same side"; one point beyond +2σ and the next beyond -2σ
    z = [0.0, 2.5, -2.5, 0.0, 1.5, -1.5, 1.5, -1.5, 0.0]
    # act
    hits = run_js([("Calc.runRules", [z, 0, 1])])[0]["hits"]
    # assert -- same-side reading sees nothing, the literal reading flags both windows
    assert hits["2"] == [] and hits["2b"] == [3, 4]
    assert hits["3"] == [] and hits["3b"] == [6, 7, 8, 9]


def test_observed_cp_msa_p26() -> None:
    # arrange -- MSA p. 26 notes: actual Cp 2 with %GRR (tolerance) 10 % / 30 % -> 1.96 / "no more than 1.71"; 60 % is
    #            printed 1.20, the formula gives 1.28 (a printed error the guide text explains)
    # act
    got = run_js([("Calc.cpObserved", ["tol", g, 2, None]) for g in (0.10, 0.30, 0.60)])
    # assert
    assert agrees_at_printed_precision(got[0]["cpo"], "1.96") and agrees_at_printed_precision(got[1]["cpo"], "1.71")
    assert got[2]["cpo"] == pytest.approx(1 / math.sqrt(0.25 + 0.36))


def test_uncertainty_budget_steel_strip(oracle: dict[str, Any]) -> None:
    # arrange -- steel strip p. 1-2: certificate 0.1 % of 1834 at k = 2, resolution 0.5 uniform, squareness 0.917 uniform,
    #            repeatability 2/d2(3) as the mean of 3 readings; printed u's 0.917, 0.289, 0.529, 0.682
    sources = [{"kind": "cert", "a": 1.834, "b": 2}, {"kind": "uni", "a": 0.5}, {"kind": "uni", "a": 0.917},
               {"kind": "range", "a": 2, "b": 1.693, "c": 3}]
    # act
    r = run_js([("Calc.uncertaintyBudget", [sources, 2])])[0]
    # assert -- the u's as printed; u_c from them is 1.292 (printed 1.264: erratum, S10-WE01 left as stated)
    for row, printed in zip(r["rows"], ("0.917", "0.289", "0.529", "0.682")):
        assert agrees_at_printed_precision(row["u"], printed)
    assert r["uc"] == pytest.approx(math.sqrt(sum(row["u"] ** 2 for row in r["rows"])))
    assert "1.264" in oracle["S10-WE01"]["stated_answers"]["u_c"]


def test_bias_test_and_msa_table_lookup() -> None:
    # arrange -- 'Theoretical background of a GRR study' p. 1: g = 5, m = 3 -> d2* 1.73857, ν 9.3; bias test MSA p. 31-32
    data = build_module().constants_data()["msa"]
    readings = [2.70, 2.50, 2.40, 2.50, 2.70, 2.30, 2.50, 2.50, 2.40, 2.40, 2.60, 2.40]
    # act
    r = run_js([("Calc.biasTest", [float(np.mean(readings)), max(readings) - min(readings), 1, 12, 1.99, 0.05, "@M"])])[0]
    # assert
    assert data["d2star"][4][1] == 1.73857 and data["nu"][4][1] == 9.3
    sr = (max(readings) - min(readings)) / r["d2"]
    assert r["t"] == pytest.approx((np.mean(readings) - 1.99) / sr * math.sqrt(12) * r["d2s"] / r["d2"])
    assert r["nu"] == 9.0 and r["significant"]


def test_discrimination_rule_msa_p30() -> None:
    # arrange -- MSA p. 30 notes: σ = MU gives UCL − LCL ≈ 3.69 MU (n = 2) and 4.36 MU (n = 3)
    # act
    n2, n3 = run_js([("Calc.discrimination", [2, 1, None, "@K"]), ("Calc.discrimination", [3, 1, None, "@K"])])
    # assert
    assert agrees_at_printed_precision(n2["width"], "3.69") and agrees_at_printed_precision(n3["width"], "4.36")
    assert len(n2["values"]) == 4 and not n2["tooCoarse"]


def test_bayes_and_beta_posterior() -> None:
    # arrange -- Naert p. 24-25 table (P(L = 2 | Rejected) = 40/60); web slides p. 55 (Beta(1, 1) + HHTHT -> Beta(4, 3),
    #            mean 0.57); Naert notes L1 p. 4 (Beta(2, 8): mode 0.125, mean 0.200, median 0.180)
    # act
    b, post, prior = run_js([("Calc.bayes", [0.46, 40 / 230, 20 / 270]), ("Calc.betaPosterior", [1, 1, 3, 5]),
                             ("Calc.betaPosterior", [2, 8, None, None])])
    # assert
    assert b["pAgivenB"] == pytest.approx(40 / 60)
    assert (post["posterior"]["a"], post["posterior"]["b"]) == (4, 3) and agrees_at_printed_precision(post["posterior"]["mean"], "0.57")
    for key, printed in (("mode", "0.125"), ("mean", "0.200"), ("median", "0.180")):
        assert agrees_at_printed_precision(prior["prior"][key], printed), key


def test_k_class_confusion_matrix_ml_p30() -> None:
    # arrange -- ML p. 30 4 × 4 matrix, rows = Predicted (CSV as printed)
    rows = [[float(r[c]) for c in ("1", "2", "3", "4")] for r in course_rows("S11_confusion_matrix_example_4_x_4_figure_rows_predicted_columns")]
    m = np.array(rows).T  # rows = actual
    # act
    r = run_js([("Calc.confusionK", [rows, False])])[0]
    # assert
    assert r["accuracy"] == pytest.approx(np.trace(m) / m.sum())
    for i, c in enumerate(r["classes"]):
        assert c["recall"] == pytest.approx(m[i, i] / m[i].sum()) and c["precision"] == pytest.approx(m[i, i] / m[:, i].sum())


# ---------- solvers in every direction (user request 2026-10-07: tables merged, calculators work both ways) ----------

def test_sigma_solver_reproduces_the_course_sigma_table_both_ways() -> None:
    # arrange -- SPC p. 21: level -> DPMO and the printed DPMO back to its level (decision 3, 1.5σ shift)
    rows = course_rows("S06_voc_vs_vop_sigma_capability_defects_per_million_opportunitie")
    calls = [("Calc.sigmaSolve", [{"level": float(r["Sigma Capability"])}]) for r in rows]
    calls += [("Calc.sigmaSolve", [{"dpmo": float(r["Defects per Million Opportunities"].replace(",", ""))}]) for r in rows]
    # act
    got = run_js(calls)
    forward, backward = got[:len(rows)], got[len(rows):]
    # assert -- at printed precision (2σ: 308,537.5 prints as 308,537 on SPC p. 21, build/README.md); back within 0.01σ
    for row, ahead, back in zip(rows, forward, backward):
        printed = row["Defects per Million Opportunities"]
        if printed != "308,537":
            assert agrees_at_printed_precision(ahead["dpmo"], printed, thousands=True), row
        assert back["level"] == pytest.approx(float(row["Sigma Capability"]), abs=0.01), row
        assert back["z"] == pytest.approx(stats.norm.isf(float(printed.replace(",", "")) / 1e6), rel=1e-9)


def test_sigma_solver_finds_the_missing_count() -> None:
    # arrange -- D, N, O give DPO; DPMO with two of D, N, O gives the third
    # act
    full, no_n, no_o = run_js([("Calc.sigmaSolve", [{"D": 15, "N": 500, "O": 6}]),
                               ("Calc.sigmaSolve", [{"D": 15, "O": 6, "dpmo": 5000}]),
                               ("Calc.sigmaSolve", [{"D": 15, "N": 500, "dpmo": 5000}])])
    # assert
    assert full["dpo"] == pytest.approx(0.005) and full["level"] == pytest.approx(stats.norm.isf(0.005) + 1.5)
    assert no_n["N"] == pytest.approx(500) and no_o["O"] == pytest.approx(6)


def test_normal_interval_every_direction_matches_scipy() -> None:
    # arrange -- µ ± kσ, a symmetric interval with its coverage, and σ from µ, one bound and the fraction outside
    mu, sigma, k = 10.0, 2.0, 1.7
    inside = float(stats.norm.cdf(k) - stats.norm.cdf(-k))
    # act
    by_k, by_bounds, by_one_bound = run_js([
        ("Calc.normalInterval", [{"mu": mu, "sigma": sigma, "k": k}]),
        ("Calc.normalInterval", [{"a": mu - k * sigma, "b": mu + k * sigma, "inside": inside}]),
        ("Calc.normalInterval", [{"mu": mu, "a": mu - k * sigma, "outside": 1 - inside}])])
    # assert
    assert by_k["inside"] == pytest.approx(inside, rel=1e-10) and by_k["b"] == pytest.approx(mu + k * sigma)
    assert by_bounds["mu"] == pytest.approx(mu) and by_bounds["sigma"] == pytest.approx(sigma, rel=1e-8)
    assert by_one_bound["sigma"] == pytest.approx(sigma, rel=1e-8) and by_one_bound["b"] == pytest.approx(mu + k * sigma, rel=1e-8)


def test_capability_solver_every_direction_matches_scipy() -> None:
    # arrange -- random processes: the forward values from scipy, then each solver direction recovers the rest
    rng = random.Random(5)
    calls, cases = [], []
    for _ in range(10):
        lsl, usl = 100.0, 100.0 + rng.uniform(5, 40)
        sigma, mean = rng.uniform(0.6, 3.0), rng.uniform(lsl + 0.3 * (usl - lsl), usl - 0.3 * (usl - lsl))
        cp, cpu, cpl = (usl - lsl) / (6 * sigma), (usl - mean) / (3 * sigma), (mean - lsl) / (3 * sigma)
        ppm = (stats.norm.sf(3 * cpu) + stats.norm.sf(3 * cpl)) * 1e6
        cases.append((sigma, mean, cp, min(cpu, cpl), ppm))
        calls += [("Calc.capabilitySolve", [{"lsl": lsl, "usl": usl, "mean": mean, "sigma": sigma}]),
                  ("Calc.capabilitySolve", [{"lsl": lsl, "usl": usl, "cp": cp, "cpk": min(cpu, cpl)}]),
                  ("Calc.capabilitySolve", [{"lsl": lsl, "usl": usl, "cp": cp, "ppm": ppm}]),
                  ("Calc.capabilitySolve", [{"usl": usl, "ppm": stats.norm.sf(3 * cpu) * 1e6}])]
    # act
    got = run_js(calls)
    # assert
    for i, (sigma, mean, cp, cpk, ppm) in enumerate(cases):
        forward, from_indices, from_ppm, one_sided = got[4 * i:4 * i + 4]
        assert forward["cp"] == pytest.approx(cp) and forward["cpk"] == pytest.approx(cpk) and forward["ppm"] == pytest.approx(ppm, rel=1e-9)
        assert from_indices["ppm"] == pytest.approx(ppm, rel=1e-9) and from_indices["sigma"] == pytest.approx(sigma)
        assert mean == pytest.approx(min(from_indices["meanOptions"], key=lambda m: abs(m - mean)))
        assert from_ppm["cpk"] == pytest.approx(cpk, rel=1e-7)
        assert one_sided["cpk"] == pytest.approx((one_sided["usl"] - mean) / (3 * sigma), rel=1e-7)


def test_capability_solver_centred_cp_2_is_2_per_billion_spc_p40() -> None:
    # arrange -- SPC p. 40: Cp = 2, centred: 0.002 ppm (study/parts/09_numbers.py); and back from the ppm
    # act
    ahead, back = run_js([("Calc.capabilitySolve", [{"lsl": 0, "usl": 12, "cp": 2, "centred": True}]),
                          ("Calc.capabilitySolve", [{"lsl": 0, "usl": 12, "ppm": 2 * stats.norm.sf(6) * 1e6, "centred": True}])])
    # assert
    assert agrees_at_printed_precision(ahead["ppm"], "0.002") and ahead["mean"] == 6
    assert back["cp"] == pytest.approx(2, rel=1e-9) and back["sigma"] == pytest.approx(1, rel=1e-9)


def test_limits_inverse_recovers_the_summary_of_the_chart() -> None:
    # arrange -- SPC p. 74 forwards with the decision 4 constants, then backwards from the limits
    n, xbb, rbar, sbar = 5, 28.46, 3.06, 1.30
    forward = run_js([("Calc.limitsSummary", [n, xbb, rbar, sbar, "@K"])])[0]
    # act
    from_r = run_js([("Calc.limitsInverse", [n, {"ucl": forward["xR"][2], "lcl": forward["xR"][0], "uclR": forward["R"][2]}, "@K"])])[0]
    from_s = run_js([("Calc.limitsInverse", [n, {"ucl": forward["xS"][2], "cl": xbb, "uclS": forward["S"][2]}, "@K"])])[0]
    # assert
    assert from_r["xbb"] == pytest.approx(xbb) and from_r["rbarFromX"] == pytest.approx(rbar) and from_r["rbarFromR"] == pytest.approx(rbar)
    assert from_r["sigmaR"] == pytest.approx(rbar / table_constant("d2", n))
    assert from_s["sbarFromX"] == pytest.approx(sbar) and from_s["sbarFromS"] == pytest.approx(sbar)


def test_sample_size_solver_every_direction() -> None:
    # arrange -- CI p. 7, 10: n 1537 for a full width of 5 % at 95 %; then width and confidence back from n
    z = stats.norm.ppf(0.975)
    # act
    n, width, confidence = run_js([("Calc.sampleSizeSolve", [0.05, 0.05, None, None, None]),
                                   ("Calc.sampleSizeSolve", [0.05, None, 1537, None, None]),
                                   ("Calc.sampleSizeSolve", [None, 0.05, 1537, None, 5])])
    # assert
    assert n["prop"]["nUp"] == 1537
    assert width["prop"]["width"] == pytest.approx(2 * z * math.sqrt(0.25 / 1537), rel=1e-12)
    assert confidence["prop"]["alpha"] == pytest.approx(2 * stats.norm.sf(0.05 * math.sqrt(1537) / (2 * 0.5)), rel=1e-10)
    assert confidence["mean"]["alpha"] == pytest.approx(2 * stats.norm.sf(0.05 * math.sqrt(1537) / (2 * 5)), rel=1e-10)


def test_detectable_shift_th_fr_p9() -> None:
    # arrange -- TH FR p. 9: σ 300, α 1 %, β 1 %: n = 195 detects the true mean 1300 against 1200 (oracle S03-WE11)
    # act
    shift = run_js([("Calc.detectableShift", [300, 195, 0.01, 0.01])])[0]
    # assert
    assert shift == pytest.approx(2 * stats.norm.ppf(0.99) * 300 / math.sqrt(195), rel=1e-12)
    assert 1200 + shift == pytest.approx(1300, abs=0.5)


def test_discrete_and_exponential_inverses_match_scipy() -> None:
    # arrange
    cases = [("Calc.binomialInv", [20, 0.1, 0.95], stats.binom.ppf(0.95, 20, 0.1)),
             ("Calc.binomialInv", [130, 0.02, 0.5], stats.binom.ppf(0.5, 130, 0.02)),
             ("Calc.poissonInv", [2.959, 0.9], stats.poisson.ppf(0.9, 2.959)),
             ("Calc.hypergeometricInv", [1000, 30, 80, 0.99], stats.hypergeom.ppf(0.99, 1000, 30, 80))]
    # act
    got = run_js([(name, args) for name, args, _ in cases] + [("Calc.exponentialSolve", [None, 10, 0.5]),
                                                               ("Calc.exponentialSolve", [0.2, None, 0.5])])
    # assert
    for (name, args, expected), value in zip(cases, got):
        assert value == expected, (name, args)
    assert got[-2]["rate"] == pytest.approx(math.log(2) / 10) and got[-1]["t"] == pytest.approx(math.log(2) / 0.2)


def test_merged_tables_keep_every_printed_number() -> None:
    # arrange -- the four sigma tables and the five constant tables became one table each (build_study.py)
    module = build_module()
    sigma_html, constants_html = module.sigma_table(), module.chart_constants()
    # act
    sigma_dpmo = [(stem, row[dpmo]) for stem, _, dpmo, _, _ in module.SIGMA_TABLES for row in course_rows(stem)]
    used = {(symbol, row[0].strip()): row[column].strip() for table in load_all() if table.source.key in set(USED_TABLE.values())
            for symbol, column in table.symbol_columns().items() if USED_TABLE.get(symbol) == table.source.key
            for row in table.rows if row[0].strip().isdigit() and row[column].strip()}
    # assert
    for stem, printed in sigma_dpmo:
        assert re.search(rf"(?<![\d,.]){re.escape(printed)}(?![\d,.])", sigma_html), (stem, printed)
    for (symbol, n), printed in used.items():
        assert f"<td>{printed}</td>" in constants_html, (symbol, n, printed)


# ---------- formulas of the calculators (study/tool_formulas.html) ----------

def calculator_blocks() -> list[str]:
    """'tool.index' of every calculator block in study/assets/tools.js, read by running it in node with a stub page."""
    script = f"""
global.Stats = require({json.dumps(str(ASSETS / "stats.js"))});
global.Calc = require({json.dumps(str(ASSETS / "calc.js"))});
global.document = {{ getElementById: () => ({{ textContent: JSON.stringify({{ chart: {{}}, msa: {{ m: [], g: [], d2: [], d2star: [], nu: [] }} }}) }}),
                    querySelectorAll: () => [] }};
global.window = {{ addEventListener() {{}} }};
global.location = {{ hash: '' }};
const source = require('fs').readFileSync({json.dumps(str(ASSETS / "tools.js"))}, 'utf8').replace('var TOOLS = {{}};', 'var TOOLS = global.TOOLS = {{}};');
eval(source);
process.stdout.write(JSON.stringify(Object.entries(global.TOOLS).flatMap(([name, blocks]) => blocks.map((_, i) => name + '.' + i))));
"""
    result = subprocess.run([shutil.which("node"), "-e", script], capture_output=True, text=True, check=True)
    return json.loads(result.stdout)


def test_every_calculator_block_has_its_formulas_and_every_symbol_a_meaning() -> None:
    # arrange
    module = build_module()
    sections = [f"{tool}.{index}" for tool, index, _ in module.formula_blocks()[0]]
    # act
    blocks = calculator_blocks()
    templates, rows, problems = module.formula_blocks()
    # assert -- one section per block and back; every \sym defined, with a meaning and a 'hoe'
    assert sorted(blocks) == sorted(sections), (set(blocks) ^ set(sections))
    assert not problems, problems
    assert all(row["betekenis"] and row["hoe"] for _, row in rows)
