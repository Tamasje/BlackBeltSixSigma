"""Numbers for Deel 11 (Meetsysteemanalyse / Gage R&R, Ottoy, Les 5) of the study guide.

Prints every value that study/parts/11_msa_grr.html marks "zelf berekend", every exercise answer
(data-answer), and re-checks every printed course result the text quotes (OK / MISMATCH lines).

Data sources (read-only):
- source/course/Les 5/20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx (sheets measurements,
  2way anova, ranges): the GRR study data and the lecturer's solution.
- source/course/Les 5/20260619_ottoy_Rheostat Knob Data.xls (xlrd): original and rounded data and the
  lecturer's solved X-bar/R charts.
- source/course/Les 5/20260619_ottoy_linearity.txt: the linearity exercise (MSA p. 33).
- inventory/constants: tabel MSA (d2*, nu, d2; tabel MSA.pdf p. 1), Table 18 and Table A
  (Les 4 ___4.1 tabellen SPC.pdf p. 1-2).
- Numbers transcribed from a page carry that page in a comment.
Nothing is taken from memory. Run from anywhere:  python3 study/parts/11_numbers.py
"""
from __future__ import annotations

import csv
import math
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.api as sm
import statsmodels.formula.api as smf
import xlrd
from scipy import stats

ROOT = Path(__file__).resolve().parents[2]
LES5 = ROOT / "source" / "course" / "Les 5"
CONSTANTS = ROOT / "inventory" / "constants"
GRR_XLSX = LES5 / "20260619_ottoy_GRR - ANOVA - avegage and range - 2.xlsx"
RHEOSTAT_XLS = LES5 / "20260619_ottoy_Rheostat Knob Data.xls"
LINEARITY_TXT = LES5 / "20260619_ottoy_linearity.txt"


# ----------------------------------------------------------------------------------------------- helpers
def show(label: str, value: float | str, note: str = "") -> None:
    """Print one result line (6 significant digits for numbers)."""
    text = f"{value:.6g}" if isinstance(value, (float, np.floating, int)) else str(value)
    print(f"  {label:<58} {text:>16}  {note}")


def check(label: str, computed: float, printed: float, tol: float) -> None:
    """Compare a computed value with a printed course value; tol is the printed precision."""
    status = "OK" if abs(computed - printed) <= tol else "MISMATCH"
    print(f"  {status:<8} {label:<52} computed {computed:.6g}  printed {printed:.6g}  (tol {tol:g})")


def header(title: str) -> None:
    """Section header."""
    print()
    print("=" * 100)
    print(title)
    print("=" * 100)


def csv_rows(stem: str) -> list[dict[str, str]]:
    """Rows of a course table in inventory/constants."""
    with (CONSTANTS / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def msa_table() -> tuple[dict[tuple[int, int], tuple[float, float]], dict[int, float]]:
    """tabel MSA.pdf p. 1: {(g, m): (nu, d2*)} and {m: d2 (g -> infinity)}."""
    rows = csv_rows("S10_values_associated_with_the_distribution_of_the_average_range")
    cells: dict[tuple[int, int], tuple[float, float]] = {}
    d2_inf: dict[int, float] = {}
    for row in rows:
        key = row["g (rows) \\ m (subgroup size, columns)"]
        for m in range(2, 21):
            nu_txt, val_txt = row[str(m)].split("/")
            if key.startswith("d2"):
                d2_inf[m] = float(nu_txt)  # last row: "d2 / cd"
            else:
                cells[(int(key), m)] = (float(nu_txt), float(val_txt))
    return cells, d2_inf


def table18() -> dict[int, dict[str, float]]:
    """Table 18 (Les 4 ___4.1 tabellen SPC.pdf p. 2), keyed by subgroup size n."""
    out = {}
    for row in csv_rows("S06_table_18_factors_for_computing_control_chart_lines"):
        if not row["Sample size n"].isdigit():  # the ">25" row holds formulas, not numbers
            continue
        out[int(row["Sample size n"])] = {k: float(v) for k, v in row.items()
                                          if k not in ("Sample size n", "source_file", "source_page", "table_name")}
    return out


def table_a() -> dict[int, dict[str, float]]:
    """Table A (Les 4 ___4.1 tabellen SPC.pdf p. 1): d2, c4, d3 by subgroup size."""
    out = {}
    for row in csv_rows("S06_table_a_bias_correction_factors_for_estimating_standard_devi"):
        out[int(row["Subgroup Size"])] = {"d2": float(row["d2"]), "d3": float(row["d3"])}
    return out


MSA, D2_INF = msa_table()
T18 = table18()
TA = table_a()


# ----------------------------------------------------------------------------------------- 1. meter unit
def meter_unit() -> None:
    """MSA - meter unit.pdf p. 1-4 and MSA p. 6-7: small checks on the length standard."""
    header("1. Meter standard (MSA - meter unit.pdf p. 1-4; MSA p. 6-7)")
    g, period = 9.81, 2.0  # meter unit p. 1 footnote 2: g about 9,81 m/s2, period of a second pendulum 2 s
    length = g * (period / (2 * math.pi)) ** 2
    show("ex-11-1 q1: second pendulum l = g (T / 2 pi)^2 [mm]", length * 1000)
    check("pendulum length 994 mm (meter unit p. 1)", length * 1000, 994, 0.5)
    kr = 1 / 1_650_763.73  # meter unit p. 3: 1 m = 1 650 763,73 wavelengths of 86Kr
    show("ex-11-1 q2: krypton wavelength 1 m / 1 650 763,73 [nm]", kr * 1e9)
    check("krypton line 605,78 nm (meter unit p. 3)", kr * 1e9, 605.78, 0.005)
    c, nu_cs = 299_792_458, 9_192_631_770  # MSA p. 7; meter unit p. 4
    lam_cs = c / nu_cs
    show("caesium wavelength c / dnu_Cs [m]", lam_cs)
    show("ex-11-1 q3: 1 m = (9 192 631 770 / 299 792 458) lambda_Cs", nu_cs / c)
    sw, am = 29.69, 28.31  # MSA p. 6 notes: Swedish foot 29,69 cm, Amsterdam foot 28,31 cm
    show("ex-11-1 q4: Vasa: Swedish - Amsterdam foot [cm]", sw - am)
    show("Vasa: difference relative to the Swedish foot [%]", 100 * (sw - am) / sw)
    show("meter unit p. 1 fn 3: polar circumference 4 x 10 000 km [km]", 4 * 1e7 / 1000)


# --------------------------------------------------------------------------------------- 2. discrimination
def discrimination() -> None:
    """MSA p. 10: rule of tens with the vernier caliper readability quoted in the notes."""
    header("2. Discrimination (MSA p. 5, 10)")
    readability = 0.02  # MSA p. 10 notes: vernier caliper readability 0,02 mm
    show("ex-11-2 q1: smallest range for which 0,02 mm is 1/10 [mm]", 10 * readability)
    show("ex-11-2 q2: kitchen scale 0,1 g -> smallest range [g]", 10 * 0.1)  # MSA p. 10: 1/10 of a gram


# ------------------------------------------------------------------------------- 3. effective resolution
def effective_resolution() -> None:
    """MSA p. 21 question: how many measurements to detect 0,1 mm when sigma = 1 mm.

    The course gives no criterion. Illustration with the course's own uncertainty rules (MSA p. 28-29):
    the expanded uncertainty of an average of n readings is U = k sigma / sqrt(n) with k = 2; ask U <= 0,1 mm.
    """
    header("3. Effective resolution (MSA p. 21; illustration with MSA p. 28-29)")
    sigma, delta, k = 1.0, 0.1, 2.0  # MSA p. 21 notes; k = 2 MSA p. 28 notes
    n = (k * sigma / delta) ** 2
    show("ex-11-5: n with k sigma / sqrt(n) <= delta", n)
    show("check: U for n = 400 [mm]", k * sigma / math.sqrt(400))


# --------------------------------------------------------------------------------- 4. probable error (PE)
def probable_error() -> None:
    """MSA p. 23: PE = 75 % quantile - mean; lower limits of the odds P(G|X)/P(B|X) (exercise in the notes).

    Assumptions (stated in the guide): no bias, normal measurement error with sd sigma, P(G) = P(B), tolerance
    much wider than sigma (only the nearest spec limit matters). Specs tightened/widened by c PE.
    Lower limit of the odds = Phi((delta - c PE)/sigma) / (1 - Phi((delta + c PE)/sigma)); the same for
    P(B|Y)/P(G|Y) by symmetry.
    """
    header("4. Probable error and decision odds (MSA p. 23)")
    pe = stats.norm.ppf(0.75)
    show("PE / sigma = z(0,75)", pe)
    check("PE about 0,67 sigma (MSA p. 23)", pe, 0.67, 0.005)
    show("share of measurements within mu +/- PE", stats.norm.cdf(pe) - stats.norm.cdf(-pe))

    def odds(delta: float, c: float) -> float:
        return stats.norm.cdf(delta - c * pe) / stats.norm.sf(delta + c * pe)

    for c in (1, 2):
        for mult in (0.0, 1.0, 2.0):
            show(f"lower limit odds, specs adjusted by {c} PE, delta = {mult:g} PE", odds(mult * pe, c))
    show("ex-11-6 q1: odds limit, PE specs, delta = PE", odds(pe, 1))
    show("ex-11-6 q2: odds limit, 2PE specs, delta = PE", odds(pe, 2))
    show("intermediate: 1 - Phi(2 PE), Phi(-PE), 1 - Phi(3 PE)",
         f"{stats.norm.sf(2 * pe):.4f} {stats.norm.cdf(-pe):.4f} {stats.norm.sf(3 * pe):.4f}")


# ------------------------------------------------------------------------------- 5. process decisions
def process_decisions() -> None:
    """MSA p. 24-26: observed vs actual Cp."""
    header("5. Effect on process decisions (MSA p. 24-26)")

    def cpo_process(cpa: float, grr: float) -> float:
        """%GRR as a fraction of observed variation: Cpo = Cpa sqrt(1 - %GRR^2) (MSA p. 24)."""
        return cpa * math.sqrt(1 - grr ** 2)

    def cpo_tolerance(cpa: float, grr: float) -> float:
        """%GRR as a fraction of the tolerance: 1/Cpo^2 = 1/Cpa^2 + %GRR^2 (MSA p. 24-25)."""
        return 1 / math.sqrt(1 / cpa ** 2 + grr ** 2)

    # exercise MSA p. 24: %GRR 50 %, observed Cp at least 1,33
    target, grr = 1.33, 0.50
    cpa_p = target / math.sqrt(1 - grr ** 2)
    cpa_t = 1 / math.sqrt(1 / target ** 2 - grr ** 2)
    show("ex-11-7 q1: Cpa needed, %GRR of process variation", cpa_p)
    show("ex-11-7 q2: Cpa needed, %GRR of tolerance", cpa_t)
    show("intermediate: 1/1,33^2 and 1/1,33^2 - 0,25", f"{1 / target ** 2:.4f} {1 / target ** 2 - grr ** 2:.4f}")
    show("check q1: Cpo at that Cpa", cpo_process(cpa_p, grr))
    show("check q2: Cpo at that Cpa (implicit p. 24 form)", cpa_t * math.sqrt(1 - (grr * cpo_tolerance(cpa_t, grr)) ** 2))
    show("tolerance basis: largest possible Cpo = 1/%GRR", 1 / grr)
    # p. 24 notes: Cp 1,33 centred -> not more than 60 ppm
    show("defects, centred Cp 1,33: 2 Phi(-3 x 1,33) [ppm]", 2e6 * stats.norm.cdf(-3 * 1.33))
    show("defects, centred Cp 4/3: 2 Phi(-4) [ppm]", 2e6 * stats.norm.cdf(-4))
    check("'not more than 60 ppm' (MSA p. 24)", 2e6 * stats.norm.cdf(-3 * 1.33), 60, 0.5)
    # examples MSA p. 26 (tolerance basis, actual Cp = 2)
    printed = {0.10: (1.96, 0.005), 0.30: (1.71, 0.006), 0.60: (1.20, 0.005)}
    for g, (val, tol) in printed.items():
        cpo = cpo_tolerance(2.0, g)
        show(f"ex-11-8: Cpo for Cpa 2, %GRR {g:.0%} of tolerance", cpo)
        check(f"observed Cp at %GRR {g:.0%} (MSA p. 26)", cpo, val, tol)
    show("%GRR of tolerance that would give Cpo 1,20 at Cpa 2", math.sqrt(1 / 1.2 ** 2 - 1 / 4))
    show("Cpo for Cpa 2, %GRR 60 % of process variation (other basis)", cpo_process(2.0, 0.60))
    show("variance ratio for %GRR 30 % (Dummies p. 178 uses variances)", 0.30 ** 2)
    # p. 26 notes, process control: process at 4,95 g, reading 4,85 g, target 5,00 g
    show("ex-11-10 q1: adjustment = 5,00 - 4,85 [g]", 5.00 - 4.85)
    show("ex-11-10 q2: new process level = 4,95 + 0,15 [g]", 4.95 + 0.15)


# --------------------------------------------------------------------------- 6. gauge performance curve
def gauge_performance_curve() -> None:
    """MSA p. 27: beta(Xr) = Phi((USL-(Xr+b))/s) - Phi((LSL-(Xr+b))/s), LSL 0,6, USL 1,0, b 0,05, s 0,05."""
    header("6. Gauge performance curve (MSA p. 27)")
    lsl, usl, b, s = 0.6, 1.0, 0.05, 0.05

    def beta(xr: float) -> float:
        return stats.norm.cdf((usl - (xr + b)) / s) - stats.norm.cdf((lsl - (xr + b)) / s)

    show("ex-11-11 q1: P(accept | Xr = 0,5)", beta(0.5))
    show("ex-11-11 q2: P(accept | Xr = 0,9)", beta(0.9))
    show("P(accept | Xr = 0,8) (middle of the specs)", beta(0.8))
    show("P(accept | Xr = 0,95)", beta(0.95))


# ------------------------------------------------------------------------------- 7. uncertainty rules
def uncertainty_rules() -> None:
    """MSA p. 29: propagation rules; numeric example x = 5 +/- 0,1."""
    header("7. Uncertainty propagation (MSA p. 29)")
    x, ux = 5.0, 0.1  # MSA p. 29 notes
    show("u(x^2) = 2 x u(x)", 2 * x * ux)
    check("x^2 = 25 +/- 1 (MSA p. 29)", 2 * x * ux, 1.0, 0.05)
    u_sqrt = math.sqrt(x) * ux / (2 * x)
    show("ex-11-12 q2: u(sqrt x) = sqrt(x) u(x) / (2x)", u_sqrt)
    show("sqrt(5)", math.sqrt(x))


# ------------------------------------------------------------------------------------- 8. steel strip
def steel_strip() -> None:
    """Calculating the uncertainty in the width of a steel strip.pdf p. 1-2."""
    header("8. Steel strip uncertainty (steel strip p. 1-2)")
    mean, rng, n = 1834.0, 2.0, 3  # p. 1: average of 3 measurements 1834 mm, range 2 mm
    d2_3 = 1.693  # p. 1: 2 / d2 = 2 / 1,693 (d2 for n = 3; Table 18 also 1.693)
    u1 = mean * 0.001 / 2
    u2 = 0.5 / math.sqrt(3)
    u3 = (mean * 0.001 / 2) / math.sqrt(3)
    sigma_single = rng / d2_3
    u4 = sigma_single / math.sqrt(n)
    correction = mean * 0.001
    show("ex-11-13 q1: u1 = 1834 x 0,1 % / 2", u1)
    check("u1 = 0,917", u1, 0.917, 0.0005)
    show("ex-11-13 q2: u2 = 0,5 / sqrt 3", u2)
    check("u2 = 0,289", u2, 0.289, 0.0005)
    show("u3 = 0,917 / sqrt 3", u3)
    check("u3 = 0,529", u3, 0.529, 0.0005)
    show("repeatability of one reading R / d2 = 2 / 1,693", sigma_single)
    show("ex-11-13 q3: u4 = (R / d2) / sqrt 3 (mean of 3, MSA p. 29)", u4)
    check("u4 = 0,682 (printed after '2 / 1,693')", u4, 0.682, 0.0005)
    check("'2 / d2 = 2 / 1,693' taken literally vs printed u4 0,682", sigma_single, 0.682, 0.0005)
    uc = math.sqrt(u1 ** 2 + u2 ** 2 + u3 ** 2 + u4 ** 2)
    uc_printed_inputs = math.sqrt(0.917 ** 2 + 0.289 ** 2 + 0.529 ** 2 + 0.682 ** 2)
    show("ex-11-13 q4: uc = sqrt(u1^2 + u2^2 + u3^2 + u4^2)", uc)
    show("uc from the printed rounded u's", uc_printed_inputs)
    show("squares 0,917^2 0,289^2 0,529^2 0,682^2 and sum",
         f"{0.917 ** 2:.4f} {0.289 ** 2:.4f} {0.529 ** 2:.4f} {0.682 ** 2:.4f} {uc_printed_inputs ** 2:.4f}")
    check("uc = 1,264 (steel strip p. 2)", uc_printed_inputs, 1.264, 0.0005)
    big_u = 2 * uc
    show("ex-11-13 q5: U = 2 uc", big_u)
    check("U = 2,5 (steel strip p. 2)", big_u, 2.5, 0.05)
    show("printed U from printed uc: 2 x 1,264", 2 * 1.264)
    show("correction for the slanted tape 0,1 % x 1834", correction)
    corrected = mean - 1.8  # the text rounds the correction to 1,8 mm
    show("ex-11-13 q6: corrected width 1834 - 1,8", corrected)
    show("ex-11-13: interval lower with U = 2,6", corrected - 2.6)
    show("ex-11-13: interval upper with U = 2,6", corrected + 2.6)
    show("interval with unrounded U", f"[{corrected - big_u:.2f} ; {corrected + big_u:.2f}]")
    print("  printed interval [1829,7 ; 1834,7] = 1832,2 -/+ 2,5")
    # uniform distribution: half-width a -> sd a / sqrt 3 (footnote 3)
    show("sd of uniform on [-0,5 ; 0,5] = sqrt(1^2/12)", math.sqrt(1 / 12))


# --------------------------------------------------------------------------- 9. inadequate discrimination
def borderline_rule() -> None:
    """MSA p. 30: borderline sigma = MU; R chart UCL = D4 d2 MU, LCL = D3 d2 MU (Table 18 constants)."""
    header("9. Inadequate discrimination: borderline condition (MSA p. 30, Table 18)")
    show("n = 2: E(R) = d2 sigma, d2", T18[2]["d2"])
    for n in range(2, 11):
        c = T18[n]
        lcl, ucl = c["D3"] * c["d2"], c["D4"] * c["d2"]
        values = [v for v in range(0, 20) if lcl <= v <= ucl]
        show(f"n = {n}: LCL, UCL, UCL-LCL [MU]; possible range values",
             f"{lcl:.3f} {ucl:.3f} {ucl - lcl:.3f}", f"{len(values)} values {values}")
    check("n = 2: UCL - LCL = 3,69 MU (MSA p. 30)", (T18[2]["D4"] - T18[2]["D3"]) * T18[2]["d2"], 3.69, 0.005)
    check("n = 3: UCL - LCL = 4,36 MU (MSA p. 30)", (T18[3]["D4"] - T18[3]["D3"]) * T18[3]["d2"], 4.36, 0.005)
    show("ex-11-15 q1: n = 4 UCL - LCL [MU]", (T18[4]["D4"] - T18[4]["D3"]) * T18[4]["d2"])
    show("ex-11-15 q2: n = 5 UCL - LCL [MU]", (T18[5]["D4"] - T18[5]["D3"]) * T18[5]["d2"])


def rheostat() -> None:
    """Rheostat Knob Data.xls: X-bar/R chart of the original (1/1000 inch) and rounded (1/100 inch) data."""
    header("10. Rheostat knob data (MSA p. 30; Rheostat Knob Data.xls)")
    book = xlrd.open_workbook(str(RHEOSTAT_XLS))
    a2, d4_t18 = T18[5]["A2"], T18[5]["D4"]
    for sheet, solved, mu, tag in (("originele data", "originele data - opgelost", 0.001, "orig"),
                                   ("afgeronde data", "afgeronde data - opgelost", 0.01, "rounded")):
        sh = book.sheet_by_name(sheet)
        data = np.array([[float(sh.cell_value(r, c)) for c in range(5)] for r in range(2, 29)])
        means, ranges = data.mean(axis=1), data.max(axis=1) - data.min(axis=1)
        xbb, rbar = means.mean(), ranges.mean()
        lcl_x, ucl_x = xbb - a2 * rbar, xbb + a2 * rbar
        ucl_r = d4_t18 * rbar
        print(f"  --- {sheet}: {data.shape[0]} subgroups of {data.shape[1]}")
        show(f"{tag}: X-double-bar", xbb)
        show(f"ex-11-14 {tag}: R-bar", rbar)
        show(f"{tag}: X-bar limits (A2 = {a2})", f"{lcl_x:.6f} {ucl_x:.6f}")
        show(f"ex-11-14 {tag}: UCL_R with D4 = {d4_t18} (Table 18)", ucl_r)
        show(f"{tag}: UCL_R with D4 = 2.114 (workbook)", 2.114 * rbar)
        out_x = [(i + 1, round(m, 4)) for i, m in enumerate(means) if m < lcl_x - 1e-12 or m > ucl_x + 1e-12]
        out_r = [(i + 1, round(r, 4)) for i, r in enumerate(ranges) if r > ucl_r + 1e-12]
        show(f"ex-11-14 {tag}: subgroups with mean outside limits", len(out_x), str([i for i, _ in out_x]))
        show(f"{tag}: subgroups with range above UCL_R", len(out_r), str([i for i, _ in out_r]))
        observed = sorted({round(r / mu) for r in ranges})
        possible = [v for v in range(0, 40) if v * mu <= ucl_r + 1e-12]
        show(f"{tag}: observed range values [MU = {mu}]", str(observed))
        show(f"ex-11-14 {tag}: possible range values within R limits", len(possible), str(possible))
        sol = book.sheet_by_name(solved)
        printed = {"X-double-bar": sol.cell_value(2, 9), "LCL_X": sol.cell_value(2, 10),
                   "UCL_X": sol.cell_value(2, 11), "R-bar": sol.cell_value(2, 13), "UCL_R": sol.cell_value(2, 15)}
        check(f"{tag}: X-double-bar (opgelost)", xbb, printed["X-double-bar"], 1e-9)
        check(f"{tag}: LCL_X (opgelost)", lcl_x, printed["LCL_X"], 1e-9)
        check(f"{tag}: UCL_X (opgelost)", ucl_x, printed["UCL_X"], 1e-9)
        check(f"{tag}: R-bar (opgelost)", rbar, printed["R-bar"], 1e-9)
        check(f"{tag}: UCL_R, Table 18 D4 2.115 (opgelost uses 2.114)", ucl_r, printed["UCL_R"], 1e-9)
        check(f"{tag}: UCL_R with D4 2.114", 2.114 * rbar, printed["UCL_R"], 1e-9)


# ----------------------------------------------------------------------------- 11. linearity and bias
def read_linearity() -> tuple[np.ndarray, np.ndarray]:
    """linearity.txt: reference values (5 parts) and a 12 x 5 matrix of measurements."""
    lines = [ln.strip() for ln in LINEARITY_TXT.read_text().replace("\r", "").split("\n") if ln.strip()]
    ref = np.array([float(v) for v in lines[1].split()[2:]])
    meas = np.array([[float(v) for v in ln.split()[1:]] for ln in lines[3:]])
    return ref, meas


def linearity() -> None:
    """MSA p. 33: regress bias on reference value; test slope = 0 and intercept = 0 (two-sided)."""
    header("11. Linearity (MSA p. 33; linearity.txt)")
    ref, meas = read_linearity()
    show("reference values", str(ref.tolist()))
    show("measurements shape (repeats x parts)", str(meas.shape))
    bias = meas - ref
    for j in range(5):
        show(f"part {j + 1}: mean measurement, mean bias, range",
             f"{meas[:, j].mean():.4f} {bias[:, j].mean():.4f} {np.ptp(meas[:, j]):.2f}")
    x = np.repeat(ref[None, :], meas.shape[0], axis=0).ravel()
    y = bias.ravel()
    fit = sm.OLS(y, sm.add_constant(x)).fit()
    a, b = fit.params
    se_a, se_b = fit.bse
    show("n points", len(y))
    show("ex-11-19 q1: slope b", b)
    show("ex-11-19 q2: intercept a", a)
    show("se(slope), se(intercept)", f"{se_b:.6f} {se_a:.6f}")
    show("ex-11-19 q3: t for slope = 0", fit.tvalues[1])
    show("p-value slope (two-sided)", fit.pvalues[1])
    show("t for intercept = 0, p-value", f"{fit.tvalues[0]:.4f} {fit.pvalues[0]:.3g}")
    show("t crit 0,975, df = n - 2", stats.t.ppf(0.975, len(y) - 2))
    ci = fit.conf_int(0.05)
    show("95 % CI intercept", f"[{ci[0][0]:.4f} ; {ci[0][1]:.4f}]")
    show("95 % CI slope", f"[{ci[1][0]:.4f} ; {ci[1][1]:.4f}]")
    show("s = sqrt(MS_E), R^2", f"{math.sqrt(fit.mse_resid):.4f} {fit.rsquared:.4f}")
    # independent cross-check with scipy
    lr = stats.linregress(x, y)
    check("slope (scipy linregress)", lr.slope, b, 1e-12)
    check("intercept (scipy linregress)", lr.intercept, a, 1e-12)
    # 95 % confidence band for E(bias) at each reference value: does it contain 0?
    pred = fit.get_prediction(sm.add_constant(ref)).summary_frame(alpha=0.05)
    for r, lo, hi in zip(ref, pred["mean_ci_lower"], pred["mean_ci_upper"]):
        show(f"band E(bias) at ref {r}", f"[{lo:.4f} ; {hi:.4f}]", "contains 0" if lo <= 0 <= hi else "0 outside")
    # a horizontal line bias = b inside the band over the whole range [1,99 ; 9,80] (MSA p. 33 notes)?
    grid = np.linspace(ref.min(), ref.max(), 2001)
    band = fit.get_prediction(sm.add_constant(grid)).summary_frame(alpha=0.05)
    lo_max, hi_min = band["mean_ci_lower"].max(), band["mean_ci_upper"].min()
    show("horizontal lines inside the band over the range: b in", f"[{lo_max:.4f} ; {hi_min:.4f}]")
    # extra checks on the U-shaped pattern (not part of the course procedure)
    show("overall mean bias (model bias = constant)", y.mean())
    t_const = y.mean() / (y.std(ddof=1) / math.sqrt(len(y)))
    show("t for constant bias = 0 (df 59), p", f"{t_const:.4f} {2 * stats.t.sf(abs(t_const), len(y) - 1):.3g}")
    f_parts, p_parts = stats.f_oneway(*[bias[:, j] for j in range(5)])
    show("one-way ANOVA bias by part (DOE p. 6-8): F, p", f"{f_parts:.4f} {p_parts:.3g}")


def bias_test() -> None:
    """MSA p. 32 applied to one part of linearity.txt (extra exercise): one subgroup (g = 1) of m = 12.

    t = bias / sigma_r * sqrt(g m) * d2* / d2, sigma_r = R / d2, nu and d2* from tabel MSA (g = 1, m = 12),
    d2 from tabel MSA last row (g -> infinity), as the Gage R&R sheet does.
    """
    header("12. Bias test (MSA p. 32) on parts of linearity.txt (extra exercise)")
    ref, meas = read_linearity()
    g, m = 1, 12
    nu, d2s = MSA[(g, m)]
    d2 = D2_INF[m]
    show("tabel MSA g = 1, m = 12: nu, d2*", f"{nu} {d2s}")
    show("tabel MSA last row: d2 for m = 12", d2)
    tq = stats.t.ppf(0.975, nu)
    show("t(0,975; nu = 9,0)", tq)
    for j in (0, 2):
        xs = meas[:, j]
        bias = xs.mean() - ref[j]
        rng = np.ptp(xs)
        sigma_r = rng / d2
        t = bias / sigma_r * math.sqrt(g * m) * d2s / d2
        half = sigma_r * d2 / (d2s * math.sqrt(g * m)) * tq
        tag = f"part {j + 1} (ref {ref[j]})"
        show(f"ex-11-18 {tag}: bias = mean - ref", bias)
        show(f"ex-11-18 {tag}: R, sigma_r = R / d2", f"{rng:.2f} {sigma_r:.6f}")
        show(f"ex-11-18 {tag}: sum, mean", f"{xs.sum():.2f} {xs.mean():.4f}")
        show(f"ex-11-18 {tag}: factors bias/sigma_r, sqrt(gm), d2*/d2",
             f"{bias / sigma_r:.4f} {math.sqrt(g * m):.4f} {d2s / d2:.4f}")
        show(f"ex-11-18 {tag}: t", t)
        show(f"ex-11-18 {tag}: 95 % CI half-width", half)
        show(f"ex-11-18 {tag}: 95 % CI", f"[{bias - half:.4f} ; {bias + half:.4f}]")


# ------------------------------------------------------------------------------------- 13. d2 and d2*
def d2_star() -> None:
    """tabel MSA.pdf: d2* lookups; relation d2*^2 = d2^2 + d3^2 / g derived from course facts
    (E(R) = d2 sigma, sd(R) = d3 sigma, var of an average of g ranges = var/g: GRR theory p. 1;
    d2*^2 = E(Rbar^2) / sigma^2: MSA p. 32 notes)."""
    header("13. d2 and d2* (tabel MSA p. 1; GRR theory p. 1; MSA p. 32)")
    show("ex-11-20 q1: d2* (g = 1, m = 3)", MSA[(1, 3)][1])
    show("ex-11-20 q2: d2* (g = 1, m = 5)", MSA[(1, 5)][1])
    show("ex-11-20 q3: nu (g = 5, m = 3)", MSA[(5, 3)][0])
    show("ex-11-20 q4: d2* (g = 5, m = 3)", MSA[(5, 3)][1])
    show("d2 (g -> infinity) for m = 2", D2_INF[2])
    check("GRR theory p. 1: g = 5, m = 3: d2* = 1,73857", MSA[(5, 3)][1], 1.73857, 5e-6)
    check("GRR theory p. 1: g = 5, m = 3: nu = 9,3", MSA[(5, 3)][0], 9.3, 0.05)
    check("GRR theory p. 1: d2 = 2,326 (sample size 5)", TA[5]["d2"], 2.326, 5e-4)
    check("GRR theory p. 1: d3 = 0,864 (sample size 5)", TA[5]["d3"], 0.864, 5e-4)
    for g, m in ((1, 2), (1, 3), (1, 5), (5, 3), (15, 2), (20, 5)):
        rel = math.sqrt(D2_INF[m] ** 2 + TA[m]["d3"] ** 2 / g)
        check(f"d2*^2 = d2^2 + d3^2/g at g = {g}, m = {m}", rel, MSA[(g, m)][1], 0.0006)
    # Monte Carlo as in the workbook sheet 'on d2 and d2star' (g = 2, m = 5), fixed seed
    rng = np.random.default_rng(11)
    samples = rng.standard_normal((200_000, 2, 5))
    rbar = np.ptp(samples, axis=2).mean(axis=1)
    show("Monte Carlo g = 2, m = 5: mean Rbar (-> d2)", rbar.mean(), f"table d2 {D2_INF[5]}")
    show("Monte Carlo g = 2, m = 5: sqrt(mean Rbar^2) (-> d2*)", math.sqrt((rbar ** 2).mean()),
         f"table d2* {MSA[(2, 5)][1]}")


# ------------------------------------------------------------------------------------------ 14. GRR study
def grr_data() -> np.ndarray:
    """GRR workbook, sheet measurements D4:F13: array [operator, part, trial] of coded values (0,0001 inch)."""
    ws = openpyxl.load_workbook(GRR_XLSX, data_only=True, read_only=True)["measurements"]
    cells = [[ws.cell(row=r, column=c).value for c in (4, 5, 6)] for r in range(4, 14)]
    arr = np.array(cells, dtype=float)  # rows: P1 t1, P1 t2, P2 t1, ... ; columns A, B, C
    return arr.reshape(5, 2, 3).transpose(2, 0, 1)  # -> [operator, part, trial]


def grr_average_range(x: np.ndarray) -> dict[str, float]:
    """Average and range method (MSA p. 34-35) with d2, d2* from tabel MSA."""
    k, n, r = x.shape
    rbar = np.ptp(x, axis=2).mean()
    op_means = x.mean(axis=(1, 2))
    part_means = x.mean(axis=(0, 2))
    xdiff = np.ptp(op_means)
    rp = np.ptp(part_means)
    d2 = D2_INF[r]
    d2_a = MSA[(1, k)][1]
    d2_p = MSA[(1, n)][1]
    ev = rbar / d2
    av2 = (xdiff / d2_a) ** 2 - ev ** 2 / (n * r)
    pv2 = (rp / d2_p) ** 2 - ev ** 2 / (k * r)
    av, pv = math.sqrt(max(av2, 0)), math.sqrt(pv2)
    grr = math.sqrt(ev ** 2 + av ** 2)
    tv = math.sqrt(grr ** 2 + pv ** 2)
    return {"Rbar": rbar, "XDIFF": xdiff, "Rp": rp, "d2": d2, "d2*A": d2_a, "d2*P": d2_p,
            "EV": ev, "EV2": ev ** 2, "AV2": av2, "AV": av, "PV2": pv2, "PV": pv, "PV_simple": rp / d2_p,
            "GRR": grr, "TV": tv, "%GRR": grr / tv,
            "op_means": op_means, "part_means": part_means}


def grr_anova(x: np.ndarray) -> dict[str, float]:
    """Two-way ANOVA (MSA p. 36): with interaction (Excel layout) and the course's additive model."""
    k, n, r = x.shape
    grand = x.mean()
    cell = x.mean(axis=2)
    op_m, part_m = x.mean(axis=(1, 2)), x.mean(axis=(0, 2))
    ss_t = ((x - grand) ** 2).sum()
    ss_a = n * r * ((op_m - grand) ** 2).sum()
    ss_p = k * r * ((part_m - grand) ** 2).sum()
    ss_cells = r * ((cell - grand) ** 2).sum()
    ss_ap = ss_cells - ss_a - ss_p
    ss_w = ss_t - ss_cells
    df_a, df_p, df_ap, df_w = k - 1, n - 1, (k - 1) * (n - 1), k * n * (r - 1)
    ms_a, ms_p, ms_ap, ms_w = ss_a / df_a, ss_p / df_p, ss_ap / df_ap, ss_w / df_w
    ss_e = ss_ap + ss_w
    df_e = n * k * (r - 1) + (n - 1) * (k - 1)  # MSA p. 36
    ms_e = ss_e / df_e
    ev2 = ms_e
    av2 = (ms_a - ev2) / (n * r)
    pv2 = (ms_p - ev2) / (k * r)
    tv2 = ev2 + av2 + pv2
    return {"SS_T": ss_t, "SS_A": ss_a, "SS_P": ss_p, "SS_AP": ss_ap, "SS_W": ss_w, "SS_E": ss_e,
            "df_A": df_a, "df_P": df_p, "df_AP": df_ap, "df_W": df_w, "df_E": df_e,
            "MS_A": ms_a, "MS_P": ms_p, "MS_AP": ms_ap, "MS_W": ms_w, "MS_E": ms_e,
            "F_A_int": ms_a / ms_w, "F_P_int": ms_p / ms_w, "F_AP": ms_ap / ms_w,
            "p_A_int": stats.f.sf(ms_a / ms_w, df_a, df_w), "p_P_int": stats.f.sf(ms_p / ms_w, df_p, df_w),
            "p_AP": stats.f.sf(ms_ap / ms_w, df_ap, df_w),
            "F_A": ms_a / ms_e, "F_P": ms_p / ms_e,
            "p_A": stats.f.sf(ms_a / ms_e, df_a, df_e), "p_P": stats.f.sf(ms_p / ms_e, df_p, df_e),
            "Fcrit_A": stats.f.isf(0.05, df_a, df_e), "Fcrit_P": stats.f.isf(0.05, df_p, df_e),
            "Fcrit_AP": stats.f.isf(0.05, df_ap, df_w),
            "EV2": ev2, "EV": math.sqrt(ev2), "AV2": av2, "AV": math.sqrt(av2), "PV2": pv2, "PV": math.sqrt(pv2),
            "GRR": math.sqrt(ev2 + av2), "TV": math.sqrt(tv2), "%GRR": math.sqrt((ev2 + av2) / tv2)}


def grr_study() -> None:
    """MSA p. 38 exercise with the lecturer's workbook as the course answer."""
    header("14. GRR study (MSA p. 38; GRR - ANOVA - average and range - 2.xlsx)")
    x = grr_data()
    k, n, r = x.shape
    show("k operators, n parts, r trials", f"{k} {n} {r}")
    for i, op in enumerate("ABC"):
        show(f"operator {op}: values per part (trial1, trial2)", str([tuple(v) for v in x[i].tolist()]))
    wb = openpyxl.load_workbook(GRR_XLSX, data_only=True, read_only=True)
    rng_ws, an_ws = wb["ranges"], wb["2way anova"]

    print("  --- average and range method")
    ar = grr_average_range(x)
    show("operator means A, B, C", str([round(v, 4) for v in ar["op_means"]]))
    show("part means P1..P5", str([round(v, 4) for v in ar["part_means"]]))
    show("ranges per operator x part", str(np.ptp(x, axis=2).tolist()))
    for key in ("Rbar", "XDIFF", "Rp", "d2", "d2*A", "d2*P", "EV2", "EV", "AV2", "AV", "PV2", "PV",
                "PV_simple", "GRR", "TV", "%GRR"):
        show(f"ex-11-23 A&R {key}", ar[key])
    show("EV in inch = EV x 0,0001", ar["EV"] * 1e-4)
    k_, n_, r_ = x.shape
    show("intermediate: sum of ranges", float(np.ptp(x, axis=2).sum()))
    show("intermediate: (XDIFF/d2*A)^2, EV^2/(nr)", f"{(ar['XDIFF'] / ar['d2*A']) ** 2:.4f} {ar['EV2'] / (n_ * r_):.4f}")
    show("intermediate: (Rp/d2*P)^2, EV^2/(kr)", f"{(ar['Rp'] / ar['d2*P']) ** 2:.4f} {ar['EV2'] / (k_ * r_):.4f}")
    show("intermediate: GRR^2 = EV^2 + AV^2", ar["EV2"] + ar["AV2"])
    show("grand mean (sum / 30)", f"{x.sum():g} / {x.size} = {x.mean():.4f}")
    printed = {"EV2": "C11", "EV": "F11", "AV2": "C12", "AV": "F12", "PV2": "C13", "PV": "F13",
               "TV": "F16", "%GRR": "C17"}
    for key, cell in printed.items():
        check(f"A&R {key} (ranges!{cell})", ar[key], float(rng_ws[cell].value), 1e-9)
    check("A&R Rbar (measurements!H17)", ar["Rbar"], 4.333333333333333, 1e-9)
    check("A&R XDIFF (measurements!H24)", ar["XDIFF"], 8.5, 1e-9)
    check("A&R Rp (measurements!H25)", ar["Rp"], 25.0, 1e-9)
    check("A&R d2* (A) 1,91155 (ranges!C6)", ar["d2*A"], float(rng_ws["C6"].value), 1e-9)
    check("A&R d2* (P) 2,48124 (ranges!C7)", ar["d2*P"], float(rng_ws["C7"].value), 1e-9)
    check("A&R d2 1,12838 (ranges!C2)", ar["d2"], float(rng_ws["C2"].value), 1e-9)

    print("  --- ANOVA method")
    an = grr_anova(x)
    for key in ("SS_T", "SS_P", "SS_A", "SS_AP", "SS_W", "SS_E", "df_E", "MS_P", "MS_A", "MS_AP", "MS_W", "MS_E",
                "F_P_int", "p_P_int", "F_A_int", "p_A_int", "F_AP", "p_AP", "Fcrit_AP",
                "F_P", "p_P", "Fcrit_P", "F_A", "p_A", "Fcrit_A",
                "EV2", "EV", "AV2", "AV", "PV2", "PV", "GRR", "TV", "%GRR"):
        show(f"ex-11-24 ANOVA {key}", an[key])
    printed_cells = {"SS_P": "C43", "SS_A": "C44", "SS_AP": "C45", "SS_W": "C46", "SS_T": "C48",
                     "F_P_int": "F43", "F_A_int": "F44", "F_AP": "F45",
                     "p_P_int": "G43", "p_A_int": "G44", "p_AP": "G45"}
    for key, cell in printed_cells.items():
        check(f"ANOVA {key} (2way anova!{cell})", an[key], float(an_ws[cell].value), 1e-6 * max(1, abs(an[key])))
    derived = {"SS_E": "K45", "MS_E": "M45", "F_P": "N43", "F_A": "N44", "p_P": "O43", "p_A": "O44",
               "EV2": "K51", "EV": "N51", "AV2": "K52", "AV": "N52", "PV2": "K53", "PV": "N53",
               "TV": "N56", "%GRR": "K57"}
    print("  derived additive table: strict (1e-9) and at the workbook's own rounding (1e-4 relative)")
    for key, cell in derived.items():
        printed_val = float(an_ws[cell].value)
        check(f"ANOVA {key} (2way anova!{cell}) strict", an[key], printed_val, 1e-9 * max(1, abs(an[key])))
        check(f"ANOVA {key} (2way anova!{cell}) 1e-4 rel", an[key], printed_val, 1e-4 * max(1, abs(an[key])))
    show("workbook K45 typed =412.5+296.667", 412.5 + 296.667)
    show("intermediate: ANOVA GRR^2 = EV^2 + AV^2", an["EV2"] + an["AV2"])
    show("intermediate: within-SS share of the range 25 (B, P4): 25^2 / 2", 25 ** 2 / 2)
    show("exact SS_W + SS_AP", an["SS_W"] + an["SS_AP"])

    # independent cross-check with statsmodels
    rows = [{"y": x[i, j, t], "op": "ABC"[i], "part": f"P{j + 1}"} for i in range(k) for j in range(n) for t in range(r)]
    df = pd.DataFrame(rows)
    full = sm.stats.anova_lm(smf.ols("y ~ C(op) * C(part)", df).fit(), typ=2)
    add = sm.stats.anova_lm(smf.ols("y ~ C(op) + C(part)", df).fit(), typ=2)
    check("statsmodels SS interaction", float(full.loc["C(op):C(part)", "sum_sq"]), an["SS_AP"], 1e-8)
    check("statsmodels SS residual (with interaction)", float(full.loc["Residual", "sum_sq"]), an["SS_W"], 1e-8)
    check("statsmodels SS residual (additive)", float(add.loc["Residual", "sum_sq"]), an["SS_E"], 1e-8)
    check("statsmodels df residual (additive)", float(add.loc["Residual", "df"]), an["df_E"], 1e-9)
    check("statsmodels F operators (additive)", float(add.loc["C(op)", "F"]), an["F_A"], 1e-8)

    print("  --- comparison and verdict (MSA p. 35: <= 10 % ok, 10-30 % maybe, > 30 % not acceptable)")
    for name, res in (("average and range", ar), ("ANOVA", an)):
        verdict = "acceptable" if res["%GRR"] <= 0.10 else ("may be acceptable" if res["%GRR"] <= 0.30 else
                                                             "not acceptable")
        show(f"%GRR {name} [%]", 100 * res["%GRR"], verdict)
        show(f"share of TV^2 due to GRR, {name} [%]", 100 * (res["GRR"] / res["TV"]) ** 2)
    check("A&R %GRR = 50,01 % (ranges!C17)", 100 * ar["%GRR"], 50.01, 0.005)
    check("ANOVA %GRR = 49,47 % (2way anova!K57)", 100 * an["%GRR"], 49.47, 0.005)

    print("  --- extra: discrimination rule (MSA p. 30) on the GRR ranges (n = 2, MU = 5 coded units)")
    d4 = T18[2]["D4"]
    ucl_r = d4 * ar["Rbar"]
    show("ex-11-26 q1: UCL_R = D4 x Rbar (D4 = 3.267, Table 18)", ucl_r)
    possible = [v for v in range(0, 100, 5) if v <= ucl_r]
    show("ex-11-26 q2: possible range values within limits", len(possible), str(possible))
    show("ranges above UCL_R", str([(op, f'P{j + 1}', float(np.ptp(x[i, j])))
                                     for i, op in enumerate('ABC') for j in range(n) if np.ptp(x[i, j]) > ucl_r]))
    show("EV / MU (sigma relative to the measurement unit)", ar["EV"] / 5)
    show("EV (ANOVA) / MU", an["EV"] / 5)


def ratios_of_tv() -> None:
    """Pitfall box 11.17: EV/TV, AV/TV and PV/TV are ratios of standard deviations and do not add up to 100 %;
    only their squares (the shares of TV^2) do, because TV^2 = EV^2 + AV^2 + PV^2 (MSA p. 35)."""
    header("15. Pitfall 11.17: ratios of standard deviations do not add up (MSA p. 35; GRR study)")
    x = grr_data()
    for name, res in (("average and range", grr_average_range(x)), ("ANOVA", grr_anova(x))):
        ratios = {comp: res[comp] / res["TV"] for comp in ("EV", "AV", "PV")}
        for comp, ratio in ratios.items():
            show(f"{name}: {comp}/TV [%]", 100 * ratio)
        show(f"{name}: EV/TV + AV/TV + PV/TV [%]", 100 * sum(ratios.values()))
        # the squared ratios are the variance shares; they must add up to exactly 100 %
        show(f"{name}: (EV/TV)^2 + (AV/TV)^2 + (PV/TV)^2 [%]", 100 * sum(r ** 2 for r in ratios.values()))


def main() -> None:
    """Run every section in the order of the guide."""
    meter_unit()
    discrimination()
    effective_resolution()
    probable_error()
    process_decisions()
    gauge_performance_curve()
    uncertainty_rules()
    steel_strip()
    borderline_rule()
    rheostat()
    linearity()
    bias_test()
    d2_star()
    grr_study()
    ratios_of_tv()
    print("\nDone.")


if __name__ == "__main__":
    main()
