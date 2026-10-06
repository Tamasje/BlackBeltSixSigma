"""Numbers for study/parts/03_normaal_sigma_dpmo.html (Deel 03: normale verdeling, sigmaniveau en DPMO).

Prints every value the fragment marks "zelf berekend", every exercise answer (data-answer) and re-checks every
printed course value the fragment quotes. Course data come from inventory/constants (rows carry file + page), from
the course workbook __NormVerdeling Excel functies.xlsx (read-only, openpyxl) or are transcribed from the page named
in the comment. Page numbers are PDF page numbers (Les 4 deck: PDF page = pptx slide number; Dummies: PDF page =
printed page + 18).

Each re-check has an expected status: "OK" (agrees after rounding to the printed precision) or "MISMATCH" (a known
printed error, listed in 03_errata.tsv). The script exits 1 if any check does not have its expected status.
Run from the project root:  python3 study/parts/03_numbers.py
"""
from __future__ import annotations

import csv
import math
import sys
from pathlib import Path

import openpyxl
from scipy import integrate, stats

ROOT = Path(__file__).resolve().parents[2]
CONSTANTS = ROOT / "inventory" / "constants"
LES4 = ROOT / "source" / "course" / "Les 4"
PHI = stats.norm.cdf          # standard normal cdf, P(Z < z)
PHI_INV = stats.norm.ppf      # inverse: z with P(Z < z) = p
FAILURES: list[str] = []


def rows(stem: str) -> list[dict[str, str]]:
    """Rows of a course table in inventory/constants."""
    with (CONSTANTS / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def num(text: str) -> float:
    """Course number as printed ('308,537' with thousands comma, '69.15%', '.0228') to float."""
    return float(text.replace(",", "").replace("%", "").strip())


def show(label: str, value: float | str, note: str = "") -> None:
    """Print one computed value."""
    text = f"{value:.6g}" if isinstance(value, float) else str(value)
    print(f"  {label:<62} {text:>16}  {note}")


def check(label: str, printed: str, computed: float, decimals: int, expected: str = "OK", note: str = "") -> None:
    """Compare a printed course value with a computation rounded to the printed number of decimals."""
    printed_value = num(printed)
    status = "OK" if round(computed, decimals) == round(printed_value, decimals) else "MISMATCH"
    if status != expected:
        FAILURES.append(f"{label}: expected {expected}, got {status}")
    flag = "" if status == expected else "   <-- UNEXPECTED"
    print(f"  {status:<8} {label:<56} printed {printed:>14}  computed {computed:.8g}  {note}{flag}")


def section(title: str) -> None:
    """Print a section header."""
    print(f"\n== {title}")


def normal_basics() -> None:
    """68-95-99.7 (SPC p. 18-19), area outside 3 sigma (p. 18), density formula (p. 27), shoe sizes (p. 19 notes)."""
    section("Normale verdeling: kengetallen (SPC p. 18-19)")
    inside = {k: PHI(k) - PHI(-k) for k in (1, 2, 3)}
    for k, p in inside.items():
        show(f"P(mu - {k} sigma < X < mu + {k} sigma)", p, "zelf berekend")
    check("SPC p. 19: +-1 sigma ~ 68 %", "68", 100 * inside[1], 0)
    check("SPC p. 19: +-2 sigma ~ 95 %", "95", 100 * inside[2], 0)
    check("SPC p. 19: +-3 sigma ~ 99.7 %", "99.7", 100 * inside[3], 1)
    check("SPC p. 18 figure: 95.4 %", "95.4", 100 * inside[2], 1)
    outside3 = 1 - inside[3]
    show("buiten +-3 sigma, totaal (%)", 100 * outside3, "zelf berekend")
    show("buiten +-3 sigma, per zijde (%)", 100 * outside3 / 2, "zelf berekend")
    check("SPC p. 18: '0,30 % totaal' buiten 3 sigma", "0.30", 100 * outside3, 2, "MISMATCH")
    check("SPC p. 18: '0,15 % links/rechts'", "0.15", 100 * outside3 / 2, 2, "MISMATCH")

    section("Dichtheidsformule SPC p. 27: oppervlakte onder de gedrukte formule")
    for sigma in (1.0, 0.5, 2.0):
        area, _ = integrate.quad(lambda x, s=sigma: math.exp(-0.5 * (x / s) ** 2) / math.sqrt(2 * math.pi),
                                 -math.inf, math.inf)
        show(f"oppervlakte zonder 1/sigma, sigma = {sigma}", area, "zelf berekend (= sigma, niet 1)")

    section("Schoenmaten (SPC p. 19, sprekersnotities): mu 42,5, sigma 2,5")
    mu, sd = 42.5, 2.5
    check("mu - 3 sigma = 35", "35", mu - 3 * sd, 0)
    check("mu + 3 sigma = 50", "50", mu + 3 * sd, 0)
    show("ex-03-2a: P(X > 45) in % (z = 1)", 100 * (1 - PHI((45 - mu) / sd)), "data-answer 15.87")
    show("ex-03-2b: P(X > 50) in % (z = 3)", 100 * (1 - PHI((50 - mu) / sd)), "data-answer 0.135")
    show("ex-03-2c: P(40 < X < 45) in % (z = -1..1)", 100 * (PHI(1) - PHI(-1)), "data-answer 68.27")
    show("1 man op ... boven maat 50", 1 / (1 - PHI(3)), "zelf berekend")


def sample_means() -> None:
    """SPC p. 14 notes / p. 28: sigma of means = sigma / sqrt(n), with the data of SPC p. 46 (sigma 0,2, n 5)."""
    section("Spreiding van gemiddelden (SPC p. 14, 28)")
    show("sqrt(5)", math.sqrt(5), "zelf berekend")
    show("ex-03-1: sigma_xbar = 0,2 / sqrt(5)", 0.2 / math.sqrt(5), "data-answer 0.0894")


def z_table() -> None:
    """Re-check the course Z table (___1.1 Ztable.pdf p. 1-2) and the lookups used in the exercises."""
    section("Z-tabel (Ztable p. 1-2): elke cel tegen Phi(z), 4 decimalen")
    table: dict[float, float] = {}
    for stem in ("S06_standard_normal_probabilities_table_entry_area_to_the_left_o",
                 "S06_standard_normal_probabilities_table_entry_area_to_the_left_o_2"):
        for row in rows(stem):
            base = num(row["z"].replace("–", "-"))
            negative = row["z"].strip().startswith(("-", "–"))
            for col in (".00", ".01", ".02", ".03", ".04", ".05", ".06", ".07", ".08", ".09"):
                z = round(base - num(col), 2) if negative else round(base + num(col), 2)
                table[z] = num(row[col])
    off = [z for z, p in table.items() if round(PHI(z), 4) != p]
    show("aantal cellen", float(len(table)))
    show("cellen die afwijken van Phi(z) op 4 decimalen", str(off) if off else "geen")
    if off:
        FAILURES.append(f"Z table cells off: {off}")
    lookups = {
        "ex-03-3a: P(Z < 1,25)": table[1.25],
        "ex-03-3b: P(Z > 1,96) = 1 - P(Z < 1,96)": 1 - table[1.96],
        "ex-03-3c: P(-1 < Z < 1) = .8413 - .1587": table[1.0] - table[-1.0],
        "ex-03-3d: P(Z < -2,00)": table[-2.0],
    }
    for label, value in lookups.items():
        show(label, round(value, 4), f"tabel; data-answer {value:.4f}")
    show("exact P(-1 < Z < 1)", PHI(1) - PHI(-1), "zelf berekend")
    show("tabel: P(Z < 2,32) / P(Z < 2,33)", f"{table[2.32]:.4f} / {table[2.33]:.4f}")
    show("ex-03-3e: z met P(Z < z) = 0,99 (tabel 2,33; exact)", PHI_INV(0.99), "data-answer 2.33")


def excel_workbook() -> None:
    """__NormVerdeling Excel functies.xlsx (Blad1): mean 20, sigma 0.5, 1 % afkeur; reverse questions."""
    section("__NormVerdeling Excel functies.xlsx (Blad1)")
    wb = openpyxl.load_workbook(LES4 / "__NormVerdeling Excel functies.xlsx", data_only=True, read_only=True)
    ws = wb["Blad1"]
    mean, sd = float(ws["B1"].value), float(ws["B2"].value)
    cached_inv, cached_dist = float(ws["B5"].value), float(ws["B6"].value)
    wb.close()
    limit = stats.norm.ppf(0.99, mean, sd)
    check("B5 NORMINV(0,99;20;0,5)", f"{cached_inv:.6f}", limit, 6)
    check("B6 NORMDIST(21,1632;20;0,5;WAAR)", f"{cached_dist:.6f}", stats.norm.cdf(21.1632, mean, sd), 6)
    show("ex-03-4a: grens voor 1 % afkeur boven (3 decimalen)", limit, "data-answer 21.163")
    show("met Z-tabel: 20 + 2,33 * 0,5", 20 + 2.33 * 0.5, "zelf berekend")
    show("ex-03-4b: z = (21,1632 - 20)/0,5", (21.1632 - mean) / sd, "data-answer 2.326")
    show("ex-03-4c: afkeur boven 21,1632 in %", 100 * (1 - stats.norm.cdf(21.1632, mean, sd)), "data-answer 1.00")
    # Reverse use (sigma or mean from a tail fraction), with the workbook's own numbers.
    z99 = PHI_INV(0.99)
    show("ex-03-5a: sigma = (21,1632 - 20) / z(0,99)", (21.1632 - mean) / z99, "data-answer 0.500")
    show("ex-03-5a met tabel-z 2,33", (21.1632 - mean) / 2.33, "zelf berekend")
    show("ex-03-5b: mu = 21,1632 - 0,5 * z(0,99)", 21.1632 - sd * z99, "data-answer 20.00")
    show("ex-03-5c: grens voor 1 % afkeur ONDER (3 decimalen)", stats.norm.ppf(0.01, mean, sd), "data-answer 18.837")


def discrete_metrics() -> None:
    """SPC p. 20 formulas with Dummies p. 152-154 and Harry & Schroeder p. 5 examples."""
    section("DPU / DPO / DPMO (Dummies p. 152-154; H&S p. 5)")
    dpu = 11 / 23
    check("Dummies p. 152: DPU = 11/23", "0.478", dpu, 3)
    check("Dummies p. 153: DPO auto = 158/14550", "0.011", 158 / 14550, 3)
    check("Dummies p. 154: DPO fiets = 2/173", "0.012", 2 / 173, 3)
    check("Dummies p. 154: DPO vliegtuig = 2/6.000.000", "0.000000333", 2 / 6e6, 9)
    show("DPMO vliegtuig = DPO * 1e6", 2 / 6e6 * 1e6, "zelf berekend")
    show("ex-03-6a: DPU leningen", dpu, "data-answer 0.478")
    show("DPO auto (4 decimalen)", 158 / 14550, "zelf berekend")
    show("DPO fiets (4 decimalen)", 2 / 173, "zelf berekend")
    show("ex-03-6b: DPMO auto", 158 / 14550 * 1e6, "data-answer 10859")
    show("ex-03-6c: DPMO fiets", 2 / 173 * 1e6, "data-answer 11561")
    d, n, o = 5, 100, 20
    check("H&S p. 5: DPU = 5/100", "0.05", d / n, 2)
    check("H&S p. 5: DPO = 0,05/20", "0.0025", d / (n * o), 4)
    check("H&S p. 5: DPMO = 2,500", "2,500", 1e6 * d / (n * o), 0)
    check("H&S p. 5: 'about 4.3 sigma' (met 1,5 shift)", "4.3", PHI_INV(1 - d / (n * o)) + 1.5, 1)
    show("ex-03-7a: DPO", d / (n * o), "data-answer 0.0025")
    show("ex-03-7b: DPMO", 1e6 * d / (n * o), "data-answer 2500")
    show("ex-03-7c: yield per opportunity = 1 - DPO (%)", 100 * (1 - d / (n * o)), "data-answer 99.75")
    show("ex-03-7d: sigmaniveau met 1,5 shift", PHI_INV(1 - d / (n * o)) + 1.5, "data-answer 4.31")
    show("ex-03-7e: Z zonder verschuiving", PHI_INV(1 - d / (n * o)), "data-answer 2.81")
    show("% naar ppm: 1 % in ppm", 0.01 * 1e6, "SPC p. 20 'factor 10.000'")


def yields() -> None:
    """Dummies p. 147-150 (Y, FTY, hidden factory, RTY), p. 38-39 dice; H&S p. 5 (RTY, NY, units)."""
    section("Yield (Dummies p. 147-150)")
    y, fty = 347 / 352, (352 - 5 - 98) / 352
    check("Dummies p. 147: Y = 347/352", "0.986", y, 3)
    check("Dummies p. 148: FTY = 249/352", "0.707", fty, 3)
    check("Dummies p. 149: verborgen fabriek 98.6 % - 70.7 % = 27.9 %", "27.9", 100 * (y - fty), 1, "MISMATCH")
    show("verborgen fabriek, onafgerond (%)", 100 * (y - fty), "zelf berekend")
    show("ex-03-8a: Y (%)", 100 * y, "data-answer 98.58")
    show("ex-03-8b: FTY (%)", 100 * fty, "data-answer 70.74")
    show("ex-03-8c: verborgen fabriek (procentpunt)", 100 * (y - fty), "data-answer 27.84")
    rty = 0.75 * 0.95 * 0.85 * 0.95 * 0.90
    check("Dummies p. 149: RTY = 0,518", "0.518", rty, 3)
    show("ex-03-9: RTY 5 stappen", rty, "data-answer 0.518")

    section("Dobbelstenen (Dummies p. 38-39): defect = een 1 gooien")
    check("2 dobbelstenen foutloos 69 %", "69", 100 * (5 / 6) ** 2, 0)
    check("3 dobbelstenen foutloos 58 %", "58", 100 * (5 / 6) ** 3, 0)
    show("100 dobbelstenen foutloos", (5 / 6) ** 100, "zelf berekend")
    show("= 1 op ... (miljoen)", 1 / (5 / 6) ** 100 / 1e6, "boek: 'less than one in 82 million'")
    show("ex-03-10: 10 dobbelstenen foutloos (%)", 100 * (5 / 6) ** 10, "data-answer 16.15")

    section("Harry & Schroeder p. 5")
    rty_hs = 0.98 * 0.93 * 0.95 * 0.98 * 0.94
    check("H&S p. 5: RTY = 0.7976", "0.7976", rty_hs, 4)
    ny = 0.368 ** (1 / 10)
    check("H&S p. 5: NY '(0.368)**(-10) = 0.9051'", "0.9051", ny, 4, "MISMATCH")
    show("NY = 0,368^(1/10)", ny, "zelf berekend; data-answer ex-03-11a 0.9049")
    show("0,368^(-10) zoals gedrukt", 0.368 ** -10, "zelf berekend (ongeveer 22 000)")
    check("H&S p. 5: herstelbaar: 1 + (1 - 0,7) = 1.3", "1.3", 1 + (1 - 0.7), 1)
    check("H&S p. 5: afgekeurd: 1/0,7 = 1.43", "1.43", 1 / 0.7, 2)
    show("ex-03-11b: eenheden per foutloze bij RTY 0,7 en afkeur", 1 / 0.7, "data-answer 1.43")


def sigma_tables() -> None:
    """Sigma level <-> DPMO: SPC p. 21, Dummies p. 41 (Table 1-2) and p. 160-161 (Table 6-3), H&S, Van Volsem."""
    section("Sigmaniveau -> DPMO met 1,5 sigma verschuiving: DPMO = 1e6 * P(Z > Z_ST - 1,5)")

    def dpmo(z_st: float) -> float:
        return 1e6 * (1 - PHI(z_st - 1.5))

    for z in (1, 2, 3, 4, 5, 6):
        show(f"Z_ST = {z}: DPMO met verschuiving / zonder (1 staart)",
             f"{dpmo(z):.1f} / {1e6 * (1 - PHI(z)):.4g}", "zelf berekend")
    for row in rows("S06_voc_vs_vop_sigma_capability_defects_per_million_opportunitie"):
        z = float(row["Sigma Capability"])
        expected = "MISMATCH" if z == 2 else "OK"
        check(f"SPC p. 21: {z:g} sigma DPMO", row["Defects per Million Opportunities"], dpmo(z),
              1 if z == 6 else 0, expected)
        decimals = len(row["% Yield"].split(".")[1].rstrip("%"))
        check(f"SPC p. 21: {z:g} sigma % yield", row["% Yield"], 100 - dpmo(z) / 1e4, decimals)
    for row in rows("S07_table_1_2_the_sigma_scale"):
        z = float(row["Sigma"])
        printed = row["Defects per Million"]
        decimals = len(printed.split(".")[1]) if "." in printed else 0
        check(f"Dummies p. 41 Table 1-2: {z:g} sigma DPM", printed, dpmo(z), decimals)
    for row in rows("S08_table_6_3_sigma_score_table_z_dpmo"):
        z = float(row["Z"])
        printed = row["DPMO"]
        decimals = len(printed.split(".")[1]) if "." in printed else 0
        check(f"Dummies p. 160-161 Table 6-3: Z = {z:g}", printed, dpmo(z), decimals)
    for row in rows("S09_the_cost_of_quality_sigma_level_vs_dpmo_vs_cost_of_quality"):
        z = float(row["Sigma Level"])
        printed = row["Defects Per Million Opportunities"].split("(")[0].strip()
        decimals = len(printed.split(".")[1]) if "." in printed else 0
        check(f"H&S samenvatting p. 2: {z:g} sigma", printed, dpmo(z), decimals,
              "MISMATCH" if z == 2 else "OK")
    show("2 sigma onafgerond", dpmo(2), "zelf berekend: 308537,5 -> 308538")
    # Van Volsem p. 7 rounds to 2-3 significant digits; 308,000 is the only row that does not round correctly.
    for printed, z, places in (("690,000", 1, -4), ("308,000", 2, -3), ("66,800", 3, -2), ("230", 5, -1)):
        value = round(dpmo(z), places)
        status = "OK" if value == num(printed) else "MISMATCH"
        expected = "MISMATCH" if z == 2 else "OK"
        if status != expected:
            FAILURES.append(f"Van Volsem {z}: expected {expected}")
        print(f"  {status:<8} Van Volsem p. 7: {z} sigma {printed:>9}  computed {dpmo(z):.1f} rounded {value:.0f}")

    section("Oefeningen sigmaniveau")
    check("Dummies p. 161: DPMO 20.000 -> 'about 3.6'", "3.6", PHI_INV(1 - 0.02) + 1.5, 1)
    show("ex-03-12a: DPMO 20000 -> Z_ST met verschuiving", PHI_INV(1 - 0.02) + 1.5, "data-answer 3.55")
    show("ex-03-12b: DPMO 20000 -> Z zonder verschuiving", PHI_INV(1 - 0.02), "data-answer 2.05")
    show("ex-03-12c: 6210 DPMO -> sigmaniveau met verschuiving", PHI_INV(1 - 0.00621) + 1.5, "data-answer 4.00")
    show("ex-03-12d: sigmaniveau 4,5 met verschuiving -> DPMO", dpmo(4.5), "data-answer 1350")
    check("Dummies p. 160: Z 4,5 -> 1,350", "1,350", dpmo(4.5), 0)
    check("H&S p. 7: 6 sigma korte termijn = 4,5 sigma lange termijn", "4.5", 6 - 1.5, 1)

    section("Harry & Schroeder p. 3: opbrengst per defectkans (product A en B)")
    a = 0.85 ** (1 / 600)
    b = 0.968 ** (1 / 48)
    check("H&S p. 3: A (0.85)^(1/600) = 99.97 %", "99.97", 100 * a, 2)
    check("H&S p. 3: A 'about 3.5 sigma' (zonder verschuiving)", "3.5", PHI_INV(a), 1)
    show("A: DPMO", 1e6 * (1 - a), "zelf berekend")
    show("A: Z met verschuiving", PHI_INV(a) + 1.5, "zelf berekend")
    dpu_a = 600 * (1 - a)
    show("A: defecten per eenheid = 600 * DPO", dpu_a, "zelf berekend")
    show("A: eenheden per defect = 1/DPU", 1 / dpu_a, "zelf berekend; boek zegt 1.2")
    show("A: 1/0,85 (eenheden per goede eenheid)", 1 / 0.85, "zelf berekend")
    check("H&S p. 3: A 'one defect for every 1.2 units'", "1.2", 1 / dpu_a, 1, "MISMATCH")
    check("H&S p. 3: B (0.968)^(1/48) = 99.97 %", "99.97", 100 * b, 2, "MISMATCH")
    check("H&S p. 3: B 'about 3.5 sigma'", "3.5", PHI_INV(b), 1, "MISMATCH")
    show("B: opbrengst per defectkans (%)", 100 * b, "zelf berekend; data-answer ex-03-13b 99.93")
    show("B: Z zonder / met verschuiving", f"{PHI_INV(b):.3f} / {PHI_INV(b) + 1.5:.3f}", "zelf berekend")
    show("ex-03-13a: A opbrengst per defectkans (%)", 100 * a, "data-answer 99.97")
    show("ex-03-13c: B Z zonder verschuiving", PHI_INV(b), "data-answer 3.20")


def main() -> int:
    """Run every block; exit 1 if a check did not have its expected status."""
    sample_means()
    normal_basics()
    z_table()
    excel_workbook()
    discrete_metrics()
    yields()
    sigma_tables()
    print()
    if FAILURES:
        print("FAILED:", *FAILURES, sep="\n  ")
        return 1
    print("All checks have their expected status (MISMATCH lines are the errata in 03_errata.tsv).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
