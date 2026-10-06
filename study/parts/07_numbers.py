"""Deel 07 (Regressie) of the study guide: every number the fragment 07_regressie.html marks "zelf berekend", every
exercise answer (data-answer), and a re-check of every printed course result the fragment quotes.

Data come from inventory/constants (each row carries its course file and page) or, for the exercise workbook,
from source/course/Les 3/20260605_devuyst_Regression_demo.xlsx (read-only, openpyxl). Nothing is taken from memory.
Lines "OK" / "MISMATCH" compare a computation with a printed course value (half-up rounding to the printed
precision, convention decision 10); "MISMATCH (erratum)" marks a known slide error listed in 07_errata.tsv.
Run from the project root:  python3 study/parts/07_numbers.py
"""
from __future__ import annotations

import csv
from decimal import ROUND_HALF_UP, Decimal
from itertools import combinations
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.api as sm
from scipy import stats

ROOT = Path(__file__).resolve().parents[2]
CONSTANTS = ROOT / "inventory" / "constants"
DEMO = ROOT / "source" / "course" / "Les 3" / "20260605_devuyst_Regression_demo.xlsx"
UNEXPECTED: list[str] = []


def rows(stem: str) -> list[dict[str, str]]:
    """Rows of a course table in inventory/constants."""
    with (CONSTANTS / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def show(label: str, value: object) -> None:
    """Print one derived number (or a short text)."""
    text = f"{value:.6g}" if isinstance(value, (float, np.floating)) else str(value)
    print(f"  {label:<58} {text}")


def check(label: str, computed: float, printed: str, erratum: bool = False) -> None:
    """Compare a computed value with a printed one after half-up rounding to the printed precision."""
    places = len(printed.split(".")[1]) if "." in printed else 0
    rounded = Decimal(repr(float(computed))).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    ok = rounded == Decimal(printed)
    status = "OK" if ok else ("MISMATCH (erratum)" if erratum else "MISMATCH")
    if not ok and not erratum:
        UNEXPECTED.append(label)
    print(f"  {status:<19} {label:<52} printed {printed:>12}  computed {float(computed):.6g}")


def check_sci(label: str, computed: float, printed: str) -> None:
    """Compare a p-value with a printed one in scientific notation (same number of mantissa digits)."""
    digits = len(printed.lower().split("e")[0].replace("-", "").replace(".", "")) - 1
    ok = f"{computed:.{digits}e}".lower() == f"{float(printed):.{digits}e}".lower()
    if not ok:
        UNEXPECTED.append(label)
    print(f"  {'OK' if ok else 'MISMATCH':<19} {label:<52} printed {printed:>12}  computed {computed:.6g}")


def check_rounded_input(label: str, r2_percent: float, n: int, k: int, printed: str) -> None:
    """Adjusted R^2 from a printed (rounded) R^2: OK when the printed value lies in the range the rounding allows."""
    half = 0.005  # R-Sq is printed with 2 decimals (in %)
    lo, hi = (100 * (1 - (1 - (r2_percent + d) / 100) * (n - 1) / (n - k - 1)) for d in (-half, half))
    ok = lo - half <= float(printed) <= hi + half
    if not ok:
        UNEXPECTED.append(label)
    print(f"  {'OK' if ok else 'MISMATCH':<19} {label:<52} printed {printed:>12}  computed {lo:.3f} to {hi:.3f}")


def ols(columns: list[np.ndarray], y: np.ndarray):
    """Ordinary least squares with intercept (statsmodels result)."""
    return sm.OLS(y, sm.add_constant(np.column_stack(columns))).fit()


def start() -> None:
    """7.1: the two notations for a critical value (REG p. 32, DOE p. 9)."""
    print("7.1 Critical values in both notations")
    check("t_0.005,18 = t quantile 0.995 (REG p. 32)", stats.t.ppf(0.995, 18), "2.88")
    check("F_0.01,3,20 = F quantile 0.99 (DOE p. 9)", stats.f.ppf(0.99, 3, 20), "4.94")


def oxygen() -> None:
    """7.4-7.7: oxygen purity (REG p. 2, 20-22, 32): fit, ANOVA, tests, intervals."""
    data = rows("S05_table_11_1_oxygen_and_hydrocarbon_levels")
    x = np.array([float(r["Hydrocarbon Level x (%)"]) for r in data])
    y = np.array([float(r["Purity y (%)"]) for r in data])
    n = len(x)
    print("7.4 Oxygen purity: least squares from summary values (REG p. 18, 20, 22, 32)")
    sx, sy, sxx_raw, sxy_raw = x.sum(), y.sum(), (x ** 2).sum(), (x * y).sum()
    for label, value in (("n", n), ("sum x", f"{sx:.2f}"), ("sum y", f"{sy:.2f}"), ("sum x^2", f"{sxx_raw:.4f}"),
                         ("sum xy", f"{sxy_raw:.4f}"), ("x-bar", f"{x.mean():.4f}"), ("y-bar", f"{y.mean():.4f}")):
        show(label, value)
    sxx = sxx_raw - n * x.mean() ** 2  # REG p. 18: Sxx = sum x^2 - n x-bar^2
    sxy = sxy_raw - n * x.mean() * y.mean()  # REG p. 18: SxY = sum xY - n x-bar Y-bar
    show("S_xy = sum xy - n x-bar y-bar", f"{sxy:.5f}")
    check("S_xx (REG p. 32)", sxx, "0.68088")
    b1 = sxy / sxx
    b0 = y.mean() - b1 * x.mean()
    check("b1 = S_xy / S_xx (REG p. 22)", b1, "14.947")
    check("b0 = y-bar - b1 x-bar (REG p. 22)", b0, "74.283")
    check("intercept of the figure caption (REG p. 21)", b0, "74.20", erratum=True)
    check("slope of the figure caption (REG p. 21)", b1, "14.97", erratum=True)
    show("answer ex-07-1: b1 from the rounded sums", f"{sxy / sxx:.3f}")
    show("answer ex-07-1: b0 with b1 rounded to 14.947", f"{y.mean() - 14.947 * x.mean():.3f}")

    print("7.5 ANOVA and R^2 (REG p. 22-28)")
    fit = b0 + b1 * x
    sst, sse = float(np.sum((y - y.mean()) ** 2)), float(np.sum((y - fit) ** 2))
    ssr, mse = sst - sse, sse / (n - 2)
    check("SS_T", sst, "173.38")
    check("SS_R", ssr, "152.13")
    check("SS_E", sse, "21.25")
    check("MS_E", mse, "1.18")
    check("S = sqrt(MS_E)", np.sqrt(mse), "1.087")
    check("R-Sq in % (REG p. 22)", 100 * ssr / sst, "87.7")
    check("R-Sq(adj) in %, (n-1)/(n-2) (REG p. 22)", 100 * (1 - (1 - ssr / sst) * (n - 1) / (n - 2)), "87.1")
    show("R-Sq(adj) with p. 56 print (n-1)/(n-k-2), k = 1, in %", f"{100 * (1 - (1 - ssr / sst) * (n - 1) / (n - 3)):.2f}")
    show("answer ex-07-2: R^2 = 152.13 / 173.38", f"{152.13 / 173.38:.4f}")
    show("answer ex-07-2: MS_E = 21.25 / 18", f"{21.25 / 18:.4f}")
    show("answer ex-07-2: S = sqrt(21.25 / 18)", f"{np.sqrt(21.25 / 18):.4f}")

    print("7.6 Tests and confidence intervals (REG p. 29-34)")
    se_b1, se_b0 = np.sqrt(mse / sxx), np.sqrt(mse * (1 / n + x.mean() ** 2 / sxx))
    check("se(b1)", se_b1, "1.317")
    check("se(b0)", se_b0, "1.593")
    check("t0 slope", b1 / se_b1, "11.35")
    check("t0 intercept", b0 / se_b0, "46.62")
    check("F0 = MS_R / MS_E", ssr / mse, "128.86")
    show("t0(slope)^2", (b1 / se_b1) ** 2)
    p_slope = 2 * stats.t.sf(b1 / se_b1, n - 2)
    print(f"  {'OK' if f'{p_slope:.2e}' == '1.23e-09' else 'MISMATCH':<19} {'P of t0 slope (REG p. 32: 1.23 x 10^-9)':<52} computed {p_slope:.3e}")
    if f"{p_slope:.2e}" != "1.23e-09":
        UNEXPECTED.append("P of t0 slope")
    t975 = stats.t.ppf(0.975, n - 2)
    show("t quantile 0.975, df 18", f"{t975:.3f}")
    show("95 % CI beta1", f"[{b1 - t975 * se_b1:.3f}, {b1 + t975 * se_b1:.3f}]")
    show("95 % CI beta0", f"[{b0 - t975 * se_b0:.3f}, {b0 + t975 * se_b0:.3f}]")
    t_ex = 14.947 / np.sqrt(1.18 / 0.68088)
    show("answer ex-07-3: t0 = 14.947 / sqrt(1.18 / 0.68088)", f"{t_ex:.3f}")
    show("answer ex-07-3: t0^2", f"{t_ex ** 2:.2f}")
    show("ex-07-3: se(b1) from rounded MS_E 1.18 / 11.35^2", f"{np.sqrt(1.18 / 0.68088):.4f} / {11.35 ** 2:.2f}")
    show("ex-07-4: t * se = 2.101 * 1.317", f"{2.101 * 1.317:.3f}")
    show("answer ex-07-4: CI beta1 with printed b1, SE 1.317, t 2.101",
         f"[{14.947 - 2.101 * 1.317:.3f}, {14.947 + 2.101 * 1.317:.3f}]")

    print("7.7 CI of the mean response and prediction interval at x0 = 1.00 (REG p. 22, 35-39)")
    x0 = 1.00
    y0 = b0 + b1 * x0
    se_fit = np.sqrt(mse * (1 / n + (x0 - x.mean()) ** 2 / sxx))
    se_pred = np.sqrt(mse * (1 + 1 / n + (x0 - x.mean()) ** 2 / sxx))
    check("Fit", y0, "89.231")
    check("SE Fit", se_fit, "0.354")
    check("95% CI lower", y0 - t975 * se_fit, "88.486")
    check("95% CI upper", y0 + t975 * se_fit, "89.975")
    check("95% PI lower", y0 - t975 * se_pred, "86.830")
    check("95% PI upper", y0 + t975 * se_pred, "91.632")
    show("se(e0) = sqrt(MS_E (1 + 1/n + (x0-xbar)^2/Sxx))", f"{se_pred:.4f}")
    # Exercise ex-07-5 with the rounded values a student has: b0, b1, MS_E 1.18, x-bar 1.196, Sxx 0.68088, t 2.101.
    y0r = 74.283 + 14.947 * x0
    sf = np.sqrt(1.18 * (1 / 20 + (x0 - 1.196) ** 2 / 0.68088))
    sp = np.sqrt(1.18 * (1 + 1 / 20 + (x0 - 1.196) ** 2 / 0.68088))
    show("answer ex-07-5: fit from rounded b0, b1", f"{y0r:.3f}")
    show("ex-07-5: (x0 - x-bar)^2 / Sxx and 1 + 1/n + that", f"{(x0 - 1.196) ** 2 / 0.68088:.4f} / {1 + 1 / 20 + (x0 - 1.196) ** 2 / 0.68088:.4f}")
    show("answer ex-07-5: se(fit) / se(e0) from rounded values", f"{sf:.4f} / {sp:.4f}")
    show("answer ex-07-5: CI", f"[{y0r - 2.101 * sf:.3f}, {y0r + 2.101 * sf:.3f}]")
    show("answer ex-07-5: PI", f"[{y0r - 2.101 * sp:.3f}, {y0r + 2.101 * sp:.3f}]")


def heating_oil() -> None:
    """7.10-7.11: oil consumption, multiple regression (REG p. 51-57)."""
    data = rows("S05_oil_consumption_example_data_for_15_houses")
    y = np.array([float(r["Oil (Gal)"].replace(",", ".")) for r in data])
    temp, ins = (np.array([float(r[c]) for r in data]) for c in ("Temp", "Insulation"))
    n, k = len(y), 2
    res = ols([temp, ins], y)
    print("7.10 Heating oil (REG p. 51-57)")
    printed = {r["Item"]: r["Value"] for r in rows("S05_oil_consumption_example_excel_regression_statistics_and_anov")}
    r2 = res.rsquared
    check("Multiple R", np.sqrt(r2), printed["Multiple R"])
    check("R Square", r2, printed["R Square"])
    check("Adjusted R Square, (n-1)/(n-k-1)", 1 - (1 - r2) * (n - 1) / (n - k - 1), printed["Adjusted R Square"])
    check("Standard Error", np.sqrt(res.mse_resid), printed["Standard Error"])
    check("Regression SS", res.ess, printed["Regression SS"])
    check("Residual SS", res.ssr, printed["Residual SS"])
    check("Total SS", res.centered_tss, printed["Total SS"])
    check("Regression MS", res.ess / 2, printed["Regression MS"])
    check("Residual MS", res.mse_resid, printed["Residual MS"])
    # The inventory CSV reads 168.47203; the slide (REG p. 52) prints 168.471203, which is used here.
    check("F (REG p. 52 prints 168.471203)", res.fvalue, "168.471203")
    check_sci("Significance F", res.f_pvalue, printed["Regression Significance F"])
    coef = {r["Variable"]: r for r in rows("S05_oil_consumption_example_excel_coefficients_table")}
    ci = res.conf_int(0.05)
    for i, name in enumerate(("Intercept", "Temp", "Insulation")):
        check(f"{name} coefficient", res.params[i], coef[name]["Coefficients"])
        check(f"{name} standard error", res.bse[i], coef[name]["Standard Error"])
        check(f"{name} t Stat", res.tvalues[i], coef[name]["t Stat"])
        check(f"{name} Lower 95%", ci[i, 0], coef[name]["Lower 95%"])
        check(f"{name} Upper 95%", ci[i, 1], coef[name]["Upper 95%"])
        check_sci(f"{name} P-value", res.pvalues[i], coef[name]["P-value"])
    show("R Square in %", f"{100 * r2:.1f}")
    show("ex-07-6 steps: 5.4365806 * 30 / 20.012321 * 6", f"{5.4365806 * 30:.3f} / {20.012321 * 6:.3f}")
    pred = float(res.params @ [1, 30, 6])
    check("prediction at 30 F, 6 inches (REG p. 57)", pred, "278.97", erratum=True)
    show("answer ex-07-6: prediction with exact coefficients", f"{pred:.2f}")
    show("prediction with 562.15, 5.44, 20.01 (p. 53 equation)", f"{562.15 - 5.44 * 30 - 20.01 * 6:.2f}")
    show("prediction with the slide's 5.43 (p. 57 line)", f"{562.15 - 5.43 * 30 - 20.01 * 6:.2f}")
    check("sigma-hat = sqrt(8120.6 / (15 - 2 - 1)) (REG p. 57)", np.sqrt(8120.6 / 12), "26.0")
    show("answer ex-07-6: sigma-hat", f"{np.sqrt(8120.6 / 12):.3f}")
    print("7.11 Adjusted R^2: the course's outputs versus the formula printed on REG p. 56")
    show("adj R^2 with (n-1)/(n-k-1) = 1 - (1 - 0.96561)*14/12", f"{1 - (1 - 0.96561) * 14 / 12:.5f}")
    show("adj R^2 with (n-1)/(n-k-2) = 1 - (1 - 0.96561)*14/11", f"{1 - (1 - 0.96561) * 14 / 11:.4f}")


def acetylene() -> None:
    """7.12: polynomial model and partial F-test (REG p. 58-62)."""
    data = rows("S05_table_6_9_the_acetylene_data")
    y = np.array([float(r["Yield Y"]) for r in data])
    t, ratio = (np.array([float(r[c]) for r in data]) for c in ("Temp T", "Ratio R"))
    tc, rc = t - 1212.5, ratio - 12.444  # centred as on REG p. 60
    print("7.12 Acetylene (REG p. 58-62)")
    show("mean T / mean R (centring values on p. 60: 1212.5, 12.444)", f"{t.mean():.4f} / {ratio.mean():.5f}")
    red = ols([tc, rc], y)
    check("first order: Constant (p. 59)", red.params[0], "36.1063")
    check("first order: Temp (p. 59)", red.params[1], "0.13396")
    check("first order: Ratio (p. 59)", red.params[2], "0.3511")
    check("first order: S (p. 59)", np.sqrt(red.mse_resid), "3.624")
    check("first order: R-Sq % (p. 59)", 100 * red.rsquared, "92.0")
    check("first order: R-Sq(adj) % (p. 59)", 100 * red.rsquared_adj, "90.7")
    check("first order: SS_E (p. 59)", red.ssr, "170.73")
    check("first order: F (p. 59)", red.fvalue, "74.35")
    full = ols([tc, rc, tc * rc, tc ** 2, rc ** 2], y)
    for i, (name, value) in enumerate((("Constant", "36.4339"), ("Temp", "0.130476"), ("Ratio", "0.48005"),
                                        ("T x R", "-0.0073346"), ("T^2", "0.00017820"), ("R^2", "-0.02367"))):
        check(f"second order: {name} (p. 62)", full.params[i], value)
    check("second order: S (p. 62)", np.sqrt(full.mse_resid), "1.066")
    check("second order: R-Sq % (p. 62)", 100 * full.rsquared, "99.5")
    check("second order: R-Sq(adj) % (p. 62)", 100 * full.rsquared_adj, "99.2")
    check("second order: SS_E (p. 62)", full.ssr, "11.37")
    check("second order: F (p. 62)", full.fvalue, "371.49")
    f0 = ((170.73 - 11.37) / 3) / (11.37 / (16 - 6))
    show("ex-07-7: numerator (170.73 - 11.37)/3 / denominator 11.37/10", f"{(170.73 - 11.37) / 3:.2f} / {11.37 / 10:.3f}")
    show("answer ex-07-7: partial F0 from printed SS_E, df (3, 10)", f"{f0:.2f}")
    show("answer ex-07-7: F crit 0.05 (3, 10)", f"{stats.f.isf(0.05, 3, 10):.3f}")
    show("p-value of the partial F0", f"{stats.f.sf(f0, 3, 10):.2e}")


def shampoo() -> None:
    """7.13: adjusted R^2 of the selected shampoo models from their R^2 (REG p. 63-66), N = 24."""
    print("7.13 Shampoo: R-Sq(adj) from R-Sq with (N-1)/(N-k-1), N = 24")
    check_rounded_input("Foam, Residue, Region: R-Sq 77.57 (p. 64)", 77.57, 24, 3, "74.20")
    check_rounded_input("Foam, Scent, Residue, Region: R-Sq 79.89 (p. 65)", 79.89, 24, 4, "75.65")
    check_rounded_input("all five: R-Sq 80.01 (p. 64)", 80.01, 24, 5, "74.45")


def demo_simple(sheet: str) -> None:
    """7.14: workbook sheets 'Simple 1' and 'Simple 2': choose a polynomial degree, predict at 2.5 and 10."""
    ws = openpyxl.load_workbook(DEMO, data_only=True, read_only=True)[sheet]
    pairs = [(r[1], r[2]) for r in ws.iter_rows(min_row=2, max_row=51, values_only=True)]
    x = np.array([p[0] for p in pairs], float)
    y = np.array([p[1] for p in pairs], float)
    print(f"7.14 Regression_demo.xlsx '{sheet}' (n = {len(x)}); x from {x.min():.2f} to {x.max():.2f}")
    fits = {}
    for degree in (1, 2, 3, 4):
        res = ols([x ** d for d in range(1, degree + 1)], y)
        fits[degree] = res
        terms = ", ".join(f"{b:.4f}" for b in res.params)
        pvals = ", ".join(f"{p:.4f}" for p in res.pvalues)
        show(f"degree {degree}: R^2 / adj R^2 / S", f"{res.rsquared:.4f} / {res.rsquared_adj:.4f} / {np.sqrt(res.mse_resid):.3f}")
        show(f"degree {degree}: coefficients b0..b{degree}", terms)
        show(f"degree {degree}: p-values of b0..b{degree}", pvals)
        new = np.array([[1] + [v ** d for d in range(1, degree + 1)] for v in (2.5, 10.0)])
        frame = res.get_prediction(new).summary_frame(alpha=0.05)
        for i, v in enumerate((2.5, 10.0)):
            row = frame.iloc[i]
            show(f"degree {degree}: x0 = {v}: fit, 95 % CI, 95 % PI",
                 f"{row['mean']:.2f}; [{row['mean_ci_lower']:.2f}, {row['mean_ci_upper']:.2f}];"
                 f" [{row['obs_ci_lower']:.2f}, {row['obs_ci_upper']:.2f}]")
    for low, high in ((1, 2), (2, 3)):
        f0 = (fits[low].ssr - fits[high].ssr) / fits[high].mse_resid
        show(f"partial F degree {low} -> {high}, df (1, {int(fits[high].df_resid)}), p",
             f"{f0:.3f}, p = {stats.f.sf(f0, 1, fits[high].df_resid):.3g}")
    b = fits[3].params
    turning = np.sort(np.roots([3 * b[3], 2 * b[2], b[1]]).real)
    show("cubic fit: x where the slope is zero (local max / min)", ", ".join(f"{t:.2f}" for t in turning))
    resid = fits[1].resid
    edges = [-5, -3, -1, 1, 3, 5]
    spread = []
    for lo, hi in zip(edges[:-1], edges[1:]):
        mask = (x >= lo) & (x < hi)
        spread.append(f"[{lo},{hi}): n {mask.sum()}, mean {resid[mask].mean():.1f}, sd {resid[mask].std(ddof=1):.1f}")
    show("straight-line residuals per x-interval", "; ".join(spread))


def demo_multiple() -> None:
    """7.14: workbook sheet 'Multiple 1': Strength from A..E, full model, subsets, prediction."""
    ws = openpyxl.load_workbook(DEMO, data_only=True, read_only=True)["Multiple 1"]
    data = pd.DataFrame([r[:6] for r in ws.iter_rows(min_row=2, max_row=31, values_only=True)],
                        columns=["A", "B", "C", "D", "E", "Strength"], dtype=float)
    print(f"7.14 Regression_demo.xlsx 'Multiple 1' (n = {len(data)})")
    for col in "ABCDE":
        show(f"range of {col}", f"{data[col].min():g} to {data[col].max():g}")
    show("range of Strength", f"{data['Strength'].min():g} to {data['Strength'].max():g}")
    corr = data.corr()
    show("correlations A-B / A-C / B-C", f"{corr.loc['A', 'B']:.3f} / {corr.loc['A', 'C']:.3f} / {corr.loc['B', 'C']:.3f}")
    total = data["A"] + data["C"]
    show("A + C: number of rows equal to 100 / min / max", f"{int((total == 100).sum())} / {total.min():g} / {total.max():g}")
    y = data["Strength"].to_numpy()

    def fit(cols: tuple[str, ...]):
        return ols([data[c].to_numpy() for c in cols], y)

    full = fit(tuple("ABCDE"))
    show("full model: R^2 / adj R^2 / S", f"{full.rsquared:.4f} / {full.rsquared_adj:.4f} / {np.sqrt(full.mse_resid):.1f}")
    show("full model: F / p-value (df 5, 24)", f"{full.fvalue:.3f} / {full.f_pvalue:.4f}")
    for i, name in enumerate(["Intercept", *"ABCDE"]):
        show(f"full model {name}: coef / se / t / p",
             f"{full.params[i]:.2f} / {full.bse[i]:.2f} / {full.tvalues[i]:.3f} / {full.pvalues[i]:.4f}")
    best = sorted(((fit(c).rsquared_adj, c) for k in range(1, 6) for c in combinations("ABCDE", k)), reverse=True)
    for adj, cols in best[:3]:
        res = fit(cols)
        show(f"subset {''.join(cols)}: adj R^2 / R^2 / F p-value", f"{adj:.4f} / {res.rsquared:.4f} / {res.f_pvalue:.4f}")
        show(f"subset {''.join(cols)}: p-values", ", ".join(f"{c} {p:.3f}" for c, p in zip(cols, res.pvalues[1:])))
    cols, step = list("ABCDE"), 1
    while cols:  # backward elimination with alpha-to-remove 0.1, as on REG p. 64
        res = fit(tuple(cols))
        pvals = dict(zip(cols, res.pvalues[1:]))
        worst = max(pvals, key=pvals.get)
        show(f"backward step {step}: model {''.join(cols)}, largest p", f"{worst} {pvals[worst]:.4f}")
        if pvals[worst] <= 0.1:
            break
        cols.remove(worst)
        step += 1
    show("backward elimination ends with", "".join(cols) or "no variable left")
    for cols in (tuple("ABCDE"), tuple("BCDE")):
        res = fit(cols)
        point = {"A": 15, "B": 220, "C": 85, "D": 580, "E": 185}
        frame = res.get_prediction(np.array([[1] + [point[c] for c in cols]])).summary_frame(alpha=0.05).iloc[0]
        show(f"model {''.join(cols)}: prediction at A=15, B=220, C=85, D=580, E=185",
             f"{frame['mean']:.0f}; 95 % PI [{frame['obs_ci_lower']:.0f}, {frame['obs_ci_upper']:.0f}]")


def main() -> None:
    """Print the numbers of Deel 07 in the order of its units."""
    start()
    oxygen()
    heating_oil()
    acetylene()
    shampoo()
    demo_simple("Simple 1")
    demo_simple("Simple 2")
    demo_multiple()
    print("unexpected mismatches:", UNEXPECTED or "none")


if __name__ == "__main__":
    main()
