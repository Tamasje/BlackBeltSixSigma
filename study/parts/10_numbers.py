"""Deel 10 (SPC en regelkaarten): every number the fragment 10_spc.html marks "zelf berekend", every exercise
answer (data-answer), and a re-check of every printed course result the text quotes (OK / MISMATCH lines).

Sources (page = PDF page; for the Les 4 deck PDF page = slide number, the footer shows one less):
- deck "2026 Lean - Six Sigma v13 - Capabiliteit - SPC.pdf" (SPC p. N);
- the lecturer's workbooks "__Gegevens oefeningen.xlsx" and "__Xbar R kaart data - berekeningen oefening 2/3/5.xlsx"
  (read read-only with openpyxl; data_only=True gives the values Excel cached when the lecturer saved them);
- constants: inventory/constants (Table 18 and Table A from "___4.1 tabellen SPC.pdf", Six Sigma Demystified from
  "Control charts - constants.pdf", Dummies Table 10-2);
- Six Sigma For Dummies (PDF page = printed page + 18): the u-chart data come from inventory/worked_examples.json
  (S08-WE21, Dummies p. 256); chart read-outs are transcribed below with their page.
Known printed errors print MISMATCH on purpose; the script still exits 0. Run from anywhere:
    python3 study/parts/10_numbers.py
"""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

ROOT = Path(__file__).resolve().parents[2]
LES4 = ROOT / "source" / "course" / "Les 4"
CONSTANTS = ROOT / "inventory" / "constants"
MISMATCHES: list[str] = []


# ----------------------------------------------------------------------------------------------- helpers
def table(stem: str, key: str) -> dict[str, dict[str, str]]:
    """Rows of a course constants table (inventory/constants), indexed by the sample-size column."""
    with (CONSTANTS / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
        return {row[key]: row for row in csv.DictReader(handle)}


T18 = table("S06_table_18_factors_for_computing_control_chart_lines", "Sample size n")
TA = table("S06_table_a_bias_correction_factors_for_estimating_standard_devi", "Subgroup Size")
SSD_X = table("S06_control_chart_constants_chart_for_average_chart_for_standard", "Observations in Sample n")
SSD_R = table("S06_control_chart_constants_chart_for_ranges_x_charts_six_sigma", "Observations in Sample n")
DUM = table("S08_table_10_2_continuous_data_control_chart_constants_a2_a3_b3", "Sample Size (n)")


def t18(name: str, n: int) -> float:
    """A Table 18 constant (decision 4: the constants the lecturer's workbooks use)."""
    return float(T18[str(n)][name])


def show(label: str, value: float | str, note: str = "") -> None:
    """Print one result line."""
    text = f"{value:.6g}" if isinstance(value, (float, np.floating)) else str(value)
    print(f"  {label:<58} {text:>16}  {note}")


def check(label: str, computed: float, printed: float, decimals: int) -> None:
    """Compare a computed value with a printed course value rounded half-up to `decimals` (decision 10)."""
    factor = 10 ** decimals
    rounded = math.floor(computed * factor + 0.5) / factor
    ok = abs(rounded - printed) < 0.5 / factor
    if not ok:
        MISMATCHES.append(label)
    print(f"{'OK      ' if ok else 'MISMATCH'} {label}: computed {computed:.8g} -> {rounded:.{decimals}f}, "
          f"printed {printed:.{decimals}f}")


def check_cell(label: str, computed: float, cached: float) -> None:
    """Compare with a value the lecturer's workbook cached (relative tolerance 1e-9)."""
    ok = math.isclose(computed, cached, rel_tol=1e-9, abs_tol=1e-15)
    if not ok:
        MISMATCHES.append(label)
    print(f"{'OK      ' if ok else 'MISMATCH'} {label}: computed {computed:.12g}, workbook {cached:.12g}")


def sheet(name: str, sheet_name: str, formulas: bool = False):
    """A worksheet of a Les 4 workbook, read-only (values cached by Excel, or the formulas)."""
    wb = openpyxl.load_workbook(LES4 / name, data_only=not formulas)
    return wb[sheet_name]


def we_rules(z: list[float]) -> dict[str, list[int]]:
    """Western Electric rules 1-4 and sensitizing rule 5 (SPC p. 68) on points given in sigma units of the plotted
    statistic (z = (point - centre line) / sigma). Returns, per rule, the 1-based numbers of the points at which the
    rule signals (the last point of the window). Rules 2-3: points on the same side (Dummies p. 246 wording);
    a point exactly on the centre line breaks a run (rule 4)."""
    out: dict[str, list[int]] = {"1": [], "2": [], "3": [], "4": [], "5": []}
    for i in range(len(z)):
        if abs(z[i]) > 3:
            out["1"].append(i + 1)
        for side in (1, -1):
            w3 = z[max(0, i - 2): i + 1]
            if len(w3) == 3 and sum(1 for v in w3 if side * v > 2) >= 2:
                out["2"].append(i + 1)
            w5 = z[max(0, i - 4): i + 1]
            if len(w5) == 5 and sum(1 for v in w5 if side * v > 1) >= 4:
                out["3"].append(i + 1)
            w8 = z[max(0, i - 7): i + 1]
            if len(w8) == 8 and all(side * v > 0 for v in w8):
                out["4"].append(i + 1)
        w6 = z[max(0, i - 5): i + 1]
        if len(w6) == 6:
            d = np.diff(w6)
            if np.all(d > 0) or np.all(d < 0):
                out["5"].append(i + 1)
    return {k: sorted(set(v)) for k, v in out.items()}


def zone(v: float) -> str:
    """Zone of a point in sigma units (SPC p. 68-69: A between 2 and 3 sigma, B 1-2, C within 1)."""
    side = "boven" if v > 0 else "onder"
    a = abs(v)
    name = "C" if a <= 1 else "B" if a <= 2 else "A" if a <= 3 else "buiten"
    return f"{name} {side}"


# ----------------------------------------------------------------------------------------------- constants
print("== Constants used (decision 4: Table 18; c4 from Table A; A3, E2 from Six Sigma Demystified)")
for n in (2, 3, 4, 5, 6, 7, 8, 9, 10):
    show(f"n={n}: A2, D3, D4, d2 (T18) | A3, E2 (SSD) | c4 (TA)",
         f"{t18('A2', n)}, {t18('D3', n)}, {t18('D4', n)}, {t18('d2', n)} | "
         f"{SSD_X[str(n)]['A3']}, {SSD_R[str(n)]['E2']} | {TA[str(n)]['c4']}")
print("Printed sources disagree in the third decimal (decision 4):")
for n in (3, 5):
    show(f"D4({n}): T18 / SSD / Dummies", f"{T18[str(n)]['D4']} / {SSD_R[str(n)]['D4']} / {DUM[str(n)]['D4']}")
show("E2(2): SSD / Dummies", f"{SSD_R['2']['E2']} / {DUM['2']['E2']}")
print("A2 = 3/(d2 sqrt n) follows from SPC p. 74 (UCL = X + 3 sigma/sqrt n, sigma = R/d2):")
for n in (2, 3, 4, 5, 6, 8, 10):
    show(f"n={n}: 3/(d2 sqrt n) vs T18 A2", f"{3 / (t18('d2', n) * math.sqrt(n)):.4f} vs {t18('A2', n):.3f}")

print("Table 18 has no A3; Table 7.2 (same file p. 2) writes X +/- A1 s sqrt((n-1)/n):")
for n in (2, 3, 4, 5, 6, 8, 10):
    show(f"n={n}: A1 sqrt((n-1)/n) vs SSD A3", f"{t18('A1', n) * math.sqrt((n - 1) / n):.4f} vs {SSD_X[str(n)]['A3']}")

# ----------------------------------------------------------------------------------------------- SPC p. 64
print("\n== SPC p. 64: X-bar chart with mu = 1.5, sigma = 0.15, n = 5 (figure values)")
mu, sig, n = 1.5, 0.15, 5
s_xbar = sig / math.sqrt(n)
check("sigma_xbar = 0.15/sqrt 5", s_xbar, 0.0671, 4)
check("UCL = 1.5 + 3 sigma_xbar (exact)", mu + 3 * s_xbar, 1.7013, 4)
check("LCL = 1.5 - 3 sigma_xbar (exact)", mu - 3 * s_xbar, 1.2987, 4)
show("UCL / LCL with the rounded sigma_xbar 0.0671", f"{mu + 3 * 0.0671:.4f} / {mu - 3 * 0.0671:.4f}",
     "-> reproduces the printed 1.7013 / 1.2987")
show("ANSWER ex-10-2: sigma_xbar", s_xbar)
show("ANSWER ex-10-2: UCL", mu + 3 * s_xbar)
show("ANSWER ex-10-2: LCL", mu - 3 * s_xbar)

# ----------------------------------------------------------------------------------------------- SPC p. 73
print("\n== SPC p. 73: Xbar-R Chart of Oil (read-outs UCL 31.59, X 28.46, LCL 25.32, R 3.06, UCL_R 7.88)")
ucl, cl, lcl, rbar, ucl_r = 31.59, 28.46, 25.32, 3.06, 7.88
show("(UCL - LCL)/(2 R) = implied A2", (ucl - lcl) / (2 * rbar))
show("UCL_R / R = implied D4", ucl_r / rbar)
for n in (2, 3, 4, 5):
    show(f"T18 n={n}: A2, D4", f"{t18('A2', n)}, {t18('D4', n)}")
show("ANSWER ex-10-3: subgroup size n", 3)
show("midpoint of UCL and LCL (the unrounded X)", (ucl + lcl) / 2)
# The read-outs are rounded by Minitab (the midpoint of UCL and LCL is 28.455), so only show the reconstruction.
show("Oil UCL, LCL from 28.46 +/- 1.023 x 3.06", f"{cl + t18('A2', 3) * rbar:.3f}, {cl - t18('A2', 3) * rbar:.3f}",
     "(printed 31.59, 25.32: rounded read-outs)")
show("Oil UCL_R = 2.575 x 3.06", t18("D4", 3) * rbar, "(printed 7.88)")

# ----------------------------------------------------------------------------------------------- oefening 2
print("\n== SPC oefening 2 (SPC p. 75; data '__Gegevens oefeningen.xlsx' sheet 'Gegevens oefening 2' B5:F24)")
ws = sheet("__Gegevens oefeningen.xlsx", "Gegevens oefening 2")
x2 = np.array([[ws.cell(r, c).value for c in range(2, 7)] for r in range(5, 25)], float)
means2 = x2.mean(axis=1)
ranges2 = x2.max(axis=1) - x2.min(axis=1)
xbb2, rbar2 = means2.mean(), ranges2.mean()
a2, d3, d4, d2 = t18("A2", 5), t18("D3", 5), t18("D4", 5), t18("d2", 5)
ucl2, lcl2 = xbb2 + a2 * rbar2, xbb2 - a2 * rbar2
uclr2, lclr2 = d4 * rbar2, d3 * rbar2
sig2 = rbar2 / d2
for i in range(20):
    show(f"subgroup {i + 1}: mean, range", f"{means2[i]:.2f}, {ranges2[i]:.1f}")
show("ANSWER X-double-bar", xbb2)
show("ANSWER R-bar", rbar2)
show("ANSWER UCL_xbar = X + A2 R", ucl2)
show("ANSWER LCL_xbar", lcl2)
show("ANSWER UCL_R = D4 R", uclr2)
show("LCL_R = D3 R", lclr2)
show("points outside X-bar limits", str([i + 1 for i in range(20) if not lcl2 <= means2[i] <= ucl2]))
show("points outside R limits", str([i + 1 for i in range(20) if not lclr2 <= ranges2[i] <= uclr2]))
show("largest / smallest subgroup mean", f"{means2.max():.2f} (#{means2.argmax() + 1}) / "
     f"{means2.min():.2f} (#{means2.argmin() + 1})")
show("largest range", f"{ranges2.max():.1f} (#{ranges2.argmax() + 1})")
s_x2 = a2 * rbar2 / 3  # sigma of the plotted means implied by the 3-sigma limits
z2 = [(m - xbb2) / s_x2 for m in means2]
show("sigma of the means = A2 R / 3", s_x2)
show("WE rules on the 20 means (signal points)", str(we_rules(z2)))
show("zones of the 20 means", str([zone(v) for v in z2]))
show("ANSWER sigma = R/d2", sig2)
show("6 sigma", 6 * sig2)
cp2 = (16.7 - 15.7) / (6 * sig2)
cpl2, cpu2 = (xbb2 - 15.7) / (3 * sig2), (16.7 - xbb2) / (3 * sig2)
show("ANSWER Cp = 1/(6 sigma)", cp2)
show("Cpk lower / upper", f"{cpl2:.6f} / {cpu2:.6f}")
show("ANSWER Cpk", min(cpl2, cpu2))
ws_b = sheet("__Xbar R kaart data - berekeningen oefening 2.xlsx", "Berekening")
check_cell("oef 2 X-double-bar (H26)", xbb2, ws_b["H26"].value)
check_cell("oef 2 R-bar (L26)", rbar2, ws_b["L26"].value)
check_cell("oef 2 UCL (B31)", ucl2, ws_b["B31"].value)
check_cell("oef 2 LCL (B32)", lcl2, ws_b["B32"].value)
check_cell("oef 2 UCL_R (B36)", uclr2, ws_b["B36"].value)
check_cell("oef 2 sigma = R/d2 (L32)", sig2, ws_b["L32"].value)
check_cell("oef 2 Cp (B44)", cp2, ws_b["B44"].value)
check_cell("oef 2 Cpk (B45)", min(cpl2, cpu2), ws_b["B45"].value)
s_all = float(np.std(x2, ddof=1))
check_cell("oef 2 STDEV of all 100 values (F26)", s_all, ws_b["F26"].value)
s_means = float(np.std(means2, ddof=1))
check_cell("oef 2 STDEV of the 20 means (R31)", s_means, ws_b["R31"].value)
check_cell("oef 2 W31 = R31*SQRT(5)", s_means * math.sqrt(5), ws_b["W31"].value)
f_w32 = sheet("__Xbar R kaart data - berekeningen oefening 2.xlsx", "Berekening", formulas=True)["W32"].value
show("W32 formula as typed", f_w32)
check_cell("oef 2 W32 cached = R31*5/2 (Excel reads (5)^1/2 as 5/2)", s_means * 5 / 2, ws_b["W32"].value)
check("oef 2 W32 'sigma individuele metingen' vs intended R31*sqrt(5)", s_means * math.sqrt(5),
      round(ws_b["W32"].value, 5), 5)
ws_g = sheet("__Xbar R kaart data - berekeningen oefening 2.xlsx", "Grafieken")
usl_c, lsl_c = 16.2 + 0.5 / math.sqrt(5), 16.2 - 0.5 / math.sqrt(5)
check_cell("oef 2 'USL corr' (Grafieken I3) = 16.2 + 0.5/sqrt 5", usl_c, ws_g["I3"].value)
check_cell("oef 2 'LSL corr' (Grafieken J3) = 16.2 - 0.5/sqrt 5", lsl_c, ws_g["J3"].value)
print("Minitab output for oefening 2 (hidden slides SPC p. 79 and 81):")
# Minitab draws X +/- 3 sigma_within/sqrt(n) (the variable-n formula of SPC p. 74), not X +/- 0.577 R.
show("UCL / LCL with A2 = 0.577 (workbook)", f"{ucl2:.5f} / {lcl2:.5f}")
check("Minitab UCL 16.5429 = X + 3 (R/d2)/sqrt 5", xbb2 + 3 * sig2 / math.sqrt(5), 16.5429, 4)
check("Minitab LCL 15.9891 = X - 3 (R/d2)/sqrt 5", xbb2 - 3 * sig2 / math.sqrt(5), 15.9891, 4)
check("Minitab UCL_R 1.015", uclr2, 1.015, 3)
check("Minitab StDev(within) 0.20636 = R/d2", sig2, 0.20636, 5)
check("Minitab StDev(overall) 0.20212", s_all, 0.20212, 5)
pp, ppk = 1 / (6 * s_all), min(xbb2 - 15.7, 16.7 - xbb2) / (3 * s_all)
check("Minitab Pp 0.82", pp, 0.82, 2)
check("Minitab Ppk 0.72", ppk, 0.72, 2)
show("ANSWER Pp (overall s)", pp)
show("ANSWER Ppk (overall s)", ppk)
out_overall = stats.norm.cdf((15.7 - xbb2) / s_all) + stats.norm.sf((16.7 - xbb2) / s_all)
check("Minitab % out of spec expected (overall) 1.84", 100 * out_overall, 1.84, 2)
s_within_mtb = 0.20982  # printed on SPC p. 81, not computed by the deck
check("Minitab Cp 0.79 with StDev(within) 0.20982", 1 / (6 * s_within_mtb), 0.79, 2)
check("Minitab Cpk 0.69 with StDev(within) 0.20982", (16.7 - xbb2) / (3 * s_within_mtb), 0.69, 2)
out_within = stats.norm.cdf((15.7 - xbb2) / s_within_mtb) + stats.norm.sf((16.7 - xbb2) / s_within_mtb)
check("Minitab % out of spec expected (within) 2.28", 100 * out_within, 2.28, 2)
sp = math.sqrt(float(np.mean(x2.var(axis=1, ddof=1))))
show("pooled sd of the 20 subgroups sqrt(mean s_i^2)", sp, "(compare 0.20982 on SPC p. 81)")
show("ANSWER % out of spec with R/d2 sigma (Cp/Cpk basis)",
     100 * (stats.norm.cdf((15.7 - xbb2) / sig2) + stats.norm.sf((16.7 - xbb2) / sig2)))

# ----------------------------------------------------------------------------------------------- oefening 3
print("\n== SPC oefening 3 (SPC p. 84; '__Gegevens oefeningen.xlsx' sheet 'Gegevens oefening 3'; coded x' = x/10000 + 0.5)")
ws3 = sheet("__Gegevens oefeningen.xlsx", "Gegevens oefening 3")
xb3 = np.array([ws3.cell(r, 2).value for r in range(2, 26)], float)  # coded (last 3 decimals)
r3 = np.array([ws3.cell(r, 4).value for r in range(2, 26)], float)


def xr_limits(xb: np.ndarray, r: np.ndarray, n: int) -> tuple[float, float, float, float, float, float]:
    """X-double-bar, R-bar and the X-bar and R chart limits (SPC p. 74, Table 18 constants)."""
    xbb, rb = float(xb.mean()), float(r.mean())
    return xbb, rb, xbb + t18("A2", n) * rb, xbb - t18("A2", n) * rb, t18("D4", n) * rb, t18("D3", n) * rb


xbb3, rb3, ucl3, lcl3, uclr3, lclr3 = xr_limits(xb3, r3, 5)
show("coded: X, R, UCL, LCL, UCL_R, LCL_R",
     f"{xbb3:.4f}, {rb3:.5f}, {ucl3:.4f}, {lcl3:.4f}, {uclr3:.4f}, {lclr3:g}")
for label, v in (("X", xbb3), ("UCL", ucl3), ("LCL", lcl3)):
    show(f"ANSWER {label} (mm) = coded/10000 + 0.5", v / 10000 + 0.5)
for label, v in (("R", rb3), ("UCL_R", uclr3)):
    show(f"ANSWER {label} (mm) = coded/10000", v / 10000)
show("points outside X-bar limits", str([i + 1 for i in range(24) if not lcl3 <= xb3[i] <= ucl3]))
show("points outside R limits", str([i + 1 for i in range(24) if not lclr3 <= r3[i] <= uclr3]))
show("sample 9: X-bar vs UCL (coded)", f"{xb3[8]} vs {ucl3:.4f}  (margin {ucl3 - xb3[8]:.4f})")
s_x3 = t18("A2", 5) * rb3 / 3
z3 = [(v - xbb3) / s_x3 for v in xb3]
show("sigma of the means A2 R/3 (coded)", s_x3)
show("2-sigma line above (coded)", xbb3 + 2 * s_x3)
show("WE rules on the 24 means (signal points)", str(we_rules(z3)))
show("zones", str([zone(v) for v in z3]))
ws3b = sheet("__Xbar R kaart data - berekeningen oefening 3.xlsx", "Berekening")
check_cell("oef 3 X (C26)", xbb3 / 10000 + 0.5, ws3b["C26"].value)
check_cell("oef 3 R (E26)", rb3 / 10000, ws3b["E26"].value)
check_cell("oef 3 UCL (B34)", ucl3 / 10000 + 0.5, ws3b["B34"].value)
check_cell("oef 3 LCL (B35)", lcl3 / 10000 + 0.5, ws3b["B35"].value)
check_cell("oef 3 UCL_R (B41)", uclr3 / 10000, ws3b["B41"].value)
show("workbook G10:G11 next to sample 9", f"{ws3b['G10'].value} {ws3b['G11'].value}")
keep = [i for i in range(24) if i != 8]
xbb3r, rb3r, ucl3r, lcl3r, uclr3r, lclr3r = xr_limits(xb3[keep], r3[keep], 5)
show("revised without sample 9, coded: X, R, UCL, LCL, UCL_R",
     f"{xbb3r:.5f}, {rb3r:.6f}, {ucl3r:.4f}, {lcl3r:.4f}, {uclr3r:.4f}")
for label, v in (("X", xbb3r), ("UCL", ucl3r), ("LCL", lcl3r)):
    show(f"ANSWER revised {label} (mm)", v / 10000 + 0.5)
show("ANSWER revised UCL_R (mm)", uclr3r / 10000)
check_cell("oef 3 revised X (K26)", xbb3r / 10000 + 0.5, ws3b["K26"].value)
check_cell("oef 3 revised R (M26)", rb3r / 10000, ws3b["M26"].value)
check_cell("oef 3 revised UCL (J34)", ucl3r / 10000 + 0.5, ws3b["J34"].value)
check_cell("oef 3 revised LCL (J35)", lcl3r / 10000 + 0.5, ws3b["J35"].value)
check_cell("oef 3 revised UCL_R (J41)", uclr3r / 10000, ws3b["J41"].value)
show("revised: remaining points outside the revised X-bar limits",
     str([i + 1 for i in keep if not lcl3r <= xb3[i] <= ucl3r]))
sig3 = rb3r / t18("d2", 5)
check_cell("oef 3 sigma = R_bar/d2 (P27, R_bar = revised M26)", sig3 / 10000, ws3b["P27"].value)
show("ANSWER sigma = R/d2 (revised, mm)", sig3 / 10000)

# ----------------------------------------------------------------------------------------------- oefening 4
print("\n== SPC oefening 4 (SPC p. 85): new subgroup size, from the revised charts of oefening 3")
show("sigma = R/d2(5), coded", sig3)
for n_new in (3, 8):
    r_new = t18("d2", n_new) * sig3  # expected range for the new n: R = d2 sigma (inverse of sigma = R/d2)
    u = xbb3r + t18("A2", n_new) * r_new
    lo = xbb3r - t18("A2", n_new) * r_new
    u_alt = xbb3r + 3 * sig3 / math.sqrt(n_new)  # SPC p. 74 variable-n formula
    show(f"n={n_new}: R_new = d2 sigma (coded)", r_new)
    show(f"n={n_new}: UCL, LCL (coded) via A2 R_new", f"{u:.4f}, {lo:.4f}")
    show(f"n={n_new}: UCL, LCL (coded) via X +/- 3 sigma/sqrt n", f"{u_alt:.5f}, {2 * xbb3r - u_alt:.5f}")
    show(f"n={n_new}: UCL, LCL (coded) via A2 R_new, 5 decimals", f"{u:.5f}, {lo:.5f}")
    show(f"ANSWER n={n_new}: UCL (mm)", u / 10000 + 0.5)
    show(f"ANSWER n={n_new}: LCL (mm)", lo / 10000 + 0.5)
    show(f"ANSWER n={n_new}: CL_R (mm)", r_new / 10000)
    show(f"ANSWER n={n_new}: UCL_R (mm)", t18("D4", n_new) * r_new / 10000)
    show(f"ANSWER n={n_new}: LCL_R (mm)", t18("D3", n_new) * r_new / 10000)
    show(f"n={n_new}: UCL_R, LCL_R (coded)", f"{t18('D4', n_new) * r_new:.4f}, {t18('D3', n_new) * r_new:.4f}")
    cell = {3: "E48", 8: "E53"}[n_new]
    check_cell(f"oef 4 workbook 'sigma' n={n_new} = R_bar/d2({n_new}) ({cell})", rb3r / t18("d2", n_new) / 10000,
               ws3b[cell].value)
    check(f"oef 4 workbook sigma n={n_new} vs sigma = R_bar/d2(5) (in 1e-5 mm)", sig3 / 10000 * 1e5,
          round(ws3b[cell].value * 1e5, 4), 4)
print("Probability that the first subgroup after a 2-sigma shift of the mean falls outside the 3-sigma limits:")
for n_new in (3, 5, 8):
    shift = 2 * math.sqrt(n_new)  # a 2 sigma shift is 2 sqrt(n) sigma_xbar (sigma_xbar = sigma/sqrt n)
    p = stats.norm.sf(3 - shift) + stats.norm.cdf(-3 - shift)
    show(f"ANSWER n={n_new}: P(signal on the next subgroup)", p)

# ----------------------------------------------------------------------------------------------- oefening 5
print("\n== SPC oefening 5 (SPC p. 86): n = 5, sum xbar = 662.50, sum R = 9.00, m = 25, specs 26.40 +/- 0.50")
xbb5, rb5 = 662.50 / 25, 9.00 / 25
sig5 = rb5 / t18("d2", 5)
usl5, lsl5 = 26.40 + 0.50, 26.40 - 0.50
cp5 = (usl5 - lsl5) / (6 * sig5)
cpu5, cpl5 = (usl5 - xbb5) / (3 * sig5), (xbb5 - lsl5) / (3 * sig5)
p_hi, p_lo = stats.norm.sf(3 * cpu5), stats.norm.sf(3 * cpl5)
p_c = 2 * stats.norm.sf((usl5 - 26.40) / sig5)
for label, v in (("X", xbb5), ("R", rb5), ("sigma = R/d2", sig5), ("Cp", cp5), ("Cpk upper", cpu5),
                 ("Cpk lower", cpl5), ("Cpk", min(cpu5, cpl5)), ("Z upper = 3 Cpk", 3 * cpu5),
                 ("fraction above USL", p_hi), ("fraction below LSL", p_lo),
                 ("% out of spec", 100 * (p_hi + p_lo)), ("% out of spec, centred at 26.40", 100 * p_c),
                 ("UCL_xbar = X + A2 R (not asked)", xbb5 + t18("A2", 5) * rb5)):
    show(f"ANSWER {label}", v)
ws5 = sheet("__Xbar R kaart data - berekeningen oefening 5.xlsx", "Berekening")
check_cell("oef 5 sigma (H4)", sig5, ws5["H4"].value)
check_cell("oef 5 Cp (B10)", cp5, ws5["B10"].value)
check_cell("oef 5 Cpk (E11)", min(cpu5, cpl5), ws5["E11"].value)
check_cell("oef 5 total out (P11)", p_hi + p_lo, ws5["P11"].value)
check_cell("oef 5 centred total out (P17)", p_c, ws5["P17"].value)

# ----------------------------------------------------------------------------------------------- oefening 6
print("\n== SPC oefening 6 (SPC p. 87): specs 100 +/- 10, X = 104, R = 9.30, n = 5")
sig6 = 9.30 / t18("d2", 5)
cp6 = 20 / (6 * sig6)
cpl6, cpu6 = (104 - 90) / (3 * sig6), (110 - 104) / (3 * sig6)
check("oef 6 sigma = 9.3/2.326 '=~ 4'", sig6, 4, 0)
check("oef 6 Cp '0,83' (printed 20/24)", cp6, 0.83, 2)
check("oef 6 Cpk lower '1,17'", cpl6, 1.17, 2)
check("oef 6 Cpk upper '0,5'", cpu6, 0.5, 1)
for label, v in (("sigma", sig6), ("6 sigma", 6 * sig6), ("Cp", cp6), ("Cpk lower", cpl6), ("Cpk upper", cpu6),
                 ("Cpk", min(cpl6, cpu6)),
                 ("% above 110 now", 100 * stats.norm.sf((110 - 104) / sig6)),
                 ("% below 90 now", 100 * stats.norm.cdf((90 - 104) / sig6)),
                 ("% out of spec now", 100 * (stats.norm.sf(6 / sig6) + stats.norm.cdf(-14 / sig6))),
                 ("% out of spec if centred at 100", 200 * stats.norm.sf(10 / sig6))):
    show(f"ANSWER {label}", v)

# ----------------------------------------------------------------------------------------------- oefening 1
print("\n== SPC oefening 1 (SPC p. 72): points read from the figure")
# Pixel rows of the gridlines and dot centres, read from the chart image embedded in SPC p. 72 (499 x 325 px):
# UCL at y = 33.5, centre line 132.5, LCL 231.5 (so 1 sigma = 33 px). Dot centres y for samples 1-25:
DOTS_Y = [167.0, 118.1, 183.1, 214.7, 198.9, 175.9, 147.0, 101.1, 69.1, 83.9, 152.0, 183.1, 168.0, 196.8, 151.0,
          119.1, 167.8, 51.3, 140.2, 151.9, 176.1, 186.2, 211.2, 224.0, 183.5]
z1 = [(132.5 - y) / 33.0 for y in DOTS_Y]
show("z (sigma units) per sample", str([round(v, 2) for v in z1]))
show("zones", str([f"{i + 1}:{zone(v)}" for i, v in enumerate(z1)]))
show("WE rules (signal points)", str(we_rules(z1)))
show("points within 0.1 sigma of a zone line (read-out uncertain)",
     str([i + 1 for i, v in enumerate(z1) if min(abs(abs(v) - k) for k in (1, 2, 3)) < 0.1]))
signs = "".join("+" if v > 0 else "-" for v in z1)
show("side of the centre line, samples 1-25", signs)

# ----------------------------------------------------------------------------------------------- Dummies
print("\n== Dummies p. 251 (Figure 10-8, I-MR chart read-outs: X 100.1, MR 1.658)")
xb, mr = 100.1, 1.658
for src, e2 in (("Dummies Table 10-2 (p. 250)", float(DUM["2"]["E2"])), ("SSD", float(SSD_R["2"]["E2"]))):
    show(f"E2 = {e2} ({src}): UCL_X, LCL_X", f"{xb + e2 * mr:.3f}, {xb - e2 * mr:.3f}")
show("ANSWER UCL_X (E2 2.659)", xb + float(DUM["2"]["E2"]) * mr)
show("ANSWER LCL_X (E2 2.659)", xb - float(DUM["2"]["E2"]) * mr)
show("ANSWER UCL_MR = D4(2) MR", t18("D4", 2) * mr)
show("ANSWER sigma = MR/d2(2)", mr / t18("d2", 2))
# Read-outs of a Minitab chart: the centre values are shown rounded, so the limits are reproduced only to the
# rounding of X (100.1) and MR (1.658); README: not a test target.
check("Dummies UCL_X 104.5", xb + float(DUM["2"]["E2"]) * mr, 104.5, 1)
show("LCL_X from the rounded X (printed 95.71)", xb - float(DUM["2"]["E2"]) * mr)
show("X that reproduces LCL 95.71 (then UCL = 104.53 -> '104.5')", 95.71 + float(DUM["2"]["E2"]) * mr)
show("UCL_MR from the rounded MR (printed 5.418)", t18("D4", 2) * mr)
show("MR that reproduces 5.418", 5.418 / t18("D4", 2))
xb_r, mr_r = 100.12, 1.6584
show("with X = 100.12, MR = 1.6584: UCL_X, LCL_X, UCL_MR",
     f"{xb_r + float(DUM['2']['E2']) * mr_r:.4f}, {xb_r - float(DUM['2']['E2']) * mr_r:.4f}, {t18('D4', 2) * mr_r:.4f}")

print("\n== Dummies p. 252 (Figure 10-9, Xbar-R read-outs: X 84.50, R 5.750)")
xb, rb = 84.50, 5.750
show("implied A2 = (UCL - X)/R", (87.82 - xb) / rb)
show("implied D4 = UCL_R/R", 12.16 / rb)
show("ANSWER n", 5)
show("ANSWER UCL", xb + t18("A2", 5) * rb)
show("ANSWER LCL", xb - t18("A2", 5) * rb)
show("ANSWER UCL_R (D4 2.115)", t18("D4", 5) * rb)
check("Dummies UCL 87.82", xb + t18("A2", 5) * rb, 87.82, 2)
check("Dummies LCL 81.18", xb - t18("A2", 5) * rb, 81.18, 2)
check("Dummies UCL_R 12.16 (D4 2.115)", t18("D4", 5) * rb, 12.16, 2)
check("Dummies UCL_R 12.16 (D4 2.114, Dummies table)", float(DUM["5"]["D4"]) * rb, 12.16, 2)

print("\n== Dummies p. 256 (Figure 10-11, u chart, data in inventory/worked_examples.json S08-WE21)")
we = {w["id"]: w for w in json.loads((ROOT / "inventory" / "worked_examples.json").read_text())["examples"]}
data = we["S08-WE21"]["given"]["data_table"]["rows"]
nn = np.array([r[1] for r in data], float)
cc = np.array([r[2] for r in data], float)
ubar = cc.sum() / nn.sum()
show("total defects / total units", f"{cc.sum():g} / {nn.sum():g}")
show("ANSWER u-bar", ubar)
check("Dummies u-bar 1.870", ubar, 1.870, 3)
ucl_u = ubar + 3 * np.sqrt(ubar / nn)
lcl_u = ubar - 3 * np.sqrt(ubar / nn)
ui = cc / nn
for i in range(20):
    flag = "ABOVE" if ui[i] > ucl_u[i] else "BELOW" if ui[i] < lcl_u[i] else ""
    show(f"subgroup {i + 1}: n, u_i, LCL_i, UCL_i", f"{nn[i]:g}, {ui[i]:.3f}, {lcl_u[i]:.3f}, {ucl_u[i]:.3f}", flag)
show("half width 3 sqrt(u/n) for n = 65", 3 * math.sqrt(ubar / 65))
show("ANSWER UCL for subgroup 20 (n = 65)", ucl_u[-1])
show("ANSWER LCL for subgroup 20 (n = 65)", lcl_u[-1])
check("Dummies '-3.0SL = 1.361' (last subgroup)", lcl_u[-1], 1.361, 3)
check("Dummies '3.0SL = 2379' (no decimal point)", ucl_u[-1], 2379, 0)
check("Dummies 3.0SL read as 2.379", ucl_u[-1], 2.379, 3)
above = [i + 1 for i in range(20) if ui[i] > ucl_u[i]]
below = [i + 1 for i in range(20) if ui[i] < lcl_u[i]]
show("ANSWER subgroups above UCL", str(above))
show("ANSWER subgroups below LCL", str(below))
show("ANSWER number of subgroups outside", len(above) + len(below))

print("\n== Dummies p. 255 (Figure 10-10, p chart read-outs: p 0.2232, limits 0.2777 / 0.1687)")
show("half width 3 sqrt(p(1-p)/n)", 0.2777 - 0.2232)
show("implied n of the last subgroup", 9 * 0.2232 * (1 - 0.2232) / (0.2777 - 0.2232) ** 2)

print()
print(f"{len(MISMATCHES)} MISMATCH line(s) (printed course values that disagree; see 10_errata.tsv): "
      + "; ".join(MISMATCHES))
